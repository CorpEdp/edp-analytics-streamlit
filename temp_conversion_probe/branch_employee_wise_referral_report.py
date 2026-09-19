"""
Auto-converted from: App-BranchEmployeeWiseReferralReport.py
Review this file before using — the converter does a best-effort wrap;
double-check indentation around any unusual control flow (loops, if/else
blocks that span large sections, etc.).
"""

import streamlit as st
import pandas as pd
import numpy as np
import calendar
import re
from io import BytesIO
from datetime import datetime
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import warnings
def normalize_branch(branch_name):
    if pd.isna(branch_name):
        return "Skip"
    key = str(branch_name).strip().lower()
    key = " ".join(key.split())
    if not key:
        return "Skip"
    if key in BRANCH_MAPPING:
        return BRANCH_MAPPING[key]
    for k, v in BRANCH_MAPPING.items():
        k_norm = " ".join(k.split())
        if k_norm and k_norm in key:
            return v
    return "Skip"
def _is_telecaller_branch(branch):
    return any(k in str(branch).lower() for k in TELECALLER_KEYWORDS)
def _fixed_rank(branch):
    branch_l = str(branch).lower()
    for i, aliases in enumerate(FIXED_BRANCH_GROUPS):
        if any(alias in branch_l for alias in aliases):
            return i
    return None
def build_branch_order(all_branches):
    all_branches = sorted(set(b for b in all_branches if pd.notna(b) and str(b).strip() != ''))
    telecaller = sorted([b for b in all_branches if _is_telecaller_branch(b)])
    remaining = [b for b in all_branches if not _is_telecaller_branch(b)]
    fixed_present = [b for b in remaining if _fixed_rank(b) is not None]
    fixed_present.sort(key=_fixed_rank)
    new_branches = sorted([b for b in remaining if _fixed_rank(b) is None])
    return fixed_present + new_branches + telecaller
def sort_branches_df(df, ordered_branches, branch_col='Branch'):
    if branch_col not in df.columns or len(df) == 0:
        return df
    order_map = {b: i for i, b in enumerate(ordered_branches)}
    df = df.copy()
    df['_branch_sort_key'] = df[branch_col].map(lambda b: order_map.get(b, len(ordered_branches)))
    df = df.sort_values('_branch_sort_key', kind='stable').drop(columns=['_branch_sort_key'])
    return df
def clean_phone_scalar(value):
    """Normalize phone: keep last 10 digits, strip country codes."""
    if pd.isna(value):
        return ""
    s = str(value).strip()
    if s.endswith(".0"):
        s = s[:-2]
    digits = re.sub(r'\D', '', s)
    if not digits:
        return ""
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    elif len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    if len(digits) > 10:
        digits = digits[-10:]
    return digits
def clean_phone(series):
    return series.apply(clean_phone_scalar)
def _parse_date_flexible(date_val):
    """Parse date into normalized pd.Timestamp (no time component)."""
    if pd.isna(date_val) or date_val == '':
        return pd.NaT
    try:
        if isinstance(date_val, (datetime, pd.Timestamp)):
            return pd.Timestamp(date_val).normalize()
        date_str = str(date_val).strip()
        if date_str == '':
            return pd.NaT
        formats = [
            '%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d',
            '%m-%d-%Y', '%m/%d/%Y',
            '%d-%m-%y', '%d/%m/%y',
            '%b %d, %Y', '%d %b %Y'
        ]
        for fmt in formats:
            try:
                return pd.Timestamp(datetime.strptime(date_str, fmt)).normalize()
            except Exception:
                continue
        return pd.to_datetime(date_str).normalize()
    except Exception:
        return pd.NaT
def clean_amount_scalar(value):
    """Convert currency string to float."""
    if pd.isna(value):
        return np.nan
    s = re.sub(r'[^\d\.]', '', str(value))
    if s == '' or s == '.':
        return np.nan
    try:
        return float(s)
    except Exception:
        return np.nan
def find_column(df, candidates):
    lookup = {str(c).strip().lower(): c for c in df.columns}
    for cand in candidates:
        key = cand.strip().lower()
        if key in lookup:
            return lookup[key]
    return None
def clean_columns(df):
    df.columns = (df.columns.astype(str).str.strip().str.replace(r"\s+", " ", regex=True))
    return df
def clean_amount(series):
    return pd.to_numeric(
        series.astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("₹", "", regex=False)
        .str.replace("Rs.", "", regex=False)
        .str.replace("Rs", "", regex=False)
        .str.strip(),
        errors="coerce"
    ).fillna(0)
def clean_text(series):
    return (series.astype(str).str.strip().replace(["nan", "None", "NaN", ""], np.nan))
def round_value(value):
    if pd.isna(value):
        return 0
    return round(value)
def round_df(df, exclude_columns=None):
    result = df.copy()
    exclude_columns = exclude_columns or []
    for col in result.columns:
        if col in exclude_columns:
            continue
        if pd.api.types.is_numeric_dtype(result[col]):
            result[col] = result[col].apply(round_value)
    return result
def classify_scheme_type(scheme_name):
    if pd.isna(scheme_name):
        return "sessional"
    text = str(scheme_name).lower().strip()
    for key in DAILY_SCHEME_KEYS:
        if key in text:
            return "daily"
    return "sessional"
def get_scheme_list(df, scheme_type="all"):
    if "Scheme" not in df.columns:
        return []
    schemes = sorted(df["Scheme"].dropna().unique().tolist())
    if scheme_type == "daily":
        return [s for s in schemes if classify_scheme_type(s) == "daily"]
    elif scheme_type == "sessional":
        return [s for s in schemes if classify_scheme_type(s) == "sessional"]
    return schemes
def apply_cell_style(cell, fill=None, font=None, alignment=None, number_format=None, border=None):
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if alignment:
        cell.alignment = alignment
    if number_format:
        cell.number_format = number_format
    if border:
        cell.border = border
def write_section_header(worksheet, row, title, max_column, color="4472C4"):
    worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_column)
    cell = worksheet.cell(row=row, column=1)
    cell.value = title
    cell.font = Font(size=14, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[row].height = 25
def format_dataframe_section(worksheet, dataframe, start_row, header_rows=1, total_identifier=None):
    if dataframe.empty:
        return
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col_idx, column_name in enumerate(dataframe.columns, start=1):
        cell = worksheet.cell(row=start_row, column=col_idx)
        cell.value = column_name
        apply_cell_style(cell, fill=header_fill, font=Font(bold=True),
                         alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
                         border=thin_border)
    for row_idx in range(start_row + 1, start_row + 1 + len(dataframe)):
        is_total = False
        if total_identifier is not None:
            fv = worksheet.cell(row=row_idx, column=1).value
            if fv is not None and total_identifier in str(fv):
                is_total = True
        for col_idx in range(1, len(dataframe.columns) + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            fill = total_fill if is_total else None
            font = Font(bold=True) if is_total else None
            number_format = '#,##0' if col_idx > 1 else None
            apply_cell_style(cell, fill=fill, font=font,
                             alignment=Alignment(horizontal="center", vertical="center"),
                             number_format=number_format, border=thin_border)
def write_projection_section(worksheet, start_row, enrollment_df, collection_df, month_label, thin_border):
    if enrollment_df.empty and collection_df.empty:
        return
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    sub_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    title = "📈 MONTHLY PROJECTION"
    if month_label:
        title = f"{title} - {month_label}"
    write_section_header(worksheet, start_row, title, 7, "ED7D31")
    sub_row = start_row + 1
    header_row = start_row + 2
    data_start_row = start_row + 3
    worksheet.merge_cells(start_row=sub_row, start_column=1, end_row=sub_row, end_column=3)
    c = worksheet.cell(row=sub_row, column=1)
    c.value = "Projection of this month's first enrollments"
    c.font = Font(bold=True, italic=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = sub_fill
    worksheet.merge_cells(start_row=sub_row, start_column=5, end_row=sub_row, end_column=7)
    c = worksheet.cell(row=sub_row, column=5)
    c.value = "Projection of this month's overall collection"
    c.font = Font(bold=True, italic=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = sub_fill
    for offset, label in enumerate(["Scheme", "Projected Count", "Projected Amount"]):
        for base_col in (1, 5):
            cell = worksheet.cell(row=header_row, column=base_col + offset)
            cell.value = label
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
    row_count = max(len(enrollment_df), len(collection_df))
    for r in range(row_count):
        row_idx = data_start_row + r
        if r < len(enrollment_df):
            erow = enrollment_df.iloc[r]
            worksheet.cell(row=row_idx, column=1).value = erow["Scheme"]
            worksheet.cell(row=row_idx, column=2).value = erow["Projected Count"]
            worksheet.cell(row=row_idx, column=3).value = erow["Projected Amount"]
        if r < len(collection_df):
            crow = collection_df.iloc[r]
            worksheet.cell(row=row_idx, column=5).value = crow["Scheme"]
            worksheet.cell(row=row_idx, column=6).value = crow["Projected Count"]
            worksheet.cell(row=row_idx, column=7).value = crow["Projected Amount"]
        for col_idx in (1, 2, 3, 5, 6, 7):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            if col_idx in (2, 3, 6, 7):
                cell.number_format = '#,##0'
def write_avg_ticket_section(worksheet, start_row, avg_ticket_data, thin_border):
    if avg_ticket_data.empty:
        return
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    write_section_header(worksheet, start_row, "📊 AVERAGE TICKET SIZE COMPARISON",
                         len(avg_ticket_data.columns), "70AD47")
    header_row = start_row + 1
    data_start_row = start_row + 2
    for col_idx, col_name in enumerate(avg_ticket_data.columns, start=1):
        cell = worksheet.cell(row=header_row, column=col_idx)
        cell.value = col_name
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    for r, row in enumerate(avg_ticket_data.values):
        row_idx = data_start_row + r
        is_total = "Total" in str(row[0]) if len(row) > 0 else False
        for col_idx, value in enumerate(row, start=1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            col_name = avg_ticket_data.columns[col_idx - 1]
            if col_name == "Scheme":
                cell.value = value
            elif col_name in ["First Enrollment Count", "Collection Count"]:
                if isinstance(value, (int, float)):
                    cell.value = value
                    cell.number_format = '#,##0'
                else:
                    cell.value = value
            elif col_name in ["First Enrollment Amount", "Collection Amount",
                              "First Enrollment Avg Ticket", "Collection Avg Ticket", "Difference"]:
                if isinstance(value, (int, float)):
                    cell.value = value
                    cell.number_format = '#,##0'
                else:
                    cell.value = value
            elif col_name == "% Change":
                if isinstance(value, (int, float)):
                    cell.value = value
                    cell.number_format = '0.00%'
                else:
                    cell.value = value
            if is_total:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
def write_sheet(workbook, sheet_data):
    worksheet = workbook.create_sheet(sheet_data["sheet_name"])
    summary_df = sheet_data["summary"]
    enrollment_df = sheet_data["enrollment"]
    collection_df = sheet_data["collection"]
    unique_df = sheet_data["unique"]
    date_range = sheet_data["date_range"]
    report_title = sheet_data["report_title"]
    avg_ticket_data = sheet_data.get("avg_ticket_data", pd.DataFrame())

    title_color = "203764"
    section_color = "4472C4"
    scheme_color = "5B9BD5"
    total_color = "C00000"
    light_blue = "D9E1F2"
    gold_color = "FFC000"

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    max_cols = max(
        len(summary_df.columns) if not summary_df.empty else 1,
        len(enrollment_df.columns) if not enrollment_df.empty else 1,
        len(collection_df.columns) if not collection_df.empty else 1,
        len(unique_df.columns) if not unique_df.empty else 1,
        len(avg_ticket_data.columns) if not avg_ticket_data.empty else 1,
        8
    )

    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_cols)
    tc = worksheet.cell(row=1, column=1)
    tc.value = report_title
    tc.font = Font(size=16, bold=True, color="FFFFFF")
    tc.fill = PatternFill(start_color=title_color, end_color=title_color, fill_type="solid")
    tc.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 30

    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_cols)
    dc = worksheet.cell(row=2, column=1)
    dc.value = f"Report Period: {date_range[0]} to {date_range[1]}"
    dc.font = Font(size=12, bold=True)
    dc.alignment = Alignment(horizontal="center", vertical="center")

    summary_section_row = 4
    write_section_header(worksheet, summary_section_row, "📊 SCHEME SUMMARY",
                         len(summary_df.columns) if not summary_df.empty else max_cols, section_color)
    summary_header_row = summary_section_row + 1
    if not summary_df.empty:
        for col_idx, col_name in enumerate(summary_df.columns, start=1):
            worksheet.cell(row=summary_header_row, column=col_idx).value = col_name
        for row_idx, row in enumerate(summary_df.values, start=summary_header_row + 1):
            for col_idx, value in enumerate(row, start=1):
                worksheet.cell(row=row_idx, column=col_idx).value = value
    format_dataframe_section(worksheet, summary_df, summary_header_row, total_identifier="Grand Total")

    avg_ticket_section_row = summary_header_row + len(summary_df) + 3
    if not avg_ticket_data.empty:
        write_avg_ticket_section(worksheet, avg_ticket_section_row, avg_ticket_data, thin_border)
        avg_ticket_last_row = avg_ticket_section_row + 2 + len(avg_ticket_data) + 1
    else:
        avg_ticket_last_row = summary_header_row + len(summary_df)

    enrollment_projection_df = sheet_data.get("enrollment_projection", pd.DataFrame())
    collection_projection_df = sheet_data.get("collection_projection", pd.DataFrame())
    projection_month_label = sheet_data.get("projection_month_label", "")
    has_projection = not enrollment_projection_df.empty or not collection_projection_df.empty
    projection_section_row = avg_ticket_last_row + 3
    if has_projection:
        write_projection_section(worksheet, projection_section_row,
                                 enrollment_projection_df, collection_projection_df,
                                 projection_month_label, thin_border)
        prc = max(len(enrollment_projection_df), len(collection_projection_df))
        projection_last_row = projection_section_row + 2 + prc
    else:
        projection_last_row = avg_ticket_last_row

    enrollment_section_row = projection_last_row + 3
    write_section_header(worksheet, enrollment_section_row, "📅 DAY-WISE NEW ENROLLMENT",
                         len(enrollment_df.columns) if not enrollment_df.empty else max_cols, section_color)
    enrollment_header_row = enrollment_section_row + 1
    if not enrollment_df.empty:
        for col_idx, col_name in enumerate(enrollment_df.columns, start=1):
            worksheet.cell(row=enrollment_header_row, column=col_idx).value = col_name
        for row_idx, row in enumerate(enrollment_df.values, start=enrollment_header_row + 1):
            for col_idx, value in enumerate(row, start=1):
                worksheet.cell(row=row_idx, column=col_idx).value = value
    for col_idx, col_name in enumerate(enrollment_df.columns if not enrollment_df.empty else [], start=1):
        cell = worksheet.cell(row=enrollment_header_row, column=col_idx)
        if str(col_name).startswith("Total"):
            fill_color, text_color = total_color, "FFFFFF"
        elif "Rate" in str(col_name):
            fill_color, text_color = scheme_color, "FFFFFF"
        else:
            fill_color, text_color = light_blue, "000000"
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        cell.font = Font(bold=True, color=text_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    for row_idx in range(enrollment_header_row + 1, enrollment_header_row + 1 + len(enrollment_df)):
        is_total_row = "Days" in str(worksheet.cell(row=row_idx, column=1).value or "")
        for col_idx in range(1, (len(enrollment_df.columns) + 1) if not enrollment_df.empty else 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            if col_idx > 1:
                cell.number_format = '#,##0'
            if is_total_row:
                cell.fill = PatternFill(start_color=gold_color, end_color=gold_color, fill_type="solid")
                cell.font = Font(bold=True)

    collection_section_row = enrollment_header_row + len(enrollment_df) + 3
    write_section_header(worksheet, collection_section_row, "💰 DAY-WISE COLLECTION",
                         len(collection_df.columns) if not collection_df.empty else max_cols, section_color)
    collection_header_row = collection_section_row + 1
    if not collection_df.empty:
        for col_idx, col_name in enumerate(collection_df.columns, start=1):
            worksheet.cell(row=collection_header_row, column=col_idx).value = col_name
        for row_idx, row in enumerate(collection_df.values, start=collection_header_row + 1):
            for col_idx, value in enumerate(row, start=1):
                worksheet.cell(row=row_idx, column=col_idx).value = value
    for col_idx, col_name in enumerate(collection_df.columns if not collection_df.empty else [], start=1):
        cell = worksheet.cell(row=collection_header_row, column=col_idx)
        if str(col_name).startswith("Total"):
            fill_color, text_color = total_color, "FFFFFF"
        else:
            fill_color, text_color = light_blue, "000000"
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        cell.font = Font(bold=True, color=text_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    for row_idx in range(collection_header_row + 1, collection_header_row + 1 + len(collection_df)):
        is_total_row = "Days" in str(worksheet.cell(row=row_idx, column=1).value or "")
        for col_idx in range(1, (len(collection_df.columns) + 1) if not collection_df.empty else 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            if col_idx > 1:
                cell.number_format = '#,##0'
            if is_total_row:
                cell.fill = PatternFill(start_color=gold_color, end_color=gold_color, fill_type="solid")
                cell.font = Font(bold=True)

    unique_section_row = collection_header_row + len(collection_df) + 3
    write_section_header(worksheet, unique_section_row, "👥 UNIQUE ENROLLMENT",
                         len(unique_df.columns) if not unique_df.empty else max_cols, section_color)
    unique_header_row = unique_section_row + 1
    if not unique_df.empty:
        for col_idx, col_name in enumerate(unique_df.columns, start=1):
            worksheet.cell(row=unique_header_row, column=col_idx).value = col_name
        for row_idx, row in enumerate(unique_df.values, start=unique_header_row + 1):
            for col_idx, value in enumerate(row, start=1):
                worksheet.cell(row=row_idx, column=col_idx).value = value
    for col_idx, col_name in enumerate(unique_df.columns if not unique_df.empty else [], start=1):
        cell = worksheet.cell(row=unique_header_row, column=col_idx)
        if str(col_name) == "Grand Total":
            fill_color, text_color = total_color, "FFFFFF"
        else:
            fill_color, text_color = light_blue, "000000"
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        cell.font = Font(bold=True, color=text_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    for row_idx in range(unique_header_row + 1, unique_header_row + 1 + len(unique_df)):
        is_total_row = "Days" in str(worksheet.cell(row=row_idx, column=1).value or "")
        for col_idx in range(1, (len(unique_df.columns) + 1) if not unique_df.empty else 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            if col_idx > 1:
                cell.number_format = '#,##0'
            if is_total_row:
                cell.fill = PatternFill(start_color=gold_color, end_color=gold_color, fill_type="solid")
                cell.font = Font(bold=True)

    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is not None:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 25)

    worksheet.column_dimensions["A"].width = 16
    worksheet.freeze_panes = "A5"
    worksheet.sheet_view.showGridLines = False
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
def _fmt_cell(val):
    if val is None:
        return "-"
    if isinstance(val, float) and pd.isna(val):
        return "-"
    if isinstance(val, (int, np.integer)):
        return val if val != 0 else "-"
    if isinstance(val, (float, np.floating)):
        if val == 0:
            return "-"
        if float(val).is_integer():
            return int(val)
        return val
    if isinstance(val, str) and val.strip() == "":
        return "-"
    return val
def read_uploaded_file(uploaded):
    if uploaded is None:
        return None
    try:
        if uploaded.name.lower().endswith(".csv"):
            return pd.read_csv(uploaded, keep_default_na=False, na_values=[])
        return pd.read_excel(uploaded, keep_default_na=False, na_values=[])
    except Exception as e:
        st.error(f"Error reading {uploaded.name}: {e}")
        return None
def _transform_branch_pretty(branch):
    """Convert raw branch name to 'Bhima Jewellery - <Title>' form."""
    if pd.isna(branch) or str(branch).strip() == "":
        return "Bhima Jewellery - Customer"
    branch_upper = str(branch).upper().strip()
    if branch_upper == "HEAD OFFICE":
        return "Bhima Jewellery - Madurai"
    elif branch_upper == "IN-TRANSIT- LOCATIONS":
        return "Bhima Jewellery - Salem"
    elif branch_upper == "APP SHOWROOM LOCATION":
        return "Bhima Jewellery - Tirunelveli"
    if "BHIMA JEWELLERY -" in branch_upper:
        parts = str(branch).split("Bhima Jewellery -", 1)
        if len(parts) > 1:
            final_branch = f"Bhima Jewellery - {parts[1].strip().title()}"
        else:
            final_branch = str(branch).title()
    elif str(branch).upper().strip().endswith("BRANCH"):
        cleaned = re.sub(r'BRANCH$', '', str(branch), flags=re.IGNORECASE).strip()
        final_branch = f"Bhima Jewellery - {cleaned.title()}"
    else:
        final_branch = f"Bhima Jewellery - {str(branch).title()}"
    final_branch = final_branch.replace("Tiruchirappalli", "Trichy")
    final_branch = final_branch.replace("TIRUCHIRAPPALLI", "Trichy")
    final_branch = final_branch.replace("tiruchirappalli", "Trichy")
    return final_branch
def build_referral_report(transactions_df, employees_df, referrals_df, bss_df=None,
                          start_date=None, end_date=None):
    """
    STRICT 3-FIELD MATCH:
      Referral.Referee Phone   == Txn.Customer Phone Number  (cleaned, 10 digits)
      Referral.Joined Date     == Txn.Date                   (same calendar day)
      Referral.Enrollment Amt  == Txn.Saved Amount           (±0.01 tolerance)
      AND Txn.Installment number == 1

    No fallbacks. No BSS. No phone-only.
    Each matching transaction counts as one row (if multiple matches, all count).

    Returns:
      (summary_pivot, daily_list, detail_df, diagnostics_dict, duplicates_df, debug_info)
    """
    empty = pd.DataFrame()
    diagnostics = {
        "total_referrals": 0, "matched_referrals": 0, "not_enrolled": 0,
        "duplicates": 0,
    }
    debug_info = {
        "referral_raw_count": 0,
        "referral_after_date_filter": 0,
        "referral_after_phone_filter": 0,
        "referrals_in_scope": 0,
        "txn_f1_count": 0,
        "txn_unique_phones": 0,
        "status_distribution": {},
        "referral_phone_sample": [],
        "txn_phone_sample": [],
        "phone_overlap": 0,
        "scheme_name_variants": [],
        "not_enrolled_breakdown": {},
    }

    if employees_df is None or referrals_df is None:
        return empty, [], empty, diagnostics, empty, debug_info
    if employees_df.empty or referrals_df.empty:
        return empty, [], empty, diagnostics, empty, debug_info

    emp = clean_columns(employees_df.copy())
    ref = clean_columns(referrals_df.copy())
    txn = transactions_df.copy()

    debug_info["referral_raw_count"] = len(ref)

    # --- Employee columns ---
    emp_ref_code_col = find_column(emp, ["Referral Code", "Referal Code", "Ref Code", "RefCode"])
    emp_branch_col = find_column(emp, ["Branch", "Branch Name", "Showroom", "Location"])
    emp_type_col = find_column(emp, ["Employee Type", "Emp Type", "Category"])
    emp_code_col = find_column(emp, ["Employee Code", "Emp Code"])

    if emp_ref_code_col is None or emp_branch_col is None:
        st.warning("⚠️ Employee file needs 'Referral Code' and 'Branch' columns.")
        return empty, [], empty, diagnostics, empty, debug_info

    # --- Referral columns ---
    ref_code_col = find_column(ref, ["Referral Code", "Referal Code", "Ref Code", "RefCode"])
    ref_phone_col = find_column(ref, ["Referee Phone", "Referee Mobile", "Referee Contact",
                                       "Referee Phone Number", "Mobile", "Phone",
                                       "Phone Number", "Mobile Number"])
    referee_name_col = find_column(ref, ["Referee Name", "Referee", "Customer Name"])
    ref_amount_col = find_column(ref, ["Enrollment Amount", "Enroll Amount", "Amount"])
    ref_status_col = find_column(ref, ["Status"])
    referrer_name_col = find_column(ref, ["Referrer Name", "Referer Name"])
    referrer_phone_col = find_column(ref, ["Referrer Phone", "Referer Phone"])
    reg_date_col = find_column(ref, ["Registered Date", "Registration Date", "Register Date"])
    join_date_col = find_column(ref, ["Joined Date", "Join Date", "Joining Date"])

    if ref_code_col is None:
        st.error("❌ Referral file needs 'Referral Code'.")
        return empty, [], empty, diagnostics, empty, debug_info
    if ref_phone_col is None:
        st.error("❌ Referral file needs 'Referee Phone'.")
        return empty, [], empty, diagnostics, empty, debug_info
    if join_date_col is None:
        st.error("❌ Referral file needs 'Joined Date' — this is essential for the strict match.")
        return empty, [], empty, diagnostics, empty, debug_info
    if ref_amount_col is None:
        st.error("❌ Referral file needs 'Enrollment Amount' — this is essential for the strict match.")
        return empty, [], empty, diagnostics, empty, debug_info

    # --- Build employee maps ---
    ref_code_to_branch = dict(zip(emp[emp_ref_code_col].astype(str).str.strip(), emp[emp_branch_col]))
    ref_code_to_empcode = {}
    if emp_code_col is not None:
        ref_code_to_empcode = dict(zip(emp[emp_ref_code_col].astype(str).str.strip(),
                                        emp[emp_code_col].astype(str)))
    ref_code_to_emptype = {}
    if emp_type_col is not None:
        ref_code_to_emptype = dict(zip(emp[emp_ref_code_col].astype(str).str.strip(),
                                        emp[emp_type_col].astype(str)))

    # --- Build referral frame ---
    fr = pd.DataFrame()
    fr['Customer Name'] = ref[referee_name_col] if referee_name_col else ""
    fr['Customer Phone'] = ref[ref_phone_col].apply(clean_phone_scalar)
    fr['Customer Enrollment Amount'] = ref[ref_amount_col].apply(clean_amount_scalar)
    fr['Status'] = ref[ref_status_col].astype(str) if ref_status_col else ""
    fr['Employee Name'] = ref[referrer_name_col] if referrer_name_col else ""
    fr['Referral Code'] = ref[ref_code_col].astype(str).str.strip()
    fr['Employee Phone'] = (
        ref[referrer_phone_col].apply(clean_phone_scalar) if referrer_phone_col else ""
    )
    fr['Employee Code'] = fr['Referral Code'].map(ref_code_to_empcode).fillna('')
    raw_branch = fr['Referral Code'].map(ref_code_to_branch)
    fr['Branch'] = raw_branch.apply(_transform_branch_pretty)
    fr['Category'] = fr['Referral Code'].map(ref_code_to_emptype).fillna('Customer')

    # --- STRICT: use Joined Date as the matching date ---
    fr['Joined Date'] = ref[join_date_col].apply(_parse_date_flexible)

    # Registered Date (for reference only)
    if reg_date_col:
        fr['Registered Date'] = ref[reg_date_col].apply(_parse_date_flexible)
    else:
        fr['Registered Date'] = pd.NaT

    # Status distribution
    if ref_status_col:
        status_counts = fr['Status'].value_counts().head(20).to_dict()
        debug_info["status_distribution"] = {str(k): int(v) for k, v in status_counts.items()}

    # --- Filter by Joined Date range ---
    if start_date is not None and end_date is not None:
        mask = pd.to_datetime(fr['Joined Date'], errors='coerce').between(
            pd.Timestamp(start_date), pd.Timestamp(end_date)
        )
        fr_filtered = fr[mask].copy()
    else:
        fr_filtered = fr.copy()

    debug_info["referral_after_date_filter"] = len(fr_filtered)

    # --- Phone filter ---
    fr_filtered = fr_filtered[fr_filtered['Customer Phone'].str.len() >= 10].copy()
    debug_info["referral_after_phone_filter"] = len(fr_filtered)
    debug_info["referrals_in_scope"] = len(fr_filtered)
    debug_info["referral_phone_sample"] = fr_filtered['Customer Phone'].head(10).tolist()

    # --- Transactions: filter to Installment #1 ---
    if 'Installment number' in txn.columns:
        txn['Installment number'] = pd.to_numeric(txn['Installment number'], errors='coerce')
        txn_f1 = txn[txn['Installment number'] == 1].copy()
    else:
        txn_f1 = txn.copy()

    txn_phone_col = find_column(txn_f1, ["Customer Phone Number", "Customer Phone", "Mobile"])
    txn_date_col = find_column(txn_f1, ["Date"])
    txn_amount_col = find_column(txn_f1, ["Saved Amount", "Amount"])
    txn_scheme_col = find_column(txn_f1, ["Scheme Name", "Scheme"])
    txn_passbook_col = find_column(txn_f1, ["Passbook number", "Passbook Number", "Doc No"])

    if txn_phone_col is None or txn_date_col is None or txn_amount_col is None:
        st.error("❌ Transactions file needs 'Customer Phone Number', 'Date', and 'Saved Amount'.")
        return empty, [], empty, diagnostics, empty, debug_info

    txn_f1['_phone'] = txn_f1[txn_phone_col].apply(clean_phone_scalar)
    txn_f1['_date'] = txn_f1[txn_date_col].apply(_parse_date_flexible)
    txn_f1['_amount'] = txn_f1[txn_amount_col].apply(clean_amount_scalar)
    txn_f1['_scheme'] = txn_f1[txn_scheme_col].astype(str) if txn_scheme_col else ''
    txn_f1['_passbook'] = txn_f1[txn_passbook_col].astype(str) if txn_passbook_col else ''

    debug_info["txn_f1_count"] = len(txn_f1)
    debug_info["txn_unique_phones"] = txn_f1['_phone'].nunique()
    debug_info["txn_phone_sample"] = txn_f1['_phone'].head(10).tolist()

    if txn_scheme_col:
        scheme_variants = txn_f1['_scheme'].value_counts().head(30).to_dict()
        debug_info["scheme_name_variants"] = {str(k): int(v) for k, v in scheme_variants.items()}

    # --- Build lookup: phone+date → list of transaction rows ---
    txn_by_phone_date = defaultdict(list)
    for idx, row in txn_f1.iterrows():
        p = row['_phone']
        d = row['_date']
        if not p or pd.isna(d):
            continue
        key = (p, d.date())
        txn_by_phone_date[key].append({
            'amount': row['_amount'],
            'scheme': row['_scheme'],
            'passbook': row['_passbook'],
            'date': d,
            'index': idx,
        })

    # Phone overlap for debug
    ref_phones = set(fr_filtered['Customer Phone'].tolist())
    txn_phones = set(txn_f1['_phone'].tolist())
    debug_info["phone_overlap"] = len(ref_phones & txn_phones)

    # --- STRICT MATCHING LOOP ---
    matched_rows = []
    not_enrolled_reasons = defaultdict(int)

    for idx, row in fr_filtered.iterrows():
        phone = row['Customer Phone']
        joined_date = row['Joined Date']
        enrollment_amount = row['Customer Enrollment Amount']

        branch = row['Branch']
        referrer = row['Employee Name']
        referee = row['Customer Name']
        emp_code = row['Employee Code']
        category = row['Category']
        reg_date = row['Registered Date']

        # Validate required fields
        if not phone:
            not_enrolled_reasons['empty phone'] += 1
            matched_rows.append({
                "Branch": branch, "Referrer Name": referrer, "Referee Name": referee,
                "Referee Phone": phone, "Employee Code": emp_code,
                "Registered Date": reg_date, "Joined Date": joined_date,
                "Transaction Date": pd.NaT, "Scheme": None, "Installment": None,
                "Enrollment Amount": float(enrollment_amount) if pd.notna(enrollment_amount) else None,
                "Paid Amount": None, "Passbook": "",
                "Match Reason": "Not Enrolled (empty phone)",
                "Category": "Not Enrolled", "Employee Category": category, "Is Duplicate": False,
            })
            continue

        if pd.isna(joined_date):
            not_enrolled_reasons['missing joined date'] += 1
            matched_rows.append({
                "Branch": branch, "Referrer Name": referrer, "Referee Name": referee,
                "Referee Phone": phone, "Employee Code": emp_code,
                "Registered Date": reg_date, "Joined Date": pd.NaT,
                "Transaction Date": pd.NaT, "Scheme": None, "Installment": None,
                "Enrollment Amount": float(enrollment_amount) if pd.notna(enrollment_amount) else None,
                "Paid Amount": None, "Passbook": "",
                "Match Reason": "Not Enrolled (missing joined date)",
                "Category": "Not Enrolled", "Employee Category": category, "Is Duplicate": False,
            })
            continue

        if pd.isna(enrollment_amount):
            not_enrolled_reasons['missing enrollment amount'] += 1
            matched_rows.append({
                "Branch": branch, "Referrer Name": referrer, "Referee Name": referee,
                "Referee Phone": phone, "Employee Code": emp_code,
                "Registered Date": reg_date, "Joined Date": joined_date,
                "Transaction Date": pd.NaT, "Scheme": None, "Installment": None,
                "Enrollment Amount": None, "Paid Amount": None, "Passbook": "",
                "Match Reason": "Not Enrolled (missing enrollment amount)",
                "Category": "Not Enrolled", "Employee Category": category, "Is Duplicate": False,
            })
            continue

        # ============================================
        # STRICT 3-FIELD MATCH
        # ============================================
        key = (phone, joined_date.date())
        candidates = txn_by_phone_date.get(key, [])

        # Filter by amount (exact, ±0.01)
        exact_matches = [
            c for c in candidates
            if pd.notna(c.get('amount')) and abs(c['amount'] - enrollment_amount) <= 0.01
        ]

        if not exact_matches:
            if not candidates:
                reason = "no txn on phone+date"
            else:
                reason = f"amount mismatch on {len(candidates)} txn(s)"
            not_enrolled_reasons[reason] += 1
            matched_rows.append({
                "Branch": branch, "Referrer Name": referrer, "Referee Name": referee,
                "Referee Phone": phone, "Employee Code": emp_code,
                "Registered Date": reg_date, "Joined Date": joined_date,
                "Transaction Date": candidates[0]['date'] if candidates else pd.NaT,
                "Scheme": candidates[0]['scheme'] if candidates else None,
                "Installment": 1,
                "Enrollment Amount": float(enrollment_amount),
                "Paid Amount": candidates[0]['amount'] if candidates else None,
                "Passbook": candidates[0]['passbook'] if candidates else "",
                "Match Reason": f"Not Enrolled ({reason})",
                "Category": "Not Enrolled", "Employee Category": category, "Is Duplicate": False,
            })
            continue

        # Success — one row per matching transaction
        for m in exact_matches:
            scheme_raw = str(m.get('scheme', '')).strip()
            passbook = str(m.get('passbook', '')).strip()
            paid_amount = m.get('amount')

            if scheme_raw:
                scheme_lower = scheme_raw.lower()
                if classify_scheme_type(scheme_raw) == "daily":
                    cat = "e-Silver" if "silver" in scheme_lower else "e-Gold"
                else:
                    cat = scheme_raw
            else:
                cat = "Not Enrolled"

            matched_rows.append({
                "Branch": branch, "Referrer Name": referrer, "Referee Name": referee,
                "Referee Phone": phone, "Employee Code": emp_code,
                "Registered Date": reg_date, "Joined Date": joined_date,
                "Transaction Date": m['date'], "Scheme": scheme_raw, "Installment": 1,
                "Enrollment Amount": float(enrollment_amount),
                "Paid Amount": float(paid_amount) if pd.notna(paid_amount) else None,
                "Passbook": passbook,
                "Match Reason": "All 3 fields matched (phone + joined date + amount)",
                "Category": cat, "Employee Category": category, "Is Duplicate": False,
            })

    detail_df = pd.DataFrame(matched_rows)
    debug_info["not_enrolled_breakdown"] = dict(not_enrolled_reasons)

    if detail_df.empty:
        return empty, [], empty, diagnostics, empty, debug_info

    # --- Duplicate detection ---
    duplicates_df = pd.DataFrame()
    has_passbook = detail_df['Passbook'].notna() & (detail_df['Passbook'].astype(str).str.strip() != '')
    if has_passbook.any():
        detail_df['_dup_key'] = (
            detail_df['Passbook'].astype(str) + '_' +
            detail_df['Referrer Name'].astype(str) + '_' +
            detail_df['Branch'].astype(str)
        )
        key_counts = detail_df[has_passbook]['_dup_key'].value_counts()
        dup_keys = key_counts[key_counts > 1].index.tolist()
        if dup_keys:
            dup_rows = []
            for grp_idx, dup_key in enumerate(dup_keys, start=1):
                grp = detail_df[detail_df['_dup_key'] == dup_key].sort_values('Joined Date')
                orig_idx = grp.index[0]
                for di in grp.index[1:]:
                    detail_df.loc[di, 'Is Duplicate'] = True
                    orig = detail_df.loc[orig_idx]
                    dup = detail_df.loc[di]
                    dup_rows.append({
                        'Group ID': f'GROUP_{grp_idx}', 'Passbook': orig['Passbook'],
                        'Original Employee': orig['Referrer Name'],
                        'Original Customer': orig['Referee Name'],
                        'Original Date': orig['Joined Date'],
                        'Original Enrollment Amount': orig['Enrollment Amount'],
                        'Duplicate Employee': dup['Referrer Name'],
                        'Duplicate Customer': dup['Referee Name'],
                        'Duplicate Date': dup['Joined Date'],
                        'Duplicate Enrollment Amount': dup['Enrollment Amount'],
                        'Branch': orig['Branch'], 'Category': orig['Category'],
                    })
            duplicates_df = pd.DataFrame(dup_rows)
            st.warning(f"⚠️ **Duplicates detected:** {len(duplicates_df)} record(s) in {len(dup_keys)} group(s).")
        detail_df = detail_df.drop(columns=['_dup_key'])

    # Diagnostics
    total = len(detail_df)
    matched = int((detail_df['Category'] != 'Not Enrolled').sum())
    not_enrolled = total - matched
    diagnostics.update({
        "total_referrals": total,
        "matched_referrals": matched,
        "not_enrolled": not_enrolled,
        "duplicates": int(detail_df['Is Duplicate'].sum()),
    })

    # Branch ordering
    branch_universe = detail_df['Branch'].dropna().unique().tolist()
    ordered_branches = build_branch_order(branch_universe)

    sessional_categories = sorted([
        c for c in detail_df['Category'].unique()
        if c not in ('e-Gold', 'e-Silver', 'Not Enrolled')
    ])
    category_order = ["e-Gold", "e-Silver"] + sessional_categories + ["Not Enrolled"]

    pivot = detail_df.pivot_table(
        index="Branch", columns="Category",
        values="Referee Phone", aggfunc="count", fill_value=0
    )
    for cat in category_order:
        if cat not in pivot.columns:
            pivot[cat] = 0
    pivot = pivot[category_order]
    for b in ordered_branches:
        if b not in pivot.index:
            pivot.loc[b] = 0
    pivot = pivot.reindex(ordered_branches)
    pivot["Grand Total"] = pivot.sum(axis=1)
    total_row = pivot.sum(axis=0)
    total_row.name = "Total"
    pivot = pd.concat([pivot, pd.DataFrame([total_row])])
    pivot = pivot.reset_index().rename(columns={"index": "Showroom"})

    # Daily pivots
    daily_list = []
    detail_df['Match Date'] = detail_df['Joined Date'].apply(
        lambda x: x.strftime('%d-%m-%Y') if pd.notna(x) else None
    )
    for date_label in sorted(
        [d for d in detail_df['Match Date'].dropna().unique()],
        key=lambda x: pd.to_datetime(x, format='%d-%m-%Y')
    ):
        day_df = detail_df[detail_df['Match Date'] == date_label]
        day_pivot = day_df.pivot_table(
            index="Branch", columns="Category",
            values="Referee Phone", aggfunc="count", fill_value=0
        )
        for cat in category_order:
            if cat not in day_pivot.columns:
                day_pivot[cat] = 0
        day_pivot = day_pivot[category_order]
        for b in ordered_branches:
            if b not in day_pivot.index:
                day_pivot.loc[b] = 0
        day_pivot = day_pivot.reindex(ordered_branches)
        day_pivot["Grand Total"] = day_pivot.sum(axis=1)
        day_total = day_pivot.sum(axis=0)
        day_total.name = "Total"
        day_pivot = pd.concat([day_pivot, pd.DataFrame([day_total])])
        day_pivot = day_pivot.reset_index().rename(columns={"index": "Showroom"})
        daily_list.append((date_label, day_pivot))

    return pivot, daily_list, detail_df, diagnostics, duplicates_df, debug_info
def write_referral_sheet(workbook, summary_df, daily_list, detail_df, duplicates_df,
                         date_range, sheet_name="Referral Report"):
    if summary_df is None or summary_df.empty:
        return
    ws = workbook.create_sheet(sheet_name)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    total_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    dup_fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
    section_color = "7030A0"

    max_cols = len(summary_df.columns)
    if daily_list:
        for _, day_df in daily_list:
            max_cols = max(max_cols, len(day_df.columns))
    if detail_df is not None and not detail_df.empty:
        max_cols = max(max_cols, len(detail_df.columns))

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_cols)
    tc = ws.cell(row=1, column=1)
    tc.value = "🎁 REFERRAL CONVERSION REPORT (Strict 3-Field Match)"
    tc.font = Font(size=16, bold=True, color="FFFFFF")
    tc.fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
    tc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_cols)
    dc = ws.cell(row=2, column=1)
    dc.value = f"Report Period: {date_range[0]} to {date_range[1]}  |  Match: Phone + Joined Date + Amount + Installment #1"
    dc.font = Font(size=12, bold=True)
    dc.alignment = Alignment(horizontal="center", vertical="center")

    write_section_header(ws, 4,
                         f"🏢 BRANCH-WISE REFERRAL SUMMARY  ({date_range[0]} to {date_range[1]})",
                         max_cols, section_color)
    summary_header_row = 5
    for c_idx, col_name in enumerate(summary_df.columns, start=1):
        cell = ws.cell(row=summary_header_row, column=c_idx)
        cell.value = col_name
        is_total_col = col_name == "Grand Total"
        cell.font = Font(bold=True, color="FFFFFF" if is_total_col else "000000")
        cell.fill = PatternFill(
            start_color="C00000" if is_total_col else "D9E1F2",
            end_color="C00000" if is_total_col else "D9E1F2",
            fill_type="solid"
        )
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    for r_idx, row in enumerate(summary_df.values, start=summary_header_row + 1):
        is_total_row = str(row[0]).strip().lower() == "total"
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.value = _fmt_cell(val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin
            if c_idx > 1 and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'
            if is_total_row:
                cell.font = Font(bold=True)
                cell.fill = total_fill

    last_row = summary_header_row + len(summary_df)

    if daily_list:
        for date_label, day_df in daily_list:
            last_row += 3
            write_section_header(ws, last_row, f"📅 REFERRAL — {date_label}",
                                 max(len(day_df.columns), max_cols), section_color)
            d_head_row = last_row + 1
            for c_idx, col_name in enumerate(day_df.columns, start=1):
                cell = ws.cell(row=d_head_row, column=c_idx)
                cell.value = col_name
                is_total_col = col_name == "Grand Total"
                cell.font = Font(bold=True, color="FFFFFF" if is_total_col else "000000")
                cell.fill = PatternFill(
                    start_color="C00000" if is_total_col else "D9E1F2",
                    end_color="C00000" if is_total_col else "D9E1F2",
                    fill_type="solid"
                )
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin
            for r_idx, row in enumerate(day_df.values, start=d_head_row + 1):
                is_total_row = str(row[0]).strip().lower() == "total"
                for c_idx, val in enumerate(row, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx)
                    cell.value = _fmt_cell(val)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin
                    if c_idx > 1 and isinstance(cell.value, (int, float)):
                        cell.number_format = '#,##0'
                    if is_total_row:
                        cell.font = Font(bold=True)
                        cell.fill = total_fill
            last_row = d_head_row + len(day_df)

    if detail_df is not None and not detail_df.empty:
        last_row += 3
        write_section_header(ws, last_row, "📋 REFERRAL MATCH DETAILS", max_cols, section_color)
        detail_header_row = last_row + 1
        display_cols = [c for c in detail_df.columns if c not in ("Referee Phone", "Match Date")]
        detail_display = detail_df[display_cols].copy()
        for col in ["Registered Date", "Joined Date", "Transaction Date"]:
            if col in detail_display.columns:
                detail_display[col] = detail_display[col].apply(
                    lambda x: x.strftime("%d-%m-%Y") if pd.notna(x) else ""
                )
        for c_idx, col_name in enumerate(detail_display.columns, start=1):
            cell = ws.cell(row=detail_header_row, column=c_idx)
            cell.value = col_name
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin
        is_dup_col_idx = None
        if "Is Duplicate" in detail_display.columns:
            is_dup_col_idx = list(detail_display.columns).index("Is Duplicate") + 1
        for r_idx, row in enumerate(detail_display.values, start=detail_header_row + 1):
            is_dup = bool(row[is_dup_col_idx - 1]) if is_dup_col_idx else False
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin
                if is_dup:
                    cell.fill = dup_fill

    if duplicates_df is not None and not duplicates_df.empty:
        dup_ws = workbook.create_sheet("Duplicate Records")
        for c_idx, col_name in enumerate(duplicates_df.columns, start=1):
            cell = dup_ws.cell(row=1, column=c_idx)
            cell.value = col_name
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin
        for r_idx, row in enumerate(duplicates_df.values, start=2):
            for c_idx, val in enumerate(row, start=1):
                cell = dup_ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin
                cell.fill = dup_fill
        for col_idx in range(1, len(duplicates_df.columns) + 1):
            letter = get_column_letter(col_idx)
            max_len = 12
            for cell in dup_ws[letter]:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            dup_ws.column_dimensions[letter].width = min(max_len + 2, 30)
        dup_ws.sheet_view.showGridLines = False

    for col_idx in range(1, max_cols + 1):
        letter = get_column_letter(col_idx)
        max_len = 12
        for cell in ws[letter]:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 28)

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
def create_formatted_excel(daily_data, sessional_data, date_range, report_title,
                           referral_summary=None, referral_daily=None,
                           referral_detail=None, referral_duplicates=None):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        workbook = writer.book
        if "Sheet" in workbook.sheetnames:
            del workbook["Sheet"]
        sheets_written = 0
        if daily_data["schemes"]:
            write_sheet(workbook, daily_data)
            sheets_written += 1
        if sessional_data["schemes"]:
            write_sheet(workbook, sessional_data)
            sheets_written += 1
        if referral_summary is not None and not referral_summary.empty:
            write_referral_sheet(
                workbook, referral_summary,
                referral_daily if referral_daily is not None else [],
                referral_detail if referral_detail is not None else pd.DataFrame(),
                referral_duplicates if referral_duplicates is not None else pd.DataFrame(),
                date_range, sheet_name="Referral Report"
            )
            sheets_written += 1
        if sheets_written == 0:
            placeholder = workbook.create_sheet("No Data")
            placeholder["A1"] = "No data matched the selected filters."
    output.seek(0)
    return output.getvalue()
def add_avg_ticket_size_comparison(summary_df):
    if summary_df.empty:
        return pd.DataFrame()
    scheme_rows = summary_df[summary_df["Scheme"] != "Grand Total"].copy()
    if scheme_rows.empty:
        return pd.DataFrame()
    result_data = []
    for _, row in scheme_rows.iterrows():
        fc = row["First Enrollment Count"]
        fa = row["First Enrollment Amount"]
        cc = row["Collection Count"]
        ca = row["Collection Amount"]
        f_avg = fa / fc if fc > 0 else 0
        c_avg = ca / cc if cc > 0 else 0
        diff = c_avg - f_avg
        if f_avg > 0:
            pct_change = round(diff / f_avg, 4)
        elif c_avg > 0:
            pct_change = "N/A"
        else:
            pct_change = 0
        result_data.append({
            "Scheme": row["Scheme"],
            "First Enrollment Count": int(fc),
            "First Enrollment Amount": int(fa),
            "First Enrollment Avg Ticket": int(round(f_avg)),
            "Collection Count": int(cc),
            "Collection Amount": int(ca),
            "Collection Avg Ticket": int(round(c_avg)),
            "Difference": int(round(diff)),
            "% Change": pct_change,
        })
    tf = sum(d["First Enrollment Count"] for d in result_data)
    ta = sum(d["First Enrollment Amount"] for d in result_data)
    cf = sum(d["Collection Count"] for d in result_data)
    ca = sum(d["Collection Amount"] for d in result_data)
    tfa = ta / tf if tf > 0 else 0
    tca = ca / cf if cf > 0 else 0
    tdiff = tca - tfa
    tpct = round(tdiff / tfa, 4) if tfa > 0 else ("N/A" if tca > 0 else 0)
    result_data.append({
        "Scheme": "Grand Total", "First Enrollment Count": int(tf),
        "First Enrollment Amount": int(ta), "First Enrollment Avg Ticket": int(round(tfa)),
        "Collection Count": int(cf), "Collection Amount": int(ca),
        "Collection Avg Ticket": int(round(tca)),
        "Difference": int(round(tdiff)), "% Change": tpct,
    })
    return pd.DataFrame(result_data)
def display_avg_ticket_comparison(avg_ticket_df):
    if avg_ticket_df.empty:
        return
    st.subheader("📊 Average Ticket Size Comparison")
    display_df = avg_ticket_df.copy()
    for col in ["First Enrollment Count", "Collection Count"]:
        if col in display_df.columns:
            display_df[col] = display_df[col].apply(
                lambda x: f"{int(x):,}" if isinstance(x, (int, float)) else x
            )
    for col in ["First Enrollment Amount", "Collection Amount",
                "First Enrollment Avg Ticket", "Collection Avg Ticket", "Difference"]:
        if col in display_df.columns:
            display_df[col] = display_df[col].apply(
                lambda x: f"{int(x):,}" if isinstance(x, (int, float)) else x
            )
    if "% Change" in display_df.columns:
        display_df["% Change"] = display_df["% Change"].apply(
            lambda x: f"{x * 100:.2f}%" if isinstance(x, (int, float)) else x
        )
    st.dataframe(display_df, use_container_width=True, hide_index=True)
def get_rate_column_name(scheme):
    if pd.isna(scheme):
        return None
    sl = str(scheme).lower()
    if "e-gold" in sl or "egold" in sl or "e gold" in sl:
        return "Gold Rate"
    elif "e-silver" in sl or "esilver" in sl or "e silver" in sl:
        return "Silver Rate"
    if "christmas" in sl or "diwali" in sl or "pongal" in sl:
        return None
    elif "akshaya tritiya" in sl:
        return "Gold Rate"
    return f"{scheme} Rate"
def generate_report_data(df, schemes):
    if not schemes or df.empty:
        return (pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), pd.DataFrame())

    installment_one = df[(df["Installment number"] == 1) & (df["Passbook number"].notna())].copy()
    installment_one = installment_one.sort_values(["Passbook number", "Date", "Id"])
    first_df = installment_one.drop_duplicates(subset=["Passbook number"], keep="first").copy()
    collection_df = df[df["Installment number"].notna() & (df["Installment number"] != 1)].copy()

    first_df["Customer Key"] = first_df["Customer Phone Number"].fillna("").astype(str).str.strip()
    first_df["Valid Phone"] = (
        first_df["Customer Key"].notna()
        & (first_df["Customer Key"] != "")
        & (first_df["Customer Key"] != "nan")
    )

    summary_rows = []
    for scheme in schemes:
        s_first = first_df[first_df["Scheme"] == scheme]
        s_coll = collection_df[collection_df["Scheme"] == scheme]
        fc = len(s_first)
        fa = s_first["Saved Amount"].sum()
        ud = s_first["Date"].nunique()
        cc = len(s_coll)
        ca = s_coll["Saved Amount"].sum()
        cd = s_coll["Date"].nunique()
        summary_rows.append({
            "Scheme": scheme, "First Enrollment Count": fc, "First Enrollment Amount": fa,
            "Avg Count/Day": fc / ud if ud > 0 else 0,
            "Avg Amount/Day": fa / ud if ud > 0 else 0,
            "Collection Count": cc, "Collection Amount": ca,
            "Avg Collection Count/Day": cc / cd if cd > 0 else 0,
            "Avg Collection Amount/Day": ca / cd if cd > 0 else 0,
        })
    if summary_rows:
        gfc = sum(r["First Enrollment Count"] for r in summary_rows)
        gfa = sum(r["First Enrollment Amount"] for r in summary_rows)
        gcc = sum(r["Collection Count"] for r in summary_rows)
        gca = sum(r["Collection Amount"] for r in summary_rows)
        td = df["Date"].nunique()
        summary_rows.append({
            "Scheme": "Grand Total", "First Enrollment Count": gfc, "First Enrollment Amount": gfa,
            "Avg Count/Day": gfc / td if td > 0 else 0,
            "Avg Amount/Day": gfa / td if td > 0 else 0,
            "Collection Count": gcc, "Collection Amount": gca,
            "Avg Collection Count/Day": gcc / td if td > 0 else 0,
            "Avg Collection Amount/Day": gca / td if td > 0 else 0,
        })
    summary_df = pd.DataFrame(summary_rows)
    avg_ticket_df = add_avg_ticket_size_comparison(summary_df)

    dates = sorted(df["Date"].dropna().unique())

    enrollment_rows = []
    for date in dates:
        row = {"Date": pd.Timestamp(date).strftime("%d-%m-%Y")}
        tc, ta = 0, 0
        for scheme in schemes:
            rc = get_rate_column_name(scheme)
            sd = df[(df["Date"] == date) & (df["Scheme"] == scheme)]
            rate = sd["Metal Rate"].iloc[0] if not sd.empty else 0
            if rc is not None:
                row[rc] = rate
            temp = first_df[(first_df["Date"] == date) & (first_df["Scheme"] == scheme)]
            cnt, amt = len(temp), temp["Saved Amount"].sum()
            row[f"{scheme} Count"] = cnt
            row[f"{scheme} Amount"] = amt
            row[f"{scheme} Avg Ticket"] = amt / cnt if cnt > 0 else 0
            tc += cnt
            ta += amt
        row["Total Count"] = tc
        row["Total Amount"] = ta
        row["Total Avg Ticket"] = ta / tc if tc > 0 else 0
        enrollment_rows.append(row)
    enrollment_df = pd.DataFrame(enrollment_rows)

    enrollment_columns = ["Date"]
    for scheme in schemes:
        rc = get_rate_column_name(scheme)
        if rc is not None:
            enrollment_columns.append(rc)
        enrollment_columns.extend([f"{scheme} Count", f"{scheme} Amount", f"{scheme} Avg Ticket"])
    enrollment_columns.extend(["Total Count", "Total Amount", "Total Avg Ticket"])
    enrollment_df = enrollment_df[[c for c in enrollment_columns if c in enrollment_df.columns]]

    if not enrollment_df.empty:
        total_row = {"Date": f"{len(enrollment_df)} Days"}
        for scheme in schemes:
            rc = get_rate_column_name(scheme)
            if rc is not None and rc in enrollment_df.columns:
                total_row[rc] = "-"
            tc_ = enrollment_df[f"{scheme} Count"].sum()
            ta_ = enrollment_df[f"{scheme} Amount"].sum()
            total_row[f"{scheme} Count"] = tc_
            total_row[f"{scheme} Amount"] = ta_
            total_row[f"{scheme} Avg Ticket"] = ta_ / tc_ if tc_ > 0 else 0
        tc_ = enrollment_df["Total Count"].sum()
        ta_ = enrollment_df["Total Amount"].sum()
        total_row["Total Count"] = tc_
        total_row["Total Amount"] = ta_
        total_row["Total Avg Ticket"] = ta_ / tc_ if tc_ > 0 else 0
        enrollment_df = pd.concat([enrollment_df, pd.DataFrame([total_row])], ignore_index=True)

    collection_rows = []
    for date in dates:
        row = {"Date": pd.Timestamp(date).strftime("%d-%m-%Y")}
        tc, ta = 0, 0
        for scheme in schemes:
            temp = collection_df[(collection_df["Date"] == date) & (collection_df["Scheme"] == scheme)]
            cnt, amt = len(temp), temp["Saved Amount"].sum()
            row[f"{scheme} Count"] = cnt
            row[f"{scheme} Amount"] = amt
            row[f"{scheme} Avg Ticket"] = amt / cnt if cnt > 0 else 0
            tc += cnt
            ta += amt
        row["Total Count"] = tc
        row["Total Amount"] = ta
        row["Total Avg Ticket"] = ta / tc if tc > 0 else 0
        collection_rows.append(row)
    collection_df_report = pd.DataFrame(collection_rows)

    collection_columns = ["Date"]
    for scheme in schemes:
        collection_columns.extend([f"{scheme} Count", f"{scheme} Amount", f"{scheme} Avg Ticket"])
    collection_columns.extend(["Total Count", "Total Amount", "Total Avg Ticket"])
    collection_df_report = collection_df_report[[c for c in collection_columns if c in collection_df_report.columns]]

    if not collection_df_report.empty:
        total_row = {"Date": f"{len(collection_df_report)} Days"}
        for scheme in schemes:
            tc_ = collection_df_report[f"{scheme} Count"].sum()
            ta_ = collection_df_report[f"{scheme} Amount"].sum()
            total_row[f"{scheme} Count"] = tc_
            total_row[f"{scheme} Amount"] = ta_
            total_row[f"{scheme} Avg Ticket"] = ta_ / tc_ if tc_ > 0 else 0
        tc_ = collection_df_report["Total Count"].sum()
        ta_ = collection_df_report["Total Amount"].sum()
        total_row["Total Count"] = tc_
        total_row["Total Amount"] = ta_
        total_row["Total Avg Ticket"] = ta_ / tc_ if tc_ > 0 else 0
        collection_df_report = pd.concat([collection_df_report, pd.DataFrame([total_row])], ignore_index=True)

    unique_rows = []
    for date in dates:
        row = {"Date": pd.Timestamp(date).strftime("%d-%m-%Y")}
        gt = 0
        for scheme in schemes:
            temp = first_df[
                (first_df["Date"] == date) & (first_df["Scheme"] == scheme) & (first_df["Valid Phone"])
            ]
            uc = temp["Customer Key"].nunique()
            row[scheme] = uc
            gt += uc
        row["Grand Total"] = gt
        unique_rows.append(row)
    unique_df = pd.DataFrame(unique_rows)

    unique_columns = ["Date"] + list(schemes) + ["Grand Total"]
    if not unique_df.empty:
        unique_df = unique_df[[c for c in unique_columns if c in unique_df.columns]]
        total_row = {"Date": f"{len(unique_df)} Days"}
        for scheme in schemes:
            if scheme in unique_df.columns:
                total_row[scheme] = unique_df[scheme].sum()
        total_row["Grand Total"] = unique_df["Grand Total"].sum()
        unique_df = pd.concat([unique_df, pd.DataFrame([total_row])], ignore_index=True)

    summary_df = round_df(summary_df)
    enrollment_df = round_df(enrollment_df)
    collection_df_report = round_df(collection_df_report)
    unique_df = round_df(unique_df)
    avg_ticket_df = round_df(avg_ticket_df, exclude_columns=["% Change"]) if not avg_ticket_df.empty else avg_ticket_df

    return summary_df, enrollment_df, collection_df_report, unique_df, avg_ticket_df
def generate_monthly_projection(full_df, schemes):
    empty = pd.DataFrame(columns=["Scheme", "Projected Count", "Projected Amount"])
    if full_df.empty or not schemes:
        return empty, empty, ""
    latest_date = full_df["Date"].max()
    month_start = latest_date.replace(day=1)
    days_in_month = calendar.monthrange(latest_date.year, latest_date.month)[1]
    days_elapsed = latest_date.day
    if days_elapsed <= 0:
        days_elapsed = 1
    month_label = (
        f"{latest_date.strftime('%B %Y')} - projected for all "
        f"{days_in_month} days ({days_elapsed} day(s) so far)"
    )
    month_df = full_df[(full_df["Date"] >= month_start) & (full_df["Date"] <= latest_date)]
    installment_one = full_df[
        (full_df["Installment number"] == 1) & (full_df["Passbook number"].notna())
    ].sort_values(["Passbook number", "Date", "Id"])
    first_df = installment_one.drop_duplicates(subset=["Passbook number"], keep="first")
    first_month_df = first_df[(first_df["Date"] >= month_start) & (first_df["Date"] <= latest_date)]
    collection_month_df = month_df[month_df["Installment number"].notna() & (month_df["Installment number"] != 1)]

    enrollment_rows, collection_rows = [], []
    for scheme in schemes:
        e = first_month_df[first_month_df["Scheme"] == scheme]
        enrollment_rows.append({
            "Scheme": scheme,
            "Projected Count": round_value((len(e) / days_elapsed) * days_in_month),
            "Projected Amount": round_value((e["Saved Amount"].sum() / days_elapsed) * days_in_month),
        })
        c = collection_month_df[collection_month_df["Scheme"] == scheme]
        collection_rows.append({
            "Scheme": scheme,
            "Projected Count": round_value((len(c) / days_elapsed) * days_in_month),
            "Projected Amount": round_value((c["Saved Amount"].sum() / days_elapsed) * days_in_month),
        })
    return pd.DataFrame(enrollment_rows), pd.DataFrame(collection_rows), month_label
def generate_clean_filename(start_date, end_date):
    if start_date and end_date:
        mn = start_date.strftime("%B")
        y = start_date.strftime("%Y")
        sd = str(start_date.day).zfill(2)
        ed = str(end_date.day).zfill(2)
        if start_date.month == end_date.month:
            return f"Daily enrollment ({mn} {sd} - {ed} {y})"
        else:
            return (
                f"Daily enrollment ({start_date.strftime('%B')} {sd} - "
                f"{end_date.strftime('%B')} {ed} {y})"
            )
    else:
        return f"Daily enrollment ({pd.Timestamp.now().strftime('%B %Y')})"

LABEL = "Branch Employee Wise Referral Report"


def run():
    warnings.filterwarnings("ignore")
    st.title("💰 Scheme Enrollment & Collection Report")
    st.caption("Strict Match: Referee Phone + Joined Date + Enrollment Amount + Installment #1")
    REQUIRED_COLUMNS = [
        "Id", "Scheme Participation Id", "Date", "Status", "Saved Amount",
        "Reward Amount", "Transaction Reference", "Installment number",
        "Metal Type", "Metal Rate", "Saved Metal Weight", "Rewards Metal Weight",
        "Benefit Metal Amount", "Benefit Metal Weight", "Benefit Metal Percentage",
        "Receipt ID", "Customer Name", "Customer Phone Number",
        "Passbook number", "Scheme Name"
    ]
    DAILY_SCHEME_KEYS = [
        "e-gold", "egold", "e gold",
        "e-silver", "esilver", "e silver",
    ]
    BRANCH_MAPPING = {
        "bhima jewellery - madurai": "MDU",
        "head office": "MDU",
        "madurai branch": "MDU",
        "bhima jewellery - telecaller": "MDU-Telecalling",
        "bhima jewellery - marthandam": "MDM",
        "bhima jewellery - salem": "SLM",
        "app showroom location": "SLM",
        "in-transit- locations": "SLM",
        "in-transit locations": "SLM",
        "n/a": "Unassigned",
        "bhima jewellery - tirunelveli": "TVL",
        "bhima jewellery - tiruchirappalli": "TCY",
        "trichy branch": "TCY",
        "bhima jewellery -  rajapalayam": "RJPM",
        "bhima jewellery - rajapalayam": "RJPM",
        "rajapalayam branch": "RJPM",
        "bhima jewellery - dindigul": "DGL",
        "dindigul branch": "DGL",
        "bhima jewellery - noida": "ND",
        "bhima jewellery - virudhunagar": "VNR",
        "bhima jewellery -anna nagar": "AN",
        "bhima jewellery - anna nagar": "AN",
        "bhima jewellery - thanjavur": "TJR",
        "bhima jewellery -thanjavur": "TJR",
    }
    BRANCH_ORDER = [
        "MDU", "MDM", "SLM", "TVL", "TCY", "RJPM",
        "DGL", "ND", "VNR", "TJR", "AN", "MDU-Telecalling",
        "Unassigned",
    ]
    FIXED_BRANCH_GROUPS = [
        ["madurai"], ["marthandam"], ["salem"], ["tirunelveli"],
        ["trichy", "tiruchirappalli", "tiruchirapalli"],
        ["rajapalayam", "rajapalaiyam", "rajapalayem"],
        ["dindigul"], ["noida"], ["virudhunagar", "virudunagar"], ["thanjavur"],
    ]
    TELECALLER_KEYWORDS = ["telecaller", "tele caller", "tellecaller"]
    st.sidebar.header("📂 File Uploads")
    uploaded_file = st.sidebar.file_uploader(
        "1️⃣ Main Transaction File (Required)",
        type=["xlsx", "xls", "csv"], key="main_file"
    )
    uploaded_employee = st.sidebar.file_uploader(
        "2️⃣ Employee Details File (For Referral Report)",
        type=["xlsx", "xls", "csv"], key="employee_file"
    )
    uploaded_referral = st.sidebar.file_uploader(
        "3️⃣ Referral Details File (For Referral Report)",
        type=["xlsx", "xls", "csv"], key="referral_file"
    )
    if uploaded_file is None:
        st.info("Please upload your raw transaction Excel/CSV file (1️⃣) from the sidebar.")
        st.stop()
    try:
        if uploaded_file.name.lower().endswith(".csv"):
            df = pd.read_csv(uploaded_file, low_memory=False)
        else:
            df = pd.read_excel(uploaded_file)
    except Exception as e:
        st.error(f"❌ Error reading file: {e}")
        st.stop()
    df = clean_columns(df)
    missing_columns = [col for col in REQUIRED_COLUMNS if col not in df.columns]
    if missing_columns:
        st.error("❌ Required columns are missing.")
        st.write("Missing columns:", missing_columns)
        st.stop()
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce", dayfirst=True)
    df["Saved Amount"] = clean_amount(df["Saved Amount"])
    df["Metal Rate"] = clean_amount(df["Metal Rate"])
    df["Installment number"] = pd.to_numeric(df["Installment number"], errors="coerce")
    df["Passbook number"] = clean_text(df["Passbook number"])
    df["Scheme Name"] = clean_text(df["Scheme Name"])
    df["Customer Phone Number"] = clean_text(df["Customer Phone Number"])
    df["Customer Name"] = clean_text(df["Customer Name"])
    df = df[df["Date"].notna()].copy()
    if df.empty:
        st.error("❌ No valid Date records found.")
        st.stop()
    df["Date"] = df["Date"].dt.normalize()
    rows_before = len(df)
    df = df[df["Scheme Name"].notna()].copy()
    if rows_before - len(df) > 0:
        st.sidebar.warning(f"⚠️ Skipped {rows_before - len(df)} row(s) with a blank 'Scheme Name'.")
    df["Scheme"] = df["Scheme Name"].astype(str).str.strip()
    st.sidebar.header("🔎 Report Filters")
    minimum_date = df["Date"].min().date()
    maximum_date = df["Date"].max().date()
    date_range = st.sidebar.date_input(
        "Date Range", value=(minimum_date, maximum_date),
        min_value=minimum_date, max_value=maximum_date
    )
    if isinstance(date_range, tuple):
        if len(date_range) == 2:
            start_date = pd.Timestamp(date_range[0])
            end_date = pd.Timestamp(date_range[1])
        else:
            start_date = pd.Timestamp(date_range[0])
            end_date = start_date
    else:
        start_date = pd.Timestamp(date_range)
        end_date = start_date
    all_schemes = get_scheme_list(df)
    selected_schemes = st.sidebar.multiselect("Select Schemes", options=all_schemes, default=all_schemes)
    filtered_df = df.copy()
    filtered_df = filtered_df[(filtered_df["Date"] >= start_date) & (filtered_df["Date"] <= end_date)].copy()
    filtered_df = filtered_df[filtered_df["Scheme"].isin(selected_schemes)].copy()
    filtered_df = filtered_df.sort_values(["Passbook number", "Date", "Id"])
    if filtered_df.empty:
        st.warning("⚠️ No records found for the selected filters.")
        st.stop()
    daily_schemes = [s for s in selected_schemes if classify_scheme_type(s) == "daily"]
    daily_df = filtered_df[filtered_df["Scheme"].isin(daily_schemes)].copy()
    sessional_schemes = [s for s in selected_schemes if classify_scheme_type(s) == "sessional"]
    sessional_df = filtered_df[filtered_df["Scheme"].isin(sessional_schemes)].copy()
    with st.spinner("Crunching daily-scheme numbers..."):
        daily_summary, daily_enrollment, daily_collection, daily_unique, daily_avg_ticket = generate_report_data(
            daily_df, daily_schemes
        )
    with st.spinner("Crunching sessional-scheme numbers..."):
        sessional_summary, sessional_enrollment, sessional_collection, sessional_unique, sessional_avg_ticket = generate_report_data(
            sessional_df, sessional_schemes
        )
    daily_enrollment_projection, daily_collection_projection, daily_projection_label = generate_monthly_projection(
        df, daily_schemes
    )
    sessional_enrollment_projection, sessional_collection_projection, sessional_projection_label = generate_monthly_projection(
        df, sessional_schemes
    )
    referral_summary = pd.DataFrame()
    referral_daily_list = []
    referral_detail = pd.DataFrame()
    referral_diagnostics = {}
    referral_duplicates = pd.DataFrame()
    referral_debug = {}
    employee_df_raw = None
    referral_df_raw = None
    referral_files_present = uploaded_employee is not None and uploaded_referral is not None
    if referral_files_present:
        employee_df_raw = read_uploaded_file(uploaded_employee)
        referral_df_raw = read_uploaded_file(uploaded_referral)

        if employee_df_raw is not None and referral_df_raw is not None:
            with st.spinner("Matching referrals (strict: phone + joined date + amount)..."):
                referral_summary, referral_daily_list, referral_detail, referral_diagnostics, referral_duplicates, referral_debug = \
                    build_referral_report(
                        filtered_df, employee_df_raw, referral_df_raw, None,
                        start_date=start_date, end_date=end_date
                    )

            if referral_diagnostics:
                st.info(
                    f"📊 **Referral Match Summary:** "
                    f"{referral_diagnostics.get('matched_referrals', 0)} matched / "
                    f"{referral_diagnostics.get('total_referrals', 0)} total "
                    f"(**{referral_diagnostics.get('not_enrolled', 0)}** not enrolled)"
                    + (f" — **{referral_diagnostics.get('duplicates', 0)}** duplicate(s) flagged"
                       if referral_diagnostics.get('duplicates') else "")
                )
    if referral_files_present and referral_debug:
        st.header("🧪 Debug — Match Analysis")

        with st.expander("📊 Referral Funnel", expanded=True):
            col1, col2, col3, col4 = st.columns(4)
            with col1:
                st.metric("Raw referral rows", f"{referral_debug.get('referral_raw_count', 0):,}")
            with col2:
                st.metric("After date filter", f"{referral_debug.get('referral_after_date_filter', 0):,}")
            with col3:
                st.metric("After phone filter", f"{referral_debug.get('referral_after_phone_filter', 0):,}")
            with col4:
                st.metric("In scope for matching", f"{referral_debug.get('referrals_in_scope', 0):,}")

            st.markdown("---")
            col1, col2, col3 = st.columns(3)
            with col1:
                st.metric("Txn rows (installment 1)", f"{referral_debug.get('txn_f1_count', 0):,}")
            with col2:
                st.metric("Unique txn phones", f"{referral_debug.get('txn_unique_phones', 0):,}")
            with col3:
                st.metric("Phone overlap (ref ∩ txn)", f"{referral_debug.get('phone_overlap', 0):,}")

        with st.expander("📋 Not Enrolled Reason Breakdown"):
            if referral_debug.get('not_enrolled_breakdown'):
                st.dataframe(
                    pd.DataFrame([
                        {"Reason": k, "Count": v}
                        for k, v in referral_debug['not_enrolled_breakdown'].items()
                    ]),
                    use_container_width=True, hide_index=True
                )

        with st.expander("🔍 Branch × Category Pivot (in-scope)"):
            if not referral_detail.empty:
                grp = referral_detail.groupby(['Branch', 'Category']).size().reset_index(name='Count')
                pivot_dbg = grp.pivot(index='Branch', columns='Category', values='Count').fillna(0).astype(int)
                st.dataframe(pivot_dbg, use_container_width=True)

        with st.expander("📋 Not Enrolled Rows (first 50)"):
            if not referral_detail.empty:
                ne = referral_detail[referral_detail['Category'] == 'Not Enrolled']
                display_cols = ['Branch', 'Referee Name', 'Joined Date', 'Enrollment Amount',
                                'Match Reason', 'Transaction Date', 'Paid Amount']
                avail = [c for c in display_cols if c in ne.columns]
                st.write(f"**Total Not Enrolled: {len(ne)}**")
                st.dataframe(ne[avail].head(50), use_container_width=True, hide_index=True)
    if referral_files_present:
        if employee_df_raw is None or referral_df_raw is None:
            st.error("❌ Could not read Employee or Referral file.")
        else:
            with st.expander("🔍 DIAGNOSTIC — Referral Data", expanded=False):
                st.write("### 📌 Column Detection")
                st.write("**Employee file columns:**", list(employee_df_raw.columns))
                st.write("**Referral file columns:**", list(referral_df_raw.columns))

                if not referral_detail.empty:
                    st.write("### 📌 Category Breakdown")
                    cs = referral_detail["Category"].value_counts().reset_index()
                    cs.columns = ["Category", "Count"]
                    st.dataframe(cs, use_container_width=True, hide_index=True)

                    st.write("### 📌 Match Reason Breakdown")
                    rs = referral_detail["Match Reason"].value_counts().reset_index()
                    rs.columns = ["Reason", "Count"]
                    st.dataframe(rs, use_container_width=True, hide_index=True)

                    if referral_diagnostics.get("duplicates", 0) > 0 and not referral_duplicates.empty:
                        st.write("### 📌 Duplicate Records")
                        st.dataframe(referral_duplicates, use_container_width=True, hide_index=True)
    st.header("📋 Formatted Report Preview")
    if daily_schemes and not daily_summary.empty:
        st.subheader("📅 Daily Schemes (e-Gold & e-Silver)")
        st.dataframe(daily_summary, use_container_width=True, hide_index=True)
        if not daily_avg_ticket.empty:
            display_avg_ticket_comparison(daily_avg_ticket)
        if not daily_enrollment_projection.empty or not daily_collection_projection.empty:
            st.markdown(f"**📈 Monthly Projection — {daily_projection_label}**")
            c5, c6 = st.columns(2)
            with c5:
                st.caption("Projection — first enrollments")
                st.dataframe(daily_enrollment_projection, use_container_width=True, hide_index=True)
            with c6:
                st.caption("Projection — overall collection")
                st.dataframe(daily_collection_projection, use_container_width=True, hide_index=True)
        c1, c2 = st.columns(2)
        with c1:
            st.dataframe(daily_enrollment, use_container_width=True, hide_index=True)
        with c2:
            st.dataframe(daily_collection, use_container_width=True, hide_index=True)
        st.dataframe(daily_unique, use_container_width=True, hide_index=True)
    if sessional_schemes and not sessional_summary.empty:
        st.subheader("🎯 Sessional Schemes")
        st.dataframe(sessional_summary, use_container_width=True, hide_index=True)
        if not sessional_avg_ticket.empty:
            display_avg_ticket_comparison(sessional_avg_ticket)
        if not sessional_enrollment_projection.empty or not sessional_collection_projection.empty:
            st.markdown(f"**📈 Monthly Projection — {sessional_projection_label}**")
            c5, c6 = st.columns(2)
            with c5:
                st.caption("Projection — first enrollments")
                st.dataframe(sessional_enrollment_projection, use_container_width=True, hide_index=True)
            with c6:
                st.caption("Projection — overall collection")
                st.dataframe(sessional_collection_projection, use_container_width=True, hide_index=True)
        c1, c2 = st.columns(2)
        with c1:
            st.dataframe(sessional_enrollment, use_container_width=True, hide_index=True)
        with c2:
            st.dataframe(sessional_collection, use_container_width=True, hide_index=True)
        st.dataframe(sessional_unique, use_container_width=True, hide_index=True)
    if referral_files_present:
        st.header("🎁 Referral Conversion Report")
        if referral_summary.empty:
            st.warning("⚠️ No referral matches found. Check the debug panel above.")
        else:
            st.subheader(
                f"🏢 Branch-wise Referral Summary — "
                f"{start_date.strftime('%d-%m-%Y')} to {end_date.strftime('%d-%m-%Y')}"
            )
            display_summary = referral_summary.copy()
            for col in display_summary.columns[1:]:
                display_summary[col] = display_summary[col].apply(
                    lambda x: "-" if (isinstance(x, (int, float)) and x == 0) else x
                )
            st.dataframe(display_summary, use_container_width=True, hide_index=True)

            if referral_daily_list:
                st.subheader("📅 Daily Referral Breakdown")
                for date_label, day_df in referral_daily_list:
                    st.markdown(f"**Referral — {date_label}**")
                    day_display = day_df.copy()
                    for col in day_display.columns[1:]:
                        day_display[col] = day_display[col].apply(
                            lambda x: "-" if (isinstance(x, (int, float)) and x == 0) else x
                        )
                    st.dataframe(day_display, use_container_width=True, hide_index=True)

            with st.expander("📋 View Referral Match Details"):
                dd = referral_detail.drop(columns=["Referee Phone", "Match Date"], errors="ignore").copy()
                for col in ["Registered Date", "Joined Date", "Transaction Date"]:
                    if col in dd.columns:
                        dd[col] = dd[col].apply(lambda x: x.strftime("%d-%m-%Y") if pd.notna(x) else "")
                st.dataframe(dd, use_container_width=True, hide_index=True)

            if not referral_duplicates.empty:
                with st.expander("🔴 View Duplicate Records"):
                    st.dataframe(referral_duplicates, use_container_width=True, hide_index=True)
    else:
        st.info("💡 Upload Employee + Referral files to see the Referral Report.")
    st.header("📥 Download Formatted Report")
    clean_filename = generate_clean_filename(start_date, end_date)
    st.info(f"📁 **Report will be saved as:** `{clean_filename}.xlsx`")
    report_title = uploaded_file.name.rsplit(".", 1)[0] or "Scheme Enrollment & Collection Report"
    daily_data = {
        "summary": daily_summary if daily_summary is not None and not daily_summary.empty else pd.DataFrame(),
        "enrollment": daily_enrollment if daily_enrollment is not None and not daily_enrollment.empty else pd.DataFrame(),
        "collection": daily_collection if daily_collection is not None and not daily_collection.empty else pd.DataFrame(),
        "unique": daily_unique if daily_unique is not None and not daily_unique.empty else pd.DataFrame(),
        "schemes": daily_schemes,
        "date_range": (start_date.strftime("%d-%m-%Y"), end_date.strftime("%d-%m-%Y")),
        "report_title": "eGold & eSilver Enrollment & Collection Report",
        "sheet_name": "eGold & eSilver",
        "enrollment_projection": daily_enrollment_projection,
        "collection_projection": daily_collection_projection,
        "projection_month_label": daily_projection_label,
        "avg_ticket_data": daily_avg_ticket if daily_avg_ticket is not None and not daily_avg_ticket.empty else pd.DataFrame(),
    }
    sessional_data = {
        "summary": sessional_summary if sessional_summary is not None and not sessional_summary.empty else pd.DataFrame(),
        "enrollment": sessional_enrollment if sessional_enrollment is not None and not sessional_enrollment.empty else pd.DataFrame(),
        "collection": sessional_collection if sessional_collection is not None and not sessional_collection.empty else pd.DataFrame(),
        "unique": sessional_unique if sessional_unique is not None and not sessional_unique.empty else pd.DataFrame(),
        "schemes": sessional_schemes,
        "date_range": (start_date.strftime("%d-%m-%Y"), end_date.strftime("%d-%m-%Y")),
        "report_title": "Sessional Scheme Enrollment & Collection Report",
        "sheet_name": "Sessional Scheme",
        "enrollment_projection": sessional_enrollment_projection,
        "collection_projection": sessional_collection_projection,
        "projection_month_label": sessional_projection_label,
        "avg_ticket_data": sessional_avg_ticket if sessional_avg_ticket is not None and not sessional_avg_ticket.empty else pd.DataFrame(),
    }
    try:
        with st.spinner("Building the formatted Excel workbook..."):
            excel_data = create_formatted_excel(
                daily_data,
                sessional_data,
                (start_date.strftime("%d-%m-%Y"), end_date.strftime("%d-%m-%Y")),
                report_title,
                referral_summary=referral_summary,
                referral_daily=referral_daily_list,
                referral_detail=referral_detail,
                referral_duplicates=referral_duplicates,
            )
        st.download_button(
            label="⬇️ Download Formatted Excel Report",
            data=excel_data,
            file_name=f"{clean_filename}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
    except Exception as e:
        st.error(f"❌ Could not build the Excel report: {e}")
