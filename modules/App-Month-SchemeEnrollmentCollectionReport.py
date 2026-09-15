import streamlit as st
import pandas as pd
import numpy as np
import calendar
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import warnings

warnings.filterwarnings("ignore")


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Scheme Enrollment & Collection Report",
    page_icon="💰",
    layout="wide"
)


# ============================================================
# TITLE
# ============================================================

st.title("💰 Scheme Enrollment & Collection Report")

st.caption(
    "Enrollment = Installment #1 | Collection = Installment #2 onward"
)


# ============================================================
# REQUIRED COLUMNS
# ============================================================

REQUIRED_COLUMNS = [
    "Id",
    "Scheme Participation Id",
    "Date",
    "Status",
    "Saved Amount",
    "Reward Amount",
    "Transaction Reference",
    "Installment number",
    "Metal Type",
    "Metal Rate",
    "Saved Metal Weight",
    "Rewards Metal Weight",
    "Benefit Metal Amount",
    "Benefit Metal Weight",
    "Benefit Metal Percentage",
    "Receipt ID",
    "Customer Name",
    "Customer Phone Number",
    "Passbook number",
    "Scheme Name"
]


# ============================================================
# DAILY SCHEME KEYS
# ============================================================

DAILY_SCHEME_KEYS = [
    "e-gold",
    "egold",
    "e gold",
    "gold",
    "e-silver",
    "esilver",
    "e silver",
    "silver"
]


# ============================================================
# BASIC HELPERS
# ============================================================

def clean_columns(df):
    """Clean column names."""

    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
        .str.replace(r"\s+", " ", regex=True)
    )

    return df


def clean_amount(series):
    """Convert amount/currency values into numeric."""

    return pd.to_numeric(
        series
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("₹", "", regex=False)
        .str.replace("Rs.", "", regex=False)
        .str.replace("Rs", "", regex=False)
        .str.strip(),
        errors="coerce"
    ).fillna(0)


def clean_text(series):
    """Clean text values."""

    return (
        series
        .astype(str)
        .str.strip()
        .replace(
            ["nan", "None", "NaN", ""],
            np.nan
        )
    )


def round_value(value):
    if pd.isna(value):
        return 0

    return round(value)


def classify_scheme_type(scheme_name):
    """
    Daily:
        e-Gold / e-Silver

    Sessional:
        Everything else.
    """

    if pd.isna(scheme_name):
        return "sessional"

    text = str(scheme_name).lower().strip()

    for key in DAILY_SCHEME_KEYS:

        if key in text:
            return "daily"

    return "sessional"


def get_scheme_list(df, scheme_type="all"):

    if "Scheme" not in df.columns:
        return []

    schemes = sorted(
        df["Scheme"]
        .dropna()
        .unique()
        .tolist()
    )

    if scheme_type == "daily":

        return [
            s for s in schemes
            if classify_scheme_type(s) == "daily"
        ]

    if scheme_type == "sessional":

        return [
            s for s in schemes
            if classify_scheme_type(s) == "sessional"
        ]

    return schemes


# ============================================================
# WEEK FUNCTIONS
# ============================================================

def get_week_number(date):
    """
    Get ISO week number.
    """

    return date.isocalendar().week


def get_week_display(date):
    """
    Display weekly period as:
        Nov Week 2, 2025
    
    Simple week numbering based on day of month:
        Days 1-7 = Week 1
        Days 8-14 = Week 2
        Days 15-21 = Week 3
        Days 22-28 = Week 4
        Days 29+ = Week 5
    """
    
    # Simple week calculation based on day of month
    day = date.day
    
    if day <= 7:
        week_of_month = 1
    elif day <= 14:
        week_of_month = 2
    elif day <= 21:
        week_of_month = 3
    elif day <= 28:
        week_of_month = 4
    else:
        week_of_month = 5
    
    month_name = date.strftime("%b")

    return (
        f"{month_name} Week "
        f"{week_of_month}, "
        f"{date.year}"
    )


# ============================================================
# CHANGE CALCULATION
# ============================================================

def calculate_change(current, previous):
    """
    Percentage change.

    Previous = 0, Current = 0
        -> 0%

    Previous = 0, Current > 0
        -> New

    Otherwise
        -> normal percentage change
    """

    current = float(current or 0)
    previous = float(previous or 0)

    if previous == 0:

        if current == 0:
            return 0

        return float("inf")

    return round(
        ((current - previous) / previous) * 100,
        2
    )


def format_change(value):

    if value is None:
        return "-"

    if value == float("inf"):
        return "New"

    if pd.isna(value):
        return "-"

    if value > 0:
        return f"+{value:.2f}%"

    if value < 0:
        return f"{value:.2f}%"

    return "0.00%"


def get_change_color(value):

    if value == float("inf"):

        return (
            "color: #006100;"
            "font-weight: bold;"
        )

    if value > 0:

        return (
            "color: #006100;"
            "font-weight: bold;"
        )

    if value < 0:

        return (
            "color: #9C0006;"
            "font-weight: bold;"
        )

    return "color: #7F7F7F;"


# ============================================================
# EXCEL HELPERS
# ============================================================

def apply_cell_style(
    cell,
    fill=None,
    font=None,
    alignment=None,
    number_format=None,
    border=None
):

    if fill:
        cell.fill = fill

    if font:
        cell.font = font

    if alignment:
        cell.alignment = alignment

    if number_format:
        cell.number_format = number_format

    if border:
        cell.border = border


def write_section_header(
    worksheet,
    row,
    title,
    max_column,
    color="4472C4"
):

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=max_column
    )

    cell = worksheet.cell(
        row=row,
        column=1
    )

    cell.value = title

    cell.font = Font(
        size=14,
        bold=True,
        color="FFFFFF"
    )

    cell.fill = PatternFill(
        start_color=color,
        end_color=color,
        fill_type="solid"
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    worksheet.row_dimensions[row].height = 25


def write_legend_section(
    worksheet,
    start_row,
    max_cols
):

    notes = [

        (
            "HOW TO READ THIS SHEET",
            True
        ),

        (
            "• Enrollment = customer's first payment only "
            "(Installment #1).",
            False
        ),

        (
            "• Each Passbook is counted only once for Enrollment.",
            False
        ),

        (
            "• Collection = Installment #2 onward only. "
            "Installment #1 is NEVER included in Collection.",
            False
        ),

        (
            "• Count = number of transactions.",
            False
        ),

        (
            "• Amount = total Saved Amount (₹).",
            False
        ),

        (
            "• Count Change % = current period count "
            "compared with previous period count.",
            False
        ),

        (
            "• Amount Change % = current period amount "
            "compared with previous period amount.",
            False
        ),

        (
            "• New = previous period was zero and "
            "current period is greater than zero.",
            False
        ),

        (
            "• TOTAL rows calculate their own "
            "count and amount changes.",
            False
        )
    ]

    note_fill = PatternFill(
        start_color="F2F2F2",
        end_color="F2F2F2",
        fill_type="solid"
    )

    for offset, (text, is_title) in enumerate(notes):

        row = start_row + offset

        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=max_cols
        )

        cell = worksheet.cell(
            row=row,
            column=1
        )

        cell.value = text

        cell.fill = note_fill

        cell.font = Font(
            bold=is_title,
            italic=not is_title,
            size=11 if is_title else 10,
            color="203764" if is_title else "595959"
        )

        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            indent=1
        )

        worksheet.row_dimensions[row].height = (
            20 if is_title else 16
        )

    return start_row + len(notes) + 1


# ============================================================
# EXCEL DATAFRAME FORMAT
# ============================================================

def format_dataframe_section(
    worksheet,
    dataframe,
    start_row,
    total_identifier=None
):

    if dataframe.empty:
        return

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    total_fill = PatternFill(
        start_color="FFF2CC",
        end_color="FFF2CC",
        fill_type="solid"
    )

    header_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid"
    )

    # Header
    for col_idx, column_name in enumerate(
        dataframe.columns,
        start=1
    ):

        cell = worksheet.cell(
            row=start_row,
            column=col_idx
        )

        cell.value = column_name

        apply_cell_style(
            cell,
            fill=header_fill,
            font=Font(bold=True),
            alignment=Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            ),
            border=thin_border
        )

    # Data
    for row_idx in range(
        start_row + 1,
        start_row + 1 + len(dataframe)
    ):

        is_total = False

        if total_identifier is not None:

            first_value = worksheet.cell(
                row=row_idx,
                column=1
            ).value

            if (
                first_value is not None
                and total_identifier in str(first_value)
            ):

                is_total = True

        for col_idx in range(
            1,
            len(dataframe.columns) + 1
        ):

            cell = worksheet.cell(
                row=row_idx,
                column=col_idx
            )

            fill = (
                total_fill
                if is_total
                else None
            )

            font = (
                Font(bold=True)
                if is_total
                else None
            )

            number_format = None

            if col_idx > 1:
                number_format = "#,##0"

            apply_cell_style(
                cell,
                fill=fill,
                font=font,
                alignment=Alignment(
                    horizontal="center",
                    vertical="center"
                ),
                number_format=number_format,
                border=thin_border
            )


# ============================================================
# AVG TICKET EXCEL
# ============================================================

def write_avg_ticket_section(
    worksheet,
    start_row,
    avg_ticket_data,
    thin_border
):

    if avg_ticket_data.empty:
        return

    header_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid"
    )

    write_section_header(
        worksheet,
        start_row,
        "📊 AVERAGE TICKET SIZE COMPARISON",
        len(avg_ticket_data.columns),
        "70AD47"
    )

    header_row = start_row + 1
    data_start_row = start_row + 2

    for col_idx, col_name in enumerate(
        avg_ticket_data.columns,
        start=1
    ):

        cell = worksheet.cell(
            row=header_row,
            column=col_idx
        )

        cell.value = col_name

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = thin_border

    for r, row in enumerate(
        avg_ticket_data.values
    ):

        row_idx = data_start_row + r

        is_total = (
            "Grand Total"
            in str(row[0])
        )

        for col_idx, value in enumerate(
            row,
            start=1
        ):

            cell = worksheet.cell(
                row=row_idx,
                column=col_idx
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            cell.border = thin_border

            col_name = avg_ticket_data.columns[
                col_idx - 1
            ]

            if col_name == "% Change":

                if isinstance(
                    value,
                    (int, float)
                ):

                    cell.value = value
                    cell.number_format = "0.00%"

                else:

                    cell.value = value

            elif isinstance(
                value,
                (int, float)
            ):

                cell.value = value
                cell.number_format = "#,##0"

            else:

                cell.value = value

            if (
                col_name == "Difference"
                and not is_total
                and isinstance(
                    value,
                    (int, float)
                )
            ):

                if value > 0:

                    cell.font = Font(
                        color="006100"
                    )

                    cell.fill = PatternFill(
                        start_color="E2EFDA",
                        end_color="E2EFDA",
                        fill_type="solid"
                    )

                elif value < 0:

                    cell.font = Font(
                        color="9C0006"
                    )

                    cell.fill = PatternFill(
                        start_color="FCE4D6",
                        end_color="FCE4D6",
                        fill_type="solid"
                    )

            if (
                col_name == "% Change"
                and not is_total
                and isinstance(
                    value,
                    (int, float)
                )
            ):

                if value > 0:

                    cell.font = Font(
                        color="006100"
                    )

                elif value < 0:

                    cell.font = Font(
                        color="9C0006"
                    )

            if is_total:

                cell.font = Font(
                    bold=True
                )

                cell.fill = PatternFill(
                    start_color="FFF2CC",
                    end_color="FFF2CC",
                    fill_type="solid"
                )

    for col_idx, col_name in enumerate(
        avg_ticket_data.columns,
        start=1
    ):

        column_letter = get_column_letter(
            col_idx
        )

        max_length = len(
            str(col_name)
        )

        for row_idx in range(
            data_start_row,
            data_start_row
            + len(avg_ticket_data)
        ):

            cell_val = worksheet.cell(
                row=row_idx,
                column=col_idx
            ).value

            if cell_val is not None:

                max_length = max(
                    max_length,
                    len(str(cell_val))
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            30
        )


# ============================================================
# LONG AGGREGATION EXCEL
# ============================================================

def write_long_aggregated_section(
    worksheet,
    start_row,
    df,
    title,
    color,
    thin_border
):

    if df.empty:
        return start_row

    header_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid"
    )

    total_fill = PatternFill(
        start_color="FFF2CC",
        end_color="FFF2CC",
        fill_type="solid"
    )

    write_section_header(
        worksheet,
        start_row,
        title,
        len(df.columns),
        color
    )

    header_row = start_row + 1
    data_start_row = start_row + 2

    # Headers
    for col_idx, col_name in enumerate(
        df.columns,
        start=1
    ):

        cell = worksheet.cell(
            row=header_row,
            column=col_idx
        )

        cell.value = col_name

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = thin_border

    scheme_col_idx = (
        df.columns.get_loc("Scheme") + 1
        if "Scheme" in df.columns
        else 2
    )

    # Data
    for r, row in enumerate(
        df.values
    ):

        row_idx = data_start_row + r

        is_total = (
            str(
                row[
                    scheme_col_idx - 1
                ]
            ).strip().upper()
            == "TOTAL"
        )

        for col_idx, value in enumerate(
            row,
            start=1
        ):

            cell = worksheet.cell(
                row=row_idx,
                column=col_idx
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

            cell.border = thin_border

            col_name = df.columns[
                col_idx - 1
            ]

            # Change %
            if "Change %" in col_name:

                if value == float("inf"):

                    cell.value = "New"

                    cell.font = Font(
                        color="006100",
                        bold=True
                    )

                elif (
                    isinstance(
                        value,
                        (int, float)
                    )
                    and not pd.isna(value)
                ):

                    cell.value = value / 100

                    cell.number_format = "0.00%"

                    if value > 0:

                        cell.font = Font(
                            color="006100",
                            bold=True
                        )

                    elif value < 0:

                        cell.font = Font(
                            color="9C0006",
                            bold=True
                        )

                    else:

                        cell.font = Font(
                            color="7F7F7F"
                        )

                else:

                    cell.value = "-"

            # Numeric
            elif isinstance(
                value,
                (int, float)
            ):

                cell.value = value
                cell.number_format = "#,##0"

            else:

                cell.value = value

            # Total
            if is_total:

                cell.font = Font(
                    bold=True
                )

                cell.fill = total_fill

    data_end_row = (
        data_start_row
        + len(df)
        - 1
    )

    # Merge period cells
    if "Period" in df.columns:

        period_col_idx = (
            df.columns.get_loc("Period")
            + 1
        )

        r = data_start_row

        while r <= data_end_row:

            span_start = r

            period_val = worksheet.cell(
                row=r,
                column=period_col_idx
            ).value

            r2 = r + 1

            while (
                r2 <= data_end_row
                and worksheet.cell(
                    row=r2,
                    column=period_col_idx
                ).value == period_val
            ):

                r2 += 1

            if r2 - span_start > 1:

                worksheet.merge_cells(
                    start_row=span_start,
                    start_column=period_col_idx,
                    end_row=r2 - 1,
                    end_column=period_col_idx
                )

                worksheet.cell(
                    row=span_start,
                    column=period_col_idx
                ).alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            r = r2

    # Width
    for col_idx, col_name in enumerate(
        df.columns,
        start=1
    ):

        column_letter = get_column_letter(
            col_idx
        )

        max_length = len(
            str(col_name)
        )

        for row_idx in range(
            data_start_row,
            data_end_row + 1
        ):

            cell_val = worksheet.cell(
                row=row_idx,
                column=col_idx
            ).value

            if cell_val is not None:

                max_length = max(
                    max_length,
                    len(str(cell_val))
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            30
        )

    worksheet.auto_filter.ref = (
        f"A{header_row}:"
        f"{get_column_letter(len(df.columns))}"
        f"{data_end_row}"
    )

    return data_end_row + 3


# ============================================================
# WRITE EXCEL SHEET
# ============================================================

def write_sheet(
    workbook,
    sheet_data
):

    worksheet = workbook.create_sheet(
        sheet_data["sheet_name"]
    )

    summary_df = sheet_data["summary"]

    week_df = sheet_data.get(
        "week_data",
        pd.DataFrame()
    )

    month_df = sheet_data.get(
        "month_data",
        pd.DataFrame()
    )

    year_df = sheet_data.get(
        "year_data",
        pd.DataFrame()
    )

    avg_ticket_data = sheet_data.get(
        "avg_ticket_data",
        pd.DataFrame()
    )

    date_range = sheet_data[
        "date_range"
    ]

    report_title = sheet_data[
        "report_title"
    ]

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    max_cols = max(
        len(summary_df.columns)
        if not summary_df.empty else 1,

        len(week_df.columns)
        if not week_df.empty else 1,

        len(month_df.columns)
        if not month_df.empty else 1,

        len(year_df.columns)
        if not year_df.empty else 1,

        len(avg_ticket_data.columns)
        if not avg_ticket_data.empty else 1,

        10
    )

    # ========================================================
    # TITLE
    # ========================================================

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max_cols
    )

    title_cell = worksheet.cell(
        row=1,
        column=1
    )

    title_cell.value = report_title

    title_cell.font = Font(
        size=16,
        bold=True,
        color="FFFFFF"
    )

    title_cell.fill = PatternFill(
        start_color="203764",
        end_color="203764",
        fill_type="solid"
    )

    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    worksheet.row_dimensions[1].height = 30

    # ========================================================
    # DATE RANGE
    # ========================================================

    worksheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=max_cols
    )

    date_cell = worksheet.cell(
        row=2,
        column=1
    )

    date_cell.value = (
        f"Report Period: "
        f"{date_range[0]} to {date_range[1]}"
    )

    date_cell.font = Font(
        size=12,
        bold=True
    )

    date_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    # ========================================================
    # LEGEND
    # ========================================================

    current_row = write_legend_section(
        worksheet,
        4,
        max_cols
    )

    current_row += 1

    # ========================================================
    # SUMMARY
    # ========================================================

    write_section_header(
        worksheet,
        current_row,
        "📊 SCHEME SUMMARY",
        len(summary_df.columns)
        if not summary_df.empty
        else max_cols
    )

    summary_header_row = current_row + 1

    if not summary_df.empty:

        for col_idx, col_name in enumerate(
            summary_df.columns,
            start=1
        ):

            worksheet.cell(
                row=summary_header_row,
                column=col_idx
            ).value = col_name

        for row_idx, row in enumerate(
            summary_df.values,
            start=summary_header_row + 1
        ):

            for col_idx, value in enumerate(
                row,
                start=1
            ):

                worksheet.cell(
                    row=row_idx,
                    column=col_idx
                ).value = value

    format_dataframe_section(
        worksheet,
        summary_df,
        summary_header_row,
        total_identifier="Grand Total"
    )

    current_row = (
        summary_header_row
        + len(summary_df)
        + 3
    )

    # ========================================================
    # AVG TICKET
    # ========================================================

    if not avg_ticket_data.empty:

        write_avg_ticket_section(
            worksheet,
            current_row,
            avg_ticket_data,
            thin_border
        )

        current_row = (
            current_row
            + 2
            + len(avg_ticket_data)
            + 3
        )

    # ========================================================
    # WEEKLY
    # ========================================================

    if not week_df.empty:

        current_row = (
            write_long_aggregated_section(
                worksheet,
                current_row,
                week_df,
                "📅 WEEKLY ENROLLMENT & COLLECTION "
                "(Count & Amount Change)",
                "ED7D31",
                thin_border
            )
        )

    # ========================================================
    # MONTHLY
    # ========================================================

    if not month_df.empty:

        current_row = (
            write_long_aggregated_section(
                worksheet,
                current_row,
                month_df,
                "📆 MONTHLY ENROLLMENT & COLLECTION "
                "(Count & Amount Change)",
                "7030A0",
                thin_border
            )
        )

    # ========================================================
    # YEARLY
    # ========================================================

    if not year_df.empty:

        current_row = (
            write_long_aggregated_section(
                worksheet,
                current_row,
                year_df,
                "📊 YEARLY ENROLLMENT & COLLECTION "
                "(Count & Amount Change)",
                "C00000",
                thin_border
            )
        )

    worksheet.freeze_panes = "A5"

    worksheet.sheet_view.showGridLines = False

    worksheet.page_setup.orientation = "landscape"

    worksheet.page_setup.fitToWidth = 1

    worksheet.page_setup.fitToHeight = 0

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True


# ============================================================
# CREATE FORMATTED EXCEL
# ============================================================

def create_formatted_excel(
    daily_data,
    sessional_data
):

    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        workbook = writer.book

        if "Sheet" in workbook.sheetnames:

            del workbook["Sheet"]

        if daily_data["schemes"]:

            write_sheet(
                workbook,
                daily_data
            )

        if sessional_data["schemes"]:

            write_sheet(
                workbook,
                sessional_data
            )

    output.seek(0)

    return output.getvalue()


# ============================================================
# AVG TICKET CALCULATION
# ============================================================

def add_avg_ticket_size_comparison(
    summary_df
):

    if summary_df.empty:
        return pd.DataFrame()

    scheme_rows = summary_df[
        summary_df["Scheme"] != "Grand Total"
    ].copy()

    if scheme_rows.empty:
        return pd.DataFrame()

    result_data = []

    for _, row in scheme_rows.iterrows():

        scheme = row["Scheme"]

        first_count = row[
            "First Enrollment Count"
        ]

        first_amount = row[
            "First Enrollment Amount"
        ]

        collection_count = row[
            "Collection Count"
        ]

        collection_amount = row[
            "Collection Amount"
        ]

        first_avg = (
            first_amount / first_count
            if first_count > 0
            else 0
        )

        collection_avg = (
            collection_amount / collection_count
            if collection_count > 0
            else 0
        )

        diff = (
            collection_avg
            - first_avg
        )

        if first_avg > 0:

            pct_change = round(
                diff / first_avg,
                4
            )

        elif collection_avg > 0:

            pct_change = "N/A"

        else:

            pct_change = 0

        result_data.append({

            "Scheme":
                scheme,

            "First Enrollment Count":
                int(first_count),

            "First Enrollment Amount":
                int(first_amount),

            "First Enrollment Avg Ticket":
                int(round(first_avg)),

            "Collection Count":
                int(collection_count),

            "Collection Amount":
                int(collection_amount),

            "Collection Avg Ticket":
                int(round(collection_avg)),

            "Difference":
                int(round(diff)),

            "% Change":
                pct_change
        })

    # ========================================================
    # GRAND TOTAL
    # ========================================================

    total_first_count = sum(
        d["First Enrollment Count"]
        for d in result_data
    )

    total_first_amount = sum(
        d["First Enrollment Amount"]
        for d in result_data
    )

    total_collection_count = sum(
        d["Collection Count"]
        for d in result_data
    )

    total_collection_amount = sum(
        d["Collection Amount"]
        for d in result_data
    )

    total_first_avg = (
        total_first_amount
        / total_first_count
        if total_first_count > 0
        else 0
    )

    total_collection_avg = (
        total_collection_amount
        / total_collection_count
        if total_collection_count > 0
        else 0
    )

    total_diff = (
        total_collection_avg
        - total_first_avg
    )

    if total_first_avg > 0:

        total_pct = round(
            total_diff / total_first_avg,
            4
        )

    elif total_collection_avg > 0:

        total_pct = "N/A"

    else:

        total_pct = 0

    result_data.append({

        "Scheme":
            "Grand Total",

        "First Enrollment Count":
            int(total_first_count),

        "First Enrollment Amount":
            int(total_first_amount),

        "First Enrollment Avg Ticket":
            int(round(total_first_avg)),

        "Collection Count":
            int(total_collection_count),

        "Collection Amount":
            int(total_collection_amount),

        "Collection Avg Ticket":
            int(round(total_collection_avg)),

        "Difference":
            int(round(total_diff)),

        "% Change":
            total_pct
    })

    return pd.DataFrame(
        result_data
    )


# ============================================================
# DISPLAY AVG TICKET
# ============================================================

def display_avg_ticket_comparison(
    avg_ticket_df
):

    if avg_ticket_df.empty:
        return

    st.subheader(
        "📊 Average Ticket Size Comparison"
    )

    display_df = avg_ticket_df.copy()

    # Counts
    for col in [
        "First Enrollment Count",
        "Collection Count"
    ]:

        if col in display_df.columns:

            display_df[col] = display_df[col].apply(
                lambda x:
                f"{int(x):,}"
                if isinstance(
                    x,
                    (int, float)
                )
                else x
            )

    # Amounts
    for col in [
        "First Enrollment Amount",
        "Collection Amount"
    ]:

        if col in display_df.columns:

            display_df[col] = display_df[col].apply(
                lambda x:
                f"{int(x):,}"
                if isinstance(
                    x,
                    (int, float)
                )
                else x
            )

    # Average ticket
    for col in [
        "First Enrollment Avg Ticket",
        "Collection Avg Ticket",
        "Difference"
    ]:

        if col in display_df.columns:

            display_df[col] = display_df[col].apply(
                lambda x:
                f"{int(x):,}"
                if isinstance(
                    x,
                    (int, float)
                )
                else x
            )

    # Percentage
    if "% Change" in display_df.columns:

        display_df["% Change"] = (
            display_df["% Change"].apply(
                lambda x:
                f"{x * 100:.2f}%"
                if isinstance(
                    x,
                    (int, float)
                )
                else x
            )
        )

    def color_pct(val):

        if (
            isinstance(val, str)
            and val.endswith("%")
        ):

            try:

                pct_val = float(
                    val.replace("%", "")
                )

                if pct_val > 0:

                    return (
                        "color: green;"
                        "font-weight: bold;"
                    )

                if pct_val < 0:

                    return (
                        "color: red;"
                        "font-weight: bold;"
                    )

            except:
                pass

        return ""

    def color_diff(val):

        if isinstance(val, str):

            try:

                num_val = float(
                    val.replace(",", "")
                )

                if num_val > 0:
                    return "color: green;"

                if num_val < 0:
                    return "color: red;"

            except:
                pass

        return ""

    styled_df = display_df.style.map(
        color_pct,
        subset=["% Change"]
    )

    styled_df = styled_df.map(
        color_diff,
        subset=["Difference"]
    )

    st.dataframe(
        styled_df,
        use_container_width=True,
        hide_index=True
    )

    # Overall interpretation
    total_row = avg_ticket_df[
        avg_ticket_df["Scheme"] == "Grand Total"
    ]

    if not total_row.empty:

        first_avg = total_row[
            "First Enrollment Avg Ticket"
        ].iloc[0]

        collection_avg = total_row[
            "Collection Avg Ticket"
        ].iloc[0]

        diff = (
            collection_avg
            - first_avg
        )

        pct = (
            diff / first_avg * 100
            if first_avg > 0
            else 0
        )

        if diff > 0:

            direction = "increase"
            emoji = "📈"

        elif diff < 0:

            direction = "decrease"
            emoji = "📉"

        else:

            direction = "no change"
            emoji = "➡️"

        st.info(
            f"{emoji} **Overall Average Ticket Size:** "
            f"{int(first_avg):,} → "
            f"{int(collection_avg):,} "
            f"({pct:.1f}% {direction})"
        )


# ============================================================
# AGGREGATED DATA
# ============================================================

def generate_aggregated_data_long(
    df,
    schemes,
    aggregation_type
):
    """
    Generate weekly/monthly/yearly report.

    ENROLLMENT:
        Installment #1 only.
        One record per Passbook.

    COLLECTION:
        Installment #2 onward only.

    CHANGE:
        Enrollment Count %
        Enrollment Amount %
        Collection Count %
        Collection Amount %
    """

    if not schemes or df.empty:

        return pd.DataFrame()

    # ========================================================
    # ENROLLMENT
    # ========================================================

    installment_one = df[
        (df["Installment number"] == 1)
        &
        (df["Passbook number"].notna())
    ].copy()

    installment_one = installment_one.sort_values(
        [
            "Passbook number",
            "Date",
            "Id"
        ]
    )

    first_df = installment_one.drop_duplicates(
        subset=[
            "Passbook number"
        ],
        keep="first"
    ).copy()

    # ========================================================
    # COLLECTION
    #
    # ONLY INSTALLMENT 2 ONWARD
    # ========================================================

    collection_df = df[
        df["Installment number"] > 1
    ].copy()

    first_work = first_df.copy()

    coll_work = collection_df.copy()

    # ========================================================
    # PERIOD
    # ========================================================

    if aggregation_type == "week":

        # Internal key remains YYYY-W##
        # so chronological comparison stays correct.
        first_work["PeriodKey"] = (
            first_work["Date"]
            .dt.strftime("%G-W%V")
        )

        coll_work["PeriodKey"] = (
            coll_work["Date"]
            .dt.strftime("%G-W%V")
        )

        period_source = df.copy()

        period_source["PeriodKey"] = (
            period_source["Date"]
            .dt.strftime("%G-W%V")
        )

        display_lookup = {}

        for period_key, group in (
            period_source
            .groupby("PeriodKey")
        ):

            # Get the first date in this ISO week
            first_date = group[
                "Date"
            ].min()
            
            # Use the actual date to determine the week number
            # Week is based on day of month (1-7=Week1, 8-14=Week2, etc.)
            display_lookup[
                period_key
            ] = get_week_display(
                first_date
            )

    elif aggregation_type == "month":

        first_work["PeriodKey"] = (
            first_work["Date"]
            .dt.strftime("%Y-%m")
        )

        coll_work["PeriodKey"] = (
            coll_work["Date"]
            .dt.strftime("%Y-%m")
        )

        display_lookup = {}

        for period_key in sorted(
            df["Date"]
            .dt.strftime("%Y-%m")
            .unique()
        ):

            year, month = period_key.split("-")

            display_lookup[
                period_key
            ] = (
                f"{calendar.month_name[int(month)]} "
                f"{year}"
            )

    else:

        first_work["PeriodKey"] = (
            first_work["Date"]
            .dt.year
            .astype(str)
        )

        coll_work["PeriodKey"] = (
            coll_work["Date"]
            .dt.year
            .astype(str)
        )

        display_lookup = {
            str(year): str(year)
            for year in sorted(
                df["Date"]
                .dt.year
                .unique()
            )
        }

    periods = sorted(
        set(first_work["PeriodKey"].dropna())
        |
        set(coll_work["PeriodKey"].dropna())
    )

    # ========================================================
    # PREVIOUS VALUES
    # ========================================================

    previous_values = {}

    rows = []

    # ========================================================
    # PERIOD LOOP
    # ========================================================

    for period in periods:

        period_display = display_lookup.get(
            period,
            period
        )

        total_enroll_count = 0
        total_enroll_amount = 0

        total_collection_count = 0
        total_collection_amount = 0

        # ====================================================
        # SCHEME LOOP
        # ====================================================

        for scheme in schemes:

            current_enrollment = first_work[
                (first_work["PeriodKey"] == period)
                &
                (first_work["Scheme"] == scheme)
            ]

            current_collection = coll_work[
                (coll_work["PeriodKey"] == period)
                &
                (coll_work["Scheme"] == scheme)
            ]

            # ================================================
            # CURRENT VALUES
            # ================================================

            enrollment_count = len(
                current_enrollment
            )

            enrollment_amount = (
                current_enrollment[
                    "Saved Amount"
                ].sum()
            )

            collection_count = len(
                current_collection
            )

            collection_amount = (
                current_collection[
                    "Saved Amount"
                ].sum()
            )

            # ================================================
            # PREVIOUS VALUES
            # ================================================

            previous = previous_values.get(
                scheme
            )

            if previous is None:

                previous_enrollment_count = None
                previous_enrollment_amount = None
                previous_collection_count = None
                previous_collection_amount = None

            else:

                (
                    previous_enrollment_count,
                    previous_enrollment_amount,
                    previous_collection_count,
                    previous_collection_amount
                ) = previous

            # ================================================
            # CHANGE VALUES
            # ================================================

            if previous is None:

                enrollment_count_change = None
                enrollment_amount_change = None
                collection_count_change = None
                collection_amount_change = None

            else:

                enrollment_count_change = (
                    calculate_change(
                        enrollment_count,
                        previous_enrollment_count
                    )
                )

                enrollment_amount_change = (
                    calculate_change(
                        enrollment_amount,
                        previous_enrollment_amount
                    )
                )

                collection_count_change = (
                    calculate_change(
                        collection_count,
                        previous_collection_count
                    )
                )

                collection_amount_change = (
                    calculate_change(
                        collection_amount,
                        previous_collection_amount
                    )
                )

            # ================================================
            # ROW
            # ================================================

            rows.append({

                "Period":
                    period_display,

                "Scheme":
                    scheme,

                "Enrollment Count":
                    enrollment_count,

                "Enrollment Amount":
                    enrollment_amount,

                "Enrollment Count Change %":
                    enrollment_count_change,

                "Enrollment Amount Change %":
                    enrollment_amount_change,

                "Collection Count":
                    collection_count,

                "Collection Amount":
                    collection_amount,

                "Collection Count Change %":
                    collection_count_change,

                "Collection Amount Change %":
                    collection_amount_change
            })

            # ================================================
            # SAVE CURRENT
            # ================================================

            previous_values[scheme] = (

                enrollment_count,
                enrollment_amount,

                collection_count,
                collection_amount
            )

            # ================================================
            # TOTAL
            # ================================================

            total_enroll_count += (
                enrollment_count
            )

            total_enroll_amount += (
                enrollment_amount
            )

            total_collection_count += (
                collection_count
            )

            total_collection_amount += (
                collection_amount
            )

        # ====================================================
        # TOTAL PREVIOUS
        # ====================================================

        previous_total = previous_values.get(
            "__TOTAL__"
        )

        if previous_total is None:

            previous_total_enrollment_count = None
            previous_total_enrollment_amount = None
            previous_total_collection_count = None
            previous_total_collection_amount = None

        else:

            (
                previous_total_enrollment_count,
                previous_total_enrollment_amount,
                previous_total_collection_count,
                previous_total_collection_amount
            ) = previous_total

        # ====================================================
        # TOTAL CHANGE
        # ====================================================

        if previous_total is None:

            total_enrollment_count_change = None
            total_enrollment_amount_change = None
            total_collection_count_change = None
            total_collection_amount_change = None

        else:

            total_enrollment_count_change = (
                calculate_change(
                    total_enroll_count,
                    previous_total_enrollment_count
                )
            )

            total_enrollment_amount_change = (
                calculate_change(
                    total_enroll_amount,
                    previous_total_enrollment_amount
                )
            )

            total_collection_count_change = (
                calculate_change(
                    total_collection_count,
                    previous_total_collection_count
                )
            )

            total_collection_amount_change = (
                calculate_change(
                    total_collection_amount,
                    previous_total_collection_amount
                )
            )

        # ====================================================
        # TOTAL ROW
        # ====================================================

        rows.append({

            "Period":
                period_display,

            "Scheme":
                "TOTAL",

            "Enrollment Count":
                total_enroll_count,

            "Enrollment Amount":
                total_enroll_amount,

            "Enrollment Count Change %":
                total_enrollment_count_change,

            "Enrollment Amount Change %":
                total_enrollment_amount_change,

            "Collection Count":
                total_collection_count,

            "Collection Amount":
                total_collection_amount,

            "Collection Count Change %":
                total_collection_count_change,

            "Collection Amount Change %":
                total_collection_amount_change
        })

        # ====================================================
        # SAVE TOTAL
        # ====================================================

        previous_values["__TOTAL__"] = (

            total_enroll_count,
            total_enroll_amount,

            total_collection_count,
            total_collection_amount
        )

    return pd.DataFrame(rows)


# ============================================================
# REPORT DATA
# ============================================================

def generate_report_data(
    df,
    schemes
):

    if not schemes or df.empty:

        return (
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame()
        )

    # ========================================================
    # FIRST INSTALLMENT
    # ========================================================

    installment_one = df[
        (df["Installment number"] == 1)
        &
        (df["Passbook number"].notna())
    ].copy()

    installment_one = installment_one.sort_values(
        [
            "Passbook number",
            "Date",
            "Id"
        ]
    )

    first_df = installment_one.drop_duplicates(
        subset=[
            "Passbook number"
        ],
        keep="first"
    ).copy()

    # ========================================================
    # COLLECTION
    # ========================================================

    collection_df = df[
        df["Installment number"] > 1
    ].copy()

    # ========================================================
    # SUMMARY
    # ========================================================

    summary_rows = []

    for scheme in schemes:

        scheme_first = first_df[
            first_df["Scheme"] == scheme
        ]

        scheme_collection = collection_df[
            collection_df["Scheme"] == scheme
        ]

        first_count = len(
            scheme_first
        )

        first_amount = (
            scheme_first[
                "Saved Amount"
            ].sum()
        )

        collection_count = len(
            scheme_collection
        )

        collection_amount = (
            scheme_collection[
                "Saved Amount"
            ].sum()
        )

        summary_rows.append({

            "Scheme":
                scheme,

            "First Enrollment Count":
                first_count,

            "First Enrollment Amount":
                first_amount,

            "Collection Count":
                collection_count,

            "Collection Amount":
                collection_amount
        })

    # ========================================================
    # GRAND TOTAL
    # ========================================================

    if summary_rows:

        summary_rows.append({

            "Scheme":
                "Grand Total",

            "First Enrollment Count":
                sum(
                    row[
                        "First Enrollment Count"
                    ]
                    for row in summary_rows
                ),

            "First Enrollment Amount":
                sum(
                    row[
                        "First Enrollment Amount"
                    ]
                    for row in summary_rows
                ),

            "Collection Count":
                sum(
                    row[
                        "Collection Count"
                    ]
                    for row in summary_rows
                ),

            "Collection Amount":
                sum(
                    row[
                        "Collection Amount"
                    ]
                    for row in summary_rows
                )
        })

    summary_df = pd.DataFrame(
        summary_rows
    )

    # ========================================================
    # AVG TICKET
    # ========================================================

    avg_ticket_df = (
        add_avg_ticket_size_comparison(
            summary_df
        )
    )

    # ========================================================
    # PERIOD DATA
    # ========================================================

    week_df = generate_aggregated_data_long(
        df,
        schemes,
        "week"
    )

    month_df = generate_aggregated_data_long(
        df,
        schemes,
        "month"
    )

    year_df = generate_aggregated_data_long(
        df,
        schemes,
        "year"
    )

    return (
        summary_df,
        avg_ticket_df,
        week_df,
        month_df,
        year_df
    )


# ============================================================
# DISPLAY AGGREGATED DATA
# ============================================================

def display_aggregated_with_comparison(
    df,
    title
):

    if df.empty:
        return

    st.subheader(title)

    display_df = df.copy()

    # ========================================================
    # CHANGE %
    # ========================================================

    change_columns = [
        col
        for col in display_df.columns
        if "Change %" in col
    ]

    for col in change_columns:

        display_df[col] = (
            display_df[col].apply(
                lambda x:
                "New"
                if x == float("inf")
                else
                (
                    f"+{x:.2f}%"
                    if isinstance(
                        x,
                        (int, float)
                    )
                    and not pd.isna(x)
                    and x > 0
                    else
                    (
                        f"{x:.2f}%"
                        if isinstance(
                            x,
                            (int, float)
                        )
                        and not pd.isna(x)
                        else "-"
                    )
                )
            )
        )

    # ========================================================
    # NUMBERS
    # ========================================================

    numeric_columns = [
        "Enrollment Count",
        "Enrollment Amount",
        "Collection Count",
        "Collection Amount"
    ]

    for col in numeric_columns:

        if col in display_df.columns:

            display_df[col] = (
                display_df[col].apply(
                    lambda x:
                    f"{int(x):,}"
                    if isinstance(
                        x,
                        (int, float)
                    )
                    else x
                )
            )

    # ========================================================
    # COLOR
    # ========================================================

    def color_change(val):

        if val == "New":

            return (
                "color: #006100;"
                "font-weight: bold;"
                "background-color: #E2EFDA;"
            )

        if (
            isinstance(val, str)
            and val.endswith("%")
        ):

            try:

                number = float(
                    val
                    .replace("%", "")
                    .replace("+", "")
                )

                if number > 0:

                    return (
                        "color: #006100;"
                        "font-weight: bold;"
                        "background-color: #E2EFDA;"
                    )

                if number < 0:

                    return (
                        "color: #9C0006;"
                        "font-weight: bold;"
                        "background-color: #FCE4D6;"
                    )

            except:
                pass

        return ""

    styled_df = display_df.style.map(
        color_change,
        subset=change_columns
    )

    st.dataframe(
        styled_df,
        use_container_width=True,
        hide_index=True
    )


# ============================================================
# FILE UPLOAD
# ============================================================

uploaded_file = st.file_uploader(
    "📂 Upload Excel / CSV File",
    type=[
        "xlsx",
        "xls",
        "csv"
    ]
)

if uploaded_file is None:

    st.info(
        "Please upload your raw transaction Excel/CSV file."
    )

    st.stop()


# ============================================================
# READ FILE
# ============================================================

try:

    if uploaded_file.name.lower().endswith(
        ".csv"
    ):

        df = pd.read_csv(
            uploaded_file,
            low_memory=False
        )

    else:

        df = pd.read_excel(
            uploaded_file
        )

except Exception as e:

    st.error(
        f"❌ Error reading file: {e}"
    )

    st.stop()


# ============================================================
# CLEAN COLUMNS
# ============================================================

df = clean_columns(df)


# ============================================================
# CHECK REQUIRED COLUMNS
# ============================================================

missing_columns = [
    col
    for col in REQUIRED_COLUMNS
    if col not in df.columns
]

if missing_columns:

    st.error(
        "❌ Required columns are missing."
    )

    st.write(
        "Missing columns:",
        missing_columns
    )

    st.stop()


# ============================================================
# CLEAN DATA
# ============================================================

df["Date"] = pd.to_datetime(
    df["Date"],
    errors="coerce",
    dayfirst=True
)

df["Saved Amount"] = clean_amount(
    df["Saved Amount"]
)

df["Metal Rate"] = clean_amount(
    df["Metal Rate"]
)

df["Installment number"] = pd.to_numeric(
    df["Installment number"],
    errors="coerce"
)

df["Passbook number"] = clean_text(
    df["Passbook number"]
)

df["Scheme Name"] = clean_text(
    df["Scheme Name"]
)

df["Customer Phone Number"] = clean_text(
    df["Customer Phone Number"]
)

df["Customer Name"] = clean_text(
    df["Customer Name"]
)


# ============================================================
# VALID DATES
# ============================================================

df = df[
    df["Date"].notna()
].copy()

if df.empty:

    st.error(
        "❌ No valid Date records found."
    )

    st.stop()


df["Date"] = (
    df["Date"]
    .dt.normalize()
)


# ============================================================
# SCHEME
# ============================================================

rows_before = len(df)

df = df[
    df["Scheme Name"].notna()
].copy()

rows_dropped = (
    rows_before - len(df)
)

if rows_dropped > 0:

    st.sidebar.warning(
        f"⚠️ Skipped {rows_dropped} row(s) "
        f"with blank Scheme Name."
    )

df["Scheme"] = (
    df["Scheme Name"]
    .astype(str)
    .str.strip()
)


# ============================================================
# SIDEBAR FILTERS
# ============================================================

st.sidebar.header(
    "🔎 Report Filters"
)

minimum_date = (
    df["Date"].min().date()
)

maximum_date = (
    df["Date"].max().date()
)

date_range = st.sidebar.date_input(
    "Date Range",
    value=(
        minimum_date,
        maximum_date
    ),
    min_value=minimum_date,
    max_value=maximum_date
)


# ============================================================
# DATE RANGE
# ============================================================

if isinstance(
    date_range,
    tuple
):

    if len(date_range) == 2:

        start_date = pd.Timestamp(
            date_range[0]
        )

        end_date = pd.Timestamp(
            date_range[1]
        )

    else:

        start_date = pd.Timestamp(
            date_range[0]
        )

        end_date = start_date

else:

    start_date = pd.Timestamp(
        date_range
    )

    end_date = start_date


# ============================================================
# SCHEME FILTER
# ============================================================

all_schemes = get_scheme_list(
    df
)

selected_schemes = st.sidebar.multiselect(
    "Select Schemes",
    options=all_schemes,
    default=all_schemes
)


# ============================================================
# FILTER DATA
# ============================================================

filtered_df = df[
    (df["Date"] >= start_date)
    &
    (df["Date"] <= end_date)
].copy()

filtered_df = filtered_df[
    filtered_df["Scheme"].isin(
        selected_schemes
    )
].copy()

filtered_df = filtered_df.sort_values(
    [
        "Passbook number",
        "Date",
        "Id"
    ]
)


if filtered_df.empty:

    st.warning(
        "⚠️ No records found for the selected filters."
    )

    st.stop()


# ============================================================
# DAILY / SESSIONAL
# ============================================================

daily_schemes = [
    s
    for s in selected_schemes
    if classify_scheme_type(s) == "daily"
]

daily_df = filtered_df[
    filtered_df["Scheme"].isin(
        daily_schemes
    )
].copy()


sessional_schemes = [
    s
    for s in selected_schemes
    if classify_scheme_type(s) == "sessional"
]

sessional_df = filtered_df[
    filtered_df["Scheme"].isin(
        sessional_schemes
    )
].copy()


# ============================================================
# REPORT DATA
# ============================================================

(
    daily_summary,
    daily_avg_ticket,
    daily_week,
    daily_month,
    daily_year
) = generate_report_data(
    daily_df,
    daily_schemes
)


(
    sessional_summary,
    sessional_avg_ticket,
    sessional_week,
    sessional_month,
    sessional_year
) = generate_report_data(
    sessional_df,
    sessional_schemes
)


# ============================================================
# PREVIEW
# ============================================================

st.header(
    "📋 Formatted Report Preview"
)


with st.expander(
    "ℹ️ How to read this report",
    expanded=False
):

    st.markdown(
        """
### Enrollment
- **Installment #1 only**
- One enrollment per **Passbook**
- Represents the customer's first payment.

### Collection
- **Installment #2 onward only**
- Installment #1 is **never included**.
- Represents additional/follow-up payments.

### Weekly Period
Weekly periods are displayed as:

**Nov Week 2, 2025**

instead of:

**Week 45, 2025**

### Change %
- **Enrollment Count Change %** = change in enrollment transaction count.
- **Enrollment Amount Change %** = change in enrollment ₹ amount.
- **Collection Count Change %** = change in collection transaction count.
- **Collection Amount Change %** = change in collection ₹ amount.
- **New** = previous period was zero and current period is greater than zero.
- `0.00%` = no change.
"""
    )


# ============================================================
# DAILY PREVIEW
# ============================================================

if daily_schemes:

    st.subheader(
        "📅 Daily Schemes (e-Gold & e-Silver)"
    )

    st.dataframe(
        daily_summary,
        use_container_width=True,
        hide_index=True
    )

    if not daily_avg_ticket.empty:

        display_avg_ticket_comparison(
            daily_avg_ticket
        )

    if not daily_week.empty:

        display_aggregated_with_comparison(
            daily_week,
            "📅 Weekly Aggregation "
            "(Week-over-Week Count & Amount Change)"
        )

    if not daily_month.empty:

        display_aggregated_with_comparison(
            daily_month,
            "📆 Monthly Aggregation "
            "(Month-over-Month Count & Amount Change)"
        )

    if not daily_year.empty:

        display_aggregated_with_comparison(
            daily_year,
            "📊 Yearly Aggregation "
            "(Year-over-Year Count & Amount Change)"
        )


# ============================================================
# SESSIONAL PREVIEW
# ============================================================

if sessional_schemes:

    st.subheader(
        "🎯 Sessional Schemes"
    )

    st.dataframe(
        sessional_summary,
        use_container_width=True,
        hide_index=True
    )

    if not sessional_avg_ticket.empty:

        display_avg_ticket_comparison(
            sessional_avg_ticket
        )

    if not sessional_week.empty:

        display_aggregated_with_comparison(
            sessional_week,
            "📅 Weekly Aggregation "
            "(Week-over-Week Count & Amount Change)"
        )

    if not sessional_month.empty:

        display_aggregated_with_comparison(
            sessional_month,
            "📆 Monthly Aggregation "
            "(Month-over-Month Count & Amount Change)"
        )

    if not sessional_year.empty:

        display_aggregated_with_comparison(
            sessional_year,
            "📊 Yearly Aggregation "
            "(Year-over-Year Count & Amount Change)"
        )


# ============================================================
# DOWNLOAD
# ============================================================

st.header(
    "📥 Download Formatted Report"
)


def generate_clean_filename(
    start_date,
    end_date
):

    if start_date and end_date:

        year = start_date.strftime(
            "%Y"
        )

        start_day = str(
            start_date.day
        ).zfill(2)

        end_day = str(
            end_date.day
        ).zfill(2)

        if (
            start_date.month
            == end_date.month
        ):

            month_name = (
                start_date.strftime(
                    "%B"
                )
            )

            return (
                f"Daily enrollment "
                f"({month_name} "
                f"{start_day} - "
                f"{end_day} "
                f"{year})"
            )

        start_month = (
            start_date.strftime(
                "%B"
            )
        )

        end_month = (
            end_date.strftime(
                "%B"
            )
        )

        return (
            f"Daily enrollment "
            f"({start_month} "
            f"{start_day} - "
            f"{end_month} "
            f"{end_day} "
            f"{year})"
        )

    return (
        f"Daily enrollment "
        f"({pd.Timestamp.now().strftime('%B %Y')})"
    )


clean_filename = generate_clean_filename(
    start_date,
    end_date
)

st.info(
    f"📁 **Report will be saved as:** "
    f"`{clean_filename}.xlsx`"
)


# ============================================================
# EXCEL DATA
# ============================================================

daily_data = {

    "summary":
        daily_summary
        if not daily_summary.empty
        else pd.DataFrame(),

    "avg_ticket_data":
        daily_avg_ticket
        if not daily_avg_ticket.empty
        else pd.DataFrame(),

    "week_data":
        daily_week
        if not daily_week.empty
        else pd.DataFrame(),

    "month_data":
        daily_month
        if not daily_month.empty
        else pd.DataFrame(),

    "year_data":
        daily_year
        if not daily_year.empty
        else pd.DataFrame(),

    "schemes":
        daily_schemes,

    "date_range": (
        start_date.strftime("%d-%m-%Y"),
        end_date.strftime("%d-%m-%Y")
    ),

    "report_title":
        "eGold & eSilver Enrollment & Collection Report",

    "sheet_name":
        "eGold & eSilver"
}


sessional_data = {

    "summary":
        sessional_summary
        if not sessional_summary.empty
        else pd.DataFrame(),

    "avg_ticket_data":
        sessional_avg_ticket
        if not sessional_avg_ticket.empty
        else pd.DataFrame(),

    "week_data":
        sessional_week
        if not sessional_week.empty
        else pd.DataFrame(),

    "month_data":
        sessional_month
        if not sessional_month.empty
        else pd.DataFrame(),

    "year_data":
        sessional_year
        if not sessional_year.empty
        else pd.DataFrame(),

    "schemes":
        sessional_schemes,

    "date_range": (
        start_date.strftime("%d-%m-%Y"),
        end_date.strftime("%d-%m-%Y")
    ),

    "report_title":
        "Sessional Scheme Enrollment & Collection Report",

    "sheet_name":
        "Sessional Scheme"
}


# ============================================================
# CREATE EXCEL
# ============================================================

excel_data = create_formatted_excel(
    daily_data,
    sessional_data
)


# ============================================================
# DOWNLOAD BUTTON
# ============================================================

st.download_button(
    label="⬇️ Download Formatted Excel Report",

    data=excel_data,

    file_name=(
        f"{clean_filename}.xlsx"
    ),

    mime=(
        "application/vnd.openxmlformats-"
        "officedocument.spreadsheetml.sheet"
    ),

    use_container_width=True
)