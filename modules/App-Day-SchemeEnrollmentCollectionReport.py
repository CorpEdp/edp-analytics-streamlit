import streamlit as st
import pandas as pd
import numpy as np
import calendar
import re
import io
from io import BytesIO
from datetime import datetime
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import warnings

warnings.filterwarnings("ignore")


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Scheme Enrollment & Collection Report",
    page_icon="💰",
    layout="wide"
)


# ============================================================
# TITLE
# ============================================================

st.title("💰 Scheme Enrollment & Collection Report")
st.caption("Enrollment = Joined Date filled within period. "
           "Not Enrolled = blank Joined Date with Registered Date inside period.")


# ============================================================
# REQUIRED COLUMNS
# ============================================================

REQUIRED_COLUMNS = [
    "Id", "Scheme Participation Id", "Date", "Status", "Saved Amount",
    "Reward Amount", "Transaction Reference", "Installment number",
    "Metal Type", "Metal Rate", "Saved Metal Weight", "Rewards Metal Weight",
    "Benefit Metal Amount", "Benefit Metal Weight", "Benefit Metal Percentage",
    "Receipt ID", "Customer Name", "Customer Phone Number",
    "Passbook number", "Scheme Name"
]


# ============================================================
# SCHEME CLASSIFICATION
# ============================================================

DAILY_SCHEME_KEYS = [
    "e-gold", "egold", "e gold",
    "e-silver", "esilver", "e silver",
]


def _scheme_priority(scheme):
    if not scheme:
        return 100
    s = str(scheme).lower().strip()
    if "e-gold" in s or "egold" in s or "e gold" in s:
        return 0
    if "e-silver" in s or "esilver" in s or "e silver" in s:
        return 1
    return 50


def classify_scheme_type(scheme_name):
    if pd.isna(scheme_name):
        return "sessional"
    text = str(scheme_name).lower().strip()
    for key in DAILY_SCHEME_KEYS:
        if key in text:
            return "daily"
    return "sessional"


# ============================================================
# SCHEME NAME NORMALISATION
# ============================================================

SCHEME_NORMALISE_MAP = {
    "valatine's day": "Valentine's Day",
    "valentine's day": "Valentine's Day",
    "valentines day": "Valentine's Day",
    "valantine's day": "Valentine's Day",
    "akshaya thrithiyai": "Akshaya Tritiya",
    "akshaya trithiya": "Akshaya Tritiya",
    "akshaya thritiya": "Akshaya Tritiya",
    "akshaya tritiya": "Akshaya Tritiya",
    "e gold": "e-Gold",
    "egold": "e-Gold",
    "e-gold": "e-Gold",
    "e silver": "e-Silver",
    "esilver": "e-Silver",
    "e-silver": "e-Silver",
}


def normalise_scheme(name):
    if name is None:
        return name
    try:
        if pd.isna(name):
            return name
    except Exception:
        pass
    key = str(name).strip().lower()
    if key == "" or key == "nan":
        return str(name).strip()
    return SCHEME_NORMALISE_MAP.get(key, str(name).strip())


SESIONAL_ORDER = {
    "pongal": 10,
    "akshaya tritiya": 20,
    "diwali": 30,
    "christmas": 40,
    "valentine's day": 50,
}


def sessional_sort_key(name):
    return SESIONAL_ORDER.get(str(name).strip().lower(), 999)


# ============================================================
# BRANCH MAPPING
# ============================================================

BRANCH_MAPPING = {
    "bhima jewellery - madurai": "MDU",
    "head office": "MDU",
    "madurai branch": "MDU",
    "bhima jewellery - telecaller": "MDU-Telecalling",
    "bhima jewellery - marthandam": "MDM",
    "bhima jewellery - salem": "SLM",
    "app showroom location": "SLM",
    "in-transit- locations": "SLM",
    "in-transit locations": "SLM",
    "n/a": "Unassigned",
    "bhima jewellery - tirunelveli": "TVL",
    "bhima jewellery - tiruchirappalli": "TCY",
    "trichy branch": "TCY",
    "bhima jewellery -  rajapalayam": "RJPM",
    "bhima jewellery - rajapalayam": "RJPM",
    "rajapalayam branch": "RJPM",
    "bhima jewellery - dindigul": "DGL",
    "dindigul branch": "DGL",
    "bhima jewellery - noida": "ND",
    "bhima jewellery - virudhunagar": "VNR",
    "bhima jewellery -anna nagar": "AN",
    "bhima jewellery - anna nagar": "AN",
    "bhima jewellery - thanjavur": "TJR",
    "bhima jewellery -thanjavur": "TJR",
}

BRANCH_ORDER = [
    "MDU", "MDM", "SLM", "TVL", "TCY", "RJPM",
    "DGL", "ND", "VNR", "TJR", "AN", "MDU-Telecalling",
    "Unassigned",
]


def normalize_branch(branch_name):
    if pd.isna(branch_name):
        return "Skip"
    key = str(branch_name).strip().lower()
    key = " ".join(key.split())
    if not key:
        return "Skip"
    if key in BRANCH_MAPPING:
        return BRANCH_MAPPING[key]
    for k, v in BRANCH_MAPPING.items():
        k_norm = " ".join(k.split())
        if k_norm and k_norm in key:
            return v
    return "Skip"


# ============================================================
# FLEXIBLE BRANCH NAME TRANSFORMATION
# ============================================================

FIXED_BRANCH_GROUPS = [
    ["madurai"], ["marthandam"], ["salem"], ["tirunelveli"],
    ["trichy", "tiruchirappalli", "tiruchirapalli"],
    ["rajapalayam", "rajapalaiyam", "rajapalayem"],
    ["dindigul"], ["noida"], ["virudhunagar", "virudunagar"], ["thanjavur"],
]
TELECALLER_KEYWORDS = ["telecaller", "tele caller", "tellecaller"]


def _is_telecaller_branch(branch):
    return any(k in str(branch).lower() for k in TELECALLER_KEYWORDS)


def _fixed_rank(branch):
    branch_l = str(branch).lower()
    for i, aliases in enumerate(FIXED_BRANCH_GROUPS):
        if any(alias in branch_l for alias in aliases):
            return i
    return None


def build_branch_order(all_branches):
    all_branches = sorted(set(b for b in all_branches if pd.notna(b) and str(b).strip() != ''))
    telecaller = sorted([b for b in all_branches if _is_telecaller_branch(b)])
    remaining = [b for b in all_branches if not _is_telecaller_branch(b)]
    fixed_present = [b for b in remaining if _fixed_rank(b) is not None]
    fixed_present.sort(key=_fixed_rank)
    new_branches = sorted([b for b in remaining if _fixed_rank(b) is None])
    return fixed_present + new_branches + telecaller


def sort_branches_df(df, ordered_branches, branch_col='Branch'):
    if branch_col not in df.columns or len(df) == 0:
        return df
    order_map = {b: i for i, b in enumerate(ordered_branches)}
    df = df.copy()
    df['_branch_sort_key'] = df[branch_col].map(lambda b: order_map.get(b, len(ordered_branches)))
    df = df.sort_values('_branch_sort_key', kind='stable').drop(columns=['_branch_sort_key'])
    return df


# ============================================================
# PHONE / DATE / AMOUNT NORMALIZATION
# ============================================================

def clean_phone_scalar(value):
    if pd.isna(value):
        return ""
    s = str(value).strip()
    if s.endswith(".0"):
        s = s[:-2]
    digits = re.sub(r'\D', '', s)
    if not digits:
        return ""
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    elif len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    if len(digits) > 10:
        digits = digits[-10:]
    return digits


def clean_phone(series):
    return series.apply(clean_phone_scalar)


def _parse_date_flexible(date_val):
    if pd.isna(date_val) or date_val == '':
        return pd.NaT
    try:
        if isinstance(date_val, (datetime, pd.Timestamp)):
            return pd.Timestamp(date_val).normalize()
        date_str = str(date_val).strip()
        if date_str == '':
            return pd.NaT
        formats = [
            '%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d',
            '%d-%m-%y', '%d/%m/%y',
            '%b %d, %Y', '%d %b %Y'
        ]
        for fmt in formats:
            try:
                return pd.Timestamp(datetime.strptime(date_str, fmt)).normalize()
            except Exception:
                continue
        return pd.to_datetime(date_str).normalize()
    except Exception:
        return pd.NaT


def clean_amount_scalar(value):
    if pd.isna(value):
        return np.nan
    s = re.sub(r'[^\d\.]', '', str(value))
    if s == '' or s == '.':
        return np.nan
    try:
        return float(s)
    except Exception:
        return np.nan


def find_column(df, candidates):
    lookup = {str(c).strip().lower(): c for c in df.columns}
    for cand in candidates:
        key = cand.strip().lower()
        if key in lookup:
            return lookup[key]
    return None


# ============================================================
# HELPERS
# ============================================================

def clean_columns(df):
    df.columns = (df.columns.astype(str).str.strip().str.replace(r"\s+", " ", regex=True))
    return df


def clean_amount(series):
    return pd.to_numeric(
        series.astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("₹", "", regex=False)
        .str.replace("Rs.", "", regex=False)
        .str.replace("Rs", "", regex=False)
        .str.strip(),
        errors="coerce"
    ).fillna(0)


def clean_text(series):
    return (series.astype(str).str.strip().replace(["nan", "None", "NaN", ""], np.nan))


def find_duplicate_transaction_references(df):
    """Return duplicate non-blank transaction references with source row details."""
    if "Transaction Reference" not in df.columns:
        return pd.DataFrame(columns=["Transaction Reference", "Occurrences", "Source Rows"])

    references = df["Transaction Reference"].astype("string").str.strip()
    source_rows = pd.Series(range(2, len(df) + 2), index=df.index)
    valid = references.notna() & references.ne("")
    duplicate_mask = valid & references.duplicated(keep=False)
    if not duplicate_mask.any():
        return pd.DataFrame(columns=["Transaction Reference", "Occurrences", "Source Rows"])

    duplicate_rows = pd.DataFrame({
        "Transaction Reference": references[duplicate_mask],
        "Source Row": source_rows[duplicate_mask],
    })
    return (
        duplicate_rows.groupby("Transaction Reference", sort=False)
        .agg(
            Occurrences=("Source Row", "size"),
            **{"Source Rows": ("Source Row", lambda rows: ", ".join(map(str, rows)))},
        )
        .reset_index()
    )


def round_value(value):
    if pd.isna(value):
        return 0
    return round(value)


def round_df(df, exclude_columns=None):
    result = df.copy()
    exclude_columns = exclude_columns or []
    for col in result.columns:
        if col in exclude_columns:
            continue
        if pd.api.types.is_numeric_dtype(result[col]):
            result[col] = result[col].apply(round_value)
    return result


def get_scheme_list(df, scheme_type="all"):
    if "Scheme" not in df.columns:
        return []
    schemes = sorted(df["Scheme"].dropna().unique().tolist())
    if scheme_type == "daily":
        return [s for s in schemes if classify_scheme_type(s) == "daily"]
    elif scheme_type == "sessional":
        return [s for s in schemes if classify_scheme_type(s) == "sessional"]
    return schemes


def get_all_transaction_schemes(df):
    if "Scheme" not in df.columns:
        return []
    schemes = df["Scheme"].dropna().astype(str).str.strip()
    schemes = schemes[schemes != ""]
    schemes = schemes.apply(normalise_scheme)
    return sorted(schemes.unique().tolist())


# ============================================================
# EXCEL FORMATTING
# ============================================================

def apply_cell_style(cell, fill=None, font=None, alignment=None, number_format=None, border=None):
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if alignment:
        cell.alignment = alignment
    if number_format:
        cell.number_format = number_format
    if border:
        cell.border = border


def write_section_header(worksheet, row, title, max_column, color="4472C4"):
    if max_column < 1:
        max_column = 1
    if max_column > 1:
        worksheet.merge_cells(start_row=row, start_column=1,
                              end_row=row, end_column=max_column)
    cell = worksheet.cell(row=row, column=1)
    cell.value = title
    cell.font = Font(size=14, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(1, max_column + 1):
        cc = worksheet.cell(row=row, column=c)
        cc.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        cc.border = Border(
            left=Side(style="thin", color=color),
            right=Side(style="thin", color=color),
            top=Side(style="thin", color=color),
            bottom=Side(style="thin", color=color)
        )
    worksheet.row_dimensions[row].height = 25


def format_dataframe_section(worksheet, dataframe, start_row, header_rows=1, total_identifier=None):
    if dataframe.empty:
        return
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col_idx, column_name in enumerate(dataframe.columns, start=1):
        cell = worksheet.cell(row=start_row, column=col_idx)
        cell.value = column_name
        apply_cell_style(cell, fill=header_fill, font=Font(bold=True),
                         alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
                         border=thin_border)
    for row_idx in range(start_row + 1, start_row + 1 + len(dataframe)):
        is_total = False
        if total_identifier is not None:
            fv = worksheet.cell(row=row_idx, column=1).value
            if fv is not None and total_identifier in str(fv):
                is_total = True
        for col_idx in range(1, len(dataframe.columns) + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            fill = total_fill if is_total else None
            font = Font(bold=True) if is_total else None
            number_format = '#,##0' if col_idx > 1 else None
            apply_cell_style(cell, fill=fill, font=font,
                             alignment=Alignment(horizontal="center", vertical="center"),
                             number_format=number_format, border=thin_border)


def write_projection_section(worksheet, start_row, enrollment_df, collection_df, month_label, thin_border):
    if enrollment_df.empty and collection_df.empty:
        return
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    sub_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    title = "📈 MONTHLY PROJECTION"
    if month_label:
        title = f"{title} - {month_label}"
    projection_width = 7
    write_section_header(worksheet, start_row, title, projection_width, "ED7D31")
    sub_row = start_row + 1
    header_row = start_row + 2
    data_start_row = start_row + 3
    worksheet.merge_cells(start_row=sub_row, start_column=1, end_row=sub_row, end_column=3)
    c = worksheet.cell(row=sub_row, column=1)
    c.value = "Projection of this month's first enrollments"
    c.font = Font(bold=True, italic=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = sub_fill
    worksheet.merge_cells(start_row=sub_row, start_column=5, end_row=sub_row, end_column=7)
    c = worksheet.cell(row=sub_row, column=5)
    c.value = "Projection of this month's overall collection"
    c.font = Font(bold=True, italic=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = sub_fill
    for offset, label in enumerate(["Scheme", "Projected Count", "Projected Amount"]):
        for base_col in (1, 5):
            cell = worksheet.cell(row=header_row, column=base_col + offset)
            cell.value = label
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
    row_count = max(len(enrollment_df), len(collection_df))
    for r in range(row_count):
        row_idx = data_start_row + r
        if r < len(enrollment_df):
            erow = enrollment_df.iloc[r]
            worksheet.cell(row=row_idx, column=1).value = erow["Scheme"]
            worksheet.cell(row=row_idx, column=2).value = erow["Projected Count"]
            worksheet.cell(row=row_idx, column=3).value = erow["Projected Amount"]
        if r < len(collection_df):
            crow = collection_df.iloc[r]
            worksheet.cell(row=row_idx, column=5).value = crow["Scheme"]
            worksheet.cell(row=row_idx, column=6).value = crow["Projected Count"]
            worksheet.cell(row=row_idx, column=7).value = crow["Projected Amount"]
        for col_idx in (1, 2, 3, 5, 6, 7):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            if col_idx in (2, 3, 6, 7):
                cell.number_format = '#,##0'


def write_avg_ticket_section(worksheet, start_row, avg_ticket_data, thin_border):
    if avg_ticket_data.empty:
        return
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    neg_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    neg_font = Font(color="C00000")
    pos_font = Font(color="006100")
    total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    write_section_header(worksheet, start_row, "📊 AVERAGE TICKET SIZE COMPARISON",
                         len(avg_ticket_data.columns), "70AD47")
    header_row = start_row + 1
    data_start_row = start_row + 2
    for col_idx, col_name in enumerate(avg_ticket_data.columns, start=1):
        cell = worksheet.cell(row=header_row, column=col_idx)
        cell.value = col_name
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for r, row in enumerate(avg_ticket_data.values):
        row_idx = data_start_row + r
        is_total = "Total" in str(row[0]) if len(row) > 0 else False
        for col_idx, value in enumerate(row, start=1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            col_name = avg_ticket_data.columns[col_idx - 1]

            is_negative_highlight = False
            is_positive_highlight = False

            if col_name == "Scheme":
                cell.value = value
            elif col_name in ["First Enrollment Count", "Collection Count"]:
                if isinstance(value, (int, float)):
                    cell.value = value
                    cell.number_format = '#,##0'
                else:
                    cell.value = value
            elif col_name in ["First Enrollment Amount", "Collection Amount",
                              "First Enrollment Avg Ticket", "Collection Avg Ticket"]:
                if isinstance(value, (int, float)):
                    cell.value = value
                    cell.number_format = '#,##0'
                else:
                    cell.value = value
            elif col_name == "Difference":
                if isinstance(value, (int, float)):
                    cell.value = value
                    cell.number_format = '#,##0'
                    if value < 0:
                        cell.fill = neg_fill
                        cell.font = neg_font
                        is_negative_highlight = True
                    elif value > 0:
                        cell.font = pos_font
                        is_positive_highlight = True
                else:
                    cell.value = value
            elif col_name == "% Change":
                if isinstance(value, (int, float)):
                    cell.value = value
                    cell.number_format = '0.00%'
                    if value < 0:
                        cell.fill = neg_fill
                        cell.font = neg_font
                        is_negative_highlight = True
                    elif value > 0:
                        cell.font = pos_font
                        is_positive_highlight = True
                else:
                    cell.value = value

            if is_total:
                if not is_negative_highlight:
                    cell.fill = total_fill
                if is_negative_highlight:
                    cell.font = Font(bold=True, color="C00000")
                elif is_positive_highlight:
                    cell.font = Font(bold=True, color="006100")
                else:
                    cell.font = Font(bold=True, color="000000")


def write_sheet(workbook, sheet_data):
    worksheet = workbook.create_sheet(sheet_data["sheet_name"])
    summary_df = sheet_data["summary"]
    enrollment_df = sheet_data["enrollment"]
    collection_df = sheet_data["collection"]
    unique_df = sheet_data["unique"]
    date_range = sheet_data["date_range"]
    report_title = sheet_data["report_title"]
    avg_ticket_data = sheet_data.get("avg_ticket_data", pd.DataFrame())

    title_color = "203764"
    section_color = "4472C4"
    scheme_color = "5B9BD5"
    total_color = "C00000"
    light_blue = "D9E1F2"
    gold_color = "FFC000"

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    max_cols = max(
        len(summary_df.columns) if not summary_df.empty else 1,
        len(enrollment_df.columns) if not enrollment_df.empty else 1,
        len(collection_df.columns) if not collection_df.empty else 1,
        len(unique_df.columns) if not unique_df.empty else 1,
        len(avg_ticket_data.columns) if not avg_ticket_data.empty else 1,
        7
    )

    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_cols)
    tc = worksheet.cell(row=1, column=1)
    tc.value = report_title
    tc.font = Font(size=16, bold=True, color="FFFFFF")
    tc.fill = PatternFill(start_color=title_color, end_color=title_color, fill_type="solid")
    tc.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 30

    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_cols)
    dc = worksheet.cell(row=2, column=1)
    dc.value = f"Report Period: {date_range[0]} to {date_range[1]}"
    dc.font = Font(size=12, bold=True)
    dc.alignment = Alignment(horizontal="center", vertical="center")

    # SCHEME SUMMARY
    summary_section_row = 4
    summary_width = len(summary_df.columns) if not summary_df.empty else max_cols
    write_section_header(worksheet, summary_section_row, "📊 SCHEME SUMMARY",
                         summary_width, section_color)
    summary_header_row = summary_section_row + 1
    if not summary_df.empty:
        for col_idx, col_name in enumerate(summary_df.columns, start=1):
            worksheet.cell(row=summary_header_row, column=col_idx).value = col_name
        for row_idx, row in enumerate(summary_df.values, start=summary_header_row + 1):
            for col_idx, value in enumerate(row, start=1):
                worksheet.cell(row=row_idx, column=col_idx).value = value
    format_dataframe_section(worksheet, summary_df, summary_header_row, total_identifier="Grand Total")

    # AVERAGE TICKET SIZE
    avg_ticket_section_row = summary_header_row + len(summary_df) + 3
    if not avg_ticket_data.empty:
        write_avg_ticket_section(worksheet, avg_ticket_section_row, avg_ticket_data, thin_border)
        avg_ticket_last_row = avg_ticket_section_row + 2 + len(avg_ticket_data) + 1
    else:
        avg_ticket_last_row = summary_header_row + len(summary_df)

    # MONTHLY PROJECTION
    enrollment_projection_df = sheet_data.get("enrollment_projection", pd.DataFrame())
    collection_projection_df = sheet_data.get("collection_projection", pd.DataFrame())
    projection_month_label = sheet_data.get("projection_month_label", "")
    has_projection = not enrollment_projection_df.empty or not collection_projection_df.empty
    projection_section_row = avg_ticket_last_row + 3
    if has_projection:
        write_projection_section(worksheet, projection_section_row,
                                 enrollment_projection_df, collection_projection_df,
                                 projection_month_label, thin_border)
        prc = max(len(enrollment_projection_df), len(collection_projection_df))
        projection_last_row = projection_section_row + 2 + prc
    else:
        projection_last_row = avg_ticket_last_row

    # DAY-WISE NEW ENROLLMENT
    enrollment_section_row = projection_last_row + 3
    enrollment_width = len(enrollment_df.columns) if not enrollment_df.empty else max_cols
    write_section_header(worksheet, enrollment_section_row, "📅 DAY-WISE NEW ENROLLMENT",
                         enrollment_width, section_color)
    enrollment_header_row = enrollment_section_row + 1
    if not enrollment_df.empty:
        for col_idx, col_name in enumerate(enrollment_df.columns, start=1):
            worksheet.cell(row=enrollment_header_row, column=col_idx).value = col_name
        for row_idx, row in enumerate(enrollment_df.values, start=enrollment_header_row + 1):
            for col_idx, value in enumerate(row, start=1):
                worksheet.cell(row=row_idx, column=col_idx).value = value
    for col_idx, col_name in enumerate(enrollment_df.columns if not enrollment_df.empty else [], start=1):
        cell = worksheet.cell(row=enrollment_header_row, column=col_idx)
        if str(col_name).startswith("Total"):
            fill_color, text_color = total_color, "FFFFFF"
        elif "Rate" in str(col_name):
            fill_color, text_color = scheme_color, "FFFFFF"
        else:
            fill_color, text_color = light_blue, "000000"
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        cell.font = Font(bold=True, color=text_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    for row_idx in range(enrollment_header_row + 1, enrollment_header_row + 1 + len(enrollment_df)):
        is_total_row = "Days" in str(worksheet.cell(row=row_idx, column=1).value or "")
        for col_idx in range(1, (len(enrollment_df.columns) + 1) if not enrollment_df.empty else 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            if col_idx > 1:
                cell.number_format = '#,##0'
            if is_total_row:
                cell.fill = PatternFill(start_color=gold_color, end_color=gold_color, fill_type="solid")
                cell.font = Font(bold=True)

    # DAY-WISE COLLECTION
    collection_section_row = enrollment_header_row + len(enrollment_df) + 3
    collection_width = len(collection_df.columns) if not collection_df.empty else max_cols
    write_section_header(worksheet, collection_section_row, "💰 DAY-WISE COLLECTION",
                         collection_width, section_color)
    collection_header_row = collection_section_row + 1
    if not collection_df.empty:
        for col_idx, col_name in enumerate(collection_df.columns, start=1):
            worksheet.cell(row=collection_header_row, column=col_idx).value = col_name
        for row_idx, row in enumerate(collection_df.values, start=collection_header_row + 1):
            for col_idx, value in enumerate(row, start=1):
                worksheet.cell(row=row_idx, column=col_idx).value = value
    for col_idx, col_name in enumerate(collection_df.columns if not collection_df.empty else [], start=1):
        cell = worksheet.cell(row=collection_header_row, column=col_idx)
        if str(col_name).startswith("Total"):
            fill_color, text_color = total_color, "FFFFFF"
        else:
            fill_color, text_color = light_blue, "000000"
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        cell.font = Font(bold=True, color=text_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    for row_idx in range(collection_header_row + 1, collection_header_row + 1 + len(collection_df)):
        is_total_row = "Days" in str(worksheet.cell(row=row_idx, column=1).value or "")
        for col_idx in range(1, (len(collection_df.columns) + 1) if not collection_df.empty else 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            if col_idx > 1:
                cell.number_format = '#,##0'
            if is_total_row:
                cell.fill = PatternFill(start_color=gold_color, end_color=gold_color, fill_type="solid")
                cell.font = Font(bold=True)

    # UNIQUE ENROLLMENT
    unique_section_row = collection_header_row + len(collection_df) + 3
    unique_width = len(unique_df.columns) if not unique_df.empty else max_cols
    write_section_header(worksheet, unique_section_row, "👥 UNIQUE ENROLLMENT",
                         unique_width, section_color)
    unique_header_row = unique_section_row + 1
    if not unique_df.empty:
        for col_idx, col_name in enumerate(unique_df.columns, start=1):
            worksheet.cell(row=unique_header_row, column=col_idx).value = col_name
        for row_idx, row in enumerate(unique_df.values, start=unique_header_row + 1):
            for col_idx, value in enumerate(row, start=1):
                worksheet.cell(row=row_idx, column=col_idx).value = value
    for col_idx, col_name in enumerate(unique_df.columns if not unique_df.empty else [], start=1):
        cell = worksheet.cell(row=unique_header_row, column=col_idx)
        if str(col_name) == "Grand Total":
            fill_color, text_color = total_color, "FFFFFF"
        else:
            fill_color, text_color = light_blue, "000000"
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        cell.font = Font(bold=True, color=text_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    for row_idx in range(unique_header_row + 1, unique_header_row + 1 + len(unique_df)):
        is_total_row = "Days" in str(worksheet.cell(row=row_idx, column=1).value or "")
        for col_idx in range(1, (len(unique_df.columns) + 1) if not unique_df.empty else 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            if col_idx > 1:
                cell.number_format = '#,##0'
            if is_total_row:
                cell.fill = PatternFill(start_color=gold_color, end_color=gold_color, fill_type="solid")
                cell.font = Font(bold=True)

    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is not None:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 25)

    worksheet.column_dimensions["A"].width = 16
    worksheet.freeze_panes = "A5"
    worksheet.sheet_view.showGridLines = False
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True


# ============================================================
# REFERRAL REPORT — PERIOD-SCOPED ENROLLMENT LOGIC
# ============================================================

def _fmt_cell(val):
    if val is None:
        return "-"
    if isinstance(val, float) and pd.isna(val):
        return "-"
    if isinstance(val, (int, np.integer)):
        return val if val != 0 else "-"
    if isinstance(val, (float, np.floating)):
        if val == 0:
            return "-"
        if float(val).is_integer():
            return int(val)
        return val
    if isinstance(val, str) and val.strip() == "":
        return "-"
    return val


def read_uploaded_file(uploaded):
    if uploaded is None:
        return None
    try:
        if uploaded.name.lower().endswith(".csv"):
            return pd.read_csv(uploaded, keep_default_na=False, na_values=[])
        return pd.read_excel(uploaded, keep_default_na=False, na_values=[])
    except Exception as e:
        st.error(f"Error reading {uploaded.name}: {e}")
        return None


def _transform_branch_pretty(branch):
    if pd.isna(branch) or str(branch).strip() == "":
        return "Bhima Jewellery - Customer"
    branch_upper = str(branch).upper().strip()
    if branch_upper == "HEAD OFFICE":
        return "Bhima Jewellery - Madurai"
    elif branch_upper == "IN-TRANSIT- LOCATIONS":
        return "Bhima Jewellery - Salem"
    elif branch_upper == "APP SHOWROOM LOCATION":
        return "Bhima Jewellery - Tirunelveli"
    if "BHIMA JEWELLERY -" in branch_upper:
        parts = str(branch).split("Bhima Jewellery -", 1)
        if len(parts) > 1:
            final_branch = f"Bhima Jewellery - {parts[1].strip().title()}"
        else:
            final_branch = str(branch).title()
    elif str(branch).upper().strip().endswith("BRANCH"):
        cleaned = re.sub(r'BRANCH$', '', str(branch), flags=re.IGNORECASE).strip()
        final_branch = f"Bhima Jewellery - {cleaned.title()}"
    else:
        final_branch = f"Bhima Jewellery - {str(branch).title()}"
    final_branch = final_branch.replace("Tiruchirappalli", "Trichy")
    final_branch = final_branch.replace("TIRUCHIRAPPALLI", "Trichy")
    final_branch = final_branch.replace("tiruchirappalli", "Trichy")
    return final_branch


def build_referral_report(transactions_df, employees_df, referrals_df, bss_df=None,
                          start_date=None, end_date=None,
                          master_scheme_list=None):
    empty = pd.DataFrame()
    diagnostics = {
        "total_referrals": 0, "matched_referrals": 0, "not_enrolled": 0,
        "duplicates": 0, "date_filter": "none", "scope_warning": None,
        "collisions_cleared": 0,
    }
    debug_info = {
        "referral_raw_count": 0,
        "referral_after_date_filter": 0,
        "referral_after_phone_filter": 0,
        "referrals_in_scope": 0,
        "dropped_out_of_period": 0,
        "txn_f1_count": 0,
        "txn_unique_phones": 0,
        "status_distribution": {},
        "referral_phone_sample": [],
        "txn_phone_sample": [],
        "phone_overlap": 0,
        "scheme_name_variants": [],
        "not_enrolled_breakdown": {},
        "date_range_used": (str(start_date), str(end_date)),
        "strict_match_count": 0,
        "weak_match_rejected": 0,
        "collisions_cleared": 0,
        "multi_scheme_same_day": 0,
        "master_scheme_count": len(master_scheme_list) if master_scheme_list else 0,
    }

    if employees_df is None or referrals_df is None:
        return empty, [], empty, diagnostics, empty, debug_info, []
    if employees_df.empty or referrals_df.empty:
        return empty, [], empty, diagnostics, empty, debug_info, []

    emp = clean_columns(employees_df.copy())
    ref = clean_columns(referrals_df.copy())
    txn = transactions_df.copy()

    debug_info["referral_raw_count"] = len(ref)

    emp_ref_code_col = find_column(emp, ["Referral Code", "Referal Code", "Ref Code", "RefCode"])
    emp_branch_col = find_column(emp, ["Branch", "Branch Name", "Showroom", "Location"])
    emp_type_col = find_column(emp, ["Employee Type", "Emp Type", "Category"])
    emp_code_col = find_column(emp, ["Employee Code", "Emp Code"])

    if emp_ref_code_col is None or emp_branch_col is None:
        st.warning("⚠️ Employee file needs 'Referral Code' and 'Branch' columns.")
        return empty, [], empty, diagnostics, empty, debug_info, []

    ref_code_col = find_column(ref, ["Referral Code", "Referal Code", "Ref Code", "RefCode"])
    ref_phone_col = find_column(ref, ["Referee Phone", "Referee Mobile", "Referee Contact",
                                       "Referee Phone Number", "Mobile", "Phone",
                                       "Phone Number", "Mobile Number"])
    referee_name_col = find_column(ref, ["Referee Name", "Referee", "Customer Name"])
    ref_amount_col = find_column(ref, ["Enrollment Amount", "Enroll Amount", "Amount"])
    ref_status_col = find_column(ref, ["Status"])
    referrer_name_col = find_column(ref, ["Referrer Name", "Referer Name"])
    referrer_phone_col = find_column(ref, ["Referrer Phone", "Referer Phone"])
    reg_date_col = find_column(ref, ["Registered Date", "Registration Date", "Register Date",
                                      "Referral Date", "Ref Date"])
    join_date_col = find_column(ref, ["Joined Date", "Join Date", "Joining Date"])
    ref_scheme_col = find_column(ref, ["Scheme", "Scheme Name", "Scheme Type"])

    if ref_code_col is None:
        st.error("❌ Referral file needs 'Referral Code'.")
        return empty, [], empty, diagnostics, empty, debug_info, []
    if ref_phone_col is None:
        st.error("❌ Referral file needs 'Referee Phone'.")
        return empty, [], empty, diagnostics, empty, debug_info, []
    if join_date_col is None:
        st.error("❌ Referral file needs 'Joined Date' — it decides Enrollment vs Not Enrolled.")
        return empty, [], empty, diagnostics, empty, debug_info, []

    ref_code_to_branch = dict(zip(emp[emp_ref_code_col].astype(str).str.strip(), emp[emp_branch_col]))
    ref_code_to_empcode = {}
    if emp_code_col is not None:
        ref_code_to_empcode = dict(zip(emp[emp_ref_code_col].astype(str).str.strip(),
                                        emp[emp_code_col].astype(str)))
    ref_code_to_emptype = {}
    if emp_type_col is not None:
        ref_code_to_emptype = dict(zip(emp[emp_ref_code_col].astype(str).str.strip(),
                                        emp[emp_type_col].astype(str)))

    fr = pd.DataFrame()
    fr['Customer Name'] = ref[referee_name_col] if referee_name_col else ""
    fr['Customer Phone'] = ref[ref_phone_col].apply(clean_phone_scalar)
    fr['Customer Enrollment Amount'] = (
        ref[ref_amount_col].apply(clean_amount_scalar) if ref_amount_col else np.nan
    )
    fr['Status'] = ref[ref_status_col].astype(str) if ref_status_col else ""
    fr['Employee Name'] = ref[referrer_name_col] if referrer_name_col else ""
    fr['Referral Code'] = ref[ref_code_col].astype(str).str.strip()
    fr['Employee Phone'] = (
        ref[referrer_phone_col].apply(clean_phone_scalar) if referrer_phone_col else ""
    )
    fr['Employee Code'] = fr['Referral Code'].map(ref_code_to_empcode).fillna('')
    raw_branch = fr['Referral Code'].map(ref_code_to_branch)
    fr['Branch'] = raw_branch.apply(_transform_branch_pretty)
    fr['Category'] = fr['Referral Code'].map(ref_code_to_emptype).fillna('Customer')

    fr['Joined Date'] = ref[join_date_col].apply(_parse_date_flexible)
    if reg_date_col:
        fr['Registered Date'] = ref[reg_date_col].apply(_parse_date_flexible)
    else:
        fr['Registered Date'] = pd.NaT
    fr['Referral Scheme'] = ref[ref_scheme_col].astype(str) if ref_scheme_col else ""

    if ref_status_col:
        status_counts = fr['Status'].value_counts().head(20).to_dict()
        debug_info["status_distribution"] = {str(k): int(v) for k, v in status_counts.items()}

    has_registered_col = reg_date_col is not None

    if start_date is not None and end_date is not None:
        start_ts = pd.Timestamp(start_date)
        end_ts = pd.Timestamp(end_date)

        joined_in_range = fr['Joined Date'].between(start_ts, end_ts)
        enrolled_mask = fr['Joined Date'].notna() & joined_in_range

        if has_registered_col:
            registered_in_range = fr['Registered Date'].between(start_ts, end_ts)
            not_enrolled_mask = fr['Joined Date'].isna() & registered_in_range
        else:
            not_enrolled_mask = fr['Joined Date'].isna()
            diagnostics["scope_warning"] = (
                "Referral file has no 'Registered Date' column — Not Enrolled rows "
                "cannot be scoped to the selected period and will be counted from the "
                "entire file."
            )

        diagnostics["date_filter"] = f"{start_date} to {end_date}"
    else:
        enrolled_mask = fr['Joined Date'].notna()
        not_enrolled_mask = fr['Joined Date'].isna()
        if not has_registered_col:
            diagnostics["scope_warning"] = (
                "Referral file has no 'Registered Date' column — Not Enrolled rows "
                "cannot be scoped to the selected period."
            )
        diagnostics["date_filter"] = "none"

    fr_filtered = fr[enrolled_mask | not_enrolled_mask].copy()

    if has_registered_col:
        fr_filtered = fr_filtered[
            fr_filtered['Joined Date'].notna() | fr_filtered['Registered Date'].notna()
        ].copy()

    debug_info["referral_after_date_filter"] = len(fr_filtered)
    debug_info["dropped_out_of_period"] = len(fr) - len(fr_filtered)
    debug_info["referral_after_phone_filter"] = len(fr_filtered)
    debug_info["referrals_in_scope"] = len(fr_filtered)
    debug_info["referral_phone_sample"] = fr_filtered['Customer Phone'].head(10).tolist()

    if 'Installment number' in txn.columns:
        txn['Installment number'] = pd.to_numeric(txn['Installment number'], errors='coerce')
        txn_f1 = txn[txn['Installment number'] == 1].copy()
    else:
        txn_f1 = txn.copy()

    txn_f1 = txn_f1.reset_index(drop=True).copy()
    txn_f1['_row_order'] = txn_f1.index

    txn_phone_col = find_column(txn_f1, ["Customer Phone Number", "Customer Phone", "Mobile"])
    txn_date_col = find_column(txn_f1, ["Date"])
    txn_amount_col = find_column(txn_f1, ["Saved Amount", "Amount"])
    txn_scheme_col = find_column(txn_f1, ["Scheme Name", "Scheme"])
    txn_passbook_col = find_column(txn_f1, ["Passbook number", "Passbook Number", "Doc No"])

    txn_lookup_ready = (
        txn_phone_col is not None and txn_date_col is not None and txn_amount_col is not None
    )

    txn_by_phone_date = defaultdict(list)
    if txn_lookup_ready:
        txn_f1['_phone'] = txn_f1[txn_phone_col].apply(clean_phone_scalar)
        txn_f1['_date'] = txn_f1[txn_date_col].apply(_parse_date_flexible)
        txn_f1['_amount'] = txn_f1[txn_amount_col].apply(clean_amount_scalar)
        txn_f1['_scheme'] = (
            txn_f1[txn_scheme_col].astype(str).apply(normalise_scheme)
            if txn_scheme_col else ''
        )
        txn_f1['_passbook'] = txn_f1[txn_passbook_col].astype(str) if txn_passbook_col else ''

        debug_info["txn_f1_count"] = len(txn_f1)
        debug_info["txn_unique_phones"] = txn_f1['_phone'].nunique()
        debug_info["txn_phone_sample"] = txn_f1['_phone'].head(10).tolist()

        if txn_scheme_col:
            scheme_variants = txn_f1['_scheme'].value_counts().head(30).to_dict()
            debug_info["scheme_name_variants"] = {str(k): int(v) for k, v in scheme_variants.items()}

        for idx, row in txn_f1.iterrows():
            p = row['_phone']
            d = row['_date']
            if not p or pd.isna(d):
                continue
            key = (p, d.date())
            txn_by_phone_date[key].append({
                'amount': row['_amount'],
                'scheme': row['_scheme'],
                'passbook': row['_passbook'],
                'date': d,
                '_row_order': row['_row_order'],
            })
    else:
        st.warning("⚠️ Transaction file missing required columns — transaction enrichment will be skipped.")

    ref_phones = set(fr_filtered['Customer Phone'].tolist())
    txn_phones = set(txn_f1['_phone'].tolist()) if txn_lookup_ready else set()
    debug_info["phone_overlap"] = len(ref_phones & txn_phones)

    matched_rows = []
    not_enrolled_reasons = defaultdict(int)
    strict_match_count = 0
    weak_match_rejected = 0
    multi_scheme_same_day = 0

    for idx, row in fr_filtered.iterrows():
        phone = row['Customer Phone']
        joined_date = row['Joined Date']
        enrollment_amount = row['Customer Enrollment Amount']
        branch = row['Branch']
        referrer = row['Employee Name']
        referee = row['Customer Name']
        emp_code = row['Employee Code']
        emp_category = row['Category']
        reg_date = row['Registered Date']
        ref_scheme = row['Referral Scheme']
        ref_scheme_norm = normalise_scheme(ref_scheme) if ref_scheme else ''

        if pd.isna(joined_date):
            not_enrolled_reasons['no joined date (registered in period)'] += 1
            matched_rows.append({
                "Branch": branch, "Referrer Name": referrer, "Referee Name": referee,
                "Referee Phone": phone, "Employee Code": emp_code,
                "Registered Date": reg_date, "Joined Date": pd.NaT,
                "Transaction Date": pd.NaT, "Scheme": None, "Installment": None,
                "Enrollment Amount": float(enrollment_amount) if pd.notna(enrollment_amount) else None,
                "Paid Amount": None, "Passbook": "",
                "Match Reason": "Not Enrolled (no joined date, registered in period)",
                "Category": "Not Enrolled",
                "Employee Category": emp_category,
                "Is Duplicate": False,
            })
            continue

        enrichment = None
        enrich_reason = "Enrolled (no transaction match)"

        if txn_lookup_ready and phone:
            key = (phone, joined_date.date())
            candidates = txn_by_phone_date.get(key, [])

            exact = [
                c for c in candidates
                if pd.notna(c.get('amount')) and pd.notna(enrollment_amount)
                and abs(c['amount'] - enrollment_amount) <= 0.01
            ]
            if exact:
                if len(exact) > 1:
                    multi_scheme_same_day += 1
                exact_sorted = sorted(
                    exact,
                    key=lambda c: (
                        _scheme_priority(c.get('scheme', '')),
                        c.get('_row_order', 10**18)
                    )
                )
                enrichment = exact_sorted[0]
                enrich_reason = "Enrolled (matched txn: phone + date + amount)"
                strict_match_count += 1
            elif candidates:
                weak_match_rejected += 1
                enrichment = None
                enrich_reason = "Enrolled (no transaction match)"

        if enrichment:
            txn_scheme = str(enrichment.get('scheme', '')).strip()
            passbook = str(enrichment.get('passbook', '')).strip()
            paid_amount = enrichment.get('amount')
            txn_date = enrichment.get('date')

            scheme_source = txn_scheme
            if not scheme_source or scheme_source.lower() == 'nan':
                scheme_source = ref_scheme_norm

            if scheme_source and str(scheme_source).strip().lower() != 'nan':
                scheme_source = normalise_scheme(scheme_source)
                if classify_scheme_type(scheme_source) == "daily":
                    cat = "e-Silver" if "silver" in str(scheme_source).lower() else "e-Gold"
                else:
                    cat = str(scheme_source).strip()
            else:
                cat = "Enrolled (Unknown Scheme)"

            matched_rows.append({
                "Branch": branch, "Referrer Name": referrer, "Referee Name": referee,
                "Referee Phone": phone, "Employee Code": emp_code,
                "Registered Date": reg_date, "Joined Date": joined_date,
                "Transaction Date": txn_date,
                "Scheme": txn_scheme or scheme_source,
                "Installment": 1,
                "Enrollment Amount": float(enrollment_amount) if pd.notna(enrollment_amount) else None,
                "Paid Amount": float(paid_amount) if pd.notna(paid_amount) else None,
                "Passbook": passbook,
                "Match Reason": enrich_reason,
                "Category": cat,
                "Employee Category": emp_category,
                "Is Duplicate": False,
            })
        else:
            fallback_cat = "Enrolled (Unknown Scheme)"
            if ref_scheme_norm and str(ref_scheme_norm).strip().lower() != 'nan':
                ref_norm = normalise_scheme(ref_scheme_norm)
                if classify_scheme_type(ref_norm) == "daily":
                    fallback_cat = "e-Silver" if "silver" in str(ref_norm).lower() else "e-Gold"
                else:
                    fallback_cat = str(ref_norm).strip()

            matched_rows.append({
                "Branch": branch, "Referrer Name": referrer, "Referee Name": referee,
                "Referee Phone": phone, "Employee Code": emp_code,
                "Registered Date": reg_date, "Joined Date": joined_date,
                "Transaction Date": pd.NaT, "Scheme": None, "Installment": None,
                "Enrollment Amount": float(enrollment_amount) if pd.notna(enrollment_amount) else None,
                "Paid Amount": None, "Passbook": "",
                "Match Reason": enrich_reason,
                "Category": fallback_cat,
                "Employee Category": emp_category,
                "Is Duplicate": False,
            })

    detail_df = pd.DataFrame(matched_rows)
    debug_info["not_enrolled_breakdown"] = dict(not_enrolled_reasons)
    debug_info["strict_match_count"] = strict_match_count
    debug_info["weak_match_rejected"] = weak_match_rejected
    debug_info["multi_scheme_same_day"] = multi_scheme_same_day

    if detail_df.empty:
        return empty, [], empty, diagnostics, empty, debug_info, []

    has_passbook_mask = (
        detail_df['Passbook'].notna()
        & (detail_df['Passbook'].astype(str).str.strip() != '')
        & (detail_df['Passbook'].astype(str).str.strip().str.lower() != 'nan')
    )
    collisions_cleared = 0
    if has_passbook_mask.any():
        claimed = detail_df[has_passbook_mask]
        pb_counts = claimed.groupby('Passbook').size()
        colliding_pbs = pb_counts[pb_counts > 1].index.tolist()

        if colliding_pbs:
            for pb in colliding_pbs:
                group = detail_df[
                    (detail_df['Passbook'].astype(str) == str(pb))
                ].sort_values('Joined Date', kind='stable')
                drop_idx = group.index[1:]
                for di in drop_idx:
                    detail_df.loc[di, 'Paid Amount'] = None
                    detail_df.loc[di, 'Passbook'] = ''
                    detail_df.loc[di, 'Transaction Date'] = pd.NaT
                    detail_df.loc[di, 'Scheme'] = None
                    detail_df.loc[di, 'Installment'] = None
                    detail_df.loc[di, 'Match Reason'] = 'Enrolled (no transaction match)'
                    detail_df.loc[di, 'Is Duplicate'] = True
                    collisions_cleared += 1

            if collisions_cleared > 0:
                st.warning(
                    f"⚠️ **Transaction collision:** {collisions_cleared} row(s) shared a "
                    f"passbook with another row and had their transaction data cleared. "
                    f"Only the first row keeps the reference."
                )

    diagnostics["collisions_cleared"] = collisions_cleared
    debug_info["collisions_cleared"] = collisions_cleared

    duplicates_df = pd.DataFrame()
    has_passbook = detail_df['Passbook'].notna() & (detail_df['Passbook'].astype(str).str.strip() != '')
    if has_passbook.any():
        detail_df['_dup_key'] = (
            detail_df['Passbook'].astype(str) + '_' +
            detail_df['Referrer Name'].astype(str) + '_' +
            detail_df['Branch'].astype(str)
        )
        key_counts = detail_df[has_passbook]['_dup_key'].value_counts()
        dup_keys = key_counts[key_counts > 1].index.tolist()
        if dup_keys:
            dup_rows = []
            for grp_idx, dup_key in enumerate(dup_keys, start=1):
                grp = detail_df[detail_df['_dup_key'] == dup_key].sort_values('Joined Date')
                orig_idx = grp.index[0]
                for di in grp.index[1:]:
                    detail_df.loc[di, 'Is Duplicate'] = True
                    orig = detail_df.loc[orig_idx]
                    dup = detail_df.loc[di]
                    dup_rows.append({
                        'Group ID': f'GROUP_{grp_idx}', 'Passbook': orig['Passbook'],
                        'Original Employee': orig['Referrer Name'],
                        'Original Customer': orig['Referee Name'],
                        'Original Date': orig['Joined Date'],
                        'Original Enrollment Amount': orig['Enrollment Amount'],
                        'Duplicate Employee': dup['Referrer Name'],
                        'Duplicate Customer': dup['Referee Name'],
                        'Duplicate Date': dup['Joined Date'],
                        'Duplicate Enrollment Amount': dup['Enrollment Amount'],
                        'Branch': orig['Branch'], 'Category': orig['Category'],
                    })
            duplicates_df = pd.DataFrame(dup_rows)
            st.warning(f"⚠️ **Duplicates detected:** {len(duplicates_df)} record(s) in {len(dup_keys)} group(s).")
        detail_df = detail_df.drop(columns=['_dup_key'])

    total = len(detail_df)
    enrolled = int((detail_df['Category'] != 'Not Enrolled').sum())
    not_enrolled = total - enrolled
    diagnostics.update({
        "total_referrals": total,
        "matched_referrals": enrolled,
        "not_enrolled": not_enrolled,
        "duplicates": int(detail_df['Is Duplicate'].sum()),
    })

    branch_universe = detail_df['Branch'].dropna().unique().tolist()
    ordered_branches = build_branch_order(branch_universe)

    known_categories = set(detail_df['Category'].dropna().unique())

    if master_scheme_list:
        for scheme in master_scheme_list:
            s = str(scheme).strip()
            if not s:
                continue
            s_norm = normalise_scheme(s)
            if classify_scheme_type(s_norm) == "daily":
                s_lower = s_norm.lower()
                if "silver" in s_lower:
                    known_categories.add("e-Silver")
                else:
                    known_categories.add("e-Gold")
            else:
                known_categories.add(s_norm)

    category_order = [c for c in ["e-Gold", "e-Silver"] if c in known_categories]

    sessional_present = sorted(
        [c for c in known_categories
         if c not in ("e-Gold", "e-Silver", "Not Enrolled", "Enrolled (Unknown Scheme)")],
        key=sessional_sort_key
    )
    category_order += sessional_present

    if "Enrolled (Unknown Scheme)" in known_categories:
        category_order.append("Enrolled (Unknown Scheme)")

    category_order.append("Not Enrolled")

    all_cats = set(detail_df['Category'].dropna().unique())
    missing = all_cats - set(category_order)
    if missing:
        category_order.extend(sorted(missing))

    pivot = detail_df.pivot_table(
        index="Branch", columns="Category",
        values="Referee Phone", aggfunc="count", fill_value=0
    )
    for cat in category_order:
        if cat not in pivot.columns:
            pivot[cat] = 0
    pivot = pivot[category_order]
    for b in ordered_branches:
        if b not in pivot.index:
            pivot.loc[b] = 0
    pivot = pivot.reindex(ordered_branches)
    pivot["Grand Total"] = pivot.sum(axis=1)
    total_row = pivot.sum(axis=0)
    total_row.name = "Total"
    pivot = pd.concat([pivot, pd.DataFrame([total_row])])
    pivot = pivot.reset_index().rename(columns={"index": "Showroom"})

    detail_df['Joined Date Str'] = detail_df['Joined Date'].apply(
        lambda x: x.strftime('%d-%m-%Y') if pd.notna(x) else None
    )
    detail_df['Registered Date Str'] = detail_df['Registered Date'].apply(
        lambda x: x.strftime('%d-%m-%Y') if pd.notna(x) else None
    )
    detail_df['Match Date'] = detail_df['Joined Date Str']

    detail_df['Display Day'] = detail_df['Joined Date Str'].where(
        detail_df['Joined Date Str'].notna(),
        detail_df['Registered Date Str']
    )

    def _build_day_pivot(day_df):
        if day_df.empty:
            return pd.DataFrame()
        p = day_df.pivot_table(
            index="Branch", columns="Category",
            values="Referee Phone", aggfunc="count", fill_value=0
        )
        for cat in category_order:
            if cat not in p.columns:
                p[cat] = 0
        p = p[category_order]
        for b in ordered_branches:
            if b not in p.index:
                p.loc[b] = 0
        p = p.reindex(ordered_branches)
        p["Grand Total"] = p.sum(axis=1)
        trow = p.sum(axis=0)
        trow.name = "Total"
        p = pd.concat([p, pd.DataFrame([trow])])
        return p.reset_index().rename(columns={"index": "Showroom"})

    daily_list = []
    all_days = sorted(
        [d for d in detail_df['Display Day'].dropna().unique()],
        key=lambda x: pd.to_datetime(x, format='%d-%m-%Y')
    )
    for date_label in all_days:
        day_df = detail_df[detail_df['Display Day'] == date_label]
        day_pivot = _build_day_pivot(day_df)
        if not day_pivot.empty:
            daily_list.append((date_label, day_pivot))

    daily_not_enrolled_list = []

    return pivot, daily_list, detail_df, diagnostics, duplicates_df, debug_info, daily_not_enrolled_list


# ============================================================
# REFERRAL EXCEL WRITER
# ============================================================

def write_referral_sheet(workbook, summary_df, daily_list, detail_df, duplicates_df,
                         date_range, sheet_name="Referral Report",
                         not_enrolled_daily=None):
    if summary_df is None or summary_df.empty:
        return
    ws = workbook.create_sheet(sheet_name)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    total_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    dup_fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
    section_color = "7030A0"

    max_cols = len(summary_df.columns)
    if daily_list:
        for _, day_df in daily_list:
            max_cols = max(max_cols, len(day_df.columns))
    if detail_df is not None and not detail_df.empty:
        max_cols = max(max_cols, len(detail_df.columns))

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_cols)
    tc = ws.cell(row=1, column=1)
    tc.value = "🎁 REFERRAL CONVERSION REPORT"
    tc.font = Font(size=16, bold=True, color="FFFFFF")
    tc.fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
    tc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_cols)
    dc = ws.cell(row=2, column=1)
    dc.value = f"Report Period: {date_range[0]} to {date_range[1]}"
    dc.font = Font(size=12, bold=True)
    dc.alignment = Alignment(horizontal="center", vertical="center")

    write_section_header(
        ws, 4,
        f"🏢 BRANCH-WISE REFERRAL SUMMARY  ({date_range[0]} to {date_range[1]})",
        len(summary_df.columns),
        section_color
    )
    summary_header_row = 5
    for c_idx, col_name in enumerate(summary_df.columns, start=1):
        cell = ws.cell(row=summary_header_row, column=c_idx)
        cell.value = col_name
        is_total_col = col_name == "Grand Total"
        is_ne_col = col_name == "Not Enrolled"
        if is_total_col:
            fill_color, text_color = "C00000", "FFFFFF"
        elif is_ne_col:
            fill_color, text_color = "FCE4D6", "000000"
        else:
            fill_color, text_color = "D9E1F2", "000000"
        cell.font = Font(bold=True, color=text_color)
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    for r_idx, row in enumerate(summary_df.values, start=summary_header_row + 1):
        is_total_row = str(row[0]).strip().lower() == "total"
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.value = _fmt_cell(val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin
            if c_idx > 1 and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'
            if is_total_row:
                cell.font = Font(bold=True)
                cell.fill = total_fill

    last_row = summary_header_row + len(summary_df)

    if daily_list:
        for date_label, day_df in daily_list:
            last_row += 3
            write_section_header(
                ws, last_row,
                f"📅 REFERRAL — {date_label}",
                len(day_df.columns),
                section_color
            )
            d_head_row = last_row + 1
            for c_idx, col_name in enumerate(day_df.columns, start=1):
                cell = ws.cell(row=d_head_row, column=c_idx)
                cell.value = col_name
                is_total_col = col_name == "Grand Total"
                is_ne_col = col_name == "Not Enrolled"
                if is_total_col:
                    fill_color, text_color = "C00000", "FFFFFF"
                elif is_ne_col:
                    fill_color, text_color = "FCE4D6", "000000"
                else:
                    fill_color, text_color = "D9E1F2", "000000"
                cell.font = Font(bold=True, color=text_color)
                cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin
            for r_idx, row in enumerate(day_df.values, start=d_head_row + 1):
                is_total_row = str(row[0]).strip().lower() == "total"
                for c_idx, val in enumerate(row, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx)
                    cell.value = _fmt_cell(val)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin
                    if c_idx > 1 and isinstance(cell.value, (int, float)):
                        cell.number_format = '#,##0'
                    if is_total_row:
                        cell.font = Font(bold=True)
                        cell.fill = total_fill
            last_row = d_head_row + len(day_df)

    if detail_df is not None and not detail_df.empty:
        last_row += 3
        display_cols = [c for c in detail_df.columns
                        if c not in ("Referee Phone", "Match Date",
                                     "Joined Date Str", "Registered Date Str",
                                     "Display Day")]
        write_section_header(
            ws, last_row,
            "📋 REFERRAL MATCH DETAILS",
            len(display_cols),
            section_color
        )
        detail_header_row = last_row + 1
        detail_display = detail_df[display_cols].copy()
        for col in ["Registered Date", "Joined Date", "Transaction Date"]:
            if col in detail_display.columns:
                detail_display[col] = detail_display[col].apply(
                    lambda x: x.strftime("%d-%m-%Y") if pd.notna(x) else ""
                )
        for c_idx, col_name in enumerate(detail_display.columns, start=1):
            cell = ws.cell(row=detail_header_row, column=c_idx)
            cell.value = col_name
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin
        is_dup_col_idx = None
        if "Is Duplicate" in detail_display.columns:
            is_dup_col_idx = list(detail_display.columns).index("Is Duplicate") + 1
        for r_idx, row in enumerate(detail_display.values, start=detail_header_row + 1):
            is_dup = bool(row[is_dup_col_idx - 1]) if is_dup_col_idx else False
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin
                if is_dup:
                    cell.fill = dup_fill

    if duplicates_df is not None and not duplicates_df.empty:
        dup_ws = workbook.create_sheet("Duplicate Records")
        for c_idx, col_name in enumerate(duplicates_df.columns, start=1):
            cell = dup_ws.cell(row=1, column=c_idx)
            cell.value = col_name
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin
        for r_idx, row in enumerate(duplicates_df.values, start=2):
            for c_idx, val in enumerate(row, start=1):
                cell = dup_ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin
                cell.fill = dup_fill
        for col_idx in range(1, len(duplicates_df.columns) + 1):
            letter = get_column_letter(col_idx)
            max_len = 12
            for cell in dup_ws[letter]:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            dup_ws.column_dimensions[letter].width = min(max_len + 2, 30)
        dup_ws.sheet_view.showGridLines = False

    for col_idx in range(1, max_cols + 1):
        letter = get_column_letter(col_idx)
        max_len = 12
        for cell in ws[letter]:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 28)

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ============================================================
# CREATE FORMATTED EXCEL
# ============================================================

def create_formatted_excel(daily_data, sessional_data, date_range, report_title,
                           referral_summary=None, referral_daily=None,
                           referral_detail=None, referral_duplicates=None,
                           referral_not_enrolled_daily=None):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        workbook = writer.book
        if "Sheet" in workbook.sheetnames:
            del workbook["Sheet"]
        sheets_written = 0
        if daily_data["schemes"]:
            write_sheet(workbook, daily_data)
            sheets_written += 1
        if sessional_data["schemes"]:
            write_sheet(workbook, sessional_data)
            sheets_written += 1
        if referral_summary is not None and not referral_summary.empty:
            write_referral_sheet(
                workbook, referral_summary,
                referral_daily if referral_daily is not None else [],
                referral_detail if referral_detail is not None else pd.DataFrame(),
                referral_duplicates if referral_duplicates is not None else pd.DataFrame(),
                date_range, sheet_name="Referral Report",
                not_enrolled_daily=[]
            )
            sheets_written += 1
        if sheets_written == 0:
            placeholder = workbook.create_sheet("No Data")
            placeholder["A1"] = "No data matched the selected filters."
    output.seek(0)
    return output.getvalue()


# ============================================================
# AVG TICKET SIZE
# ============================================================

def add_avg_ticket_size_comparison(summary_df):
    if summary_df.empty:
        return pd.DataFrame()
    scheme_rows = summary_df[summary_df["Scheme"] != "Grand Total"].copy()
    if scheme_rows.empty:
        return pd.DataFrame()
    result_data = []
    for _, row in scheme_rows.iterrows():
        fc = row["First Enrollment Count"]
        fa = row["First Enrollment Amount"]
        cc = row["Collection Count"]
        ca = row["Collection Amount"]
        f_avg = fa / fc if fc > 0 else 0
        c_avg = ca / cc if cc > 0 else 0
        diff = c_avg - f_avg
        if f_avg > 0:
            pct_change = round(diff / f_avg, 4)
        elif c_avg > 0:
            pct_change = "N/A"
        else:
            pct_change = 0
        result_data.append({
            "Scheme": row["Scheme"],
            "First Enrollment Count": int(fc),
            "First Enrollment Amount": int(fa),
            "First Enrollment Avg Ticket": int(round(f_avg)),
            "Collection Count": int(cc),
            "Collection Amount": int(ca),
            "Collection Avg Ticket": int(round(c_avg)),
            "Difference": int(round(diff)),
            "% Change": pct_change,
        })
    tf = sum(d["First Enrollment Count"] for d in result_data)
    ta = sum(d["First Enrollment Amount"] for d in result_data)
    cf = sum(d["Collection Count"] for d in result_data)
    ca = sum(d["Collection Amount"] for d in result_data)
    tfa = ta / tf if tf > 0 else 0
    tca = ca / cf if cf > 0 else 0
    tdiff = tca - tfa
    tpct = round(tdiff / tfa, 4) if tfa > 0 else ("N/A" if tca > 0 else 0)
    result_data.append({
        "Scheme": "Grand Total", "First Enrollment Count": int(tf),
        "First Enrollment Amount": int(ta), "First Enrollment Avg Ticket": int(round(tfa)),
        "Collection Count": int(cf), "Collection Amount": int(ca),
        "Collection Avg Ticket": int(round(tca)),
        "Difference": int(round(tdiff)), "% Change": tpct,
    })
    return pd.DataFrame(result_data)


def display_avg_ticket_comparison(avg_ticket_df):
    if avg_ticket_df.empty:
        return
    st.subheader("📊 Average Ticket Size Comparison")
    display_df = avg_ticket_df.copy()
    for col in ["First Enrollment Count", "Collection Count"]:
        if col in display_df.columns:
            display_df[col] = display_df[col].apply(
                lambda x: f"{int(x):,}" if isinstance(x, (int, float)) else x
            )
    for col in ["First Enrollment Amount", "Collection Amount",
                "First Enrollment Avg Ticket", "Collection Avg Ticket", "Difference"]:
        if col in display_df.columns:
            display_df[col] = display_df[col].apply(
                lambda x: f"{int(x):,}" if isinstance(x, (int, float)) else x
            )
    if "% Change" in display_df.columns:
        display_df["% Change"] = display_df["% Change"].apply(
            lambda x: f"{x * 100:.2f}%" if isinstance(x, (int, float)) else x
        )
    st.dataframe(display_df, width="stretch", hide_index=True)


# ============================================================
# FILE UPLOAD
# ============================================================

st.sidebar.header("📂 File Uploads")

uploaded_file = st.sidebar.file_uploader(
    "1️⃣ Main Transaction File (Required)",
    type=["xlsx", "xls", "csv"], key="main_file"
)
uploaded_employee = st.sidebar.file_uploader(
    "2️⃣ Employee Details File (For Referral Report)",
    type=["xlsx", "xls", "csv"], key="employee_file"
)
uploaded_referral = st.sidebar.file_uploader(
    "3️⃣ Referral Details File (For Referral Report)",
    type=["xlsx", "xls", "csv"], key="referral_file"
)

if uploaded_file is None:
    st.info("Please upload your raw transaction Excel/CSV file (1️⃣) from the sidebar.")
    st.stop()


# ============================================================
# READ MAIN FILE
# ============================================================

try:
    if uploaded_file.name.lower().endswith(".csv"):
        df = pd.read_csv(uploaded_file, low_memory=False)
    else:
        df = pd.read_excel(uploaded_file)
except Exception as e:
    st.error(f"❌ Error reading file: {e}")
    st.stop()

df = clean_columns(df)
missing_columns = [col for col in REQUIRED_COLUMNS if col not in df.columns]
if missing_columns:
    st.error("❌ Required columns are missing.")
    st.write("Missing columns:", missing_columns)
    st.stop()

duplicate_references = find_duplicate_transaction_references(df)
if not duplicate_references.empty:
    st.error(
        "❌ Duplicate Transaction Reference values were detected in the Main Transaction File. "
        "Please correct the source file before generating the report."
    )
    st.dataframe(duplicate_references, width="stretch", hide_index=True)
    st.stop()

df["Date"] = pd.to_datetime(df["Date"], errors="coerce", dayfirst=True)
df["Saved Amount"] = clean_amount(df["Saved Amount"])
df["Metal Rate"] = clean_amount(df["Metal Rate"])
df["Installment number"] = pd.to_numeric(df["Installment number"], errors="coerce")
df["Passbook number"] = clean_text(df["Passbook number"])
df["Scheme Name"] = clean_text(df["Scheme Name"])
df["Customer Phone Number"] = clean_text(df["Customer Phone Number"])
df["Customer Name"] = clean_text(df["Customer Name"])

df = df[df["Date"].notna()].copy()
if df.empty:
    st.error("❌ No valid Date records found.")
    st.stop()
df["Date"] = df["Date"].dt.normalize()

rows_before = len(df)
df = df[df["Scheme Name"].notna()].copy()
if rows_before - len(df) > 0:
    st.sidebar.warning(f"⚠️ Skipped {rows_before - len(df)} row(s) with a blank 'Scheme Name'.")

df["Scheme"] = df["Scheme Name"].astype(str).str.strip().apply(normalise_scheme)


# ============================================================
# FILTERS
# ============================================================

st.sidebar.header("🔎 Report Filters")
minimum_date = df["Date"].min().date()
maximum_date = df["Date"].max().date()

date_range = st.sidebar.date_input(
    "Date Range", value=(minimum_date, maximum_date),
    min_value=minimum_date, max_value=maximum_date
)

if isinstance(date_range, tuple):
    if len(date_range) == 2:
        start_date = pd.Timestamp(date_range[0])
        end_date = pd.Timestamp(date_range[1])
    else:
        start_date = pd.Timestamp(date_range[0])
        end_date = start_date
else:
    start_date = pd.Timestamp(date_range)
    end_date = start_date

all_schemes = get_scheme_list(df)
selected_schemes = st.sidebar.multiselect("Select Schemes", options=all_schemes, default=all_schemes)

filtered_df = df.copy()
filtered_df = filtered_df[(filtered_df["Date"] >= start_date) & (filtered_df["Date"] <= end_date)].copy()
filtered_df = filtered_df[filtered_df["Scheme"].isin(selected_schemes)].copy()
filtered_df = filtered_df.sort_values(["Passbook number", "Date", "Id"])

if filtered_df.empty:
    st.warning("⚠️ No records found for the selected filters.")
    st.stop()


# ============================================================
# RATE COLUMN NAME
# ============================================================

def get_rate_column_name(scheme):
    if pd.isna(scheme):
        return None
    sl = str(scheme).lower()
    if "e-gold" in sl or "egold" in sl or "e gold" in sl:
        return "Gold Rate"
    elif "e-silver" in sl or "esilver" in sl or "e silver" in sl:
        return "Silver Rate"
    if "christmas" in sl or "diwali" in sl or "pongal" in sl:
        return None
    elif "akshaya tritiya" in sl:
        return "Gold Rate"
    return f"{scheme} Rate"


# ============================================================
# GENERATE REPORT DATA
# ============================================================
# Averages (per-scheme AND Grand Total) are computed against the
# number of distinct calendar days in the selected report range, so
# every scheme is directly comparable.
# ============================================================

@st.cache_data(show_spinner=False)
def generate_report_data(df, schemes, start_date=None, end_date=None):
    if not schemes or df.empty:
        return (pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), pd.DataFrame())

    installment_one = df[(df["Installment number"] == 1) & (df["Passbook number"].notna())].copy()
    installment_one = installment_one.sort_values(["Passbook number", "Date", "Id"])
    first_df = installment_one.drop_duplicates(subset=["Passbook number"], keep="first").copy()
    collection_df = df[df["Installment number"].notna() & (df["Installment number"] != 1)].copy()

    first_df["Customer Key"] = first_df["Customer Phone Number"].fillna("").astype(str).str.strip()
    first_df["Valid Phone"] = (
        first_df["Customer Key"].notna()
        & (first_df["Customer Key"] != "")
        & (first_df["Customer Key"] != "nan")
    )

    # ✅ Period length in days (calendar days, inclusive) — the SINGLE
    # denominator used for every scheme AND the Grand Total.
    if start_date is not None and end_date is not None:
        report_day_count = (pd.Timestamp(end_date) - pd.Timestamp(start_date)).days + 1
    else:
        report_day_count = df["Date"].nunique()
    if report_day_count <= 0:
        report_day_count = 1

    summary_rows = []
    for scheme in schemes:
        s_first = first_df[first_df["Scheme"] == scheme]
        s_coll = collection_df[collection_df["Scheme"] == scheme]
        fc = len(s_first)
        fa = s_first["Saved Amount"].sum()
        cc = len(s_coll)
        ca = s_coll["Saved Amount"].sum()
        total_count = fc + cc
        total_amount = fa + ca
        summary_rows.append({
            "Scheme": scheme,
            "First Enrollment Count": fc,
            "First Enrollment Amount": fa,
            "Avg Count/Day": fc / report_day_count,
            "Avg Amount/Day": fa / report_day_count,
            "Collection Count": cc,
            "Collection Amount": ca,
            "Avg Collection Count/Day": cc / report_day_count,
            "Avg Collection Amount/Day": ca / report_day_count,
            "Total Count": total_count,
            "Total Amount": total_amount,
            "Total Avg Ticket": total_amount / total_count if total_count > 0 else 0,
        })

    if summary_rows:
        gfc = sum(r["First Enrollment Count"] for r in summary_rows)
        gfa = sum(r["First Enrollment Amount"] for r in summary_rows)
        gcc = sum(r["Collection Count"] for r in summary_rows)
        gca = sum(r["Collection Amount"] for r in summary_rows)
        total_count = gfc + gcc
        total_amount = gfa + gca

        summary_rows.append({
            "Scheme": "Grand Total",
            "First Enrollment Count": gfc,
            "First Enrollment Amount": gfa,
            "Avg Count/Day": gfc / report_day_count,
            "Avg Amount/Day": gfa / report_day_count,
            "Collection Count": gcc,
            "Collection Amount": gca,
            "Avg Collection Count/Day": gcc / report_day_count,
            "Avg Collection Amount/Day": gca / report_day_count,
            "Total Count": total_count,
            "Total Amount": total_amount,
            "Total Avg Ticket": total_amount / total_count if total_count > 0 else 0,
        })

    summary_df = pd.DataFrame(summary_rows)
    avg_ticket_df = add_avg_ticket_size_comparison(summary_df)

    dates = sorted(df["Date"].dropna().unique())

    enrollment_rows = []
    for date in dates:
        row = {"Date": pd.Timestamp(date).strftime("%d-%m-%Y")}
        tc, ta = 0, 0
        for scheme in schemes:
            rc = get_rate_column_name(scheme)
            sd = df[(df["Date"] == date) & (df["Scheme"] == scheme)]
            rate = sd["Metal Rate"].iloc[0] if not sd.empty else 0
            if rc is not None:
                row[rc] = rate
            temp = first_df[(first_df["Date"] == date) & (first_df["Scheme"] == scheme)]
            cnt, amt = len(temp), temp["Saved Amount"].sum()
            row[f"{scheme} Count"] = cnt
            row[f"{scheme} Amount"] = amt
            row[f"{scheme} Avg Ticket"] = amt / cnt if cnt > 0 else 0
            tc += cnt
            ta += amt
        row["Total Count"] = tc
        row["Total Amount"] = ta
        row["Total Avg Ticket"] = ta / tc if tc > 0 else 0
        enrollment_rows.append(row)
    enrollment_df = pd.DataFrame(enrollment_rows)

    enrollment_columns = ["Date"]
    for scheme in schemes:
        rc = get_rate_column_name(scheme)
        if rc is not None:
            enrollment_columns.append(rc)
        enrollment_columns.extend([f"{scheme} Count", f"{scheme} Amount", f"{scheme} Avg Ticket"])
    enrollment_columns.extend(["Total Count", "Total Amount", "Total Avg Ticket"])
    enrollment_df = enrollment_df[[c for c in enrollment_columns if c in enrollment_df.columns]]

    if not enrollment_df.empty:
        total_row = {"Date": f"{len(enrollment_df)} Days"}
        for scheme in schemes:
            rc = get_rate_column_name(scheme)
            if rc is not None and rc in enrollment_df.columns:
                total_row[rc] = "-"
            tc_ = enrollment_df[f"{scheme} Count"].sum()
            ta_ = enrollment_df[f"{scheme} Amount"].sum()
            total_row[f"{scheme} Count"] = tc_
            total_row[f"{scheme} Amount"] = ta_
            total_row[f"{scheme} Avg Ticket"] = ta_ / tc_ if tc_ > 0 else 0
        tc_ = enrollment_df["Total Count"].sum()
        ta_ = enrollment_df["Total Amount"].sum()
        total_row["Total Count"] = tc_
        total_row["Total Amount"] = ta_
        total_row["Total Avg Ticket"] = ta_ / tc_ if tc_ > 0 else 0
        enrollment_df = pd.concat([enrollment_df, pd.DataFrame([total_row])], ignore_index=True)

    collection_rows = []
    for date in dates:
        row = {"Date": pd.Timestamp(date).strftime("%d-%m-%Y")}
        tc, ta = 0, 0
        for scheme in schemes:
            temp = collection_df[(collection_df["Date"] == date) & (collection_df["Scheme"] == scheme)]
            cnt, amt = len(temp), temp["Saved Amount"].sum()
            row[f"{scheme} Count"] = cnt
            row[f"{scheme} Amount"] = amt
            row[f"{scheme} Avg Ticket"] = amt / cnt if cnt > 0 else 0
            tc += cnt
            ta += amt
        row["Total Count"] = tc
        row["Total Amount"] = ta
        row["Total Avg Ticket"] = ta / tc if tc > 0 else 0
        collection_rows.append(row)
    collection_df_report = pd.DataFrame(collection_rows)

    collection_columns = ["Date"]
    for scheme in schemes:
        collection_columns.extend([f"{scheme} Count", f"{scheme} Amount", f"{scheme} Avg Ticket"])
    collection_columns.extend(["Total Count", "Total Amount", "Total Avg Ticket"])
    collection_df_report = collection_df_report[[c for c in collection_columns if c in collection_df_report.columns]]

    if not collection_df_report.empty:
        total_row = {"Date": f"{len(collection_df_report)} Days"}
        for scheme in schemes:
            tc_ = collection_df_report[f"{scheme} Count"].sum()
            ta_ = collection_df_report[f"{scheme} Amount"].sum()
            total_row[f"{scheme} Count"] = tc_
            total_row[f"{scheme} Amount"] = ta_
            total_row[f"{scheme} Avg Ticket"] = ta_ / tc_ if tc_ > 0 else 0
        tc_ = collection_df_report["Total Count"].sum()
        ta_ = collection_df_report["Total Amount"].sum()
        total_row["Total Count"] = tc_
        total_row["Total Amount"] = ta_
        total_row["Total Avg Ticket"] = ta_ / tc_ if tc_ > 0 else 0
        collection_df_report = pd.concat([collection_df_report, pd.DataFrame([total_row])], ignore_index=True)

    unique_rows = []
    for date in dates:
        row = {"Date": pd.Timestamp(date).strftime("%d-%m-%Y")}
        gt = 0
        for scheme in schemes:
            temp = first_df[
                (first_df["Date"] == date) & (first_df["Scheme"] == scheme) & (first_df["Valid Phone"])
            ]
            uc = temp["Customer Key"].nunique()
            row[scheme] = uc
            gt += uc
        row["Grand Total"] = gt
        unique_rows.append(row)
    unique_df = pd.DataFrame(unique_rows)

    unique_columns = ["Date"] + list(schemes) + ["Grand Total"]
    if not unique_df.empty:
        unique_df = unique_df[[c for c in unique_columns if c in unique_df.columns]]
        total_row = {"Date": f"{len(unique_df)} Days"}
        for scheme in schemes:
            if scheme in unique_df.columns:
                total_row[scheme] = unique_df[scheme].sum()
        total_row["Grand Total"] = unique_df["Grand Total"].sum()
        unique_df = pd.concat([unique_df, pd.DataFrame([total_row])], ignore_index=True)

    summary_df = round_df(summary_df)
    enrollment_df = round_df(enrollment_df)
    collection_df_report = round_df(collection_df_report)
    unique_df = round_df(unique_df)
    avg_ticket_df = round_df(avg_ticket_df, exclude_columns=["% Change"]) if not avg_ticket_df.empty else avg_ticket_df

    return summary_df, enrollment_df, collection_df_report, unique_df, avg_ticket_df


# ============================================================
# MONTHLY PROJECTION
# ============================================================

@st.cache_data(show_spinner=False)
def generate_monthly_projection(full_df, schemes, reference_date=None):
    empty = pd.DataFrame(columns=["Scheme", "Projected Count", "Projected Amount"])
    if full_df.empty or not schemes:
        return empty, empty, ""
    if reference_date is not None:
        latest_date = pd.Timestamp(reference_date)
    else:
        latest_date = full_df["Date"].max()
    month_start = latest_date.replace(day=1)
    days_in_month = calendar.monthrange(latest_date.year, latest_date.month)[1]
    days_elapsed = latest_date.day
    if days_elapsed <= 0:
        days_elapsed = 1
    month_label = (
        f"{latest_date.strftime('%B %Y')} - projected for all "
        f"{days_in_month} days ({days_elapsed} day(s) so far)"
    )
    month_df = full_df[(full_df["Date"] >= month_start) & (full_df["Date"] <= latest_date)]
    installment_one = full_df[
        (full_df["Installment number"] == 1) & (full_df["Passbook number"].notna())
    ].sort_values(["Passbook number", "Date", "Id"])
    first_df = installment_one.drop_duplicates(subset=["Passbook number"], keep="first")
    first_month_df = first_df[(first_df["Date"] >= month_start) & (first_df["Date"] <= latest_date)]
    collection_month_df = month_df[month_df["Installment number"].notna() & (month_df["Installment number"] != 1)]

    enrollment_rows, collection_rows = [], []
    for scheme in schemes:
        e = first_month_df[first_month_df["Scheme"] == scheme]
        enrollment_rows.append({
            "Scheme": scheme,
            "Projected Count": round_value((len(e) / days_elapsed) * days_in_month),
            "Projected Amount": round_value((e["Saved Amount"].sum() / days_elapsed) * days_in_month),
        })
        c = collection_month_df[collection_month_df["Scheme"] == scheme]
        collection_rows.append({
            "Scheme": scheme,
            "Projected Count": round_value((len(c) / days_elapsed) * days_in_month),
            "Projected Amount": round_value((c["Saved Amount"].sum() / days_elapsed) * days_in_month),
        })
    return pd.DataFrame(enrollment_rows), pd.DataFrame(collection_rows), month_label


# ============================================================
# GENERATE MAIN REPORTS
# ============================================================

daily_schemes = [s for s in selected_schemes if classify_scheme_type(s) == "daily"]
daily_df = filtered_df[filtered_df["Scheme"].isin(daily_schemes)].copy()

sessional_schemes = [s for s in selected_schemes if classify_scheme_type(s) == "sessional"]
sessional_df = filtered_df[filtered_df["Scheme"].isin(sessional_schemes)].copy()

with st.spinner("Crunching daily-scheme numbers..."):
    daily_summary, daily_enrollment, daily_collection, daily_unique, daily_avg_ticket = generate_report_data(
        daily_df, daily_schemes, start_date=start_date, end_date=end_date
    )

with st.spinner("Crunching sessional-scheme numbers..."):
    sessional_summary, sessional_enrollment, sessional_collection, sessional_unique, sessional_avg_ticket = generate_report_data(
        sessional_df, sessional_schemes, start_date=start_date, end_date=end_date
    )

daily_enrollment_projection, daily_collection_projection, daily_projection_label = generate_monthly_projection(
    df, daily_schemes, reference_date=end_date
)
sessional_enrollment_projection, sessional_collection_projection, sessional_projection_label = generate_monthly_projection(
    df, sessional_schemes, reference_date=end_date
)


# ============================================================
# MASTER SCHEME LIST
# ============================================================

all_txn_schemes = get_all_transaction_schemes(df)
st.sidebar.caption(f"📋 Master scheme list: {len(all_txn_schemes)} scheme(s) → "
                   + ", ".join(all_txn_schemes[:8])
                   + (" ..." if len(all_txn_schemes) > 8 else ""))


# ============================================================
# REFERRAL REPORT
# ============================================================

referral_summary = pd.DataFrame()
referral_daily_list = []
referral_not_enrolled_daily_list = []
referral_detail = pd.DataFrame()
referral_diagnostics = {}
referral_duplicates = pd.DataFrame()
referral_debug = {}
employee_df_raw = None
referral_df_raw = None

referral_files_present = uploaded_employee is not None and uploaded_referral is not None

if referral_files_present:
    employee_df_raw = read_uploaded_file(uploaded_employee)
    referral_df_raw = read_uploaded_file(uploaded_referral)

    if employee_df_raw is not None and referral_df_raw is not None:
        with st.spinner("Processing referrals (period-scoped enrollment)..."):
            (referral_summary,
             referral_daily_list,
             referral_detail,
             referral_diagnostics,
             referral_duplicates,
             referral_debug,
             referral_not_enrolled_daily_list) = build_referral_report(
                filtered_df, employee_df_raw, referral_df_raw, None,
                start_date=start_date, end_date=end_date,
                master_scheme_list=all_txn_schemes,
            )

        if referral_diagnostics:
            st.info(
                f"📅 **Report Period:** {referral_diagnostics.get('date_filter', 'none')}  |  "
                f"📊 {referral_diagnostics.get('matched_referrals', 0)} enrolled / "
                f"{referral_diagnostics.get('total_referrals', 0)} total "
                f"(**{referral_diagnostics.get('not_enrolled', 0)}** not enrolled)"
                + (f" — **{referral_diagnostics.get('duplicates', 0)}** duplicate(s) flagged"
                   if referral_diagnostics.get('duplicates') else "")
                + (f" — **{referral_diagnostics.get('collisions_cleared', 0)}** txn collision(s) cleared"
                   if referral_diagnostics.get('collisions_cleared') else "")
            )
            if referral_diagnostics.get("scope_warning"):
                st.warning(f"⚠️ {referral_diagnostics['scope_warning']}")


# ============================================================
# DEBUG PANEL
# ============================================================

if referral_files_present and referral_debug:
    st.header("🧪 Debug — Match Analysis")

    with st.expander("📊 Referral Funnel", expanded=True):
        c1, c2, c3, c4, c5 = st.columns(5)
        with c1:
            st.metric("Raw referral rows", f"{referral_debug.get('referral_raw_count', 0):,}")
        with c2:
            st.metric("Dropped (outside period)", f"{referral_debug.get('dropped_out_of_period', 0):,}")
        with c3:
            st.metric("In-scope for report", f"{referral_debug.get('referral_after_date_filter', 0):,}")
        with c4:
            st.metric("Enrolled", f"{referral_diagnostics.get('matched_referrals', 0):,}")
        with c5:
            st.metric("Not Enrolled (in period)", f"{referral_diagnostics.get('not_enrolled', 0):,}")

    with st.expander("🎯 Strict Match Analysis", expanded=True):
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            st.metric("✅ Strict match (phone+date+amount)",
                      f"{referral_debug.get('strict_match_count', 0):,}")
        with c2:
            st.metric("❌ Weak match rejected (phone+date only)",
                      f"{referral_debug.get('weak_match_rejected', 0):,}")
        with c3:
            st.metric("⚠️ Txn collisions cleared",
                      f"{referral_debug.get('collisions_cleared', 0):,}")
        with c4:
            st.metric("🔀 Multi-scheme same day (first picked)",
                      f"{referral_debug.get('multi_scheme_same_day', 0):,}")

    with st.expander("📋 Not Enrolled Reason Breakdown"):
        if referral_debug.get('not_enrolled_breakdown'):
            st.dataframe(
                pd.DataFrame([
                    {"Reason": k, "Count": v}
                    for k, v in referral_debug['not_enrolled_breakdown'].items()
                ]),
                width="stretch", hide_index=True
            )

    with st.expander("🔍 Branch × Category Pivot (in-scope)"):
        if not referral_detail.empty:
            grp = referral_detail.groupby(['Branch', 'Category']).size().reset_index(name='Count')
            pivot_dbg = grp.pivot(index='Branch', columns='Category', values='Count').fillna(0).astype(int)
            st.dataframe(pivot_dbg, width="stretch")


# ============================================================
# PREVIEW
# ============================================================

st.header("📋 Formatted Report Preview")

if daily_schemes and not daily_summary.empty:
    st.subheader("📅 Daily Schemes (e-Gold & e-Silver)")
    st.dataframe(daily_summary, width="stretch", hide_index=True)
    if not daily_avg_ticket.empty:
        display_avg_ticket_comparison(daily_avg_ticket)
    if not daily_enrollment_projection.empty or not daily_collection_projection.empty:
        st.markdown(f"**📈 Monthly Projection — {daily_projection_label}**")
        c5, c6 = st.columns(2)
        with c5:
            st.caption("Projection — first enrollments")
            st.dataframe(daily_enrollment_projection, width="stretch", hide_index=True)
        with c6:
            st.caption("Projection — overall collection")
            st.dataframe(daily_collection_projection, width="stretch", hide_index=True)
    c1, c2 = st.columns(2)
    with c1:
        st.dataframe(daily_enrollment, width="stretch", hide_index=True)
    with c2:
        st.dataframe(daily_collection, width="stretch", hide_index=True)
    st.dataframe(daily_unique, width="stretch", hide_index=True)

if sessional_schemes and not sessional_summary.empty:
    st.subheader("🎯 Seasonal Schemes")
    st.dataframe(sessional_summary, width="stretch", hide_index=True)
    if not sessional_avg_ticket.empty:
        display_avg_ticket_comparison(sessional_avg_ticket)
    if not sessional_enrollment_projection.empty or not sessional_collection_projection.empty:
        st.markdown(f"**📈 Monthly Projection — {sessional_projection_label}**")
        c5, c6 = st.columns(2)
        with c5:
            st.caption("Projection — first enrollments")
            st.dataframe(sessional_enrollment_projection, width="stretch", hide_index=True)
        with c6:
            st.caption("Projection — overall collection")
            st.dataframe(sessional_collection_projection, width="stretch", hide_index=True)
    c1, c2 = st.columns(2)
    with c1:
        st.dataframe(sessional_enrollment, width="stretch", hide_index=True)
    with c2:
        st.dataframe(sessional_collection, width="stretch", hide_index=True)
    st.dataframe(sessional_unique, width="stretch", hide_index=True)


# ============================================================
# REFERRAL PREVIEW
# ============================================================

if referral_files_present:
    st.header("🎁 Referral Conversion Report")
    if referral_summary.empty:
        st.warning("⚠️ No referral data found. Check the debug panel above.")
    else:
        st.subheader(
            f"🏢 Branch-wise Referral Summary — "
            f"{start_date.strftime('%d-%m-%Y')} to {end_date.strftime('%d-%m-%Y')}"
        )
        display_summary = referral_summary.copy()
        for col in display_summary.columns[1:]:
            display_summary[col] = display_summary[col].apply(
                lambda x: "-" if (isinstance(x, (int, float)) and x == 0) else x
            )
        st.dataframe(display_summary, width="stretch", hide_index=True)

        if referral_daily_list:
            st.subheader("📅 Daily Referral Breakdown")
            for date_label, day_df in referral_daily_list:
                st.markdown(f"**📅 REFERRAL — {date_label}**")
                day_display = day_df.copy()
                for col in day_display.columns[1:]:
                    day_display[col] = day_display[col].apply(
                        lambda x: "-" if (isinstance(x, (int, float)) and x == 0) else x
                    )
                st.dataframe(day_display, width="stretch", hide_index=True)

        st.markdown("---")
        st.subheader("📋 Referral Match Details — All Rows")

        dd = referral_detail.copy()
        for col in ["Registered Date", "Joined Date", "Transaction Date"]:
            if col in dd.columns:
                dd[col] = dd[col].apply(
                    lambda x: x.strftime("%d-%m-%Y") if pd.notna(x) else ""
                )

        fcol1, fcol2, fcol3, fcol4 = st.columns(4)
        with fcol1:
            status_filter = st.selectbox(
                "Match Status",
                ["All", "Enrolled Only", "Not Enrolled Only"],
                key="ref_detail_status_filter"
            )
        with fcol2:
            branches_in_detail = sorted(dd['Branch'].dropna().unique().tolist())
            branch_filter = st.multiselect(
                "Filter by Branch",
                options=branches_in_detail,
                default=[],
                key="ref_detail_branch_filter",
                help="Leave empty for ALL branches"
            )
        with fcol3:
            categories_in_detail = sorted(dd['Category'].dropna().unique().tolist())
            category_filter = st.multiselect(
                "Filter by Category",
                options=categories_in_detail,
                default=[],
                key="ref_detail_category_filter",
                help="Leave empty for ALL categories"
            )
        with fcol4:
            search_text = st.text_input(
                "🔎 Search (name / phone / passbook)",
                key="ref_detail_search"
            )

        dd_filtered = dd.copy()
        if status_filter == "Enrolled Only":
            dd_filtered = dd_filtered[dd_filtered['Category'] != 'Not Enrolled']
        elif status_filter == "Not Enrolled Only":
            dd_filtered = dd_filtered[dd_filtered['Category'] == 'Not Enrolled']
        if branch_filter:
            dd_filtered = dd_filtered[dd_filtered['Branch'].isin(branch_filter)]
        if category_filter:
            dd_filtered = dd_filtered[dd_filtered['Category'].isin(category_filter)]
        if search_text:
            term = search_text.strip().lower()
            mask = (
                dd_filtered['Referee Name'].astype(str).str.lower().str.contains(term, na=False) |
                dd_filtered['Referee Phone'].astype(str).str.lower().str.contains(term, na=False) |
                dd_filtered['Referrer Name'].astype(str).str.lower().str.contains(term, na=False) |
                dd_filtered['Passbook'].astype(str).str.lower().str.contains(term, na=False)
            )
            dd_filtered = dd_filtered[mask]

        total_rows = len(referral_detail)
        enrolled_rows = len(referral_detail[referral_detail['Category'] != 'Not Enrolled'])
        not_enrolled_rows = total_rows - enrolled_rows
        filtered_rows = len(dd_filtered)

        mcol1, mcol2, mcol3, mcol4 = st.columns(4)
        with mcol1:
            st.metric("Total in Report", f"{total_rows:,}")
        with mcol2:
            st.metric("✅ Enrolled", f"{enrolled_rows:,}")
        with mcol3:
            st.metric("❌ Not Enrolled", f"{not_enrolled_rows:,}")
        with mcol4:
            st.metric("Showing (after filters)", f"{filtered_rows:,}")

        dd_display = dd_filtered.drop(
            columns=["Match Date", "Joined Date Str", "Registered Date Str", "Display Day"],
            errors="ignore"
        ).copy()

        def highlight_rows(row):
            if row.get('Category') == 'Not Enrolled':
                return ['background-color: #ffe6e6'] * len(row)
            else:
                return ['background-color: #e6ffe6'] * len(row)

        if len(dd_display) <= 5000:
            styled_dd = dd_display.style.apply(highlight_rows, axis=1)
            st.dataframe(styled_dd, width="stretch", height=600)
        else:
            st.info(f"ℹ️ Showing {len(dd_display):,} rows without highlighting for performance")
            st.dataframe(dd_display, width="stretch", height=600)

        csv_buffer = io.StringIO()
        dd_display.to_csv(csv_buffer, index=False)
        st.download_button(
            label="📥 Download Filtered Match Details (CSV)",
            data=csv_buffer.getvalue(),
            file_name=f"referral_match_details_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.csv",
            mime="text/csv",
            width="stretch"
        )

        if not_enrolled_rows > 0:
            with st.expander(f"❌ View {not_enrolled_rows} Not Enrolled Rows (blank Joined Date, Registered in period)"):
                ne_df = referral_detail[referral_detail['Category'] == 'Not Enrolled'].copy()
                for col in ["Registered Date", "Joined Date", "Transaction Date"]:
                    if col in ne_df.columns:
                        ne_df[col] = ne_df[col].apply(
                            lambda x: x.strftime("%d-%m-%Y") if pd.notna(x) else ""
                        )
                ne_cols = [
                    'Branch', 'Referee Name', 'Referee Phone',
                    'Registered Date', 'Joined Date', 'Enrollment Amount',
                    'Referrer Name', 'Employee Code'
                ]
                ne_cols = [c for c in ne_cols if c in ne_df.columns]
                st.dataframe(ne_df[ne_cols], width="stretch", height=500)

        if not referral_duplicates.empty:
            with st.expander("🔴 View Duplicate Records"):
                st.dataframe(referral_duplicates, width="stretch", hide_index=True)
else:
    st.info("💡 Upload Employee + Referral files to see the Referral Report.")


# ============================================================
# DOWNLOAD EXCEL
# ============================================================

st.header("📥 Download Formatted Report")


def generate_clean_filename(start_date, end_date):
    if start_date and end_date:
        mn = start_date.strftime("%B")
        y = start_date.strftime("%Y")
        sd = str(start_date.day).zfill(2)
        ed = str(end_date.day).zfill(2)
        if start_date.month == end_date.month:
            return f"Daily enrollment ({mn} {sd} - {ed} {y})"
        else:
            return (
                f"Daily enrollment ({start_date.strftime('%B')} {sd} - "
                f"{end_date.strftime('%B')} {ed} {y})"
            )
    else:
        return f"Daily enrollment ({pd.Timestamp.now().strftime('%B %Y')})"


clean_filename = generate_clean_filename(start_date, end_date)
st.info(f"📁 **Report will be saved as:** `{clean_filename}.xlsx`")

report_title = uploaded_file.name.rsplit(".", 1)[0] or "Scheme Enrollment & Collection Report"

daily_data = {
    "summary": daily_summary if daily_summary is not None and not daily_summary.empty else pd.DataFrame(),
    "enrollment": daily_enrollment if daily_enrollment is not None and not daily_enrollment.empty else pd.DataFrame(),
    "collection": daily_collection if daily_collection is not None and not daily_collection.empty else pd.DataFrame(),
    "unique": daily_unique if daily_unique is not None and not daily_unique.empty else pd.DataFrame(),
    "schemes": daily_schemes,
    "date_range": (start_date.strftime("%d-%m-%Y"), end_date.strftime("%d-%m-%Y")),
    "report_title": "eGold & eSilver Enrollment & Collection Report",
    "sheet_name": "eGold & eSilver",
    "enrollment_projection": daily_enrollment_projection,
    "collection_projection": daily_collection_projection,
    "projection_month_label": daily_projection_label,
    "avg_ticket_data": daily_avg_ticket if daily_avg_ticket is not None and not daily_avg_ticket.empty else pd.DataFrame(),
}

sessional_data = {
    "summary": sessional_summary if sessional_summary is not None and not sessional_summary.empty else pd.DataFrame(),
    "enrollment": sessional_enrollment if sessional_enrollment is not None and not sessional_enrollment.empty else pd.DataFrame(),
    "collection": sessional_collection if sessional_collection is not None and not sessional_collection.empty else pd.DataFrame(),
    "unique": sessional_unique if sessional_unique is not None and not sessional_unique.empty else pd.DataFrame(),
    "schemes": sessional_schemes,
    "date_range": (start_date.strftime("%d-%m-%Y"), end_date.strftime("%d-%m-%Y")),
    "report_title": "Seasonal Scheme Enrollment & Collection Report",
    "sheet_name": "Seasonal Scheme",
    "enrollment_projection": sessional_enrollment_projection,
    "collection_projection": sessional_collection_projection,
    "projection_month_label": sessional_projection_label,
    "avg_ticket_data": sessional_avg_ticket if sessional_avg_ticket is not None and not sessional_avg_ticket.empty else pd.DataFrame(),
}

try:
    with st.spinner("Building the formatted Excel workbook..."):
        excel_data = create_formatted_excel(
            daily_data,
            sessional_data,
            (start_date.strftime("%d-%m-%Y"), end_date.strftime("%d-%m-%Y")),
            report_title,
            referral_summary=referral_summary,
            referral_daily=referral_daily_list,
            referral_detail=referral_detail,
            referral_duplicates=referral_duplicates,
            referral_not_enrolled_daily=referral_not_enrolled_daily_list,
        )
    st.download_button(
        label="⬇️ Download Formatted Excel Report",
        data=excel_data,
        file_name=f"{clean_filename}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch"
    )
except Exception as e:
    st.error(f"❌ Could not build the Excel report: {e}")


# ============================================================
# END
# ============================================================