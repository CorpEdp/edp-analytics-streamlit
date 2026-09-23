import streamlit as st
import pandas as pd
import numpy as np
from datetime import datetime
import io
import re
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter
import plotly.express as px
import plotly.graph_objects as go

# Increase pandas Styler max elements limit to handle large dataframes
pd.set_option("styler.render.max_elements", 500000)

REQUIRED_COLUMNS = {
    "Employees Report": ["Referral Code", "Employee Code", "Branch", "Employee Type"],
    "Referrals Report": [
        "Referee Name", "Referee Phone", "Enrollment Amount", "Status",
        "Referrer Name", "Referral Code", "Referrer Phone"
    ],
    "Transactions Report": [
        "Customer Phone Number", "Date", "Saved Amount", "Installment number"
    ],
    "BSS Joining Report (optional)": ["Mobileno", "Date", "Online", "Scheme", "Doc No"],
}


# ---------------------------------------------------------------------------
# Custom branch ordering
# ---------------------------------------------------------------------------
FIXED_BRANCH_GROUPS = [
    ["Madurai"],
    ["Marthandam"],
    ["Salem"],
    ["Tirunelveli"],
    ["Trichy", "Tiruchirappalli", "Tiruchirapalli"],
    ["Rajapalayam", "Rajapalaiyam", "Rajapalayem"],
    ["Dindigul"],
    ["Noida"],
    ["Virudhunagar", "Virudunagar"],
    ["Thanjavur"],
]
TELECALLER_KEYWORDS = ["telecaller", "tele caller", "tellecaller"]


def _is_telecaller_branch(branch):
    branch_l = str(branch).lower()
    return any(k in branch_l for k in TELECALLER_KEYWORDS)


def _fixed_rank(branch):
    branch_l = str(branch).lower()
    for i, aliases in enumerate(FIXED_BRANCH_GROUPS):
        if any(alias.lower() in branch_l for alias in aliases):
            return i
    return None


def build_branch_order(all_branches):
    all_branches = sorted(set(b for b in all_branches if pd.notna(b) and str(b).strip() != ''))

    telecaller_branches = sorted([b for b in all_branches if _is_telecaller_branch(b)])
    remaining = [b for b in all_branches if not _is_telecaller_branch(b)]

    fixed_present = [b for b in remaining if _fixed_rank(b) is not None]
    fixed_present.sort(key=_fixed_rank)
    new_branches = sorted([b for b in remaining if _fixed_rank(b) is None])

    return fixed_present + new_branches + telecaller_branches


def sort_branches_df(df, ordered_branches, branch_col='Branch'):
    order_map = {b: i for i, b in enumerate(ordered_branches)}
    df = df.copy()
    df['_branch_sort_key'] = df[branch_col].map(lambda b: order_map.get(b, len(ordered_branches)))
    df = df.sort_values('_branch_sort_key', kind='stable').drop(columns=['_branch_sort_key'])
    return df


def reindex_all_branches(df, ordered_branches, branch_col='Branch'):
    if branch_col not in df.columns or len(df) == 0:
        if branch_col not in df.columns:
            return df

    existing = set(df[branch_col].tolist())
    missing = [b for b in ordered_branches if b not in existing]

    if missing:
        placeholder_rows = []
        for b in missing:
            row = {}
            for col in df.columns:
                if col == branch_col:
                    row[col] = b
                elif pd.api.types.is_numeric_dtype(df[col]):
                    row[col] = 0
                else:
                    row[col] = ''
            placeholder_rows.append(row)
        df = pd.concat([df, pd.DataFrame(placeholder_rows)], ignore_index=True)

    return sort_branches_df(df, ordered_branches, branch_col)


# ---------------------------------------------------------------------------
# Core report generation logic
# ---------------------------------------------------------------------------
class ReportGenerator:
    def __init__(self, employees_df, referrals_df, transactions_df, bss_df=None):
        self.employees_df = employees_df
        self.referrals_df = referrals_df
        self.transactions_df = transactions_df
        self.bss_df = bss_df
        self.duplicate_indices = set()
        self.duplicate_details = []

    def clean_phone_number(self, phone):
        if pd.isna(phone):
            return ""
        phone_str = str(phone).strip()
        phone_digits = re.sub(r'\D', '', phone_str)
        if phone_digits.startswith('91') and len(phone_digits) == 12:
            phone_digits = phone_digits[2:]
        elif phone_digits.startswith('0') and len(phone_digits) == 11:
            phone_digits = phone_digits[1:]
        return phone_digits

    def clean_currency_amount(self, amount):
        if pd.isna(amount):
            return np.nan
        amount_str = str(amount).strip()
        amount_str = re.sub(r'[^\d\.]', '', amount_str)
        try:
            return float(amount_str)
        except Exception:
            return np.nan

    def parse_date_flexible(self, date_val):
        if pd.isna(date_val):
            return np.nan
        try:
            if isinstance(date_val, (datetime, pd.Timestamp)):
                return date_val.date()
            date_str = str(date_val).strip()
            formats = [
                '%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d',
                '%m-%d-%Y', '%m/%d/%Y', '%d-%m-%y',
                '%d/%m/%y', '%b %d, %Y', '%d %b %Y'
            ]
            for fmt in formats:
                try:
                    return datetime.strptime(date_str, fmt).date()
                except Exception:
                    continue
            return pd.to_datetime(date_str).date()
        except Exception:
            return np.nan

    def get_updated_date(self, row):
        status = str(row['Status']).lower().strip() if pd.notna(row['Status']) else ""
        if 'joined' in status or 'scheme joined' in status:
            return self.parse_date_flexible(row.get('Joined Date', np.nan))
        elif 'registered' in status or 'customer register' in status:
            return self.parse_date_flexible(row.get('Registered Date', np.nan))
        else:
            if pd.notna(row.get('Joined Date', np.nan)):
                return self.parse_date_flexible(row['Joined Date'])
            else:
                return self.parse_date_flexible(row.get('Registered Date', np.nan))

    def transform_branch(self, branch):
        if pd.isna(branch) or branch == "":
            return "Bhima Jewellery - Customer"

        branch_upper = str(branch).upper().strip()

        if branch_upper == "HEAD OFFICE":
            return "Bhima Jewellery - Madurai"
        elif branch_upper == "IN-TRANSIT- LOCATIONS":
            return "Bhima Jewellery - Salem"
        elif branch_upper == "APP SHOWROOM LOCATION":
            return "Bhima Jewellery - Tirunelveli"

        if "BHIMA JEWELLERY -" in branch_upper:
            parts = str(branch).split("Bhima Jewellery -", 1)
            if len(parts) > 1:
                suffix = parts[1].strip().title()
                final_branch = f"Bhima Jewellery - {suffix}"
            else:
                final_branch = str(branch).title()
        elif str(branch).upper().strip().endswith("BRANCH"):
            cleaned = re.sub(r'BRANCH$', '', str(branch), flags=re.IGNORECASE).strip()
            cleaned = cleaned.title()
            final_branch = f"Bhima Jewellery - {cleaned}"
        else:
            branch_proper = str(branch).title()
            final_branch = f"Bhima Jewellery - {branch_proper}"

        final_branch = final_branch.replace("Tiruchirappalli", "Trichy")
        final_branch = final_branch.replace("TIRUCHIRAPPALLI", "Trichy")
        final_branch = final_branch.replace("tiruchirappalli", "Trichy")

        if " - " in final_branch:
            prefix, suffix = final_branch.split(" - ", 1)
            suffix = suffix.title()
            final_branch = f"{prefix} - {suffix}"
        else:
            final_branch = final_branch.title()

        return final_branch

    def get_branch_universe(self, extra_branches=None):
        branches = set()
        if 'Branch' in self.employees_df.columns:
            branches.update(
                self.employees_df['Branch'].apply(self.transform_branch).dropna().unique().tolist()
            )
        if extra_branches is not None:
            branches.update([b for b in extra_branches if pd.notna(b) and str(b).strip() != ''])
        return build_branch_order(branches)

    def detect_and_separate_duplicates(self, final_report):
        if final_report is None or len(final_report) == 0:
            return final_report, pd.DataFrame(), set(), final_report

        report_copy = final_report.copy()
        report_copy['Is Duplicate'] = False
        report_copy['Duplicate Group'] = ''

        has_passbook = report_copy['Scheme Passbook Number'].notna() & (report_copy['Scheme Passbook Number'] != '')

        if not has_passbook.any():
            st.info("ℹ️ No records with Scheme Passbook Number found")
            return report_copy, pd.DataFrame(), set(), report_copy

        report_copy['Duplicate Key'] = (
            report_copy['Scheme Passbook Number'].astype(str) + '_' +
            report_copy['Employee Name'].astype(str) + '_' +
            report_copy['Employee Code'].astype(str) + '_' +
            report_copy['Branch'].astype(str)
        )

        key_counts = report_copy[has_passbook]['Duplicate Key'].value_counts()
        duplicate_keys = key_counts[key_counts > 1].index.tolist()

        if len(duplicate_keys) == 0:
            st.success("✅ No duplicate records found based on Scheme Passbook Number and Employee Reference")
            report_copy = report_copy.drop('Duplicate Key', axis=1)
            return report_copy, pd.DataFrame(), set(), report_copy

        st.warning(f"⚠️ Found {len(duplicate_keys)} duplicate groups")

        group_counter = 1
        duplicate_indices = set()
        duplicate_details_list = []

        for dup_key in duplicate_keys:
            group_records = report_copy[report_copy['Duplicate Key'] == dup_key].copy()
            group_records = group_records.sort_values('Updated Date', ascending=True)
            original_idx = group_records.index[0]
            report_copy.loc[original_idx, 'Duplicate Group'] = f'GROUP_{group_counter}_ORIGINAL'

            for idx in group_records.index[1:]:
                report_copy.loc[idx, 'Is Duplicate'] = True
                report_copy.loc[idx, 'Duplicate Group'] = f'GROUP_{group_counter}_DUPLICATE'
                duplicate_indices.add(idx)

            if len(group_records.index[1:]) > 0:
                original_row = report_copy.loc[original_idx]
                for dup_idx in group_records.index[1:]:
                    duplicate_row = report_copy.loc[dup_idx]
                    duplicate_details_list.append({
                        'Group ID': f'GROUP_{group_counter}',
                        'Scheme Passbook Number': original_row['Scheme Passbook Number'],
                        'Original Employee Name': original_row['Employee Name'],
                        'Original Employee Code': original_row['Employee Code'],
                        'Original Customer Name': original_row['Customer Name'],
                        'Original Customer Phone': original_row['Customer Phone'],
                        'Original Updated Date': original_row['Updated Date'],
                        'Original Enrollment Amount': original_row['Customer Enrollment Amount'],
                        'Original Status': original_row['Status'],
                        'Duplicate Employee Name': duplicate_row['Employee Name'],
                        'Duplicate Employee Code': duplicate_row['Employee Code'],
                        'Duplicate Customer Name': duplicate_row['Customer Name'],
                        'Duplicate Customer Phone': duplicate_row['Customer Phone'],
                        'Duplicate Updated Date': duplicate_row['Updated Date'],
                        'Duplicate Enrollment Amount': duplicate_row['Customer Enrollment Amount'],
                        'Duplicate Status': duplicate_row['Status'],
                        'Branch': original_row['Branch'],
                        'Category': original_row['Category'],
                        'Original Row Index': original_idx,
                        'Duplicate Row Index': dup_idx
                    })

            group_counter += 1

        duplicates_df = pd.DataFrame(duplicate_details_list)
        self.duplicate_indices = duplicate_indices
        self.duplicate_details = duplicate_details_list

        non_duplicate_report = report_copy[report_copy['Is Duplicate'] == False].copy()
        report_copy = report_copy.drop('Duplicate Key', axis=1)
        if 'Duplicate Key' in non_duplicate_report.columns:
            non_duplicate_report = non_duplicate_report.drop('Duplicate Key', axis=1)

        total_duplicate_records = len(duplicate_indices)
        total_groups = len(duplicate_keys)

        st.warning(f"⚠️ **DUPLICATES IDENTIFIED:** {total_duplicate_records} duplicate records in {total_groups} groups")
        st.info("📊 **KEEPING ORIGINALS ONLY** for all reports except Consolidated Report")
        st.info("📋 **DUPLICATES SHEET:** A separate 'Duplicate Records' sheet will show original vs duplicate comparison")

        return report_copy, duplicates_df, duplicate_indices, non_duplicate_report

    def apply_duplicate_highlighting(self, workbook, sheet_row_numbers, sheet_name='Consolidated Report'):
        if not sheet_row_numbers:
            return

        ws = workbook[sheet_name]
        highlight_fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
        max_row = ws.max_row

        for excel_row in sheet_row_numbers:
            if excel_row <= max_row:
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=excel_row, column=col)
                    cell.fill = highlight_fill
                    thin_border = Border(
                        left=Side(style='thin', color='FF0000'),
                        right=Side(style='thin', color='FF0000'),
                        top=Side(style='thin', color='FF0000'),
                        bottom=Side(style='thin', color='FF0000')
                    )
                    cell.border = thin_border

        ws.insert_rows(1)
        note_cell = ws.cell(row=1, column=1)
        note_cell.value = "🔴 HIGHLIGHTED ROWS (PINK) ARE DUPLICATE RECORDS - Duplicate means same Scheme Passbook Number & Employee Reference"
        note_cell.font = Font(bold=True, color="FF0000", size=12)
        note_cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws.max_column)

    def match_with_bss_report(self, final_report):
        if self.bss_df is None or len(self.bss_df) == 0:
            st.warning("⚠️ BSS Joining Report not provided. Some scheme details may remain unmatched.")
            return final_report

        st.info("🔍 Attempting to match remaining records with BSS Joining Report...")

        bss_clean = self.bss_df.copy()

        required_bss_cols = {
            'Mobileno': 'Customer Phone', 'Date': 'Date', 'Online': 'Online Amount',
            'Scheme': 'Scheme Name', 'Doc No': 'Passbook Number'
        }

        missing_cols = [c for c in required_bss_cols if c not in bss_clean.columns]
        if missing_cols:
            st.warning(f"⚠️ BSS Joining Report missing columns: {', '.join(missing_cols)}")

        if 'Mobileno' in bss_clean.columns:
            bss_clean['Clean Phone'] = bss_clean['Mobileno'].apply(self.clean_phone_number)
        else:
            st.error("❌ 'Mobileno' column not found in BSS Joining Report")
            return final_report

        if 'Date' in bss_clean.columns:
            bss_clean['Clean Date'] = bss_clean['Date'].apply(self.parse_date_flexible)
            st.info("📊 BSS Joining Report kept in original order (newest to oldest)")
        else:
            st.warning("⚠️ 'Date' column not found in BSS Joining Report")
            bss_clean['Clean Date'] = None

        if 'Online' in bss_clean.columns:
            bss_clean['Online Amount'] = bss_clean['Online'].apply(self.clean_currency_amount)
        else:
            st.warning("⚠️ 'Online' column not found in BSS Joining Report")
            bss_clean['Online Amount'] = np.nan

        bss_lookup = defaultdict(list)
        for idx, row in bss_clean.iterrows():
            phone = row['Clean Phone']
            date = row['Clean Date']
            amount = row['Online Amount']
            if phone and phone != "":
                key = phone
                if date:
                    key = f"{phone}_{date}"
                bss_lookup[key].append({
                    'amount': amount, 'scheme_name': row.get('Scheme', ''),
                    'doc_no': row.get('Doc No', ''), 'date': date, 'index': idx
                })

        need_bss_matching = final_report[
            (final_report['Not Enrolled'] == True) &
            (final_report['Status'].str.lower().str.contains('joined', na=False))
        ].copy()

        st.info(f"📊 Found {len(need_bss_matching)} records that need BSS Joining Report matching")

        if len(need_bss_matching) == 0:
            st.success("✅ No records need BSS Joining Report matching")
            return final_report

        matched_count = 0
        for idx in need_bss_matching.index:
            row = final_report.loc[idx]
            customer_phone = row['Customer Phone']
            updated_date = row['Updated Date']
            enrollment_amount = row['Customer Enrollment Amount']

            if not customer_phone or customer_phone == "":
                continue

            matched = False

            if pd.notna(updated_date):
                exact_key = f"{customer_phone}_{updated_date}"
                if exact_key in bss_lookup:
                    for match in reversed(bss_lookup[exact_key]):
                        if pd.notna(enrollment_amount) and pd.notna(match['amount']):
                            if abs(match['amount'] - enrollment_amount) <= 0.01:
                                final_report.loc[idx, 'Customer Payment'] = match['amount']
                                final_report.loc[idx, 'Scheme Name'] = match['scheme_name']
                                final_report.loc[idx, 'Scheme Passbook Number'] = match['doc_no']
                                final_report.loc[idx, 'True/False'] = True
                                final_report.loc[idx, 'Not Enrolled'] = False
                                matched_count += 1
                                matched = True
                                break

                    if not matched and len(bss_lookup[exact_key]) > 0:
                        oldest_match = bss_lookup[exact_key][-1]
                        final_report.loc[idx, 'Customer Payment'] = oldest_match['amount'] if pd.notna(oldest_match['amount']) else enrollment_amount
                        final_report.loc[idx, 'Scheme Name'] = oldest_match['scheme_name']
                        final_report.loc[idx, 'Scheme Passbook Number'] = oldest_match['doc_no']
                        final_report.loc[idx, 'True/False'] = True
                        final_report.loc[idx, 'Not Enrolled'] = False
                        matched_count += 1
                        matched = True

            if not matched:
                phone_key = customer_phone
                if phone_key in bss_lookup:
                    for match in reversed(bss_lookup[phone_key]):
                        if pd.notna(enrollment_amount) and pd.notna(match['amount']):
                            if abs(match['amount'] - enrollment_amount) <= 0.01:
                                final_report.loc[idx, 'Customer Payment'] = match['amount']
                                final_report.loc[idx, 'Scheme Name'] = match['scheme_name']
                                final_report.loc[idx, 'Scheme Passbook Number'] = match['doc_no']
                                final_report.loc[idx, 'True/False'] = True
                                final_report.loc[idx, 'Not Enrolled'] = False
                                matched_count += 1
                                matched = True
                                break

                    if not matched and len(bss_lookup[phone_key]) > 0:
                        oldest_match = bss_lookup[phone_key][-1]
                        final_report.loc[idx, 'Customer Payment'] = oldest_match['amount'] if pd.notna(oldest_match['amount']) else enrollment_amount
                        final_report.loc[idx, 'Scheme Name'] = oldest_match['scheme_name']
                        final_report.loc[idx, 'Scheme Passbook Number'] = oldest_match['doc_no']
                        final_report.loc[idx, 'True/False'] = True
                        final_report.loc[idx, 'Not Enrolled'] = False
                        matched_count += 1
                        matched = True

        st.success(f"✅ Successfully matched {matched_count} records using BSS Joining Report")

        still_unmatched = final_report[
            (final_report['Not Enrolled'] == True) &
            (final_report['Status'].str.lower().str.contains('joined', na=False))
        ].shape[0]

        if still_unmatched > 0:
            st.warning(f"⚠️ {still_unmatched} joined scheme records remain unmatched even after BSS Joining Report matching")

        return final_report

    def generate_report(self):
        final_report = pd.DataFrame()

        has_joined_date = 'Joined Date' in self.referrals_df.columns
        has_registered_date = 'Registered Date' in self.referrals_df.columns

        if has_joined_date or has_registered_date:
            final_report['Updated Date'] = self.referrals_df.apply(self.get_updated_date, axis=1)
            st.info("📅 Date Logic: Using 'Joined Date' for 'Scheme Joined' status and 'Registered Date' for 'Customer Register' status")
        else:
            st.error("❌ Either 'Joined Date' or 'Registered Date' column must exist in Referrals Report")
            return None

        if 'Referee Name' in self.referrals_df.columns:
            final_report['Customer Name'] = self.referrals_df['Referee Name']
        else:
            st.error("❌ 'Referee Name' column not found in Referrals Report")
            return None

        if 'Referee Phone' in self.referrals_df.columns:
            final_report['Customer Phone'] = self.referrals_df['Referee Phone'].apply(self.clean_phone_number)
        else:
            st.error("❌ 'Referee Phone' column not found in Referrals Report")
            return None

        if 'Enrollment Amount' in self.referrals_df.columns:
            final_report['Customer Enrollment Amount'] = self.referrals_df['Enrollment Amount'].apply(self.clean_currency_amount)
        else:
            st.error("❌ 'Enrollment Amount' column not found in Referrals Report")
            return None

        if 'Status' in self.referrals_df.columns:
            final_report['Status'] = self.referrals_df['Status']
        else:
            st.error("❌ 'Status' column not found in Referrals Report")
            return None

        if 'Referrer Name' in self.referrals_df.columns:
            final_report['Employee Name'] = self.referrals_df['Referrer Name']
        else:
            st.error("❌ 'Referrer Name' column not found in Referrals Report")
            return None

        if 'Referral Code' in self.referrals_df.columns:
            final_report['Referral Code'] = self.referrals_df['Referral Code'].astype(str)
        else:
            st.error("❌ 'Referral Code' column not found in Referrals Report")
            return None

        if 'Referrer Phone' in self.referrals_df.columns:
            final_report['Employee Phone'] = self.referrals_df['Referrer Phone'].apply(self.clean_phone_number)
        else:
            st.error("❌ 'Referrer Phone' column not found in Referrals Report")
            return None

        if 'Referral Code' in self.employees_df.columns and 'Employee Code' in self.employees_df.columns:
            emp_code_dict = {}
            for idx, row in self.employees_df.iterrows():
                referral_code = str(row['Referral Code']).strip()
                employee_code = str(row['Employee Code']).strip() if pd.notna(row['Employee Code']) else ''
                emp_code_dict[referral_code] = employee_code
            final_report['Employee Code'] = final_report['Referral Code'].astype(str).map(emp_code_dict)
            final_report['Employee Code'] = final_report['Employee Code'].fillna('')
        else:
            st.warning("⚠️ 'Referral Code' or 'Employee Code' missing in Employees Report")
            final_report['Employee Code'] = ''

        if 'Referral Code' in self.employees_df.columns and 'Branch' in self.employees_df.columns:
            branch_dict = dict(zip(self.employees_df['Referral Code'].astype(str), self.employees_df['Branch']))
            final_report['Raw Branch'] = final_report['Referral Code'].astype(str).map(branch_dict)
            final_report['Branch'] = final_report['Raw Branch'].apply(self.transform_branch)
            final_report.drop('Raw Branch', axis=1, inplace=True)
            st.success("✅ Branch transformation logic applied successfully (Proper Case)")
        else:
            st.warning("⚠️ 'Referral Code' or 'Branch' missing in Employees Report")
            final_report['Branch'] = "Bhima Jewellery - Customer"

        if 'Referral Code' in self.employees_df.columns and 'Employee Type' in self.employees_df.columns:
            emp_type_dict = dict(zip(self.employees_df['Referral Code'].astype(str), self.employees_df['Employee Type']))
            final_report['Category'] = final_report['Referral Code'].astype(str).map(emp_type_dict)
            final_report['Category'] = final_report['Category'].fillna('Customer')
        else:
            st.warning("⚠️ 'Referral Code' or 'Employee Type' missing in Employees Report")
            final_report['Category'] = 'Customer'

        if 'Installment number' in self.transactions_df.columns:
            self.transactions_df['Installment number'] = pd.to_numeric(self.transactions_df['Installment number'], errors='coerce')
            trans_filtered = self.transactions_df[self.transactions_df['Installment number'] == 1].copy()
            st.info(f"📊 Transactions with Installment number = 1: {len(trans_filtered)} out of {len(self.transactions_df)} total transactions")
        else:
            st.warning("⚠️ 'Installment number' column not found, using all transactions")
            trans_filtered = self.transactions_df.copy()

        st.info("📊 Transactions kept in original Excel order (newest to oldest)")

        if 'Customer Phone Number' in trans_filtered.columns:
            trans_filtered['Clean Phone'] = trans_filtered['Customer Phone Number'].apply(self.clean_phone_number)
        else:
            st.error("❌ 'Customer Phone Number' column not found in Transactions Report")
            return None

        if 'Date' in trans_filtered.columns:
            trans_filtered['Date Clean'] = trans_filtered['Date'].apply(self.parse_date_flexible)
        else:
            st.error("❌ 'Date' column not found in Transactions Report")
            return None

        if 'Saved Amount' in trans_filtered.columns:
            trans_filtered['Saved Amount Clean'] = trans_filtered['Saved Amount'].apply(self.clean_currency_amount)
        else:
            st.warning("⚠️ 'Saved Amount' column not found in Transactions Report")
            trans_filtered['Saved Amount Clean'] = np.nan

        trans_filtered['Match Key'] = trans_filtered['Clean Phone'] + "_" + trans_filtered['Date Clean'].astype(str)

        transaction_lookup = defaultdict(list)
        for idx, row in trans_filtered.iterrows():
            key = row['Match Key']
            transaction_lookup[key].append({
                'saved_amount': row['Saved Amount Clean'] if 'Saved Amount Clean' in row else np.nan,
                'scheme_name': row.get('Scheme Name', '') if 'Scheme Name' in row else '',
                'passbook_no': row.get('Passbook number', '') if 'Passbook number' in row else '',
                'paid_date': row.get('Date Clean', ''),
                'customer_phone': row['Clean Phone'],
                'index': idx
            })

        def get_transaction_details(row):
            customer_phone = row['Customer Phone']
            updated_date = row['Updated Date']
            enrollment_amount = row['Customer Enrollment Amount']

            if pd.isna(updated_date) or customer_phone == "":
                return np.nan, "", ""

            match_key = customer_phone + "_" + str(updated_date)

            if match_key in transaction_lookup:
                transactions = transaction_lookup[match_key]
                for trans in reversed(transactions):
                    if pd.notna(enrollment_amount) and pd.notna(trans['saved_amount']):
                        if abs(trans['saved_amount'] - enrollment_amount) <= 0.01:
                            return trans['saved_amount'], trans['scheme_name'], trans['passbook_no']
                oldest_trans = transactions[-1]
                return oldest_trans['saved_amount'], oldest_trans['scheme_name'], oldest_trans['passbook_no']
            else:
                return np.nan, "", ""

        transaction_details = final_report.apply(get_transaction_details, axis=1, result_type='expand')
        final_report['Customer Payment'] = transaction_details[0]
        final_report['Scheme Name'] = transaction_details[1]
        final_report['Scheme Passbook Number'] = transaction_details[2]

        final_report['True/False'] = np.where(
            abs(final_report['Customer Enrollment Amount'] - final_report['Customer Payment']) <= 0.01,
            True, False
        )

        final_report.loc[final_report['True/False'] == False, 'Customer Payment'] = np.nan
        final_report.loc[final_report['True/False'] == False, 'Scheme Name'] = ''
        final_report.loc[final_report['True/False'] == False, 'Scheme Passbook Number'] = ''

        final_report['Not Enrolled'] = final_report['Customer Payment'].isna()

        final_report = self.match_with_bss_report(final_report)

        final_report['Month'] = final_report['Updated Date'].apply(
            lambda x: x.strftime('%B').lower() if pd.notna(x) else ''
        )

        final_columns = [
            'Updated Date', 'Customer Name', 'Customer Phone', 'Customer Enrollment Amount',
            'Status', 'Employee Name', 'Referral Code', 'Employee Phone', 'Employee Code',
            'Branch', 'Customer Payment', 'True/False', 'Scheme Name', 'Scheme Passbook Number',
            'Category', 'Month', 'Not Enrolled'
        ]

        for col in final_columns:
            if col not in final_report.columns:
                final_report[col] = np.nan

        matched_count = final_report['Customer Payment'].notna().sum()
        not_enrolled_count = final_report['Not Enrolled'].sum()
        report_count = len(final_report)
        if report_count > 0:
            st.info(f"📊 Final Match Results: {matched_count} out of {report_count} records matched ({matched_count/report_count*100:.2f}%)")
            st.info(f"📊 Not Enrolled Customers: {not_enrolled_count} out of {report_count} ({not_enrolled_count/report_count*100:.2f}%)")

        total_enrollment = final_report['Customer Enrollment Amount'].sum()
        total_payment = final_report['Customer Payment'].sum()
        if pd.notna(total_enrollment) and pd.notna(total_payment):
            st.info(f"💰 Total Enrollment Amount: {format_inr(total_enrollment)}")
            st.info(f"💰 Total Payment Amount: {format_inr(total_payment)}")

        category_counts = final_report['Category'].value_counts()
        st.info(f"📊 Category Distribution: {dict(category_counts)}")

        branch_counts = final_report['Branch'].value_counts().head(10)
        st.info(f"📊 Top 10 Branches: {dict(branch_counts)}")

        return final_report[final_columns]

    def generate_branch_wise_scheme_report(self, final_report):
        if 'Is Duplicate' in final_report.columns:
            report_data = final_report[
                (final_report['Is Duplicate'] == False) &
                (final_report['Scheme Name'].notna() & (final_report['Scheme Name'] != ''))
            ].copy()
        else:
            report_data = final_report[final_report['Scheme Name'].notna() & (final_report['Scheme Name'] != '')].copy()

        if len(report_data) == 0:
            branch_scheme_report = pd.DataFrame(columns=[
                'Branch', 'Scheme Name', 'Number of Customers', 'Total Enrollment Amount',
                'Total Payment Received', 'Number of Matched Payments', 'Unique Employees',
                'Unique Referral Codes', 'Match Rate (%)', 'Pending Amount'
            ])
        else:
            branch_scheme_report = report_data.groupby(['Branch', 'Scheme Name']).agg({
                'Customer Name': 'count', 'Customer Enrollment Amount': 'sum',
                'Customer Payment': 'sum', 'True/False': 'sum',
                'Employee Name': 'nunique', 'Referral Code': 'nunique'
            }).reset_index()

            branch_scheme_report.columns = [
                'Branch', 'Scheme Name', 'Number of Customers', 'Total Enrollment Amount',
                'Total Payment Received', 'Number of Matched Payments', 'Unique Employees', 'Unique Referral Codes'
            ]

            branch_scheme_report['Match Rate (%)'] = np.where(
                branch_scheme_report['Number of Customers'] > 0,
                branch_scheme_report['Number of Matched Payments'] / branch_scheme_report['Number of Customers'] * 100, 0
            )
            branch_scheme_report['Pending Amount'] = (
                branch_scheme_report['Total Enrollment Amount'] - branch_scheme_report['Total Payment Received']
            )

        ordered_branches = self.get_branch_universe(branch_scheme_report['Branch'].unique() if len(branch_scheme_report) else None)
        present_branches = set(branch_scheme_report['Branch'].tolist())
        missing_branches = [b for b in ordered_branches if b not in present_branches]
        if missing_branches:
            placeholder_rows = [{
                'Branch': b, 'Scheme Name': 'No Scheme', 'Number of Customers': 0,
                'Total Enrollment Amount': 0, 'Total Payment Received': 0,
                'Number of Matched Payments': 0, 'Unique Employees': 0,
                'Unique Referral Codes': 0, 'Match Rate (%)': 0, 'Pending Amount': 0
            } for b in missing_branches]
            branch_scheme_report = pd.concat([branch_scheme_report, pd.DataFrame(placeholder_rows)], ignore_index=True)

        order_map = {b: i for i, b in enumerate(ordered_branches)}
        branch_scheme_report['_branch_order'] = branch_scheme_report['Branch'].map(lambda b: order_map.get(b, len(ordered_branches)))
        branch_scheme_report = branch_scheme_report.sort_values(
            ['_branch_order', 'Number of Customers'], ascending=[True, False]
        ).drop(columns=['_branch_order']).reset_index(drop=True)
        return branch_scheme_report

    def generate_branch_employee_wise_scheme_report(self, final_report):
        if 'Is Duplicate' in final_report.columns:
            report_data = final_report[
                (final_report['Is Duplicate'] == False) &
                (final_report['Scheme Name'].notna() & (final_report['Scheme Name'] != ''))
            ].copy()
        else:
            report_data = final_report[final_report['Scheme Name'].notna() & (final_report['Scheme Name'] != '')].copy()

        if len(report_data) == 0:
            return pd.DataFrame()

        report_data['Employee Code'] = report_data['Employee Code'].astype(str).fillna('')

        emp_scheme_report = report_data.groupby(['Branch', 'Employee Name', 'Employee Code', 'Referral Code', 'Scheme Name']).agg({
            'Customer Name': 'count', 'Customer Enrollment Amount': 'sum',
            'Customer Payment': 'sum', 'True/False': 'sum', 'Customer Phone': 'nunique'
        }).reset_index()

        emp_scheme_report.columns = [
            'Branch', 'Employee Name', 'Employee Code', 'Referral Code', 'Scheme Name',
            'Number of Customers', 'Total Enrollment Amount', 'Total Payment Received',
            'Number of Matched Payments', 'Unique Customers'
        ]

        emp_scheme_report['Match Rate (%)'] = np.where(
            emp_scheme_report['Number of Customers'] > 0,
            emp_scheme_report['Number of Matched Payments'] / emp_scheme_report['Number of Customers'] * 100, 0
        )
        emp_scheme_report['Average Enrollment Amount'] = np.where(
            emp_scheme_report['Number of Customers'] > 0,
            emp_scheme_report['Total Enrollment Amount'] / emp_scheme_report['Number of Customers'], 0
        )
        emp_scheme_report['Pending Amount'] = (
            emp_scheme_report['Total Enrollment Amount'] - emp_scheme_report['Total Payment Received']
        )

        ordered_branches = self.get_branch_universe(emp_scheme_report['Branch'].unique())
        order_map = {b: i for i, b in enumerate(ordered_branches)}
        emp_scheme_report['_branch_order'] = emp_scheme_report['Branch'].map(lambda b: order_map.get(b, len(ordered_branches)))
        emp_scheme_report = emp_scheme_report.sort_values(
            ['_branch_order', 'Employee Name', 'Number of Customers'], ascending=[True, True, False]
        ).drop(columns=['_branch_order']).reset_index(drop=True)
        return emp_scheme_report

    def generate_branch_summary_report(self, final_report):
        if 'Is Duplicate' in final_report.columns:
            branch_summary = final_report[final_report['Is Duplicate'] == False].copy()
        else:
            branch_summary = final_report.copy()

        if len(branch_summary) == 0:
            branch_summary = pd.DataFrame(columns=[
                'Branch', 'Total Customers', 'Total Enrollment Amount', 'Total Payment Received',
                'Matched Payments', 'Unique Employees', 'Unique Referral Codes', 'Schemes Sold',
                'Not Enrolled', 'Match Rate (%)', 'Enrollment Rate (%)', 'Pending Amount',
                'Average Enrollment Amount'
            ])
        else:
            branch_summary = branch_summary.groupby('Branch').agg({
                'Customer Name': 'count', 'Customer Enrollment Amount': 'sum',
                'Customer Payment': 'sum', 'True/False': 'sum',
                'Employee Name': 'nunique', 'Referral Code': 'nunique',
                'Scheme Name': lambda x: (x.notna() & (x != '')).sum(), 'Not Enrolled': 'sum'
            }).reset_index()

            branch_summary.columns = [
                'Branch', 'Total Customers', 'Total Enrollment Amount', 'Total Payment Received',
                'Matched Payments', 'Unique Employees', 'Unique Referral Codes', 'Schemes Sold', 'Not Enrolled'
            ]

            branch_summary['Match Rate (%)'] = np.where(
                branch_summary['Total Customers'] > 0,
                branch_summary['Matched Payments'] / branch_summary['Total Customers'] * 100, 0
            )
            branch_summary['Enrollment Rate (%)'] = np.where(
                branch_summary['Total Customers'] > 0,
                (branch_summary['Total Customers'] - branch_summary['Not Enrolled']) / branch_summary['Total Customers'] * 100, 0
            )
            branch_summary['Pending Amount'] = (
                branch_summary['Total Enrollment Amount'] - branch_summary['Total Payment Received']
            )
            branch_summary['Average Enrollment Amount'] = np.where(
                branch_summary['Total Customers'] > 0,
                branch_summary['Total Enrollment Amount'] / branch_summary['Total Customers'], 0
            )

        ordered_branches = self.get_branch_universe(branch_summary['Branch'].unique() if len(branch_summary) else None)
        branch_summary = reindex_all_branches(branch_summary, ordered_branches, 'Branch').reset_index(drop=True)
        return branch_summary

    def generate_employee_performance_report(self, final_report):
        if 'Is Duplicate' in final_report.columns:
            final_report_copy = final_report[final_report['Is Duplicate'] == False].copy()
        else:
            final_report_copy = final_report.copy()

        if len(final_report_copy) == 0:
            return pd.DataFrame()

        final_report_copy['Employee Code'] = final_report_copy['Employee Code'].astype(str).fillna('')

        emp_performance = final_report_copy.groupby(['Employee Name', 'Employee Code', 'Referral Code', 'Branch', 'Category']).agg({
            'Customer Name': 'count', 'Customer Enrollment Amount': 'sum',
            'Customer Payment': 'sum', 'True/False': 'sum',
            'Scheme Name': lambda x: (x.notna() & (x != '')).sum(), 'Not Enrolled': 'sum'
        }).reset_index()

        emp_performance.columns = [
            'Employee Name', 'Employee Code', 'Referral Code', 'Branch', 'Category',
            'Total Customers', 'Total Enrollment Amount', 'Total Payment Received',
            'Matched Payments', 'Schemes Sold', 'Not Enrolled'
        ]

        emp_performance['Match Rate (%)'] = np.where(
            emp_performance['Total Customers'] > 0,
            emp_performance['Matched Payments'] / emp_performance['Total Customers'] * 100, 0
        )
        emp_performance['Enrollment Rate (%)'] = np.where(
            emp_performance['Total Customers'] > 0,
            (emp_performance['Total Customers'] - emp_performance['Not Enrolled']) / emp_performance['Total Customers'] * 100, 0
        )
        emp_performance['Average Enrollment Amount'] = np.where(
            emp_performance['Total Customers'] > 0,
            emp_performance['Total Enrollment Amount'] / emp_performance['Total Customers'], 0
        )
        emp_performance['Pending Amount'] = (
            emp_performance['Total Enrollment Amount'] - emp_performance['Total Payment Received']
        )

        ordered_branches = self.get_branch_universe(emp_performance['Branch'].unique())
        order_map = {b: i for i, b in enumerate(ordered_branches)}
        emp_performance['_branch_order'] = emp_performance['Branch'].map(lambda b: order_map.get(b, len(ordered_branches)))
        emp_performance = emp_performance.sort_values(
            ['_branch_order', 'Total Customers'], ascending=[True, False]
        ).drop(columns=['_branch_order']).reset_index(drop=True)
        return emp_performance

    def generate_registration_analysis_report(self, final_report):
        if 'Is Duplicate' in final_report.columns:
            reg_analysis = final_report[final_report['Is Duplicate'] == False].copy()
        else:
            reg_analysis = final_report.copy()

        def get_registration_status(row):
            if pd.notna(row['Customer Payment']):
                return 'Enrolled (Payment Made)'
            elif 'joined' in str(row['Status']).lower() or 'scheme joined' in str(row['Status']).lower():
                return 'Registered but Not Enrolled'
            elif 'registered' in str(row['Status']).lower() or 'customer register' in str(row['Status']).lower():
                return 'Registered but Not Enrolled'
            else:
                return 'Not Enrolled'

        if len(reg_analysis) == 0:
            return pd.DataFrame(), pd.DataFrame()

        reg_analysis['Registration Status'] = reg_analysis.apply(get_registration_status, axis=1)

        branch_reg_summary = reg_analysis.groupby(['Branch', 'Registration Status']).agg({
            'Customer Name': 'count', 'Customer Enrollment Amount': 'sum',
            'Customer Payment': 'sum', 'Referral Code': 'nunique'
        }).reset_index()

        branch_reg_summary.columns = [
            'Branch', 'Registration Status', 'Customer Count',
            'Total Enrollment Amount', 'Total Payment Received', 'Unique Referral Codes'
        ]

        branch_reg_pivot = branch_reg_summary.pivot_table(
            index='Branch', columns='Registration Status', values='Customer Count', fill_value=0
        ).reset_index()

        for col_name, out_name in [
            ('Enrolled (Payment Made)', 'Total Enrolled'),
            ('Registered but Not Enrolled', 'Total Registered'),
            ('Not Enrolled', 'Total Not Enrolled'),
        ]:
            branch_reg_pivot[out_name] = branch_reg_pivot[col_name] if col_name in branch_reg_pivot.columns else 0

        branch_reg_pivot['Total Customers'] = (
            branch_reg_pivot['Total Enrolled'] + branch_reg_pivot['Total Registered'] + branch_reg_pivot['Total Not Enrolled']
        )
        branch_reg_pivot['Enrollment Rate (%)'] = np.where(
            branch_reg_pivot['Total Customers'] > 0,
            branch_reg_pivot['Total Enrolled'] / branch_reg_pivot['Total Customers'] * 100, 0
        )

        ordered_branches = self.get_branch_universe(branch_reg_pivot['Branch'].unique())
        branch_reg_pivot = reindex_all_branches(branch_reg_pivot, ordered_branches, 'Branch').reset_index(drop=True)

        if 'Is Duplicate' in final_report.columns:
            not_enrolled_customers = final_report[(final_report['Not Enrolled'] == True) & (final_report['Is Duplicate'] == False)].copy()
        else:
            not_enrolled_customers = final_report[final_report['Not Enrolled'] == True].copy()

        not_enrolled_customers = not_enrolled_customers[[
            'Customer Name', 'Customer Phone', 'Employee Name', 'Employee Code',
            'Referral Code', 'Branch', 'Status', 'Customer Enrollment Amount', 'Updated Date'
        ]]

        return branch_reg_pivot, not_enrolled_customers

    def generate_branch_employee_referral_report(self, final_report, start_date=None, end_date=None):
        if 'Is Duplicate' in final_report.columns:
            filtered_report = final_report[final_report['Is Duplicate'] == False].copy()
        else:
            filtered_report = final_report.copy()

        filtered_report['Employee Code'] = filtered_report['Employee Code'].astype(str).fillna('')

        if start_date and end_date:
            filtered_report['Updated Date'] = pd.to_datetime(filtered_report['Updated Date'], errors='coerce')
            mask = (filtered_report['Updated Date'] >= pd.to_datetime(start_date)) & (filtered_report['Updated Date'] <= pd.to_datetime(end_date))
            filtered_report = filtered_report[mask]

        enrolled_data = filtered_report[filtered_report['Not Enrolled'] == False].copy()

        all_schemes = enrolled_data['Scheme Name'].dropna().unique()
        all_schemes = sorted([s for s in all_schemes if s != ''])
        if len(all_schemes) == 0:
            all_schemes = ['No Scheme']

        branch_summary = []
        branches = self.get_branch_universe(filtered_report['Branch'].unique())

        grand_totals = {
            'scheme_counts': {scheme: 0 for scheme in all_schemes},
            'scheme_amounts': {scheme: 0 for scheme in all_schemes},
            'total_enrolled_count': 0, 'total_enrolled_amount': 0, 'total_not_enrolled': 0
        }

        for branch in branches:
            branch_data = filtered_report[filtered_report['Branch'] == branch]
            branch_enrolled = branch_data[branch_data['Not Enrolled'] == False]

            branch_row = {'Branch': branch, 'Not Enrolled Count': len(branch_data[branch_data['Not Enrolled'] == True])}

            branch_total_count = 0
            branch_total_amount = 0

            for scheme in all_schemes:
                scheme_data = branch_enrolled[branch_enrolled['Scheme Name'] == scheme]
                scheme_count = len(scheme_data)
                scheme_amount = scheme_data['Customer Enrollment Amount'].sum() if len(scheme_data) > 0 else 0
                branch_row[f'{scheme} Count'] = scheme_count
                branch_row[f'{scheme} Amount'] = scheme_amount
                branch_total_count += scheme_count
                branch_total_amount += scheme_amount
                grand_totals['scheme_counts'][scheme] += scheme_count
                grand_totals['scheme_amounts'][scheme] += scheme_amount

            branch_row['Total Enrolled Count'] = branch_total_count
            branch_row['Total Enrolled Amount'] = branch_total_amount

            grand_totals['total_enrolled_count'] += branch_total_count
            grand_totals['total_enrolled_amount'] += branch_total_amount
            grand_totals['total_not_enrolled'] += branch_row['Not Enrolled Count']

            branch_summary.append(branch_row)

        branch_df = pd.DataFrame(branch_summary)

        if len(branch_df) == 0:
            return pd.DataFrame(), pd.DataFrame(), all_schemes

        branch_df['Count %'] = (
            (branch_df['Total Enrolled Count'] / grand_totals['total_enrolled_count'] * 100).round(1).astype(str) + '%'
            if grand_totals['total_enrolled_count'] > 0 else '0%'
        )
        branch_df['Amount %'] = (
            (branch_df['Total Enrolled Amount'] / grand_totals['total_enrolled_amount'] * 100).round(1).astype(str) + '%'
            if grand_totals['total_enrolled_amount'] > 0 else '0%'
        )

        grand_total_row = {'Branch': 'Grand Total'}
        for scheme in all_schemes:
            grand_total_row[f'{scheme} Count'] = grand_totals['scheme_counts'][scheme]
            grand_total_row[f'{scheme} Amount'] = grand_totals['scheme_amounts'][scheme]
        grand_total_row['Total Enrolled Count'] = grand_totals['total_enrolled_count']
        grand_total_row['Total Enrolled Amount'] = grand_totals['total_enrolled_amount']
        grand_total_row['Not Enrolled Count'] = grand_totals['total_not_enrolled']
        grand_total_row['Count %'] = '100%'
        grand_total_row['Amount %'] = '100%'

        branch_df = pd.concat([branch_df, pd.DataFrame([grand_total_row])], ignore_index=True)

        amount_columns = [col for col in branch_df.columns if 'Amount' in col and col != 'Amount %']
        for col in amount_columns:
            branch_df[col] = pd.to_numeric(branch_df[col], errors='coerce').fillna(0)

        scheme_cols = []
        for scheme in all_schemes:
            scheme_cols.append(f'{scheme} Count')
            scheme_cols.append(f'{scheme} Amount')
        desired_col_order = (
            ['Branch', 'Count %', 'Amount %'] + scheme_cols +
            ['Total Enrolled Count', 'Total Enrolled Amount', 'Not Enrolled Count']
        )
        desired_col_order = [c for c in desired_col_order if c in branch_df.columns]
        remaining_cols = [c for c in branch_df.columns if c not in desired_col_order]
        branch_df = branch_df[desired_col_order + remaining_cols]

        employee_details = []
        for branch in branches:
            branch_data = filtered_report[filtered_report['Branch'] == branch]
            branch_enrolled = branch_data[branch_data['Not Enrolled'] == False]

            employees_with_enrolled = branch_enrolled['Employee Name'].dropna().unique()
            not_enrolled_employees = branch_data[branch_data['Not Enrolled'] == True]['Employee Name'].dropna().unique()
            all_employees = list(set(list(employees_with_enrolled) + list(not_enrolled_employees)))

            if len(all_employees) == 0 and len(branch_enrolled) == 0:
                emp_row = {
                    'Branch': branch, 'Employee Name': '', 'Employee Code': '', 'Referral Code': '',
                    'Total Enrolled Count': 0, 'Total Enrolled Amount': 0,
                    'Not Enrolled Count': len(branch_data[branch_data['Not Enrolled'] == True])
                }
                for scheme in all_schemes:
                    emp_row[f'{scheme} Count'] = 0
                    emp_row[f'{scheme} Amount'] = 0
                employee_details.append(emp_row)
            else:
                for employee in all_employees:
                    emp_data = branch_data[branch_data['Employee Name'] == employee]
                    emp_enrolled = emp_data[emp_data['Not Enrolled'] == False]

                    emp_code = emp_data['Employee Code'].iloc[0] if len(emp_data) > 0 and pd.notna(emp_data['Employee Code'].iloc[0]) else ''
                    referral_code = emp_data['Referral Code'].iloc[0] if len(emp_data) > 0 and pd.notna(emp_data['Referral Code'].iloc[0]) else ''

                    employee_row = {
                        'Branch': branch,
                        'Employee Name': employee if employee and str(employee) != 'nan' else 'Unassigned',
                        'Employee Code': str(emp_code) if emp_code else '',
                        'Referral Code': str(referral_code) if referral_code else '',
                    }

                    emp_total_count = 0
                    emp_total_amount = 0
                    for scheme in all_schemes:
                        scheme_data = emp_enrolled[emp_enrolled['Scheme Name'] == scheme]
                        scheme_count = len(scheme_data)
                        scheme_amount = scheme_data['Customer Enrollment Amount'].sum() if len(scheme_data) > 0 else 0
                        employee_row[f'{scheme} Count'] = scheme_count
                        employee_row[f'{scheme} Amount'] = scheme_amount
                        emp_total_count += scheme_count
                        emp_total_amount += scheme_amount

                    employee_row['Total Enrolled Count'] = emp_total_count
                    employee_row['Total Enrolled Amount'] = emp_total_amount
                    employee_row['Not Enrolled Count'] = len(emp_data[emp_data['Not Enrolled'] == True])
                    employee_details.append(employee_row)

                unassigned_data = branch_enrolled[branch_enrolled['Employee Name'].isna() | (branch_enrolled['Employee Name'] == '')]
                if len(unassigned_data) > 0:
                    unassigned_row = {'Branch': branch, 'Employee Name': 'Unassigned', 'Employee Code': '', 'Referral Code': ''}
                    unassigned_total_count = 0
                    unassigned_total_amount = 0
                    for scheme in all_schemes:
                        scheme_data = unassigned_data[unassigned_data['Scheme Name'] == scheme]
                        scheme_count = len(scheme_data)
                        scheme_amount = scheme_data['Customer Enrollment Amount'].sum() if len(scheme_data) > 0 else 0
                        unassigned_row[f'{scheme} Count'] = scheme_count
                        unassigned_row[f'{scheme} Amount'] = scheme_amount
                        unassigned_total_count += scheme_count
                        unassigned_total_amount += scheme_amount
                    unassigned_row['Total Enrolled Count'] = unassigned_total_count
                    unassigned_row['Total Enrolled Amount'] = unassigned_total_amount
                    unassigned_row['Not Enrolled Count'] = 0
                    employee_details.append(unassigned_row)

        employee_df = pd.DataFrame(employee_details)
        if len(employee_df) > 0:
            branch_order_map = {b: i for i, b in enumerate(branches)}
            employee_df['_branch_order'] = employee_df['Branch'].map(lambda b: branch_order_map.get(b, len(branches)))
            employee_df = employee_df.sort_values(
                ['_branch_order', 'Total Enrolled Count'], ascending=[True, False]
            ).drop(columns=['_branch_order']).reset_index(drop=True)

        amount_columns_emp = [col for col in employee_df.columns if 'Amount' in col]
        for col in amount_columns_emp:
            if col in employee_df.columns:
                employee_df[col] = pd.to_numeric(employee_df[col], errors='coerce').fillna(0)

        if len(employee_df) > 0:
            emp_desired_order = (
                ['Branch', 'Employee Name', 'Employee Code', 'Referral Code'] + scheme_cols +
                ['Total Enrolled Count', 'Total Enrolled Amount', 'Not Enrolled Count']
            )
            emp_desired_order = [c for c in emp_desired_order if c in employee_df.columns]
            emp_remaining_cols = [c for c in employee_df.columns if c not in emp_desired_order]
            employee_df = employee_df[emp_desired_order + emp_remaining_cols]

        return branch_df, employee_df, all_schemes


# ---------------------------------------------------------------------------
# Cached helpers
# ---------------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def cached_read_file(uploaded_file):
    if uploaded_file.name.lower().endswith('.csv'):
        return pd.read_csv(uploaded_file, keep_default_na=False, na_values=[])
    return pd.read_excel(uploaded_file, keep_default_na=False, na_values=[])


def validate_columns(df, required_cols, label):
    missing = [c for c in required_cols if c not in df.columns]
    return missing


@st.cache_data(show_spinner=False)
def cached_generate_pipeline(employees_file, referrals_file, transactions_file, bss_file):
    employees_df = cached_read_file(employees_file)
    referrals_df = cached_read_file(referrals_file)
    transactions_df = cached_read_file(transactions_file)
    bss_df = cached_read_file(bss_file) if bss_file else None

    generator = ReportGenerator(employees_df, referrals_df, transactions_df, bss_df)
    final_report = generator.generate_report()

    if final_report is None:
        return None

    final_report_with_dup_flag, duplicates_df, duplicate_indices, non_duplicate_report = \
        generator.detect_and_separate_duplicates(final_report)

    return {
        "employees_df": employees_df,
        "referrals_df": referrals_df,
        "transactions_df": transactions_df,
        "bss_df": bss_df,
        "final_report": final_report_with_dup_flag,
        "duplicates_df": duplicates_df,
        "duplicate_indices": duplicate_indices,
        "non_duplicate_report": non_duplicate_report,
    }


# ===========================================================================
# Single-page Excel export (filter-aware)
# ===========================================================================
def format_inr(value, decimals=0):
    """Format number in Indian numbering system: 1,23,45,678"""
    if pd.isna(value):
        return ""
    try:
        value = float(value)
    except (TypeError, ValueError):
        return str(value)

    if value == 0:
        return "-"

    negative = value < 0
    value = abs(value)
    s = f"{value:.{decimals}f}" if decimals else f"{value:.0f}"
    if '.' in s:
        int_part, dec_part = s.split('.')
    else:
        int_part, dec_part = s, ''

    if len(int_part) > 3:
        last3 = int_part[-3:]
        rest = int_part[:-3]
        groups = []
        while len(rest) > 2:
            groups.insert(0, rest[-2:])
            rest = rest[:-2]
        if rest:
            groups.insert(0, rest)
        formatted_int = ','.join(groups + [last3])
    else:
        formatted_int = int_part

    formatted = f"{formatted_int}.{dec_part}" if dec_part else formatted_int
    return f"-{formatted}" if negative else formatted


def create_single_page_referral_excel(
    branch_df,
    employee_df,
    schemes,
    start_date=None,
    end_date=None,
    selected_branches=None,
    output=None,
):
    """
    Single-sheet Excel with:
      Row 1 : Title (dark blue)
      Row 2 : Date range (medium blue)
      Row 3 : Selected branches note (light blue)
      Row 4 : thin spacer
      Row 5 : BRANCH-WISE SUMMARY section band
      Row 6 : Branch header (blue)
      Row 7+: Branch rows + Grand Total (gold)
      -- spacer --
             EMPLOYEE-WISE DETAILS section band
             Employee header
             Employee rows
       Employee Grand Total (gold)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Branch & Employee Referral"

    TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    SUBTITLE_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    BRANCH_NOTE_FONT = Font(name="Calibri", size=10, italic=True, color="1F4E78")
    HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    BODY_FONT = Font(name="Calibri", size=10)
    GRAND_FONT = Font(name="Calibri", size=10, bold=True)
    SEC_FONT = Font(name="Calibri", size=12, bold=True, color="1F4E78")

    TITLE_FILL = PatternFill("solid", start_color="1F4E78")
    SUBTITLE_FILL = PatternFill("solid", start_color="2E75B6")
    BRANCH_NOTE_FILL = PatternFill("solid", start_color="E7F0FB")
    HEADER_FILL = PatternFill("solid", start_color="4472C4")
    GRAND_FILL = PatternFill("solid", start_color="FFD966")
    ALT_FILL = PatternFill("solid", start_color="F2F7FB")
    SECTION_FILL = PatternFill("solid", start_color="D9E1F2")

    thin = Side(style="thin", color="B4C6E7")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT = Alignment(horizontal="left", vertical="center")
    RIGHT = Alignment(horizontal="right", vertical="center")

    scheme_pairs = []
    for s in schemes:
        scheme_pairs.append(f"{s} Count")
        scheme_pairs.append(f"{s} Amount")

    branch_cols = ["Branch", "Count %", "Amount %"] + scheme_pairs + \
                  ["Total Enrolled Count", "Total Enrolled Amount", "Not Enrolled Count"]

    employee_cols = ["Branch", "Employee Code", "Referral Code"] + scheme_pairs + \
                    ["Total Enrolled Count", "Total Enrolled Amount", "Not Enrolled Count"]

    n_cols = max(len(branch_cols), len(employee_cols))

    # Row 1: Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(row=1, column=1, value="Branch & Employee - Wise Referral Report")
    c.font = TITLE_FONT
    c.fill = TITLE_FILL
    c.alignment = CENTER
    ws.row_dimensions[1].height = 30

    # Row 2: Date range
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    if start_date and end_date:
        date_str = f"Date :- {pd.to_datetime(start_date).strftime('%d-%m-%Y')} to {pd.to_datetime(end_date).strftime('%d-%m-%Y')}"
    else:
        date_str = "Date :- All Data"
    c = ws.cell(row=2, column=1, value=date_str)
    c.font = SUBTITLE_FONT
    c.fill = SUBTITLE_FILL
    c.alignment = CENTER
    ws.row_dimensions[2].height = 22

    # Row 3: Branches note
    if selected_branches:
        branches_str = ", ".join(selected_branches)
        branch_note = f"Branches : {branches_str}"
    else:
        branch_note = "Branches : All"

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n_cols)
    c = ws.cell(row=3, column=1, value=branch_note)
    c.font = BRANCH_NOTE_FONT
    c.fill = BRANCH_NOTE_FILL
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 22

    # Row 4: spacer
    ws.row_dimensions[4].height = 6

    # Row 5: BRANCH-WISE SUMMARY
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=n_cols)
    c = ws.cell(row=5, column=1, value="BRANCH-WISE SUMMARY")
    c.font = SEC_FONT
    c.fill = SECTION_FILL
    c.alignment = CENTER
    ws.row_dimensions[5].height = 22

    # Row 6: Branch header
    for idx, col_name in enumerate(branch_cols, start=1):
        cell = ws.cell(row=6, column=idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[6].height = 32

    # Branch data
    current_row = 7
    for i, (_, row) in enumerate(branch_df.iterrows()):
        is_grand = str(row.get("Branch", "")).strip().lower() == "grand total"
        for idx, col_name in enumerate(branch_cols, start=1):
            value = row.get(col_name, "")

            if "Amount" in col_name and col_name != "Amount %":
                try:
                    value = format_inr(float(value)) if pd.notna(value) else "-"
                except (TypeError, ValueError):
                    value = ""

            cell = ws.cell(row=current_row, column=idx, value=value)
            cell.font = GRAND_FONT if is_grand else BODY_FONT
            cell.border = BORDER

            if col_name == "Branch":
                cell.alignment = LEFT
            elif "%" in col_name:
                cell.alignment = CENTER
            elif "Count" in col_name or "Amount" in col_name:
                cell.alignment = RIGHT
            else:
                cell.alignment = CENTER

            if is_grand:
                cell.fill = GRAND_FILL
            elif i % 2 == 1:
                cell.fill = ALT_FILL
        current_row += 1

    # Blank spacer row
    current_row += 1

    # EMPLOYEE-WISE DETAILS
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=n_cols)
    c = ws.cell(row=current_row, column=1, value="EMPLOYEE-WISE DETAILS")
    c.font = SEC_FONT
    c.fill = SECTION_FILL
    c.alignment = CENTER
    ws.row_dimensions[current_row].height = 22
    current_row += 1

    # Employee header
    header_row_emp = current_row
    for idx, col_name in enumerate(employee_cols, start=1):
        cell = ws.cell(row=header_row_emp, column=idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[header_row_emp].height = 32
    current_row += 1

    # Employee data
    for i, (_, row) in enumerate(employee_df.iterrows()):
        for idx, col_name in enumerate(employee_cols, start=1):
            value = row.get(col_name, "")

            if "Amount" in col_name and col_name != "Amount %":
                try:
                    value = format_inr(float(value)) if pd.notna(value) else "-"
                except (TypeError, ValueError):
                    value = ""

            cell = ws.cell(row=current_row, column=idx, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER

            if col_name in ("Branch", "Employee Code", "Referral Code"):
                cell.alignment = LEFT
            elif "Count" in col_name or "Amount" in col_name:
                cell.alignment = RIGHT
            else:
                cell.alignment = CENTER

            if i % 2 == 1:
                cell.fill = ALT_FILL
        current_row += 1

    # Employee Grand Total
    # Always calculate from the employee rows actually supplied to this export,
    # so branch/date/employee filters are reflected in the downloaded report.
    if len(employee_df) > 0:
        employee_total = {
            "Branch": "Grand Total",
            "Employee Code": "",
            "Referral Code": "",
        }
        for col_name in employee_cols:
            if col_name in ("Branch", "Employee Code", "Referral Code"):
                continue
            if "Count" in col_name or ("Amount" in col_name and col_name != "Amount %"):
                employee_total[col_name] = pd.to_numeric(
                    employee_df[col_name], errors="coerce"
                ).fillna(0).sum() if col_name in employee_df.columns else 0

        current_row += 1
        for idx, col_name in enumerate(employee_cols, start=1):
            value = employee_total.get(col_name, "")
            if "Amount" in col_name and col_name != "Amount %":
                value = format_inr(float(value)) if pd.notna(value) else "-"
            cell = ws.cell(row=current_row, column=idx, value=value)
            cell.font = GRAND_FONT
            cell.fill = GRAND_FILL
            cell.border = BORDER
            if col_name == "Branch":
                cell.alignment = LEFT
            elif "Count" in col_name or "Amount" in col_name:
                cell.alignment = RIGHT
            else:
                cell.alignment = CENTER

    # Column widths
    for idx in range(1, n_cols + 1):
        col_letter = get_column_letter(idx)
        if idx == 1:
            ws.column_dimensions[col_letter].width = 32
        else:
            ws.column_dimensions[col_letter].width = 14

    ws.freeze_panes = "A7"

    if output is None:
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
    else:
        wb.save(output)
        return output


def create_excel_report(final_report, generator, duplicates_df=None, duplicate_indices=None, start_date=None, end_date=None, selected_branches=None, selected_employee_names=None):
    """Create the master Excel file (all reports) with duplicate highlighting."""
    excel_buffer = io.BytesIO()

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    def style_header_and_widths(ws, df):
        if ws.max_row == 0 or ws.max_column == 0:
            return
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.freeze_panes = "A2"
        for i, col in enumerate(df.columns, start=1):
            try:
                sample = df[col].astype(str).head(200)
                max_len = max([len(str(col))] + [len(v) for v in sample])
            except Exception:
                max_len = len(str(col))
            ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 10), 45)

    with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
        final_report_for_excel = final_report.reset_index(drop=False)
        original_index_col = final_report_for_excel.columns[0]
        final_report_for_excel = final_report_for_excel.drop(columns=[original_index_col])
        final_report_for_excel.to_excel(writer, sheet_name='Consolidated Report', index=False)
        style_header_and_widths(writer.sheets['Consolidated Report'], final_report_for_excel)

        if duplicate_indices and len(duplicate_indices) > 0:
            positions = final_report.index.get_indexer(list(duplicate_indices))
            valid_positions = [p for p in positions if p != -1]
            if valid_positions:
                excel_rows = [p + 2 for p in valid_positions]
                generator.apply_duplicate_highlighting(writer.book, excel_rows, 'Consolidated Report')

        if duplicates_df is not None and len(duplicates_df) > 0:
            present_indices = set(final_report.index)
            filtered_duplicates = duplicates_df[
                duplicates_df['Duplicate Row Index'].isin(present_indices) |
                duplicates_df['Original Row Index'].isin(present_indices)
            ].copy()

            if len(filtered_duplicates) > 0:
                filtered_duplicates.to_excel(writer, sheet_name='Duplicate Records', index=False)
                style_header_and_widths(writer.sheets['Duplicate Records'], filtered_duplicates)
                st.info(f"📊 Added 'Duplicate Records' sheet with {len(filtered_duplicates)} duplicate entries")
            else:
                empty_dup_df = pd.DataFrame({'Note': ['No duplicate records in the filtered date range']})
                empty_dup_df.to_excel(writer, sheet_name='Duplicate Records', index=False)

        filter_info = {
            'Report Generated On': [datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
            'Date Filter Applied': [f"{start_date} to {end_date}" if start_date and end_date else "No Filter"],
            'Total Records in Report': [len(final_report)],
            'Total Enrollment Amount': [format_inr(final_report['Customer Enrollment Amount'].sum()) if 'Customer Enrollment Amount' in final_report.columns else "-"],
            'Total Payment Received': [format_inr(final_report['Customer Payment'].sum()) if 'Customer Payment' in final_report.columns else "-"],
        }
        filter_df = pd.DataFrame(filter_info)
        filter_df.to_excel(writer, sheet_name='Report Info', index=False)
        style_header_and_widths(writer.sheets['Report Info'], filter_df)

        sheet_specs = [
            ('Branch-wise Scheme', generator.generate_branch_wise_scheme_report(final_report)),
            ('Branch-Employee Scheme', generator.generate_branch_employee_wise_scheme_report(final_report)),
            ('Branch Summary', generator.generate_branch_summary_report(final_report)),
            ('Employee Performance', generator.generate_employee_performance_report(final_report)),
        ]
        for sheet_name, df in sheet_specs:
            if len(df) > 0:
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                style_header_and_widths(writer.sheets[sheet_name], df)

        reg_pivot, not_enrolled = generator.generate_registration_analysis_report(final_report)
        if len(reg_pivot) > 0:
            reg_pivot.to_excel(writer, sheet_name='Registration Analysis', index=False)
            style_header_and_widths(writer.sheets['Registration Analysis'], reg_pivot)
        if len(not_enrolled) > 0:
            not_enrolled.to_excel(writer, sheet_name='Not Enrolled Customers', index=False)
            style_header_and_widths(writer.sheets['Not Enrolled Customers'], not_enrolled)

        branch_ref, employee_ref, referral_schemes = generator.generate_branch_employee_referral_report(final_report, start_date, end_date)

        # Apply the SAME branch/employee filters used on screen to the
        # "Download All Reports" workbook.
        if selected_branches is not None and len(selected_branches) > 0:
            branch_ref = branch_ref[
                branch_ref['Branch'].isin(selected_branches) |
                (branch_ref['Branch'].astype(str).str.strip().str.lower() == 'grand total')
            ].copy()
            employee_ref = employee_ref[employee_ref['Branch'].isin(selected_branches)].copy()

        if selected_employee_names:
            employee_ref = employee_ref[employee_ref['Employee Name'].isin(selected_employee_names)].copy()

        if len(branch_ref) > 0:
            # Recompute branch grand total after branch filtering.
            branch_data = branch_ref[
                branch_ref['Branch'].astype(str).str.strip().str.lower() != 'grand total'
            ].copy()
            if len(branch_data) > 0:
                branch_gt = {'Branch': 'Grand Total', 'Count %': '100%', 'Amount %': '100%'}
                for col in branch_data.columns:
                    if col == 'Branch' or col in ('Count %', 'Amount %'):
                        continue
                    if 'Count' in col or 'Amount' in col:
                        branch_gt[col] = pd.to_numeric(branch_data[col], errors='coerce').fillna(0).sum()
                branch_ref = pd.concat([branch_data, pd.DataFrame([branch_gt])], ignore_index=True)
            branch_ref.to_excel(writer, sheet_name='Branch Referral Summary', index=False)
            style_header_and_widths(writer.sheets['Branch Referral Summary'], branch_ref)

        if len(employee_ref) > 0:
            employee_ref.to_excel(writer, sheet_name='Employee Referral Details', index=False)
            style_header_and_widths(writer.sheets['Employee Referral Details'], employee_ref)

    return excel_buffer


def get_week_number(date):
    if pd.isna(date):
        return None
    try:
        date_obj = pd.to_datetime(date)
        return f"{date_obj.year}-W{date_obj.isocalendar()[1]:02d}"
    except Exception:
        return None


def get_month_name(date):
    if pd.isna(date):
        return None
    try:
        return pd.to_datetime(date).strftime('%B %Y')
    except Exception:
        return None


def format_inr_compact(value):
    if pd.isna(value):
        return "₹0"
    try:
        value = float(value)
    except (TypeError, ValueError):
        return "₹0"

    negative = value < 0
    value = abs(value)
    sign = "-" if negative else ""

    if value >= 1_00_00_000:
        return f"{sign}₹{value/1_00_00_000:.2f} Cr"
    elif value >= 1_00_000:
        return f"{sign}₹{value/1_00_000:.2f} L"
    elif value >= 1_000:
        return f"{sign}₹{value/1_000:.1f} K"
    else:
        return f"{sign}₹{value:.0f}"


def format_currency(value):
    if pd.isna(value) or value == 0:
        return "-"
    return f"₹{format_inr(value)}"


def render_dashboard(display_report):
    if len(display_report) == 0:
        st.info("ℹ️ No data available for the current filter selection.")
        return

    non_dup = display_report[display_report['Is Duplicate'] == False] if 'Is Duplicate' in display_report.columns else display_report

    col1, col2, col3, col4, col5 = st.columns(5)
    total_customers = len(non_dup)
    total_enrollment = non_dup['Customer Enrollment Amount'].sum()
    total_payment = non_dup['Customer Payment'].sum()
    matched = non_dup['True/False'].sum() if 'True/False' in non_dup.columns else 0
    not_enrolled = non_dup['Not Enrolled'].sum() if 'Not Enrolled' in non_dup.columns else 0

    col1.metric("Total Customers", f"{total_customers:,}")
    col2.metric("Total Enrollment", format_inr_compact(total_enrollment))
    col3.metric("Total Payments", format_inr_compact(total_payment))
    col4.metric("Match Rate", f"{(matched/total_customers*100) if total_customers else 0:.1f}%")
    col5.metric("Not Enrolled", f"{int(not_enrolled):,}")

    st.markdown("---")
    chart_col1, chart_col2 = st.columns(2)

    with chart_col1:
        st.markdown("#### 📈 Enrollment Trend Over Time")
        trend_data = non_dup.dropna(subset=['Updated Date']).copy()
        if len(trend_data) > 0:
            trend_data['Updated Date'] = pd.to_datetime(trend_data['Updated Date'])
            daily = trend_data.groupby(trend_data['Updated Date'].dt.date).agg(
                Customers=('Customer Name', 'count'),
                Amount=('Customer Enrollment Amount', 'sum')
            ).reset_index()
            fig = px.line(daily, x='Updated Date', y='Customers', markers=True,
                          title=None, labels={'Updated Date': 'Date', 'Customers': 'Customers'})
            fig.update_layout(margin=dict(l=10, r=10, t=10, b=10), height=320)
            st.plotly_chart(fig, width="stretch")
        else:
            st.info("No dated records to chart.")

    with chart_col2:
        st.markdown("#### 🏷️ Category Distribution")
        if 'Category' in non_dup.columns and len(non_dup) > 0:
            cat_counts = non_dup['Category'].value_counts().reset_index()
            cat_counts.columns = ['Category', 'Count']
            fig = px.pie(cat_counts, names='Category', values='Count', hole=0.45)
            fig.update_layout(margin=dict(l=10, r=10, t=10, b=10), height=320)
            st.plotly_chart(fig, width="stretch")
        else:
            st.info("No category data available.")

    chart_col3, chart_col4 = st.columns(2)

    with chart_col3:
        st.markdown("#### 🏢 Top 10 Branches by Customers")
        if 'Branch' in non_dup.columns and len(non_dup) > 0:
            branch_counts = non_dup['Branch'].value_counts().head(10).reset_index()
            branch_counts.columns = ['Branch', 'Customers']
            fig = px.bar(branch_counts, x='Customers', y='Branch', orientation='h')
            fig.update_layout(margin=dict(l=10, r=10, t=10, b=10), height=360, yaxis={'categoryorder': 'total ascending'})
            st.plotly_chart(fig, width="stretch")
        else:
            st.info("No branch data available.")

    with chart_col4:
        st.markdown("#### ⭐ Top 10 Employees by Enrollment Amount")
        if 'Employee Name' in non_dup.columns and len(non_dup) > 0:
            top_emp = non_dup.groupby('Employee Name')['Customer Enrollment Amount'].sum().sort_values(ascending=False).head(10).reset_index()
            fig = px.bar(top_emp, x='Customer Enrollment Amount', y='Employee Name', orientation='h')
            fig.update_layout(margin=dict(l=10, r=10, t=10, b=10), height=360, yaxis={'categoryorder': 'total ascending'})
            st.plotly_chart(fig, width="stretch")
        else:
            st.info("No employee data available.")


def main():
    st.set_page_config(
        page_title="Report Generator System",
        layout="wide",
        page_icon="📊",
        initial_sidebar_state="expanded"
    )

    st.markdown("""
    <style>
    .main-header {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        padding: 1.5rem; border-radius: 15px; color: white; text-align: center;
        margin-bottom: 2rem; box-shadow: 0 4px 6px rgba(0, 0, 0, 0.1);
    }
    .main-header h1 { margin: 0; font-size: 2rem; font-weight: 600; }
    .main-header p { margin: 0.5rem 0 0 0; opacity: 0.9; }
    .metric-card {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        padding: 1rem; border-radius: 10px; text-align: center; color: white;
        box-shadow: 0 2px 4px rgba(0, 0, 0, 0.1);
    }
    .metric-card h3 { margin: 0; font-size: 0.9rem; opacity: 0.9; }
    .metric-card .value { font-size: 1.8rem; font-weight: bold; margin: 0.5rem 0; }
    .success-box { background: linear-gradient(135deg, #84fab0 0%, #8fd3f4 100%); padding: 1rem; border-radius: 10px; margin: 1rem 0; }
    .warning-box { background: linear-gradient(135deg, #ffe259 0%, #ffa751 100%); padding: 1rem; border-radius: 10px; margin: 1rem 0; }
    .info-box { background: linear-gradient(135deg, #a8edea 0%, #fed6e3 100%); padding: 1rem; border-radius: 10px; margin: 1rem 0; }
    .error-box { background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%); padding: 1rem; border-radius: 10px; margin: 1rem 0; color: white; }
    .stTabs [data-baseweb="tab-list"] { gap: 8px; background-color: #f8f9fa; padding: 0.5rem; border-radius: 10px; }
    .stTabs [data-baseweb="tab"] { border-radius: 8px; padding: 0.5rem 1rem; font-weight: 500; }
    .stTabs [aria-selected="true"] { background: linear-gradient(135deg, #667eea 0%, #764ba2 100%); color: white !important; }
    .stButton > button { background: linear-gradient(135deg, #667eea 0%, #764ba2 100%); color: white; border: none; border-radius: 8px; font-weight: 500; }
    </style>
    """, unsafe_allow_html=True)

    st.markdown("""
    <div class="main-header">
        <h1>📊 Automated Report Generator</h1>
        <p>Generate consolidated reports from Employees, Referrals, Transactions, and BSS Joining data</p>
    </div>
    """, unsafe_allow_html=True)

    with st.sidebar:
        st.markdown("### 📁 Upload Files")
        employees_file = st.file_uploader("Employees Report *", type=['xlsx', 'xls', 'csv'], key="emp_file")
        referrals_file = st.file_uploader("Referrals Report *", type=['xlsx', 'xls', 'csv'], key="ref_file")
        transactions_file = st.file_uploader("Transactions Report *", type=['xlsx', 'xls', 'csv'], key="trans_file")
        bss_file = st.file_uploader("BSS Joining Report (optional)", type=['xlsx', 'xls', 'csv'], key="bss_file")
        st.caption("* Required files")

    if not (employees_file and referrals_file and transactions_file):
        st.markdown("""
        <div class="info-box">
            👈 <strong>Getting Started</strong><br>
            Please upload all three required files (Employees, Referrals, and Transactions reports) in the sidebar to generate the consolidated report.
            The BSS Joining Report is optional but recommended for better matching accuracy.
        </div>
        """, unsafe_allow_html=True)
        return

    try:
        with st.spinner("📂 Validating files..."):
            emp_preview = cached_read_file(employees_file)
            ref_preview = cached_read_file(referrals_file)
            trans_preview = cached_read_file(transactions_file)
            bss_preview = cached_read_file(bss_file) if bss_file else None
    except Exception as e:
        st.markdown(f'<div class="error-box">❌ Could not read one of the uploaded files: {str(e)}</div>', unsafe_allow_html=True)
        st.info("💡 Make sure the files are valid .xlsx, .xls, or .csv files and are not password-protected.")
        return

    has_joined_or_registered = 'Joined Date' in ref_preview.columns or 'Registered Date' in ref_preview.columns
    validation_issues = {}
    for label, df, cols in [
        ("Employees Report", emp_preview, REQUIRED_COLUMNS["Employees Report"]),
        ("Referrals Report", ref_preview, REQUIRED_COLUMNS["Referrals Report"]),
        ("Transactions Report", trans_preview, REQUIRED_COLUMNS["Transactions Report"]),
    ]:
        missing = validate_columns(df, cols, label)
        if not has_joined_or_registered and label == "Referrals Report":
            missing.append("Joined Date OR Registered Date")
        if missing:
            validation_issues[label] = missing
    if bss_preview is not None:
        missing = validate_columns(bss_preview, REQUIRED_COLUMNS["BSS Joining Report (optional)"], "BSS Joining Report")
        if missing:
            validation_issues["BSS Joining Report (optional)"] = missing

    if validation_issues:
        st.markdown('<div class="error-box">❌ Some required columns are missing. Please check your files:</div>', unsafe_allow_html=True)
        for label, missing in validation_issues.items():
            st.warning(f"**{label}** is missing: {', '.join(missing)}")
        st.stop()

    try:
        with st.spinner("🔄 Generating consolidated report..."):
            result = cached_generate_pipeline(employees_file, referrals_file, transactions_file, bss_file)

        if result is None:
            st.markdown('<div class="error-box">❌ Failed to generate report. Please check your data.</div>', unsafe_allow_html=True)
            return

        final_report = result["final_report"]
        duplicates_df = result["duplicates_df"]
        duplicate_indices = result["duplicate_indices"]
        non_duplicate_report = result["non_duplicate_report"]
        generator = ReportGenerator(result["employees_df"], result["referrals_df"], result["transactions_df"], result["bss_df"])

        st.markdown('<div class="success-box">✅ Report generated successfully!</div>', unsafe_allow_html=True)

        if len(duplicates_df) > 0:
            st.markdown(f"""
            <div class="warning-box">
                ⚠️ <strong>Duplicates Found:</strong> {len(duplicates_df)} duplicate records detected in {len(duplicates_df['Group ID'].unique())} groups.
                <br><br>
                <strong>✅ ORIGINAL RECORDS KEPT</strong> - Only original records are used in all reports (except Consolidated Report).
                <br><br>
                <strong>🔴 HIGHLIGHTED IN PINK</strong> - Duplicate rows are highlighted in pink in the Consolidated Report sheet.
                <br><br>
                <strong>📋 DUPLICATE RECORDS SHEET</strong> - A separate sheet shows original vs duplicate comparison.
                <br><br>
                <strong>⚠️ DUPLICATES EXCLUDED</strong> - All other sheets (Branch-wise, Employee Performance, etc.) EXCLUDE duplicate records for accurate calculations.
            </div>
            """, unsafe_allow_html=True)

            with st.expander("📋 Preview Duplicate Records (Original vs Duplicate)"):
                st.dataframe(duplicates_df, width='stretch')

            col1, col2, col3, col4 = st.columns(4)
            with col1:
                st.metric("Total Duplicate Groups", len(duplicates_df['Group ID'].unique()))
            with col2:
                st.metric("Total Duplicate Records", len(duplicates_df))
            with col3:
                st.metric("Total Records (All)", len(final_report))
            with col4:
                non_dup_count = len(final_report[final_report['Is Duplicate'] == False]) if 'Is Duplicate' in final_report.columns else len(final_report)
                st.metric("Non-Duplicate Records", non_dup_count)

            st.markdown("#### 📊 Impact of Duplicates on Totals")
            col1, col2, col3 = st.columns(3)
            with col1:
                all_enrollment = final_report['Customer Enrollment Amount'].sum()
                non_dup_enrollment = non_duplicate_report['Customer Enrollment Amount'].sum() if len(non_duplicate_report) > 0 else 0
                dup_enrollment = all_enrollment - non_dup_enrollment
                st.metric("Enrollment Amount Impact", format_currency(dup_enrollment),
                          delta=f"Duplicates add {dup_enrollment/all_enrollment*100:.1f}%" if all_enrollment else "0%")
            with col2:
                all_payment = final_report['Customer Payment'].sum() if 'Customer Payment' in final_report.columns else 0
                non_dup_payment = non_duplicate_report['Customer Payment'].sum() if len(non_duplicate_report) > 0 and 'Customer Payment' in non_duplicate_report.columns else 0
                dup_payment = (all_payment - non_dup_payment) if pd.notna(all_payment) and pd.notna(non_dup_payment) else 0
                st.metric("Payment Amount Impact", format_currency(dup_payment),
                          delta=f"Duplicates add {dup_payment/all_payment*100:.1f}%" if all_payment else "0%")
            with col3:
                all_count = len(final_report)
                non_dup_count2 = len(non_duplicate_report)
                dup_count = all_count - non_dup_count2
                st.metric("Record Count Impact", f"{dup_count} records",
                          delta=f"{dup_count/all_count*100:.1f}% of total" if all_count else "0%")
        else:
            st.markdown('<div class="success-box">✅ No duplicate records found. All data is clean.</div>', unsafe_allow_html=True)
            non_duplicate_report = final_report.copy()

        st.markdown("---")
        st.markdown("### 📅 Date Range Filter for Reports")
        st.markdown("Apply filters to view data for specific time periods. **All reports will show ONLY filtered data.**")

        col1, col2, col3, col4 = st.columns([2, 2, 1, 1])

        if 'filter_start_date' not in st.session_state:
            st.session_state.filter_start_date = None
        if 'filter_end_date' not in st.session_state:
            st.session_state.filter_end_date = None
        if 'filter_applied' not in st.session_state:
            st.session_state.filter_applied = False

        with col1:
            start_date = st.date_input("Start Date", value=st.session_state.filter_start_date, key="global_start_date")
        with col2:
            end_date = st.date_input("End Date", value=st.session_state.filter_end_date, key="global_end_date")
        with col3:
            if 'Updated Date' in final_report.columns:
                final_report['Week Number'] = final_report['Updated Date'].apply(get_week_number)
                unique_weeks = sorted([w for w in final_report['Week Number'].unique() if w is not None], reverse=True)
                week_options = ['All'] + unique_weeks
                selected_week = st.selectbox("Select Week", week_options, key="week_select")

                if selected_week != 'All' and selected_week in unique_weeks:
                    if st.session_state.get('_last_week_applied') != selected_week:
                        week_data = final_report[final_report['Week Number'] == selected_week]
                        if len(week_data) > 0:
                            min_date = week_data['Updated Date'].min()
                            max_date = week_data['Updated Date'].max()
                            if pd.notna(min_date) and pd.notna(max_date):
                                st.session_state.filter_start_date = min_date.date() if hasattr(min_date, 'date') else min_date
                                st.session_state.filter_end_date = max_date.date() if hasattr(max_date, 'date') else max_date
                                st.session_state._last_week_applied = selected_week
                                st.session_state._last_month_applied = None
                                st.rerun()
        with col4:
            if 'Updated Date' in final_report.columns:
                final_report['Month Name'] = final_report['Updated Date'].apply(get_month_name)
                unique_months = sorted([m for m in final_report['Month Name'].unique() if m is not None], reverse=True)
                month_options = ['All'] + unique_months
                selected_month = st.selectbox("Select Month", month_options, key="month_select")

                if selected_month != 'All' and selected_month in unique_months:
                    if st.session_state.get('_last_month_applied') != selected_month:
                        month_data = final_report[final_report['Month Name'] == selected_month]
                        if len(month_data) > 0:
                            min_date = month_data['Updated Date'].min()
                            max_date = month_data['Updated Date'].max()
                            if pd.notna(min_date) and pd.notna(max_date):
                                st.session_state.filter_start_date = min_date.date() if hasattr(min_date, 'date') else min_date
                                st.session_state.filter_end_date = max_date.date() if hasattr(max_date, 'date') else max_date
                                st.session_state._last_month_applied = selected_month
                                st.session_state._last_week_applied = None
                                st.rerun()

        col5, col6 = st.columns([1, 1])
        with col5:
            if st.button("🔍 Apply Date Filter", width='content') and start_date and end_date:
                st.session_state.filter_start_date = start_date
                st.session_state.filter_end_date = end_date
                st.session_state.filter_applied = True
                st.rerun()
        with col6:
            if st.button("🗑️ Clear Filters", width='content'):
                st.session_state.filter_start_date = None
                st.session_state.filter_end_date = None
                st.session_state.filter_applied = False
                st.session_state._last_week_applied = None
                st.session_state._last_month_applied = None
                st.rerun()

        filter_start_date = st.session_state.filter_start_date
        filter_end_date = st.session_state.filter_end_date

        if filter_start_date and filter_end_date:
            filtered_data = final_report[
                (pd.to_datetime(final_report['Updated Date']) >= pd.to_datetime(filter_start_date)) &
                (pd.to_datetime(final_report['Updated Date']) <= pd.to_datetime(filter_end_date))
            ].copy()
            st.success(f"✅ Filter applied: {filter_start_date} to {filter_end_date} - Showing {len(filtered_data)} records")
        else:
            st.info("ℹ️ No date filter applied - showing all data")
            filtered_data = final_report.copy()

        st.markdown("---")

        tab0, tab1, tab2, tab3, tab4, tab5, tab6, tab7 = st.tabs([
            "🚀 **Dashboard**",
            "📋 **Consolidated Report**",
            "🏢 **Branch-wise Scheme**",
            "👥 **Branch & Employee Scheme**",
            "📊 **Branch Summary**",
            "⭐ **Employee Performance**",
            "📝 **Registration Analysis**",
            "📈 **Branch & Employee Referral**"
        ])

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

        with tab0:
            st.markdown("### 🚀 Overview Dashboard")
            render_dashboard(filtered_data)

        with tab1:
            st.markdown("### 📈 Final Consolidated Report")
            st.markdown("🔴 <span style='color:red;font-weight:bold'>PINK HIGHLIGHTED</span> rows are duplicate records (same Scheme Passbook Number & Employee Reference)", unsafe_allow_html=True)

            search_term = st.text_input("🔎 Search by customer name, phone, or employee name", key="consolidated_search")

            display_report = filtered_data.copy()
            if search_term:
                term = search_term.strip().lower()
                mask = (
                    display_report['Customer Name'].astype(str).str.lower().str.contains(term, na=False) |
                    display_report['Customer Phone'].astype(str).str.lower().str.contains(term, na=False) |
                    display_report['Employee Name'].astype(str).str.lower().str.contains(term, na=False)
                )
                display_report = display_report[mask]
                st.caption(f"Showing {len(display_report):,} matching records")

            total_cells = len(display_report) * len(display_report.columns)
            use_styler = total_cells <= 250000

            def highlight_duplicates(row):
                if row.name in duplicate_indices:
                    return ['background-color: #FFB6C1'] * len(row)
                return [''] * len(row)

            if 'Duplicate Group' in display_report.columns and use_styler:
                styled_df = display_report.style.apply(highlight_duplicates, axis=1)
                st.dataframe(styled_df, width='stretch', height=400)
            else:
                if not use_styler and 'Duplicate Group' in display_report.columns:
                    st.info(f"💡 Displaying {len(display_report):,} records without row highlighting for better performance")
                st.dataframe(display_report, width='stretch', height=400)

            st.markdown("### 📊 Report Statistics")
            col1, col2, col3, col4, col5, col6 = st.columns(6)

            with col1:
                st.markdown('<div class="metric-card"><h3>Total Records</h3><div class="value">{:,}</div></div>'.format(len(display_report)), unsafe_allow_html=True)
            with col2:
                total_amount = display_report['Customer Enrollment Amount'].sum()
                st.markdown('<div class="metric-card"><h3>Total Enrollment Amount</h3><div class="value">{}</div></div>'.format(format_currency(total_amount) if pd.notna(total_amount) else "₹0.00"), unsafe_allow_html=True)
            with col3:
                total_payment = display_report['Customer Payment'].sum()
                st.markdown('<div class="metric-card"><h3>Total Payments</h3><div class="value">{}</div></div>'.format(format_currency(total_payment) if pd.notna(total_payment) else "₹0.00"), unsafe_allow_html=True)
            with col4:
                matched = display_report['True/False'].sum() if 'True/False' in display_report.columns else 0
                st.markdown('<div class="metric-card"><h3>Matched Records</h3><div class="value">{:,}</div></div>'.format(int(matched)), unsafe_allow_html=True)
            with col5:
                match_percent = (matched / len(display_report) * 100) if len(display_report) > 0 else 0
                st.markdown('<div class="metric-card"><h3>Match Rate</h3><div class="value">{:.2f}%</div></div>'.format(match_percent), unsafe_allow_html=True)
            with col6:
                not_enrolled = display_report['Not Enrolled'].sum() if 'Not Enrolled' in display_report.columns else 0
                st.markdown('<div class="metric-card"><h3>Not Enrolled</h3><div class="value">{:,}</div></div>'.format(int(not_enrolled)), unsafe_allow_html=True)

            st.markdown("#### 💾 Download This Report")
            csv_buffer = io.StringIO()
            display_report.to_csv(csv_buffer, index=False)
            st.download_button(
                label="📥 Download Consolidated Report (CSV)",
                data=csv_buffer.getvalue(),
                file_name=f"consolidated_report_{timestamp}.csv",
                mime="text/csv",
                width='content'
            )

        with tab2:
            st.markdown("### 🏢 Branch-wise Scheme Consolidation Report")
            st.info("📊 **Note:** This report EXCLUDES duplicate records for accurate calculations. All branches (including zero-referral branches) are listed in the fixed branch order, with Telecaller last.")
            branch_scheme_report = generator.generate_branch_wise_scheme_report(filtered_data)

            if len(branch_scheme_report) > 0:
                col1, col2 = st.columns(2)
                with col1:
                    branches = ['All'] + list(dict.fromkeys(branch_scheme_report['Branch'].tolist()))
                    selected_branch = st.selectbox("🏢 Filter by Branch", branches, key="branch_scheme_filter")
                with col2:
                    schemes = ['All'] + sorted(branch_scheme_report['Scheme Name'].unique().tolist())
                    selected_scheme = st.selectbox("📋 Filter by Scheme", schemes, key="scheme_filter")

                filtered_report_data = branch_scheme_report.copy()
                if selected_branch != 'All':
                    filtered_report_data = filtered_report_data[filtered_report_data['Branch'] == selected_branch]
                if selected_scheme != 'All':
                    filtered_report_data = filtered_report_data[filtered_report_data['Scheme Name'] == selected_scheme]

                st.dataframe(filtered_report_data, width='stretch')

                st.markdown("#### 💾 Download This Report")
                csv_buffer = io.StringIO()
                filtered_report_data.to_csv(csv_buffer, index=False)
                st.download_button(
                    label="📥 Download Branch-wise Scheme Report (CSV)",
                    data=csv_buffer.getvalue(),
                    file_name=f"branch_wise_scheme_{timestamp}.csv",
                    mime="text/csv",
                    width='content'
                )
            else:
                st.info("ℹ️ No scheme data available")

        with tab3:
            st.markdown("### 👥 Branch & Employee-wise Scheme Report")
            st.info("📊 **Note:** This report EXCLUDES duplicate records for accurate calculations")
            emp_scheme_report = generator.generate_branch_employee_wise_scheme_report(filtered_data)

            if len(emp_scheme_report) > 0:
                col1, col2, col3 = st.columns(3)
                with col1:
                    branches = ['All'] + list(dict.fromkeys(emp_scheme_report['Branch'].tolist()))
                    selected_branch = st.selectbox("🏢 Filter by Branch", branches, key="emp_branch_filter")
                with col2:
                    if selected_branch != 'All':
                        employees = ['All'] + sorted(emp_scheme_report[emp_scheme_report['Branch'] == selected_branch]['Employee Name'].unique().tolist())
                    else:
                        employees = ['All'] + sorted(emp_scheme_report['Employee Name'].unique().tolist())
                    selected_employee = st.selectbox("👤 Filter by Employee", employees, key="emp_filter")
                with col3:
                    schemes = ['All'] + sorted(emp_scheme_report['Scheme Name'].unique().tolist())
                    selected_scheme = st.selectbox("📋 Filter by Scheme", schemes, key="emp_scheme_filter")

                filtered_report_data = emp_scheme_report.copy()
                if selected_branch != 'All':
                    filtered_report_data = filtered_report_data[filtered_report_data['Branch'] == selected_branch]
                if selected_employee != 'All':
                    filtered_report_data = filtered_report_data[filtered_report_data['Employee Name'] == selected_employee]
                if selected_scheme != 'All':
                    filtered_report_data = filtered_report_data[filtered_report_data['Scheme Name'] == selected_scheme]

                display_columns = ['Branch', 'Employee Name', 'Employee Code', 'Referral Code', 'Scheme Name',
                                    'Number of Customers', 'Total Enrollment Amount', 'Match Rate (%)']
                st.dataframe(filtered_report_data[display_columns], width='stretch')

                if selected_employee != 'All':
                    st.markdown("---")
                    st.markdown(f"### 📋 Detailed Referral Records for: {selected_employee}")

                    if 'Is Duplicate' in filtered_data.columns:
                        employee_records = filtered_data[
                            (filtered_data['Employee Name'] == selected_employee) &
                            (filtered_data['Is Duplicate'] == False)
                        ].copy()
                    else:
                        employee_records = filtered_data[filtered_data['Employee Name'] == selected_employee].copy()

                    if selected_branch != 'All':
                        employee_records = employee_records[employee_records['Branch'] == selected_branch]
                    if selected_scheme != 'All':
                        employee_records = employee_records[employee_records['Scheme Name'] == selected_scheme]

                    if len(employee_records) > 0:
                        col_a, col_b, col_c, col_d, col_e = st.columns(5)
                        with col_a:
                            st.metric("Total Referrals", len(employee_records))
                        with col_b:
                            total_amount = employee_records['Customer Enrollment Amount'].sum()
                            st.metric("Total Enrollment Amount", format_currency(total_amount))
                        with col_c:
                            matched = employee_records['True/False'].sum() if 'True/False' in employee_records.columns else 0
                            st.metric("Matched Payments", int(matched))
                        with col_d:
                            not_enrolled = employee_records['Not Enrolled'].sum() if 'Not Enrolled' in employee_records.columns else 0
                            st.metric("Not Enrolled", int(not_enrolled))
                        with col_e:
                            match_rate = (matched / len(employee_records) * 100) if len(employee_records) > 0 else 0
                            st.metric("Match Rate", f"{match_rate:.1f}%")

                        detail_columns = [
                            'Updated Date', 'Customer Name', 'Customer Phone', 'Customer Enrollment Amount',
                            'Status', 'Scheme Name', 'Scheme Passbook Number', 'Customer Payment', 'True/False',
                            'Not Enrolled', 'Branch', 'Category'
                        ]
                        available_detail_cols = [col for col in detail_columns if col in employee_records.columns]
                        st.dataframe(employee_records[available_detail_cols], width='stretch', height=300)

                        st.markdown("#### 💾 Download Employee Detail Report")
                        emp_timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                        emp_clean_name = selected_employee.replace(" ", "_")

                        download_columns = [
                            'Employee Name', 'Employee Code', 'Referral Code', 'Employee Phone',
                            'Updated Date', 'Customer Name', 'Customer Phone', 'Customer Enrollment Amount',
                            'Status', 'Scheme Name', 'Scheme Passbook Number', 'Customer Payment', 'True/False',
                            'Not Enrolled', 'Branch', 'Category'
                        ]
                        available_download_cols = [col for col in download_columns if col in employee_records.columns]
                        download_df = employee_records[available_download_cols].copy()

                        if 'Customer Enrollment Amount' in download_df.columns:
                            download_df['Customer Enrollment Amount'] = download_df['Customer Enrollment Amount'].apply(
                                lambda x: format_currency(x) if pd.notna(x) else "")
                        if 'Customer Payment' in download_df.columns:
                            download_df['Customer Payment'] = download_df['Customer Payment'].apply(
                                lambda x: format_currency(x) if pd.notna(x) else "")

                        csv_emp_buffer = io.StringIO()
                        download_df.to_csv(csv_emp_buffer, index=False)

                        excel_emp_buffer = io.BytesIO()
                        with pd.ExcelWriter(excel_emp_buffer, engine='openpyxl') as writer:
                            download_df.to_excel(writer, sheet_name='Employee Details', index=False)
                            summary_data = {
                                'Employee Name': [selected_employee],
                                'Employee Code': [employee_records['Employee Code'].iloc[0] if len(employee_records) > 0 and pd.notna(employee_records['Employee Code'].iloc[0]) else ''],
                                'Referral Code': [employee_records['Referral Code'].iloc[0] if len(employee_records) > 0 and pd.notna(employee_records['Referral Code'].iloc[0]) else ''],
                                'Employee Phone': [employee_records['Employee Phone'].iloc[0] if len(employee_records) > 0 and pd.notna(employee_records['Employee Phone'].iloc[0]) else ''],
                                'Total Referrals': [len(employee_records)],
                                'Total Enrollment Amount': [format_currency(employee_records['Customer Enrollment Amount'].sum())],
                                'Total Payment Received': [format_currency(employee_records['Customer Payment'].sum()) if 'Customer Payment' in employee_records.columns else '₹0.00'],
                                'Match Rate': [f"{(employee_records['True/False'].sum() / len(employee_records) * 100):.1f}%" if 'True/False' in employee_records.columns and len(employee_records) > 0 else '0%'],
                                'Date Filter Applied': [f"{filter_start_date} to {filter_end_date}" if filter_start_date and filter_end_date else "No Filter"],
                                'Report Generated': [datetime.now().strftime('%Y-%m-%d %H:%M:%S')]
                            }
                            pd.DataFrame([summary_data]).to_excel(writer, sheet_name='Summary', index=False)

                            if 'Scheme Name' in employee_records.columns:
                                scheme_breakdown = employee_records.groupby('Scheme Name').agg({
                                    'Customer Name': 'count',
                                    'Customer Enrollment Amount': 'sum',
                                    'Customer Payment': 'sum' if 'Customer Payment' in employee_records.columns else 'count'
                                }).reset_index()
                                scheme_breakdown.columns = ['Scheme Name', 'Number of Customers', 'Total Amount', 'Total Payment']
                                scheme_breakdown['Total Amount'] = scheme_breakdown['Total Amount'].apply(format_currency)
                                scheme_breakdown['Total Payment'] = scheme_breakdown['Total Payment'].apply(format_currency)
                                scheme_breakdown.to_excel(writer, sheet_name='Scheme Breakdown', index=False)

                        col_download1, col_download2 = st.columns(2)
                        with col_download1:
                            st.download_button(
                                label=f"📥 Download {selected_employee} Report (CSV)",
                                data=csv_emp_buffer.getvalue(),
                                file_name=f"employee_{emp_clean_name}_details_{emp_timestamp}.csv",
                                mime="text/csv",
                                width='stretch'
                            )
                        with col_download2:
                            st.download_button(
                                label=f"📥 Download {selected_employee} Report (Excel)",
                                data=excel_emp_buffer.getvalue(),
                                file_name=f"employee_{emp_clean_name}_details_{emp_timestamp}.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                width='stretch'
                            )

                        st.markdown("#### 📊 Scheme-wise Breakdown")
                        scheme_breakdown = employee_records.groupby('Scheme Name').agg({
                            'Customer Name': 'count',
                            'Customer Enrollment Amount': 'sum',
                            'Customer Payment': 'sum' if 'Customer Payment' in employee_records.columns else 'count'
                        }).reset_index()
                        scheme_breakdown.columns = ['Scheme Name', 'Number of Customers', 'Total Amount', 'Total Payment']
                        st.dataframe(scheme_breakdown, width='stretch')

                        if len(employee_records['Branch'].unique()) > 1:
                            st.markdown("#### 📊 Branch-wise Breakdown")
                            branch_breakdown = employee_records.groupby('Branch').agg({
                                'Customer Name': 'count',
                                'Customer Enrollment Amount': 'sum'
                            }).reset_index()
                            branch_breakdown.columns = ['Branch', 'Number of Customers', 'Total Amount']
                            st.dataframe(branch_breakdown, width='stretch')
                    else:
                        st.info(f"ℹ️ No records found for {selected_employee} with the current filters")

                st.markdown("#### 💾 Download This Report")
                csv_buffer = io.StringIO()
                filtered_report_data[display_columns].to_csv(csv_buffer, index=False)
                st.download_button(
                    label="📥 Download Branch-Employee Scheme Report (CSV)",
                    data=csv_buffer.getvalue(),
                    file_name=f"branch_employee_scheme_{timestamp}.csv",
                    mime="text/csv",
                    width='content'
                )
            else:
                st.info("ℹ️ No employee scheme data available")

        with tab4:
            st.markdown("### 📊 Branch Summary Report")
            st.info("📊 **Note:** This report EXCLUDES duplicate records for accurate calculations. Every branch is listed (including zero-referral branches) in the fixed branch order, with Telecaller last.")
            branch_summary = generator.generate_branch_summary_report(filtered_data)

            if len(branch_summary) > 0:
                st.dataframe(branch_summary, width='stretch')

                col1, col2 = st.columns(2)
                with col1:
                    st.markdown("#### Branches by Customers (fixed branch order)")
                    top_branches = branch_summary.copy()
                    fig = px.bar(top_branches, x='Branch', y='Total Customers')
                    fig.update_xaxes(categoryorder='array', categoryarray=top_branches['Branch'].tolist())
                    fig.update_layout(margin=dict(l=10, r=10, t=10, b=10), height=320)
                    st.plotly_chart(fig, width="stretch")
                with col2:
                    st.markdown("#### Enrollment Rate by Branch")
                    fig = px.bar(top_branches, x='Branch', y='Enrollment Rate (%)')
                    fig.update_xaxes(categoryorder='array', categoryarray=top_branches['Branch'].tolist())
                    fig.update_layout(margin=dict(l=10, r=10, t=10, b=10), height=320)
                    st.plotly_chart(fig, width="stretch")

                st.markdown("#### 💾 Download This Report")
                csv_buffer = io.StringIO()
                branch_summary.to_csv(csv_buffer, index=False)
                st.download_button(
                    label="📥 Download Branch Summary Report (CSV)",
                    data=csv_buffer.getvalue(),
                    file_name=f"branch_summary_{timestamp}.csv",
                    mime="text/csv",
                    width='content'
                )
            else:
                st.info("ℹ️ No branch summary data available")

        with tab5:
            st.markdown("### ⭐ Employee Performance Report")
            st.info("📊 **Note:** This report EXCLUDES duplicate records for accurate calculations")
            emp_performance = generator.generate_employee_performance_report(filtered_data)

            if len(emp_performance) > 0:
                col1, col2 = st.columns(2)
                with col1:
                    categories = ['All'] + sorted(emp_performance['Category'].unique().tolist())
                    selected_category = st.selectbox("🏷️ Filter by Category", categories)
                with col2:
                    branches = ['All'] + list(dict.fromkeys(emp_performance['Branch'].tolist()))
                    selected_branch = st.selectbox("🏢 Filter by Branch", branches, key="perf_branch_filter")

                filtered_performance = emp_performance.copy()
                if selected_category != 'All':
                    filtered_performance = filtered_performance[filtered_performance['Category'] == selected_category]
                if selected_branch != 'All':
                    filtered_performance = filtered_performance[filtered_performance['Branch'] == selected_branch]

                display_columns = ['Employee Name', 'Referral Code', 'Branch', 'Total Customers',
                                    'Total Enrollment Amount', 'Enrollment Rate (%)', 'Match Rate (%)']
                st.dataframe(filtered_performance[display_columns], width='stretch')

                st.markdown("#### 💾 Download This Report")
                csv_buffer = io.StringIO()
                filtered_performance[display_columns].to_csv(csv_buffer, index=False)
                st.download_button(
                    label="📥 Download Employee Performance Report (CSV)",
                    data=csv_buffer.getvalue(),
                    file_name=f"employee_performance_{timestamp}.csv",
                    mime="text/csv",
                    width='content'
                )
            else:
                st.info("ℹ️ No employee performance data available")

        with tab6:
            st.markdown("### 📝 Registration Analysis Report")
            st.info("📊 **Note:** This report EXCLUDES duplicate records for accurate calculations. Every branch is listed (including zero-referral branches) in the fixed branch order, with Telecaller last.")
            reg_pivot, not_enrolled_customers = generator.generate_registration_analysis_report(filtered_data)

            st.markdown("#### Registration Summary")
            col1, col2, col3, col4 = st.columns(4)
            with col1:
                total_customers = len(filtered_data[filtered_data['Is Duplicate'] == False]) if 'Is Duplicate' in filtered_data.columns else len(filtered_data)
                st.metric("Total Customers", f"{total_customers:,}")
            with col2:
                if 'Is Duplicate' in filtered_data.columns:
                    enrolled = len(filtered_data[(filtered_data['Not Enrolled'] == False) & (filtered_data['Is Duplicate'] == False)])
                else:
                    enrolled = len(filtered_data[filtered_data['Not Enrolled'] == False])
                st.metric("Enrolled", f"{enrolled:,}")
            with col3:
                st.metric("Not Enrolled", f"{len(not_enrolled_customers):,}")
            with col4:
                enrollment_rate = (enrolled / total_customers * 100) if total_customers > 0 else 0
                st.metric("Enrollment Rate", f"{enrollment_rate:.2f}%")

            st.markdown("#### Branch-wise Registration Status")
            st.dataframe(reg_pivot, width='stretch')

            st.markdown("#### 💾 Download Registration Analysis")
            csv_buffer = io.StringIO()
            reg_pivot.to_csv(csv_buffer, index=False)
            st.download_button(
                label="📥 Download Registration Analysis Report (CSV)",
                data=csv_buffer.getvalue(),
                file_name=f"registration_analysis_{timestamp}.csv",
                mime="text/csv",
                width='content'
            )

            if len(not_enrolled_customers) > 0:
                st.markdown("#### Not Enrolled Customers")
                st.dataframe(not_enrolled_customers, width='stretch')

                csv_buffer = io.StringIO()
                not_enrolled_customers.to_csv(csv_buffer, index=False)
                st.download_button(
                    label="📥 Download Not Enrolled Customers Report (CSV)",
                    data=csv_buffer.getvalue(),
                    file_name=f"not_enrolled_customers_{timestamp}.csv",
                    mime="text/csv",
                    width='content'
                )

        with tab7:
            st.markdown("### 📈 Branch & Employee Wise Referral Report")
            st.info("📊 **Note:** This report EXCLUDES duplicate records for accurate calculations. Every branch is listed (including zero-referral branches) in the fixed branch order, with Telecaller last.")
            st.markdown("Shows branch and employee-wise scheme distribution for ALL schemes")

            branch_df, employee_df, schemes = generator.generate_branch_employee_referral_report(
                filtered_data, filter_start_date, filter_end_date
            )

            if len(branch_df) > 0:
                st.info(f"📊 **Schemes found:** {', '.join(schemes)}")

                st.markdown("#### 🔍 Filter by Branch (Applies to both sections below)")
                all_branches = branch_df[branch_df['Branch'] != 'Grand Total']['Branch'].tolist()

                # ---- FIX: no 'Select All' magic string; empty = all branches ----
                selected_branches = st.multiselect(
                    "🏢 Select Branches (leave empty to include ALL branches)",
                    options=all_branches,
                    default=[],
                    key="branch_ref_filter_sync"
                )

                if not selected_branches:
                    selected_branches = all_branches
                    st.caption("ℹ️ No branches selected → showing ALL branches. Pick specific branches to filter.")
                else:
                    st.caption(f"✅ Filtering to {len(selected_branches)} branch(es): {', '.join(selected_branches)}")

                # ---- FILTER BRANCH DATA ----
                branch_filter_mask = branch_df['Branch'].isin(selected_branches)
                grand_total_mask = branch_df['Branch'] == 'Grand Total'
                filtered_branch_df = branch_df[branch_filter_mask | grand_total_mask].copy()

                branch_data_only = branch_df[branch_df['Branch'] != 'Grand Total'].copy()
                branch_data_filtered = branch_data_only[branch_data_only['Branch'].isin(selected_branches)].copy()

                # ---- RECOMPUTE GRAND TOTAL FOR FILTERED BRANCHES ----
                if len(branch_data_filtered) > 0:
                    for col in branch_data_filtered.columns:
                        if 'Count' in col and col != 'Count %':
                            branch_data_filtered[col] = pd.to_numeric(branch_data_filtered[col], errors='coerce').fillna(0)
                        if 'Amount' in col and col != 'Amount %':
                            branch_data_filtered[col] = pd.to_numeric(branch_data_filtered[col], errors='coerce').fillna(0)

                    new_grand_total = {'Branch': 'Grand Total'}
                    for scheme in schemes:
                        new_grand_total[f'{scheme} Count'] = 0
                        new_grand_total[f'{scheme} Amount'] = 0

                    new_grand_total['Total Enrolled Count'] = int(branch_data_filtered['Total Enrolled Count'].sum())
                    new_grand_total['Total Enrolled Amount'] = float(branch_data_filtered['Total Enrolled Amount'].sum())
                    new_grand_total['Not Enrolled Count'] = int(branch_data_filtered['Not Enrolled Count'].sum())

                    for scheme in schemes:
                        count_col = f'{scheme} Count'
                        amount_col = f'{scheme} Amount'
                        if count_col in branch_data_filtered.columns:
                            new_grand_total[count_col] = int(branch_data_filtered[count_col].sum())
                        if amount_col in branch_data_filtered.columns:
                            new_grand_total[amount_col] = float(branch_data_filtered[amount_col].sum())

                    total_count = new_grand_total['Total Enrolled Count']
                    total_amount = new_grand_total['Total Enrolled Amount']

                    new_grand_total['Count %'] = '100%' if total_count > 0 else '0%'
                    new_grand_total['Amount %'] = '100%' if total_amount > 0 else '0%'

                    for idx, row in branch_data_filtered.iterrows():
                        branch_name = row['Branch']
                        count_val = row['Total Enrolled Count']
                        amount_val = row['Total Enrolled Amount']

                        if total_count > 0:
                            count_pct = (count_val / total_count) * 100
                            filtered_branch_df.loc[filtered_branch_df['Branch'] == branch_name, 'Count %'] = f"{count_pct:.1f}%"
                        else:
                            filtered_branch_df.loc[filtered_branch_df['Branch'] == branch_name, 'Count %'] = "0%"

                        if total_amount > 0:
                            amount_pct = (amount_val / total_amount) * 100
                            filtered_branch_df.loc[filtered_branch_df['Branch'] == branch_name, 'Amount %'] = f"{amount_pct:.1f}%"
                        else:
                            filtered_branch_df.loc[filtered_branch_df['Branch'] == branch_name, 'Amount %'] = "0%"

                    filtered_branch_df = filtered_branch_df[filtered_branch_df['Branch'] != 'Grand Total']
                    filtered_branch_df = sort_branches_df(filtered_branch_df, selected_branches, 'Branch')
                    filtered_branch_df = pd.concat([filtered_branch_df, pd.DataFrame([new_grand_total])], ignore_index=True)
                    _col_order = [c for c in branch_df.columns if c in filtered_branch_df.columns]
                    _remaining = [c for c in filtered_branch_df.columns if c not in _col_order]
                    filtered_branch_df = filtered_branch_df[_col_order + _remaining]

                filtered_employee_df = employee_df[employee_df['Branch'].isin(selected_branches)].copy()
                filtered_employee_df = sort_branches_df(filtered_employee_df, selected_branches, 'Branch') if len(filtered_employee_df) > 0 else filtered_employee_df

                st.markdown("#### Branch-wise Summary")
                st.dataframe(filtered_branch_df, width='stretch')

                st.markdown("#### Employee-wise Details")

                if len(filtered_employee_df) > 0:
                    available_employees = sorted(filtered_employee_df['Employee Name'].unique().tolist())
                    selected_employee_names = st.multiselect(
                        "👤 Filter by Employees (Optional - leave empty for all)",
                        options=available_employees,
                        default=[],
                        key="emp_ref_emp_sync"
                    )

                    if selected_employee_names:
                        filtered_employee_df = filtered_employee_df[
                            filtered_employee_df['Employee Name'].isin(selected_employee_names)
                        ]

                    display_emp_df = filtered_employee_df.copy()
                    amount_cols = [col for col in display_emp_df.columns if 'Amount' in col]
                    for col in amount_cols:
                        if col in display_emp_df.columns:
                            display_emp_df[col] = display_emp_df[col].apply(format_currency)

                    st.dataframe(display_emp_df, width='stretch')

                    # Employee-wise Grand Total - respects BOTH branch and employee filters.
                    employee_grand_total = {
                        'Branch': 'Grand Total',
                        'Employee Code': '',
                        'Referral Code': '',
                    }
                    for col in filtered_employee_df.columns:
                        if col in ('Branch', 'Employee Code', 'Referral Code'):
                            continue
                        if 'Count' in col or ('Amount' in col and col != 'Amount %'):
                            employee_grand_total[col] = pd.to_numeric(
                                filtered_employee_df[col], errors='coerce'
                            ).fillna(0).sum()

                    employee_gt_df = pd.DataFrame([employee_grand_total])
                    for col in employee_gt_df.columns:
                        if 'Amount' in col and col != 'Amount %':
                            employee_gt_df[col] = employee_gt_df[col].apply(format_currency)
                    st.markdown("##### 🏆 Employee-wise Grand Total")
                    st.dataframe(employee_gt_df, width='stretch', hide_index=True)

                    st.markdown("#### 📊 Summary for Selected Branches")
                    col_a, col_b, col_c, col_d = st.columns(4)

                    branch_data = filtered_branch_df[filtered_branch_df['Branch'] != 'Grand Total'].copy()
                    for col in branch_data.columns:
                        if 'Count' in col and col != 'Count %':
                            branch_data[col] = pd.to_numeric(branch_data[col], errors='coerce').fillna(0)
                        if 'Amount' in col and col != 'Amount %':
                            branch_data[col] = pd.to_numeric(branch_data[col], errors='coerce').fillna(0)

                    with col_a:
                        st.metric("Selected Branches", len(branch_data))
                    with col_b:
                        total_enrolled = branch_data['Total Enrolled Count'].sum() if 'Total Enrolled Count' in branch_data.columns else 0
                        st.metric("Total Enrolled", f"{int(total_enrolled):,}")
                    with col_c:
                        total_not_enrolled = branch_data['Not Enrolled Count'].sum() if 'Not Enrolled Count' in branch_data.columns else 0
                        st.metric("Not Enrolled", f"{int(total_not_enrolled):,}")
                    with col_d:
                        total_amt = branch_data['Total Enrolled Amount'].sum() if 'Total Enrolled Amount' in branch_data.columns else 0
                        st.metric("Total Amount", format_currency(total_amt))

                    st.markdown("---")
                    st.markdown("#### 💾 Download Both Reports Together (Single Sheet)")

                    # ==================================================
                    # REBUILD EXPORT FRAMES FROM SCRATCH (guaranteed filter)
                    # ==================================================
                    branch_export_df = branch_df[
                        branch_df['Branch'].isin(selected_branches)
                    ].copy()
                    branch_export_df = branch_export_df[
                        branch_export_df['Branch'] != 'Grand Total'
                    ].copy()

                    # Recompute Grand Total from filtered branches
                    gt_row = {'Branch': 'Grand Total', 'Count %': '100%', 'Amount %': '100%'}
                    for scheme in schemes:
                        cnt_col = f'{scheme} Count'
                        amt_col = f'{scheme} Amount'
                        if cnt_col in branch_export_df.columns:
                            gt_row[cnt_col] = int(pd.to_numeric(branch_export_df[cnt_col], errors='coerce').fillna(0).sum())
                        if amt_col in branch_export_df.columns:
                            gt_row[amt_col] = float(pd.to_numeric(branch_export_df[amt_col], errors='coerce').fillna(0).sum())

                    gt_row['Total Enrolled Count'] = int(pd.to_numeric(
                        branch_export_df['Total Enrolled Count'] if 'Total Enrolled Count' in branch_export_df.columns else pd.Series([0]),
                        errors='coerce'
                    ).fillna(0).sum())
                    gt_row['Total Enrolled Amount'] = float(pd.to_numeric(
                        branch_export_df['Total Enrolled Amount'] if 'Total Enrolled Amount' in branch_export_df.columns else pd.Series([0]),
                        errors='coerce'
                    ).fillna(0).sum())
                    gt_row['Not Enrolled Count'] = int(pd.to_numeric(
                        branch_export_df['Not Enrolled Count'] if 'Not Enrolled Count' in branch_export_df.columns else pd.Series([0]),
                        errors='coerce'
                    ).fillna(0).sum())

                    branch_export_df = sort_branches_df(branch_export_df, selected_branches, 'Branch')
                    branch_export_df = pd.concat([branch_export_df, pd.DataFrame([gt_row])], ignore_index=True)

                    _col_order_export = [c for c in branch_df.columns if c in branch_export_df.columns]
                    _remaining_export = [c for c in branch_export_df.columns if c not in _col_order_export]
                    branch_export_df = branch_export_df[_col_order_export + _remaining_export]

                    # Employee export
                    employee_export_df = employee_df[employee_df['Branch'].isin(selected_branches)].copy()
                    if selected_employee_names:
                        employee_export_df = employee_export_df[
                            employee_export_df['Employee Name'].isin(selected_employee_names)
                        ].copy()

                    # Debug expander
                    with st.expander("🔎 Confirm what will be exported (click to expand)", expanded=False):
                        st.write(f"**Selected branches ({len(selected_branches)}):** {', '.join(selected_branches)}")
                        st.write(f"**Branches in Branch table:** {sorted(branch_export_df[branch_export_df['Branch'] != 'Grand Total']['Branch'].unique().tolist())}")
                        st.write(f"**Branches in Employee table:** {sorted(employee_export_df['Branch'].unique().tolist())}")
                        if len(employee_export_df) > 0:
                            st.write(f"**Unique employees:** {employee_export_df['Employee Name'].nunique()}")

                    # Build Excel
                    combined_excel = create_single_page_referral_excel(
                        branch_df=branch_export_df,
                        employee_df=employee_export_df,
                        schemes=schemes,
                        start_date=filter_start_date,
                        end_date=filter_end_date,
                        selected_branches=selected_branches,
                    )

                    excel_bytes = combined_excel.getvalue()

                    timestamp_local = datetime.now().strftime('%Y%m%d_%H%M%S')
                    if filter_start_date and filter_end_date:
                        date_part = f"{filter_start_date}_to_{filter_end_date}"
                    else:
                        date_part = "all_data"

                    all_branches_list = branch_df[branch_df['Branch'] != 'Grand Total']['Branch'].tolist()
                    is_all_selected = set(selected_branches) == set(all_branches_list)

                    if is_all_selected or len(selected_branches) > 3:
                        branch_part = "all_branches"
                    else:
                        branch_part = "_".join(
                            b.replace("Bhima Jewellery - ", "").replace(" ", "")
                            for b in selected_branches[:3]
                        )

                    filename = f"branch_employee_referral_{branch_part}_{date_part}_{timestamp_local}.xlsx"

                    filter_signature = (
                        f"{len(selected_branches)}_"
                        f"{abs(hash(tuple(selected_branches)))}_"
                        f"{len(selected_employee_names) if selected_employee_names else 'all'}_"
                        f"{filter_start_date}_{filter_end_date}"
                    )

                    col1, col2, col3 = st.columns([1, 2, 1])
                    with col2:
                        st.download_button(
                            label="📥 Download Both Reports (Single Sheet Excel)",
                            data=excel_bytes,
                            file_name=filename,
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            width='stretch',
                            key=f"download_referral_{filter_signature}",
                            help="Download Branch-wise Summary and Employee-wise Details on a single formatted Excel sheet, filtered by your branch selection"
                        )

                    with st.expander("📎 Additional Download Options (CSV)"):
                        st.markdown("**Download individual reports as CSV:**")
                        col_csv1, col_csv2 = st.columns(2)
                        with col_csv1:
                            csv_buffer_branch = io.StringIO()
                            branch_export_df.to_csv(csv_buffer_branch, index=False)
                            st.download_button(
                                label="📥 Branch Summary (CSV)",
                                data=csv_buffer_branch.getvalue(),
                                file_name=f"branch_summary_{branch_part}_{timestamp_local}.csv",
                                mime="text/csv",
                                width='stretch',
                                key=f"csv_branch_{filter_signature}"
                            )
                        with col_csv2:
                            csv_buffer_employee = io.StringIO()
                            employee_export_df.to_csv(csv_buffer_employee, index=False)
                            st.download_button(
                                label="📥 Employee Details (CSV)",
                                data=csv_buffer_employee.getvalue(),
                                file_name=f"employee_details_{branch_part}_{timestamp_local}.csv",
                                mime="text/csv",
                                width='stretch',
                                key=f"csv_employee_{filter_signature}"
                            )
                else:
                    st.info("ℹ️ No employee data available for the selected branches")
            else:
                st.info("ℹ️ No data available for referral report")

        st.markdown("---")
        st.markdown("### 💾 Export All Reports")

        col1, col2 = st.columns(2)
        with col1:
            excel_buffer = create_excel_report(
                filtered_data, generator, duplicates_df, duplicate_indices,
                filter_start_date, filter_end_date,
                selected_branches=selected_branches if 'selected_branches' in locals() else None,
                selected_employee_names=selected_employee_names if 'selected_employee_names' in locals() else None,
            )

            if filter_start_date and filter_end_date:
                filename = f"complete_reports_{filter_start_date}_to_{filter_end_date}_{timestamp}.xlsx"
            else:
                filename = f"complete_reports_all_data_{timestamp}.xlsx"

            all_report_signature = (
                f"{filter_start_date}_{filter_end_date}_"
                f"{abs(hash(tuple(selected_branches))) if 'selected_branches' in locals() else 'all'}_"
                f"{abs(hash(tuple(selected_employee_names))) if 'selected_employee_names' in locals() else 'all'}"
            )
            st.download_button(
                label="📥 Download All Reports (Excel - Complete)",
                data=excel_buffer.getvalue(),
                file_name=filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                width='stretch',
                key=f"download_all_reports_{all_report_signature}"
            )

            if len(duplicates_df) > 0:
                st.info("📊 Excel file includes:")
                st.info("   • ALL records kept in Consolidated Report (duplicates HIGHLIGHTED IN PINK)")
                st.info("   • 'Duplicate Records' sheet with duplicate entries showing original vs duplicate")
                st.info("   • All other reports EXCLUDE duplicates for accurate calculations")
                st.info("   • Branch-wise Scheme, Employee Performance, etc. show data WITHOUT duplicates, ordered Madurai → Marthandam → Salem → Tirunelveli → Trichy → Rajapalayam → Dindigul → Noida → Virudhunagar → Thanjavur → (any new branch) → Telecaller")

        with col2:
            if filter_start_date and filter_end_date:
                st.info(f"💡 **Filter Applied:** {filter_start_date} to {filter_end_date}\n\nThe Excel file will include filtered data only.")
            else:
                st.info("💡 **Tip:** Use the date filters above to generate reports for specific time periods.\n\nYou can filter by week, month, or custom date range.")

    except Exception as e:
        st.markdown(f'<div class="error-box">❌ Error: {str(e)}</div>', unsafe_allow_html=True)
        st.info("💡 **Tip:** Please check that your files have the correct column names and data formats.")


if __name__ == "__main__":
    main()