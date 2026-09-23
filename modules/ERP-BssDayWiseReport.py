"""
New Joining Scheme Report - Streamlit application.

PROCESSING ORDER
----------------
1.  Upload joining register
2.  Read and clean data
3.  Detect branch
4.  Sales Man blank -> e-Gold App
5.  Normalize scheme names
6.  Look every scheme up in the Scheme Master
7.  NEW scheme (not in the master) -> STOP and ASK for its Report Scheme Name
    and its Scheme Type; the answer is saved permanently in the Scheme Master
8.  Validate the COMPLETE uploaded file
9.  Cash Scheme with booked weight -> STOP (or IGNORE, see CASH_WEIGHT_POLICY)
10. Weight Scheme with zero weight -> FLAG FOR REVIEW
11. Booked Wt is kept ONLY on Weight Scheme rows (all other rows -> 0)
12. Apply filters
13. Cash / Weight scheme reports + Joining Register
14. Export Excel
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date
from enum import Enum
from io import BytesIO
from typing import Dict, List, Optional, Tuple

import pandas as pd
import streamlit as st

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="New Joining Scheme Report",
    page_icon="📊",
    layout="wide",
)


# ============================================================
# CONFIGURATION
# ============================================================

MAPPING_FILE = "branch_mapping.csv"
SCHEME_MASTER_FILE = "scheme_master_v2.csv"

CASH_SCHEME = "Cash Scheme"
WEIGHT_SCHEME = "Weight Scheme"

EGOLD_APP_BRANCH = "e-Gold App"
EGOLD_APP_CODE = "EGOLD"

CASH_WEIGHT_POLICY = "STOP"
PENDING_SCHEME = "Pending"

SELECT_OPTION = "— select —"
NEW_REPORT_OPTION = "➕ New report scheme name"


DEFAULT_BRANCHES: Dict[str, str] = {
    "AN": "Anna Nagar Madurai",
    "MDU": "Madurai",
    "SLM": "Salem",
    "TVL": "Tirunelveli",
    "TCY": "Trichy",
    "DGL": "Dindigul",
    "RJ": "Rajapalayam",
    "MDM": "Marthandam",
    "EG": "App",
    "ES": "App",
    "ND": "Noida",
    "VNR": "Virudhunagar",
    "TJ": "Thanjavur",
}


BRANCH_ORDER: List[str] = [
    "MDU",
    "AN",
    "MDM",
    "SLM",
    "TVL",
    "TCY",
    "RJ",
    "DGL",
    "ND",
    "VNR",
    "TJ",
    EGOLD_APP_CODE,
    "EG",
    "ES",
]


# ============================================================
# COLUMN ALIASES
# ============================================================

COLUMN_ALIASES: Dict[str, List[str]] = {
    "srno": ["Srno", "Sr No", "Sr. No", "S.No", "Sl No"],
    "date": ["Date", "Joining Date", "Join Date"],
    "doc": ["Doc No", "DocNo", "Document No", "Document Number"],
    "scheme": ["Scheme", "Scheme Name"],
    "customer": ["Customer", "Customer Name"],
    "mobile": ["Mobileno", "Mobile No", "Mobile", "Mobile Number"],
    "board_rate": ["Board Rate"],
    "weight": [
        "Booked Wt",
        "Booked Weight",
        "Booked Wt (g)",
        "Weight",
        "Gross Wt",
        "Booked Grams",
        "Weight (g)",
        "Wt",
    ],
    "cash": ["Cash Amt"],
    "card": ["Card Amt"],
    "cheque": ["Cheque Amt"],
    "neft": ["Neft Amt", "NEFT"],
    "rtgs": ["Rtgs Amt", "RTGS"],
    "online": ["Online"],
    "voucher": ["Vou Ref Amt"],
    "gift": ["Gift Voucher"],
    "gift2": ["Gift Voucher.1", "Gift Voucher_1"],
    "salesman": ["Sales Man", "Salesman", "Sales Man Name", "Salesman Name"],
    "maturity": ["Maturity Dt.", "Maturity Dt", "Maturity Date"],
    "installment": ["Installment Amount", "Installment"],
}


# ============================================================
# SCHEME NAME NORMALIZATION
# ============================================================

_PAREN_RE = re.compile(r"\(\s*(.*?)\s*\)")
_HYPHEN_RE = re.compile(r"\s*-\s*")


def normalize_scheme_name(name: object) -> str:
    if pd.isna(name):
        return ""

    s = str(name).strip()

    if not s:
        return ""

    s = s.upper()

    s = _PAREN_RE.sub(lambda m: "(" + re.sub(r"\s+", " ", m.group(1)).strip() + ")", s)
    s = re.sub(r"\s*\(", " (", s)

    s = _HYPHEN_RE.sub(" - ", s)

    s = re.sub(r"\s+", " ", s).strip()

    return s


# ============================================================
# SCHEME MASTER - COMPLETE INITIAL MAPPING
# ============================================================

SCHEME_GROUPS: List[Tuple[str, str, List[str]]] = [
    (
        "FUTURE PLUS",
        CASH_SCHEME,
        [
            "FUTURE PLUS CASH",
            "02.BHIMA GOLD TREE FPL 1000",
            "03.BHIMA GOLD TREE FPL 2000",
            "05.BHIMA GOLD TREE FPL 10000",
            "FPL NEW 1000",
            "FPL NEW 2000",
            "BHIMA GOLD TREE FPL COIN 1000",
            "BHIMA GOLD TREE FPL COIN 2000",
            "FPL NEW 500",
            "04.BHIMA GOLD TREE FPL 5000",
            "FPL NEW 10000",
            "FUTURE PLUS 1000",
            "FUTURE PLUS 1000 ( 11 )",
            "BSP 11-1000",
            "FPL NEW 5000",
            "01.BHIMA GOLD TREE FPL 500",
            "FUTURE PLUS 5000",
            "BHIMA GOLD TREE FPL COIN 5000",
            "BHIMA GOLD TREE FPL COIN 10000",
            "Future Plus",
            "FUTURE PLUS 10000",
        ],
    ),
    (
        "FUTURE PLUS 6 M",
        CASH_SCHEME,
        [
            "DIA FUTURE PLUS 10000 ( 6 )",
            "DIA FUTURE PLUS 5000 ( 6 )",
        ],
    ),
    (
        "FUTURE PLUS 18 M",
        CASH_SCHEME,
        [
            "DIA FUTURE PLUS 10000 ( 18 )",
            "DIA FUTURE PLUS 2500 ( 18 )",
            "FUTURE PLUS 500 (18 )",
        ],
    ),
    (
        "BHIMA GOLDEN SIX",
        WEIGHT_SCHEME,
        [
            "BHIMA GOLDEN SIX 6 M - 10000",
            "BHIMA GOLDEN SIX 6 M -10000",
            "BHIMA GOLDEN SIX 6 M - 1000",
            "BHIMA GOLDEN SIX 6 M 10000",
            "BHIMA GOLDEN SIX 6 M - 5000",
            "BHIMA GOLDEN SIX 6 M - 2500",
            "BHIMA GOLDEN SIX 6 M -5000",
            "BHIMA GOLDEN SIX 6 M -25000",
            "BHIMA GOLDEN SIX 6 M-2500",
            "BHIMA GOLDEN SIX 6 M -1000",
            "BHIMA GOLDEN SIX 6 M- 10000",
            "BHIMA GOLDEN SIX 6 M - 25000",
            "BHIMA GOLDEN SIX 6 M -100000",
            "BHIMA GOLDEN SIX 6 M -50000",
            "BHIMA GOLDEN SIX 6 M- 50000",
            "BHIMA GOLDEN SIX 6 M -2500",
            "BHIMA GOLDEN SIX 6 M-25000",
        ],
    ),
    (
        "LUCKY DRAW 18 M",
        CASH_SCHEME,
        [
            "LUCKY DRAW 1000 ( 18 )",
            "LUCKY DRAW 500 ( 18 )",
        ],
    ),
    (
        "GOLD TREE 18 M",
        WEIGHT_SCHEME,
        [
            "GOLD TREE FU GOLD - 1000 ( 18 )",
            "GOLD TREE FU GOLD - 2500 ( 18 )",
            "GOLD TREE FU GOLD - 5000 ( 18 )",
            "GOLD TREE FU GOLD - 10000 ( 18 )",
            "GOLD TREE FU GOLD - 500 ( 18 )",
        ],
    ),
    (
        "GOLD TREE 15 M",
        WEIGHT_SCHEME,
        [
            "GOLD TREE 15 M - 1000",
            "GOLD TREE 15 M - 10000",
            "GOLD TREE 15 M - 2500",
            "GOLD TREE 15 M - 5000",
            "Gold Tree 15 M",
            "GOLD TREE FU GOLD - 1000 ( 15 )",
            "GOLD TREE FU GOLD - 10000 ( 15 )",
            "GOLD TREE FU GOLD - 2500 ( 15 )",
            "GOLD TREE FU GOLD - 5000 ( 15 )",
            "GOLD TREE - 10000 15 M",
            "GOLD TREE - 1000 15M",
        ],
    ),
    (
        "GOLD TREE 6 M",
        WEIGHT_SCHEME,
        [
            "GOLD TREE 6 M -WT 10000",
            "GOLD TREE 6 M - 10000",
            "GOLD TREE FU GOLD 10000 (6)",
            "GOLD TREE 6 M - 1000",
            "GOLD TREE 6 M - 5000",
            "GOLD TREE 6 MONTH -1000",
            "GOLD TREE FU GOLD - 2500 ( 6 )",
            "GOLD TREE 6 MONTH -5000",
            "GOLD TREE 6 M -WT 1000",
            "GOLD TREE 6 MONTH -2500",
            "GOLD TREE 6 MONTH -10000",
            "GOLD TREE 6 MONTH -25000",
            "Gold Tree 6 M",
            "GOLD TREE 6 MONTH - 10000",
        ],
    ),
    (
        "GOLD TREE",
        WEIGHT_SCHEME,
        [
            "GOLD TREE",
            "GOLD TREE FUTURE GOLD - 10000",
            "GOLD TREE FUTURE GOLD - 5000",
            "BGT GOLD TREE WT 1000",
            "BGT GOLD TREE WT 10000",
            "BGT GOLD TREE WT 2000",
            "BGT GOLD TREE WT 25000",
            "BGT GOLD TREE WT 5000",
            "BGT GOLD TREE WT 50000",
            "BGT GOLD TREE WT 100000",
            "BGT GOLD TREE WT 500",
            "GOLD TREE FUTURE GOLD - 1000",
            "GOLD TREE FUTURE GOLD - 2500",
            "GOLD TREE FUTURE GOLD - 25000",
            "GOLD TREE FUTURE GOLD - 50000",
            "GOLD TREE FUTURE",
        ],
    ),
    (
        "ONE TIME DEPOSIT",
        WEIGHT_SCHEME,
        [
            "ONE TIME DEPOSIT SCHEMES",
            "ONE TIME DEPOSIT",
        ],
    ),
    (
        "SILVER TREE PLUS",
        CASH_SCHEME,
        [
            "SILVER TREE 1000",
            "SILVER TREE 500",
            "SILVER TREE 5000",
            "SILVER TREE PLUS 1000",
            "SILVER TREE PLUS 500",
            "SILVER TREE PLUS 2500",
            "SILVER TREE PLUS 5000",
        ],
    ),
    (
        "FUTURE SPARK +",
        CASH_SCHEME,
        [
            "FUTURE SPARK +",
        ],
    ),
]


def _build_default_master() -> Tuple[Dict[str, Tuple[str, str]], List[str]]:
    master: Dict[str, Tuple[str, str]] = {}
    conflicts: List[str] = []

    for report, scheme_type, raw_names in SCHEME_GROUPS:
        for raw in raw_names:
            key = normalize_scheme_name(raw)

            if not key:
                continue

            existing = master.get(key)

            if existing and existing != (scheme_type, report):
                conflicts.append(
                    f"{key}: {existing[1]} ({existing[0]}) vs {report} ({scheme_type})"
                )
                continue

            master[key] = (scheme_type, report)

    return master, conflicts


DEFAULT_SCHEME_MASTER, SCHEME_MASTER_CONFLICTS = _build_default_master()


# ============================================================
# SCHEME MASTER LOAD / SAVE
# ============================================================

def load_scheme_master() -> Dict[str, Tuple[str, str]]:
    default_reports = {
        scheme: report for scheme, (_type, report) in DEFAULT_SCHEME_MASTER.items()
    }

    if os.path.exists(SCHEME_MASTER_FILE):
        try:
            master_df = pd.read_csv(SCHEME_MASTER_FILE)

            if "Scheme" in master_df.columns and "Type" in master_df.columns:
                has_report = "Report Name" in master_df.columns

                loaded: Dict[str, Tuple[str, str]] = {}

                for _, row in master_df.iterrows():
                    scheme = normalize_scheme_name(row["Scheme"])
                    scheme_type = str(row["Type"]).strip()

                    if not scheme:
                        continue

                    if scheme_type not in (CASH_SCHEME, WEIGHT_SCHEME):
                        continue

                    report = row["Report Name"] if has_report else None

                    if report is None or pd.isna(report) or not str(report).strip():
                        report = default_reports.get(scheme, scheme)

                    loaded[scheme] = (scheme_type, str(report).strip())

                if loaded:
                    return loaded

        except Exception as e:
            st.warning(f"⚠️ Could not read {SCHEME_MASTER_FILE}: {e}")

    return dict(DEFAULT_SCHEME_MASTER)


def save_scheme_master(master: Dict[str, Tuple[str, str]]) -> Tuple[bool, str]:
    try:
        cleaned: Dict[str, Tuple[str, str]] = {}

        for scheme, (scheme_type, report) in master.items():
            normalized = normalize_scheme_name(scheme)

            if not normalized:
                continue

            if scheme_type not in (CASH_SCHEME, WEIGHT_SCHEME):
                continue

            cleaned[normalized] = (scheme_type, str(report).strip() or normalized)

        master_df = pd.DataFrame(
            [
                {"Scheme": scheme, "Report Name": report, "Type": scheme_type}
                for scheme, (scheme_type, report) in sorted(cleaned.items())
            ]
        )

        master_df.to_csv(SCHEME_MASTER_FILE, index=False)

        return True, "Scheme master saved."

    except Exception as e:
        return False, f"Could not save scheme master: {e}"


def reset_scheme_master() -> Tuple[bool, str]:
    return save_scheme_master(dict(DEFAULT_SCHEME_MASTER))


# ============================================================
# SCHEME CLASSIFICATION
# ============================================================

def _master_is_current(master: object) -> bool:
    if not isinstance(master, dict) or not master:
        return False

    return all(
        isinstance(value, (tuple, list)) and len(value) == 2
        for value in master.values()
    )


def report_registry(master: Dict[str, Tuple[str, str]]) -> Dict[str, str]:
    registry: Dict[str, str] = {}

    for scheme_type, report in master.values():
        registry.setdefault(report, scheme_type)

    return registry


def suggest_report_name(
    scheme: str,
    master: Dict[str, Tuple[str, str]],
) -> Optional[str]:
    words = scheme.split()

    best_count = 0
    best_report: Optional[str] = None

    for known, (_type, report) in master.items():
        count = 0

        for a, b in zip(words, known.split()):
            if a != b:
                break

            count += 1

        if count >= 2 and count > best_count:
            best_count = count
            best_report = report

    return best_report


def classify_from_master(
    df: pd.DataFrame,
    scheme_col: Optional[str],
    master: Dict[str, Tuple[str, str]],
) -> pd.DataFrame:
    df = df.copy()

    if scheme_col is None or scheme_col not in df.columns:
        df["Scheme Type"] = "N/A"
        df["Scheme Group"] = pd.NA
        df["Scheme Label"] = pd.NA
        return df

    norm = df[scheme_col].astype(str).map(normalize_scheme_name)

    entry = norm.map(master)

    df["Scheme Type"] = entry.map(
        lambda e: e[0] if isinstance(e, (tuple, list)) else PENDING_SCHEME
    )
    df["Scheme Group"] = entry.map(
        lambda e: e[1] if isinstance(e, (tuple, list)) else None
    )
    df["Scheme Label"] = norm

    return df


# ============================================================
# NEW SCHEME MAPPER
# ============================================================

def render_new_scheme_mapper(
    df: pd.DataFrame,
    scheme_col: Optional[str],
    raw_col: str,
    doc_col: str,
    master: Dict[str, Tuple[str, str]],
) -> bool:
    if scheme_col is None or "Scheme Type" not in df.columns:
        return False

    pending = df[df["Scheme Type"] == PENDING_SCHEME]

    if pending.empty:
        return False

    new_schemes = list(pending[scheme_col].drop_duplicates())

    registry = report_registry(master)
    report_names = sorted(registry, key=lambda name: name.upper())

    st.warning(
        f"⚠️ {len(new_schemes)} new scheme(s) detected. Report processing is "
        f"paused until each one is given a Report Scheme Name."
    )

    st.subheader("🆕 New Scheme → Report Scheme Name")

    st.caption(
        "A new scheme is NEVER classified automatically - not by Booked Wt, "
        "not by its name. Choose the existing Report Scheme it belongs to "
        "(its Cash / Weight type is taken from it), or add a new Report "
        "Scheme Name and pick its type. The answer is saved in the Scheme "
        "Master, so you are never asked about the same scheme again."
    )

    for scheme_name in new_schemes:
        rows = pending[pending[scheme_col] == scheme_name]

        display = scheme_name

        if raw_col in rows.columns:
            first_raw = rows[raw_col].dropna()

            if not first_raw.empty:
                display = str(first_raw.iloc[0]).strip() or scheme_name

        example_doc = rows[doc_col].iloc[0] if doc_col in rows.columns else ""

        st.markdown(
            f"**New Scheme Detected:** `{display}`  \n"
            f"**Which Report Scheme Name should this be reported as?**  \n"
            f"{len(rows)} row(s) · Example Doc No `{example_doc}`"
        )

        suggestion = suggest_report_name(scheme_name, master)

        options = [SELECT_OPTION] + report_names + [NEW_REPORT_OPTION]

        default_index = options.index(suggestion) if suggestion in options else 0

        if suggestion:
            st.caption(
                f"Suggestion only: **{suggestion}** (similar to an existing "
                f"scheme) - please confirm or change it."
            )

        choice = st.selectbox(
            f"Report {display} as",
            options,
            index=default_index,
            key=f"newscheme_pick_{scheme_name}",
            format_func=lambda o, reg=registry: (
                f"{o}  —  {reg[o]}" if o in reg else o
            ),
        )

        report_name: Optional[str] = None
        scheme_type: Optional[str] = None

        if choice == NEW_REPORT_OPTION:
            typed = st.text_input(
                f"New Report Scheme Name for {display}",
                key=f"newscheme_name_{scheme_name}",
            )

            typed = re.sub(r"\s+", " ", typed or "").strip()

            type_choice = st.selectbox(
                "Scheme Type of this new Report Scheme Name",
                [SELECT_OPTION, CASH_SCHEME, WEIGHT_SCHEME],
                key=f"newscheme_type_{scheme_name}",
            )

            if typed:
                existing = {name.upper(): name for name in registry}.get(typed.upper())

                if existing:
                    report_name = existing
                    scheme_type = registry[existing]

                    st.info(
                        f"'{existing}' already exists ({scheme_type}); "
                        f"that Report Scheme will be used."
                    )

                elif type_choice in (CASH_SCHEME, WEIGHT_SCHEME):
                    report_name = typed
                    scheme_type = type_choice

        elif choice != SELECT_OPTION:
            report_name = choice
            scheme_type = registry[choice]

        if st.button(
            f"💾 Save mapping for {display}",
            key=f"newscheme_save_{scheme_name}",
        ):
            if not report_name or not scheme_type:
                st.error(
                    "Please choose the Report Scheme Name "
                    "(and the Scheme Type, for a brand new name) first."
                )

            else:
                master[scheme_name] = (scheme_type, report_name)

                ok, msg = save_scheme_master(master)

                st.session_state["scheme_master"] = master

                if ok:
                    st.success(
                        f"✅ {display} → {report_name} ({scheme_type}) saved "
                        f"permanently in the Scheme Master."
                    )
                    st.rerun()
                else:
                    st.error(f"⚠️ {msg}")

        st.divider()

    return True


# ============================================================
# BOOKED WT RULE - WEIGHT SCHEMES ONLY
# ============================================================

def enforce_weight_only_for_weight_schemes(
    df: pd.DataFrame,
    weight_col: Optional[str],
) -> pd.DataFrame:
    if weight_col is None or weight_col not in df.columns:
        return df

    if "Scheme Type" not in df.columns:
        return df

    df = df.copy()

    df[weight_col] = pd.to_numeric(df[weight_col], errors="coerce")
    df.loc[df["Scheme Type"] != WEIGHT_SCHEME, weight_col] = 0.0

    return df


def weight_scheme_total(
    df: pd.DataFrame,
    weight_col: Optional[str],
) -> float:
    if not weight_col or weight_col not in df.columns:
        return 0.0

    if "Scheme Type" not in df.columns:
        return 0.0

    mask = df["Scheme Type"] == WEIGHT_SCHEME

    total = pd.to_numeric(df.loc[mask, weight_col], errors="coerce").sum()

    return 0.0 if pd.isna(total) else float(total)


# ============================================================
# VALIDATION
# ============================================================

@dataclass
class ValidationIssue:
    row_index: int
    doc_no: str
    branch: str
    scheme: str
    scheme_type: str
    weight: float
    issue: str
    message: str


def validate_scheme_types(
    df: pd.DataFrame,
    scheme_col: Optional[str],
    weight_col: Optional[str],
    doc_col: str,
    branch_col: str = "Branch",
) -> List[ValidationIssue]:
    if scheme_col is None or scheme_col not in df.columns:
        return []

    if "Scheme Type" not in df.columns:
        return []

    issues: List[ValidationIssue] = []

    has_weight = weight_col is not None and weight_col in df.columns

    for idx, row in df.iterrows():
        scheme_type = row["Scheme Type"]

        if scheme_type not in (CASH_SCHEME, WEIGHT_SCHEME):
            continue

        label = row.get("Scheme Label")
        scheme = "" if pd.isna(label) else str(label).strip()

        if not scheme:
            continue

        weight = 0.0

        if has_weight:
            raw_weight = row.get(weight_col)

            if pd.notna(raw_weight):
                try:
                    weight = float(raw_weight)
                except (TypeError, ValueError):
                    weight = 0.0

        doc_no = str(row.get(doc_col, "")).strip()

        branch = ""
        if branch_col in df.columns:
            branch = str(row.get(branch_col, "")).strip()

        if scheme_type == CASH_SCHEME and weight > 0:
            issues.append(
                ValidationIssue(
                    row_index=idx,
                    doc_no=doc_no,
                    branch=branch,
                    scheme=scheme,
                    scheme_type=CASH_SCHEME,
                    weight=weight,
                    issue="CASH_WITH_WEIGHT",
                    message=(
                        f"CASH SCHEME VALIDATION FAILED: '{scheme}' has "
                        f"booked weight {weight:,.3f} g. Cash Scheme must "
                        f"have ZERO booked weight."
                    ),
                )
            )

        elif scheme_type == WEIGHT_SCHEME and weight <= 0:
            issues.append(
                ValidationIssue(
                    row_index=idx,
                    doc_no=doc_no,
                    branch=branch,
                    scheme=scheme,
                    scheme_type=WEIGHT_SCHEME,
                    weight=weight,
                    issue="WEIGHT_WITHOUT_WEIGHT",
                    message=(
                        f"WEIGHT SCHEME REVIEW: '{scheme}' has booked "
                        f"weight = 0. Please verify the source booking."
                    ),
                )
            )

    return issues


def has_cash_weight_errors(issues: List[ValidationIssue]) -> bool:
    return any(issue.issue == "CASH_WITH_WEIGHT" for issue in issues)


def has_hard_validation_errors(issues: List[ValidationIssue]) -> bool:
    return has_cash_weight_errors(issues)


# ============================================================
# BRANCH MAPPING
# ============================================================

def load_branch_mapping() -> Dict[str, str]:
    if os.path.exists(MAPPING_FILE):
        try:
            mapping_df = pd.read_csv(MAPPING_FILE)

            if "Code" in mapping_df.columns and "Branch Name" in mapping_df.columns:
                return dict(
                    zip(
                        mapping_df["Code"].astype(str).str.upper().str.strip(),
                        mapping_df["Branch Name"].astype(str).str.strip(),
                    )
                )
        except Exception:
            pass

    return DEFAULT_BRANCHES.copy()


def save_branch_mapping(mapping: Dict[str, str]) -> Tuple[bool, str]:
    try:
        mapping_df = pd.DataFrame(
            [{"Code": code, "Branch Name": branch} for code, branch in mapping.items()]
        ).sort_values("Code")

        mapping_df.to_csv(MAPPING_FILE, index=False)

        return True, "Mapping saved."

    except Exception as e:
        return False, f"Could not save mapping: {e}"


def detect_branches_vectorized(
    doc_series: pd.Series,
    mapping: Dict[str, str],
) -> Tuple[pd.Series, pd.Series]:
    if not mapping:
        empty = pd.Series([None] * len(doc_series), index=doc_series.index)
        return empty.copy(), empty.copy()

    codes_sorted = sorted(mapping.keys(), key=len, reverse=True)

    pattern = "|".join(re.escape(code) for code in codes_sorted)

    normalized = doc_series.astype(str).str.upper().str.strip()

    detected_code = normalized.str.extract(f"({pattern})", expand=False)

    branch = detected_code.map(mapping)

    return branch, detected_code


def assign_branches(
    df: pd.DataFrame,
    doc_col: str,
    mapping: Dict[str, str],
) -> None:
    branch_series, code_series = detect_branches_vectorized(df[doc_col], mapping)

    df["Branch"] = branch_series
    df["Detected Code"] = code_series


def extract_unknown_code(
    doc_no: object,
    mapping: Dict[str, str],
) -> Optional[str]:
    if pd.isna(doc_no):
        return None

    doc_no = str(doc_no).strip().upper()

    if not doc_no:
        return None

    for code in mapping.keys():
        if code in doc_no:
            return None

    match = re.match(r"^[A-Z]+(?=[\d/_\-\s])", doc_no)

    if match:
        return match.group(0)

    match = re.match(r"^[A-Z]+", doc_no)

    if match:
        return match.group(0)

    return None


# ============================================================
# SALESMAN BLANK -> e-GOLD APP
# ============================================================

def apply_salesman_blank_rule(
    df: pd.DataFrame,
    salesman_col: Optional[str],
    egold_branch: str = EGOLD_APP_BRANCH,
) -> pd.DataFrame:
    if "Branch" not in df.columns:
        df["Branch"] = None

    if "Branch Source" not in df.columns:
        df["Branch Source"] = "Doc No Prefix"

    if salesman_col is None or salesman_col not in df.columns:
        return df

    salesman_str = df[salesman_col].astype(str).str.strip()

    salesman_blank = (
        df[salesman_col].isna()
        | (salesman_str == "")
        | (salesman_str.str.lower() == "nan")
        | (salesman_str.str.lower() == "none")
    )

    df.loc[salesman_blank, "Branch"] = egold_branch
    df.loc[salesman_blank, "Branch Source"] = "Salesman Blank → e-Gold App"

    return df


# ============================================================
# BRANCH ORDERING
# ============================================================

def sort_branches_by_codes(codes: List[str]) -> List[str]:
    seen = set()
    ordered: List[str] = []

    for code in BRANCH_ORDER:
        if code in codes and code not in seen:
            ordered.append(code)
            seen.add(code)

    leftovers = sorted(code for code in codes if code not in seen)

    return ordered + leftovers


def _build_name_to_code_lookup() -> Dict[str, str]:
    name_to_code: Dict[str, str] = {}

    mapping = st.session_state.get("branch_mapping", DEFAULT_BRANCHES)

    for code, name in mapping.items():
        name_to_code.setdefault(name, code)

    for code, name in DEFAULT_BRANCHES.items():
        name_to_code.setdefault(name, code)

    name_to_code[EGOLD_APP_BRANCH] = EGOLD_APP_CODE

    return name_to_code


def sort_branches_by_df(df: pd.DataFrame) -> List[str]:
    if "Branch" not in df.columns:
        return []

    present_names = list(df["Branch"].dropna().unique())

    if not present_names:
        return []

    name_to_code = _build_name_to_code_lookup()

    names_with_code = [n for n in present_names if n in name_to_code]

    present_codes = [name_to_code[n] for n in names_with_code]

    ordered_codes = sort_branches_by_codes(present_codes)

    code_to_name: Dict[str, str] = {}

    for name in names_with_code:
        code_to_name.setdefault(name_to_code[name], name)

    ordered_names: List[str] = []
    seen: set = set()

    for code in ordered_codes:
        name = code_to_name.get(code)

        if name and name not in seen:
            ordered_names.append(name)
            seen.add(name)

    for name in sorted(n for n in present_names if n not in seen):
        ordered_names.append(name)

    return ordered_names


# ============================================================
# GENERAL HELPERS
# ============================================================

def make_unique_columns(columns) -> List[str]:
    counts: Dict[str, int] = {}
    new_columns: List[str] = []

    for col in columns:
        col = str(col).strip()

        if col not in counts:
            counts[col] = 0
            new_columns.append(col)
        else:
            counts[col] += 1
            new_columns.append(f"{col}_{counts[col]}")

    return new_columns


def find_column(
    df: pd.DataFrame,
    possible_names: List[str],
) -> Optional[str]:
    normalized = {str(c).strip().lower(): c for c in df.columns}

    for name in possible_names:
        key = str(name).strip().lower()

        if key in normalized:
            return normalized[key]

    return None


def clean_numeric_column(df: pd.DataFrame, col: Optional[str]) -> None:
    if col is None or col not in df.columns:
        return

    cleaned = (
        df[col]
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("₹", "", regex=False)
        .str.replace(r"^Rs\.?\s*", "", regex=True)
        .str.strip()
        .replace({"": None, "nan": None, "None": None, "NaN": None})
    )

    df[col] = pd.to_numeric(cleaned, errors="coerce")


# ============================================================
# COLUMN MAP
# ============================================================

@dataclass(frozen=True)
class ColumnMap:
    srno: Optional[str]
    date: Optional[str]
    doc: Optional[str]
    scheme: Optional[str]
    customer: Optional[str]
    mobile: Optional[str]
    board_rate: Optional[str]
    weight: Optional[str]
    cash: Optional[str]
    card: Optional[str]
    cheque: Optional[str]
    neft: Optional[str]
    rtgs: Optional[str]
    online: Optional[str]
    voucher: Optional[str]
    gift: Optional[str]
    gift2: Optional[str]
    salesman: Optional[str]
    maturity: Optional[str]
    installment: Optional[str]

    @classmethod
    def detect(cls, df: pd.DataFrame) -> "ColumnMap":
        return cls(
            **{attr: find_column(df, names) for attr, names in COLUMN_ALIASES.items()}
        )

    @property
    def numeric(self) -> List[str]:
        candidates = [
            self.board_rate,
            self.weight,
            self.installment,
            self.cash,
            self.card,
            self.cheque,
            self.neft,
            self.rtgs,
            self.online,
            self.voucher,
            self.gift,
        ]

        return list(dict.fromkeys(c for c in candidates if c is not None))


# ============================================================
# EXCEL STYLES
# ============================================================

TITLE_BG = "1F4E79"
SECTION_BG = "2E75B6"

TYPE_CASH_BG = "FCE4D6"
TYPE_WEIGHT_BG = "E2EFDA"

SCHEME_CASH_BG = "F8CBAD"
SCHEME_WEIGHT_BG = "C6E0B4"

BRANCH_HDR_BG = "4472C4"
COL_HDR_BG = "D9E1F2"

WEIGHT_COL_BG = "E2EFDA"
AMOUNT_COL_BG = "FFEB9C"
JOINING_COL_BG = "D9E1F2"
BRANCH_COL_BG = "E2EFDA"

PCT_JOINING_BG = "EDF3FB"
PCT_WEIGHT_BG = "F2F9EF"
PCT_AMOUNT_BG = "FFF8E1"

TOTAL_BG = "FFF2CC"
GRAND_BG = "C00000"

EGOLD_COL_BG = "FCE4D6"
DAY_BANNER_BG = "8EA9DB"

WHITE = "FFFFFF"
BLACK = "000000"
BORDER_COLOR = "B7B7B7"

_THIN = Side(style="thin", color=BORDER_COLOR)
_MEDIUM = Side(style="medium", color=BLACK)

BORDER_THIN = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
BORDER_THICK = Border(left=_MEDIUM, right=_MEDIUM, top=_MEDIUM, bottom=_MEDIUM)

FONT_TITLE = Font(bold=True, size=16, color=WHITE)
FONT_SECTION = Font(bold=True, size=14, color=WHITE)
FONT_DAY_BANNER = Font(bold=True, size=12, color=BLACK)
FONT_TYPE = Font(bold=True, size=12, color=BLACK)
FONT_SCHEME = Font(bold=True, size=11, color=BLACK)
FONT_BRANCH_HDR = Font(bold=True, size=10, color=WHITE)
FONT_COL_HDR = Font(bold=True, size=10, color=BLACK)
FONT_TOTAL = Font(bold=True, size=10, color=BLACK)
FONT_GRAND = Font(bold=True, size=11, color=WHITE)
FONT_DATA = Font(size=10, color=BLACK)
FONT_LABEL = Font(bold=True, size=10, color=BLACK)
FONT_PERIOD = Font(bold=True, size=11, color=BLACK)
FONT_PCT = Font(size=9, italic=True, color="595959")


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


FILL_TITLE = _fill(TITLE_BG)
FILL_SECTION = _fill(SECTION_BG)
FILL_DAY_BANNER = _fill(DAY_BANNER_BG)
FILL_TYPE_CASH = _fill(TYPE_CASH_BG)
FILL_TYPE_WEIGHT = _fill(TYPE_WEIGHT_BG)
FILL_SCHEME_CASH = _fill(SCHEME_CASH_BG)
FILL_SCHEME_WEIGHT = _fill(SCHEME_WEIGHT_BG)
FILL_BRANCH_HDR = _fill(BRANCH_HDR_BG)
FILL_WEIGHT_COL = _fill(WEIGHT_COL_BG)
FILL_AMOUNT_COL = _fill(AMOUNT_COL_BG)
FILL_JOINING_COL = _fill(JOINING_COL_BG)
FILL_BRANCH_COL = _fill(BRANCH_COL_BG)
FILL_PCT_JOINING = _fill(PCT_JOINING_BG)
FILL_PCT_WEIGHT = _fill(PCT_WEIGHT_BG)
FILL_PCT_AMOUNT = _fill(PCT_AMOUNT_BG)
FILL_TOTAL = _fill(TOTAL_BG)
FILL_GRAND = _fill(GRAND_BG)
FILL_EGOLD = _fill(EGOLD_COL_BG)
FILL_WARN = _fill("FFC7CE")
FILL_WARN_HDR = _fill("C00000")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_LEFT_INDENT = Alignment(horizontal="left", vertical="center", indent=2)

FMT_COUNT = '#,##0;-#,##0;"-"'
FMT_WEIGHT = '#,##0.000;-#,##0.000;""'
FMT_AMOUNT = '"₹"#,##0.00;-"₹"#,##0.00;""'
FMT_PERCENT = '0.00"%"'


class ColKind(str, Enum):
    COUNT = "count"
    WEIGHT = "weight"
    AMOUNT = "amount"
    PCT_COUNT = "pct_count"
    PCT_WEIGHT = "pct_weight"
    PCT_AMOUNT = "pct_amount"


KIND_FILL = {
    ColKind.COUNT: FILL_JOINING_COL,
    ColKind.WEIGHT: FILL_WEIGHT_COL,
    ColKind.AMOUNT: FILL_AMOUNT_COL,
    ColKind.PCT_COUNT: FILL_PCT_JOINING,
    ColKind.PCT_WEIGHT: FILL_PCT_WEIGHT,
    ColKind.PCT_AMOUNT: FILL_PCT_AMOUNT,
}

KIND_NUMBER_FORMAT = {
    ColKind.COUNT: FMT_COUNT,
    ColKind.WEIGHT: FMT_WEIGHT,
    ColKind.AMOUNT: FMT_AMOUNT,
    ColKind.PCT_COUNT: FMT_PERCENT,
    ColKind.PCT_WEIGHT: FMT_PERCENT,
    ColKind.PCT_AMOUNT: FMT_PERCENT,
}

STYLE_DATA = "data"
STYLE_TOTAL = "total"
STYLE_GRAND = "grand"
STYLE_DAY = "day"

ColumnSpec = Tuple[str, str, ColKind]

COLUMNS_CASH: Tuple[ColumnSpec, ...] = (
    ("Joinings", "Joinings", ColKind.COUNT),
    ("Joinings %", "Joinings %", ColKind.PCT_COUNT),
    ("Amount", "Amount", ColKind.AMOUNT),
    ("Amount %", "Amount %", ColKind.PCT_AMOUNT),
)

COLUMNS_WEIGHT: Tuple[ColumnSpec, ...] = (
    ("Joinings", "Joinings", ColKind.COUNT),
    ("Joinings %", "Joinings %", ColKind.PCT_COUNT),
    ("Booked Wt", "Booked Wt (g)", ColKind.WEIGHT),
    ("Booked Wt %", "Booked Wt %", ColKind.PCT_WEIGHT),
    ("Amount", "Amount", ColKind.AMOUNT),
    ("Amount %", "Amount %", ColKind.PCT_AMOUNT),
)


# ============================================================
# EXCEL STYLE HELPERS
# ============================================================

def _label_style(style: str):
    if style == STYLE_GRAND:
        return FONT_GRAND, FILL_GRAND, BORDER_THICK

    if style == STYLE_TOTAL:
        return FONT_TOTAL, FILL_TOTAL, BORDER_THIN

    if style == STYLE_DAY:
        return FONT_LABEL, FILL_JOINING_COL, BORDER_THIN

    return FONT_LABEL, FILL_BRANCH_COL, BORDER_THIN


def _value_style(style: str, kind: ColKind):
    if style == STYLE_GRAND:
        return FONT_GRAND, FILL_GRAND, BORDER_THICK

    if style == STYLE_TOTAL:
        return FONT_TOTAL, FILL_TOTAL, BORDER_THIN

    if style == STYLE_DAY:
        return FONT_TOTAL, FILL_JOINING_COL, BORDER_THIN

    if kind in (ColKind.PCT_COUNT, ColKind.PCT_WEIGHT, ColKind.PCT_AMOUNT):
        return FONT_PCT, KIND_FILL[kind], BORDER_THIN

    return FONT_DATA, KIND_FILL[kind], BORDER_THIN


# ============================================================
# JOINING REGISTER
# ============================================================

REGISTER_COLUMNS: List[str] = [
    "Srno",
    "Date",
    "Doc No",
    "Scheme",
    "Customer",
    "Mobileno",
    "Board Rate",
    "Booked Wt",
    "Cash Amt",
    "Card Amt",
    "Cheque Amt",
    "Neft Amt",
    "Rtgs Amt",
    "Online",
    "Vou Ref Amt",
    "Installment Amount",
    "Gift Voucher",
    "Gift Voucher (2)",
    "Sales Man",
    "Maturity Dt.",
]

REGISTER_TOTAL_COLUMNS: List[str] = [
    "Booked Wt",
    "Cash Amt",
    "Card Amt",
    "Cheque Amt",
    "Neft Amt",
    "Rtgs Amt",
    "Online",
    "Vou Ref Amt",
    "Installment Amount",
]


def build_register_df(df: pd.DataFrame, cols: ColumnMap) -> pd.DataFrame:
    df = df.reset_index(drop=True)
    n = len(df)

    def pick(col: Optional[str]) -> pd.Series:
        if col and col in df.columns:
            return df[col]
        return pd.Series([None] * n, dtype="object")

    if "Scheme Type" in df.columns:
        is_weight = df["Scheme Type"].eq(WEIGHT_SCHEME)
    else:
        is_weight = pd.Series([False] * n)

    if cols.weight and cols.weight in df.columns:
        wt = pd.to_numeric(df[cols.weight], errors="coerce")
        booked = wt.where(is_weight & (wt > 0))
    else:
        booked = pd.Series([None] * n, dtype="object")

    scheme_src = "Scheme (Original)" if "Scheme (Original)" in df.columns else cols.scheme

    srno = pick(cols.srno) if cols.srno else pd.Series(range(1, n + 1))

    maturity = pick(cols.maturity)
    parsed = pd.to_datetime(maturity, errors="coerce", dayfirst=True)
    if parsed.notna().sum() >= maturity.notna().sum():
        maturity = parsed

    out = pd.DataFrame(
        {
            "Srno": srno,
            "Date": pick(cols.date),
            "Doc No": pick(cols.doc),
            "Scheme": pick(scheme_src),
            "Customer": pick(cols.customer),
            "Mobileno": pick(cols.mobile),
            "Board Rate": pick(cols.board_rate),
            "Booked Wt": booked,
            "Cash Amt": pick(cols.cash),
            "Card Amt": pick(cols.card),
            "Cheque Amt": pick(cols.cheque),
            "Neft Amt": pick(cols.neft),
            "Rtgs Amt": pick(cols.rtgs),
            "Online": pick(cols.online),
            "Vou Ref Amt": pick(cols.voucher),
            "Installment Amount": pick(cols.installment),
            "Gift Voucher": pick(cols.gift),
            "Gift Voucher (2)": pick(cols.gift2),
            "Sales Man": pick(cols.salesman),
            "Maturity Dt.": maturity,
        }
    )

    return out[REGISTER_COLUMNS]


def write_register_sheet(wb: Workbook, register_df: pd.DataFrame):
    ws = wb.create_sheet("Joining Register")
    ws.sheet_view.showGridLines = False

    formats = {
        "Date": "dd-mm-yyyy",
        "Maturity Dt.": "dd-mm-yyyy",
        "Board Rate": "#,##0.00",
        "Booked Wt": FMT_WEIGHT,
        "Gift Voucher": FMT_AMOUNT,
    }
    for name in REGISTER_TOTAL_COLUMNS[1:]:
        formats[name] = FMT_AMOUNT

    for c, name in enumerate(REGISTER_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name.replace(" (2)", ""))
        cell.font = FONT_BRANCH_HDR
        cell.fill = FILL_BRANCH_HDR
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN

    ws.row_dimensions[1].height = 28

    def clean(v):
        if v is None:
            return None
        if not isinstance(v, str) and pd.isna(v):
            return None
        if isinstance(v, pd.Timestamp):
            return v.to_pydatetime()
        if hasattr(v, "item") and not isinstance(v, str):
            try:
                return v.item()
            except Exception:
                return v
        return v

    for r, row in enumerate(register_df.itertuples(index=False), start=2):
        for c, (name, v) in enumerate(zip(REGISTER_COLUMNS, row), start=1):
            cell = ws.cell(row=r, column=c, value=clean(v))
            cell.font = FONT_DATA
            cell.border = BORDER_THIN

            if name in formats:
                cell.number_format = formats[name]
                cell.alignment = ALIGN_RIGHT
            else:
                cell.alignment = ALIGN_LEFT

    last = len(register_df) + 1
    total_row = last + 1

    if len(register_df) > 0:
        for c, name in enumerate(REGISTER_COLUMNS, start=1):
            cell = ws.cell(row=total_row, column=c)
            cell.font = FONT_TOTAL
            cell.fill = FILL_TOTAL
            cell.border = BORDER_THIN

            if c == 1:
                cell.value = "TOTAL"
            elif name in REGISTER_TOTAL_COLUMNS:
                letter = get_column_letter(c)
                cell.value = f"=SUM({letter}2:{letter}{last})"
                cell.number_format = formats[name]
                cell.alignment = ALIGN_RIGHT

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(REGISTER_COLUMNS))}{last}"

    for c, name in enumerate(REGISTER_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = (
            26 if name in ("Scheme", "Customer") else 16
        )

    return ws


# ============================================================
# EXCEL BUILDER
# ============================================================

def build_scheme_joining_excel(
    filtered_df: pd.DataFrame,
    report_start: pd.Timestamp,
    report_end: pd.Timestamp,
    date_col: str,
    scheme_col: Optional[str],
    weight_col: Optional[str],
    installment_col: Optional[str],
    validation_issues: Optional[List[ValidationIssue]] = None,
    register_df: Optional[pd.DataFrame] = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Scheme Joining Report"
    ws.sheet_view.showGridLines = False

    # --------------------------------------------------------
    # SCHEME ORDER
    # --------------------------------------------------------

    SchemeKey = Tuple[str, bool]

    if (
        scheme_col is None
        or "Scheme Group" not in filtered_df.columns
        or "Scheme Type" not in filtered_df.columns
    ):
        schemes_ordered: List[SchemeKey] = []
    else:
        pairs = (
            filtered_df[["Scheme Group", "Scheme Type"]]
            .dropna()
            .drop_duplicates()
        )

        schemes_ordered = sorted(
            {(str(lbl), typ == WEIGHT_SCHEME) for lbl, typ in pairs.itertuples(index=False)},
            key=lambda k: (k[1], k[0].upper()),
        )

    scheme_is_weight = {key: key[1] for key in schemes_ordered}

    scheme_col_plan = [
        (key, COLUMNS_WEIGHT if key[1] else COLUMNS_CASH)
        for key in schemes_ordered
    ]

    total_cols = 1 + sum(len(cols) for _, cols in scheme_col_plan)

    if total_cols == 1:
        total_cols = 2

    has_any_cash = any(not w for w in scheme_is_weight.values())
    has_any_weight = any(scheme_is_weight.values())

    # --------------------------------------------------------
    # SUMMARIZATION
    # --------------------------------------------------------

    def summarize(sch: SchemeKey, sch_df: pd.DataFrame) -> Dict[str, float]:
        result = {
            "Joinings": int(len(sch_df)),
            "Booked Wt": 0.0,
            "Amount": 0.0,
        }

        if (
            sch[1]
            and weight_col is not None
            and weight_col in sch_df.columns
        ):
            result["Booked Wt"] = float(sch_df[weight_col].fillna(0).sum())

        if installment_col is not None and installment_col in sch_df.columns:
            result["Amount"] = float(sch_df[installment_col].fillna(0).sum())

        return result

    def summarize_all(data: pd.DataFrame) -> Dict[SchemeKey, Dict[str, float]]:
        if not schemes_ordered:
            return {}

        if "Scheme Group" not in data.columns or "Scheme Type" not in data.columns:
            return {}

        return {
            key: summarize(
                key,
                data[
                    (data["Scheme Group"] == key[0])
                    & (data["Scheme Type"] == (WEIGHT_SCHEME if key[1] else CASH_SCHEME))
                ],
            )
            for key in schemes_ordered
        }

    # --------------------------------------------------------
    # PERCENTAGE HELPER
    # --------------------------------------------------------

    def make_pct_attacher(
        base_totals: Dict[SchemeKey, Dict[str, float]]
    ):
        def attach(
            values: Dict[SchemeKey, Dict[str, float]]
        ) -> Dict[SchemeKey, Dict[str, float]]:
            out: Dict[SchemeKey, Dict[str, float]] = {}

            for key in schemes_ordered:
                merged = dict(values.get(key, {}))

                base = base_totals.get(key, {})

                j = float(merged.get("Joinings", 0) or 0)
                w = float(merged.get("Booked Wt", 0) or 0)
                a = float(merged.get("Amount", 0) or 0)

                j_base = float(base.get("Joinings", 0) or 0)
                w_base = float(base.get("Booked Wt", 0) or 0)
                a_base = float(base.get("Amount", 0) or 0)

                merged["Joinings %"] = (j / j_base * 100.0) if j_base > 0 else 0.0
                merged["Booked Wt %"] = (w / w_base * 100.0) if w_base > 0 else 0.0
                merged["Amount %"] = (a / a_base * 100.0) if a_base > 0 else 0.0

                out[key] = merged

            return out

        return attach

    # --------------------------------------------------------
    # ROW WRITERS
    # --------------------------------------------------------

    def write_title(row: int, text: str, span: int) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)

        c = ws.cell(row=row, column=1, value=text)
        c.font = FONT_TITLE
        c.fill = FILL_TITLE
        c.alignment = ALIGN_CENTER

        for cc in range(2, span + 1):
            ws.cell(row=row, column=cc).fill = FILL_TITLE

        ws.row_dimensions[row].height = 30

    def write_section(row: int, text: str, span: int) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)

        c = ws.cell(row=row, column=1, value=text)
        c.font = FONT_SECTION
        c.fill = FILL_SECTION
        c.alignment = ALIGN_CENTER

        for cc in range(2, span + 1):
            ws.cell(row=row, column=cc).fill = FILL_SECTION

        ws.row_dimensions[row].height = 24

    def write_merged_group(
        row: int,
        start: int,
        span: int,
        text: str,
        font,
        fill,
    ) -> int:
        end = start + span - 1

        if span > 1:
            ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)

        c = ws.cell(row=row, column=start, value=text)
        c.font = font
        c.fill = fill
        c.alignment = ALIGN_CENTER

        for cc in range(start, end + 1):
            ws.cell(row=row, column=cc).border = BORDER_THIN
            ws.cell(row=row, column=cc).fill = fill

        return end + 1

    def write_corner_cell(row: int, text: str = "") -> None:
        c = ws.cell(row=row, column=1, value=text)

        if text:
            c.font = FONT_BRANCH_HDR

        c.fill = FILL_BRANCH_HDR
        c.border = BORDER_THIN
        c.alignment = ALIGN_CENTER

    def write_type_group_row(row: int) -> None:
        write_corner_cell(row)

        col = 2

        for label, fill in (
            (CASH_SCHEME, FILL_TYPE_CASH),
            (WEIGHT_SCHEME, FILL_TYPE_WEIGHT),
        ):
            span = sum(
                len(cols)
                for scheme, cols in scheme_col_plan
                if (WEIGHT_SCHEME if scheme_is_weight[scheme] else CASH_SCHEME) == label
            )

            if span > 0:
                col = write_merged_group(row, col, span, label, FONT_TYPE, fill)

        ws.row_dimensions[row].height = 22

    def write_scheme_header_row(row: int, first_label: str = "Branch") -> None:
        c = ws.cell(row=row, column=1, value=first_label)
        c.font = FONT_BRANCH_HDR
        c.fill = FILL_BRANCH_HDR
        c.alignment = ALIGN_CENTER
        c.border = BORDER_THIN

        col = 2

        for scheme, cols in scheme_col_plan:
            fill = FILL_SCHEME_WEIGHT if scheme_is_weight[scheme] else FILL_SCHEME_CASH

            col = write_merged_group(row, col, len(cols), scheme[0], FONT_SCHEME, fill)

        ws.row_dimensions[row].height = 24

    def write_column_header_row(row: int, first_label: str = "") -> None:
        write_corner_cell(row, first_label)

        col = 2

        for _, cols in scheme_col_plan:
            for _key, hdr_text, kind in cols:
                cell = ws.cell(row=row, column=col, value=hdr_text)
                cell.font = FONT_COL_HDR
                cell.alignment = ALIGN_CENTER
                cell.border = BORDER_THIN
                cell.fill = KIND_FILL[kind]

                col += 1

        ws.row_dimensions[row].height = 32

    def write_data_row(
        row: int,
        label: str,
        values: Dict[SchemeKey, Dict[str, float]],
        style: str = STYLE_DATA,
        indent: bool = False,
        highlight_egold: bool = False,
        is_total_row: bool = False,
    ) -> None:
        """Write a data / total / grand-total row.

        Percentage rule (applies to every % column):

        * Data rows  -> actual % (0 -> 0.00 %)
        * Total rows -> 100.00 % IF the scheme has any base value, otherwise 0.00 %
        """
        font, fill, border = _label_style(style)

        if highlight_egold and style == STYLE_DATA:
            fill = FILL_EGOLD

        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.alignment = ALIGN_LEFT_INDENT if indent else ALIGN_LEFT
        label_cell.border = border
        label_cell.font = font
        label_cell.fill = fill

        col = 2

        for scheme, cols in scheme_col_plan:
            scheme_vals = values.get(scheme, {})

            for key, _hdr, kind in cols:
                is_pct = kind in (
                    ColKind.PCT_COUNT,
                    ColKind.PCT_WEIGHT,
                    ColKind.PCT_AMOUNT,
                )

                if is_pct:
                    if is_total_row:
                        base_key = {
                            ColKind.PCT_COUNT: "Joinings",
                            ColKind.PCT_WEIGHT: "Booked Wt",
                            ColKind.PCT_AMOUNT: "Amount",
                        }[kind]

                        base_val = float(scheme_vals.get(base_key, 0) or 0)

                        value = 100.0 if base_val > 0 else 0.0
                    else:
                        value = scheme_vals.get(key, None)

                        if value is None:
                            value = 0.0

                        value = float(value)
                else:
                    value = scheme_vals.get(key, 0)

                    if value == 0 and kind in (ColKind.WEIGHT, ColKind.AMOUNT):
                        value = None

                cell = ws.cell(row=row, column=col, value=value)

                vfont, vfill, vborder = _value_style(style, kind)

                if highlight_egold and style == STYLE_DATA:
                    vfill = FILL_EGOLD

                cell.border = vborder
                cell.alignment = ALIGN_RIGHT
                cell.number_format = KIND_NUMBER_FORMAT[kind]
                cell.font = vfont
                cell.fill = vfill

                col += 1

    def write_day_banner(row: int, text: str, span: int) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)

        c = ws.cell(row=row, column=1, value=text)
        c.font = FONT_DAY_BANNER
        c.fill = FILL_DAY_BANNER
        c.alignment = ALIGN_LEFT_INDENT
        c.border = BORDER_THIN

        for cc in range(2, span + 1):
            cell = ws.cell(row=row, column=cc)
            cell.fill = FILL_DAY_BANNER
            cell.border = BORDER_THIN

        ws.row_dimensions[row].height = 22

    def write_day_table_headers(row: int) -> int:
        if has_any_cash and has_any_weight:
            write_type_group_row(row)
            row += 1

        write_scheme_header_row(row, first_label="📅 Date / Branch")
        row += 1

        write_column_header_row(row)
        row += 1

        return row

    def write_table_headers(row: int, first_label: str) -> int:
        if has_any_cash and has_any_weight:
            write_type_group_row(row)
            row += 1

        write_scheme_header_row(row, first_label=first_label)
        row += 1

        write_column_header_row(row)
        row += 1

        return row

    # --------------------------------------------------------
    # BRANCH LIST
    # --------------------------------------------------------

    branches = sort_branches_by_df(filtered_df)

    # ========================================================
    # MAIN SHEET
    # ========================================================

    row = 1

    write_title(row, "📊 NEW JOINING SCHEME REPORT", total_cols)

    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_cols)

    period_text = (
        f"Report Period: {report_start.strftime('%d-%m-%Y')} "
        f"to {report_end.strftime('%d-%m-%Y')}"
    )

    c = ws.cell(row=row, column=1, value=period_text)
    c.font = FONT_PERIOD
    c.alignment = ALIGN_CENTER

    row += 2

    if not schemes_ordered:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_cols)

        msg = ws.cell(
            row=row,
            column=1,
            value="⚠️ No scheme column was found in the upload.",
        )
        msg.font = FONT_PERIOD
        msg.alignment = ALIGN_CENTER

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output.getvalue()

    # --------------------------------------------------------
    # BRANCH-WISE SECTION
    # --------------------------------------------------------

    total_values = summarize_all(filtered_df)

    grand_totals_per_scheme: Dict[SchemeKey, Dict[str, float]] = {
        key: {
            "Joinings": float(total_values.get(key, {}).get("Joinings", 0) or 0),
            "Booked Wt": float(total_values.get(key, {}).get("Booked Wt", 0) or 0),
            "Amount": float(total_values.get(key, {}).get("Amount", 0) or 0),
        }
        for key in schemes_ordered
    }

    attach_pct_period = make_pct_attacher(grand_totals_per_scheme)

    row = write_table_headers(row, first_label="Branch")

    for branch in branches:
        branch_values = summarize_all(filtered_df[filtered_df["Branch"] == branch])

        write_data_row(
            row,
            branch,
            attach_pct_period(branch_values),
            highlight_egold=(branch == EGOLD_APP_BRANCH),
        )

        row += 1

    # Branch-wise TOTAL row
    write_data_row(
        row,
        "📌 TOTAL (Branch-wise)",
        attach_pct_period(total_values),
        STYLE_TOTAL,
        is_total_row=True,
    )
    row += 1

    write_data_row(
        row,
        "🔴 GRAND TOTAL",
        attach_pct_period(total_values),
        STYLE_GRAND,
        is_total_row=True,
    )
    row += 2

    # --------------------------------------------------------
    # DAY-WISE SECTION
    # --------------------------------------------------------

    write_section(row, "📅 DAY-WISE BSS REPORT", total_cols)

    row += 1

    if date_col in filtered_df.columns and len(filtered_df) > 0:
        day_keys = pd.to_datetime(filtered_df[date_col]).dt.date

        unique_dates = sorted(day_keys.unique())

        for d in unique_dates:
            day_df = filtered_df[day_keys == d]

            day_name = pd.Timestamp(d).strftime("%A")

            date_label = f"📅 {d.strftime('%d-%m-%Y')} ({day_name})"
            write_day_banner(row, date_label, total_cols)
            row += 1

            # ---- THIS DAY's totals become the denominator ----
            day_totals = summarize_all(day_df)

            day_grand_totals: Dict[SchemeKey, Dict[str, float]] = {
                key: {
                    "Joinings": float(day_totals.get(key, {}).get("Joinings", 0) or 0),
                    "Booked Wt": float(day_totals.get(key, {}).get("Booked Wt", 0) or 0),
                    "Amount": float(day_totals.get(key, {}).get("Amount", 0) or 0),
                }
                for key in schemes_ordered
            }

            attach_pct_day = make_pct_attacher(day_grand_totals)

            row = write_day_table_headers(row)

            for branch in sort_branches_by_df(day_df):
                branch_values = summarize_all(day_df[day_df["Branch"] == branch])

                write_data_row(
                    row,
                    branch,
                    attach_pct_day(branch_values),
                    indent=True,
                    highlight_egold=(branch == EGOLD_APP_BRANCH),
                )

                row += 1

            # Per-date subtotal
            write_data_row(
                row,
                f"📌 {d.strftime('%d-%m-%Y')} TOTAL",
                attach_pct_day(day_totals),
                STYLE_TOTAL,
                is_total_row=True,
            )

            row += 2

    # --------------------------------------------------------
    # COLUMN WIDTHS
    # --------------------------------------------------------

    ws.column_dimensions["A"].width = 28

    for c_idx in range(2, total_cols + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 15

    # --------------------------------------------------------
    # PAGE SETUP
    # --------------------------------------------------------

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.50
    ws.page_margins.bottom = 0.50

    ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    # ========================================================
    # JOINING REGISTER SHEET
    # ========================================================

    if register_df is not None:
        write_register_sheet(wb, register_df)

    # ========================================================
    # VALIDATION SHEET
    # ========================================================

    if validation_issues:
        ws2 = wb.create_sheet("⚠️ Validation Issues")
        ws2.sheet_view.showGridLines = False

        v_headers = [
            "#",
            "Issue Type",
            "Doc No",
            "Branch",
            "Scheme",
            "Scheme Type",
            "Weight (g)",
            "Message",
        ]

        for col_idx, header in enumerate(v_headers, start=1):
            cell = ws2.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, color=WHITE)
            cell.fill = FILL_WARN_HDR
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN

        for i, issue in enumerate(validation_issues, start=1):
            r = i + 1

            values = [
                i,
                issue.issue,
                issue.doc_no,
                issue.branch,
                issue.scheme,
                issue.scheme_type,
                issue.weight,
                issue.message,
            ]

            for col_idx, value in enumerate(values, start=1):
                cell = ws2.cell(row=r, column=col_idx, value=value)
                cell.border = BORDER_THIN
                cell.fill = FILL_WARN
                cell.alignment = ALIGN_LEFT if col_idx in (2, 8) else ALIGN_CENTER

                if col_idx == 7:
                    cell.number_format = FMT_WEIGHT

        for letter, width in zip("ABCDEFGH", (6, 25, 18, 22, 32, 16, 14, 75)):
            ws2.column_dimensions[letter].width = width

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output.getvalue()


# ============================================================
# FILE READING
# ============================================================

@st.cache_data(show_spinner=False)
def read_uploaded_file(file_bytes: bytes, file_name: str) -> pd.DataFrame:
    buffer = BytesIO(file_bytes)

    if file_name.lower().endswith(".csv"):
        try:
            return pd.read_csv(buffer)
        except UnicodeDecodeError:
            buffer.seek(0)
            return pd.read_csv(buffer, encoding="latin-1")

    return pd.read_excel(buffer)


def load_dataframe(uploaded_file) -> pd.DataFrame:
    try:
        return read_uploaded_file(uploaded_file.getvalue(), uploaded_file.name)

    except ImportError as ie:
        st.error(
            f"❌ Missing Excel engine: {ie}\n\n"
            "Install the required engine:\n"
            "- .xlsx: pip install openpyxl\n"
            "- .xls: pip install xlrd"
        )
        st.stop()

    except Exception as e:
        st.error(f"❌ Unable to read file: {e}")
        st.stop()


# ============================================================
# DATA PREPARATION
# ============================================================

def prepare_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()

    df.columns = make_unique_columns(df.columns)

    empty_columns = [col for col in df.columns if df[col].isna().all()]

    if empty_columns:
        df = df.drop(columns=empty_columns)

    return df


def require_columns(df: pd.DataFrame, cols: ColumnMap) -> None:
    missing = []

    if cols.date is None:
        missing.append("Date")

    if cols.doc is None:
        missing.append("Doc No")

    if missing:
        st.error("❌ Required column(s) not found: " + ", ".join(missing))
        st.write("Columns detected:", list(df.columns))
        st.stop()


def convert_dates(df: pd.DataFrame, date_col: str) -> pd.DataFrame:
    df = df.copy()

    df[date_col] = pd.to_datetime(df[date_col], errors="coerce", dayfirst=True)

    invalid_dates = int(df[date_col].isna().sum())

    if invalid_dates > 0:
        st.warning(f"⚠️ {invalid_dates} invalid date row(s) removed.")
        df = df.dropna(subset=[date_col])

    if df.empty:
        st.error("❌ No rows with a valid date were found.")
        st.stop()

    return df


# ============================================================
# UI HEADER
# ============================================================

def render_header() -> None:
    st.title("📊 NEW JOINING SCHEME REPORT")

    st.caption(
        "Daily + Cumulative + Branch-wise + Scheme-wise Report | "
        "Scheme Master is the final source of truth | "
        "Booked Wt from Weight Schemes only | "
        "Sales Man blank → e-Gold App"
    )


def render_welcome() -> None:
    st.info("📁 Please upload your Excel or CSV file.")

    st.markdown(
        """
### Features

- 📅 Daily joining report
- 📈 Cumulative joining report
- 🏢 Branch-wise report
- 📋 Scheme-wise Cash / Weight columns
- 📊 **Percentage column for EVERY metric** (Joinings %, Booked Wt %, Amount %)
- 📅 **Day-wise BSS Report uses each day's own totals as the denominator** (each day sums to 100%)
- 🔢 **Empty TOTAL rows show 0.00 %, not 100 %**
- 🗂️ **Scheme Master is the final source of truth**
- 🚨 **Complete-file validation before report processing**
- ⚖️ **Booked Wt taken from Weight Schemes only**
- 💰 Every Cash Scheme checked for booked weight
- 🗂️ Every raw scheme is reported under ONE Report Scheme Name
- 🆕 **A new scheme is never guessed** — you are asked for its Report
  Scheme Name + Type, and the answer is saved permanently
- 🔤 Scheme names normalized (spacing / case only — the tier marker
  `( 6 )` / `( 18 )` / `( 11 )` is always preserved)
- 🧑‍💼 Sales Man blank → e-Gold App
- 📋 Joining Register in exact column order
- 📗 Formatted Excel export with **day-wise banner headers per date**
- ⬜ Zero weight/amount shown as blank
- `-` shown for zero counts
"""
    )


# ============================================================
# UNKNOWN BRANCH MAPPER
# ============================================================

def render_unknown_code_mapper(
    df: pd.DataFrame,
    doc_col: str,
    mapping: Dict[str, str],
) -> None:
    candidate_df = df

    if "Branch Source" in df.columns:
        candidate_df = df[df["Branch Source"] == "Doc No Prefix"]

    unknown_df = candidate_df[candidate_df["Branch"].isna()].copy()

    if not unknown_df.empty:
        unknown_df["Unknown Code"] = unknown_df[doc_col].apply(
            lambda x: extract_unknown_code(x, mapping)
        )

        unknown_df = unknown_df[unknown_df["Unknown Code"].notna()].drop_duplicates(
            subset=["Unknown Code"]
        )

    if unknown_df.empty:
        return

    def save_code(code: str, branch_name: str) -> None:
        mapping[code] = branch_name

        ok, msg = save_branch_mapping(mapping)

        st.session_state["branch_mapping"] = mapping

        if ok:
            st.success(f"✅ {code} → {branch_name} saved.")
            st.rerun()
        else:
            st.error(f"⚠️ {msg}")

    st.warning(f"⚠️ {len(unknown_df)} new Doc No code(s) detected.")

    st.subheader("🆕 New Doc No → Branch Mapping")

    branch_options = ["➕ Add New Branch"] + sorted(set(mapping.values()))

    for _, unknown_row in unknown_df.iterrows():
        code = unknown_row["Unknown Code"]
        example_doc = unknown_row[doc_col]

        st.markdown(f"**New Code:** `{code}`  \n**Example Doc No:** `{example_doc}`")

        selected_branch = st.selectbox(
            f"Select Branch for {code}",
            branch_options,
            key=f"branch_{code}",
        )

        if selected_branch == "➕ Add New Branch":
            new_branch = st.text_input(
                f"New Branch Name for {code}",
                key=f"new_branch_{code}",
            )

            if st.button(f"💾 Save {code}", key=f"save_new_{code}"):
                if new_branch.strip():
                    save_code(code, new_branch.strip())

        else:
            if st.button(
                f"💾 Save {code} → {selected_branch}",
                key=f"save_{code}",
            ):
                save_code(code, selected_branch)


# ============================================================
# SCHEME TYPES SIDEBAR
# ============================================================

def render_scheme_types_sidebar(df: pd.DataFrame) -> None:
    if "Scheme Group" not in df.columns:
        return

    data = df.dropna(subset=["Scheme Group"])

    if data.empty:
        return

    st.sidebar.subheader("🧠 Scheme Classification")

    for title, scheme_type in (
        ("💰 Cash Schemes", CASH_SCHEME),
        ("⚖️ Weight Schemes", WEIGHT_SCHEME),
    ):
        part = data[data["Scheme Type"] == scheme_type]

        if part.empty:
            continue

        groups = sorted(part["Scheme Group"].unique(), key=lambda g: str(g).upper())

        st.sidebar.markdown(f"**{title} ({len(groups)}):**")

        for group in groups:
            n_schemes = part.loc[part["Scheme Group"] == group, "Scheme Label"].nunique()

            detail = f" — {n_schemes} schemes" if n_schemes > 1 else ""

            st.sidebar.markdown(f"- {group}{detail}")


# ============================================================
# SCHEME MASTER EDITOR
# ============================================================

def render_scheme_master_editor(master: Dict[str, Tuple[str, str]]) -> None:
    with st.expander("🗂️ Edit Scheme Master", expanded=False):
        st.caption(
            "The Scheme Master is the FINAL SOURCE OF TRUTH. It decides how "
            "every raw scheme is reported: the Report Scheme Name it is "
            "grouped under and its Cash / Weight type."
        )

        if SCHEME_MASTER_CONFLICTS:
            st.error(
                "⚠️ The built-in mapping lists the same raw scheme under two "
                "different Report Scheme Names:\n\n- "
                + "\n- ".join(SCHEME_MASTER_CONFLICTS)
            )

        master_df = pd.DataFrame(
            [
                {"Scheme": scheme, "Report Name": report, "Type": scheme_type}
                for scheme, (scheme_type, report) in sorted(master.items())
            ]
        )

        edited_df = st.data_editor(
            master_df,
            width="stretch",
            hide_index=True,
            column_config={
                "Scheme": st.column_config.TextColumn(required=True, width="large"),
                "Report Name": st.column_config.TextColumn(required=True, width="large"),
                "Type": st.column_config.SelectboxColumn(
                    options=[CASH_SCHEME, WEIGHT_SCHEME],
                    required=True,
                ),
            },
            num_rows="dynamic",
            key="scheme_master_editor",
        )

        left, right = st.columns(2)

        if left.button("💾 Save Scheme Master", key="save_scheme_master"):
            new_master: Dict[str, Tuple[str, str]] = {}

            for _, row in edited_df.iterrows():
                scheme = normalize_scheme_name(row["Scheme"])
                scheme_type = str(row["Type"]).strip()

                report = row["Report Name"]
                report = "" if pd.isna(report) else str(report).strip()

                if scheme and scheme_type in (CASH_SCHEME, WEIGHT_SCHEME):
                    new_master[scheme] = (scheme_type, report or scheme)

            ok, msg = save_scheme_master(new_master)

            if ok:
                st.session_state["scheme_master"] = new_master
                st.success("✅ Scheme master saved.")
                st.rerun()
            else:
                st.error(f"⚠️ {msg}")

        if right.button(
            "♻️ Reset to built-in mapping",
            key="reset_scheme_master",
            help="Discard every saved change and restore the complete "
                 "built-in scheme list.",
        ):
            ok, msg = reset_scheme_master()

            if ok:
                st.session_state["scheme_master"] = dict(DEFAULT_SCHEME_MASTER)
                st.success("✅ Scheme master reset to the built-in mapping.")
                st.rerun()
            else:
                st.error(f"⚠️ {msg}")

        st.download_button(
            "⬇️ Download Scheme Master (CSV)",
            master_df.to_csv(index=False),
            file_name="scheme_master.csv",
            mime="text/csv",
            key="download_scheme_master",
        )


# ============================================================
# VALIDATION PANEL
# ============================================================

def render_validation_panel(issues: List[ValidationIssue]) -> None:
    if not issues:
        st.success("✅ COMPLETE FILE VALIDATION PASSED")
        st.success("Every Cash Scheme has ZERO booked weight.")
        return

    cash_issues = [i for i in issues if i.issue == "CASH_WITH_WEIGHT"]
    weight_issues = [i for i in issues if i.issue == "WEIGHT_WITHOUT_WEIGHT"]

    if cash_issues:
        st.error("🛑 CASH SCHEME VALIDATION FAILED")
        st.error(f"{len(cash_issues)} Cash Scheme booking(s) have booked weight.")
        st.warning(
            "Every Cash Scheme must have ZERO booked weight. "
            "Booked Wt is taken from Weight Schemes only."
        )

    if weight_issues:
        st.warning(
            f"⚠️ {len(weight_issues)} Weight Scheme booking(s) have zero booked weight."
        )

    c1, c2 = st.columns(2)

    c1.metric("🛑 Cash w/ Weight", len(cash_issues))
    c2.metric("⚠️ Weight w/o Weight", len(weight_issues))

    issues_df = pd.DataFrame(
        [
            {
                "Issue": i.issue,
                "Doc No": i.doc_no,
                "Branch": i.branch,
                "Scheme": i.scheme,
                "Scheme Type": i.scheme_type,
                "Weight": i.weight,
                "Message": i.message,
            }
            for i in issues
        ]
    )

    st.dataframe(issues_df, width="stretch", hide_index=True)

    st.download_button(
        "⬇️ Download Validation Issues (CSV)",
        issues_df.to_csv(index=False),
        file_name="scheme_validation_issues.csv",
        mime="text/csv",
        key="download_validation",
    )


# ============================================================
# SIDEBAR FILTERS
# ============================================================

def apply_sidebar_filters(
    df: pd.DataFrame,
    cols: ColumnMap,
) -> Tuple[pd.DataFrame, date, date]:
    st.sidebar.subheader("📅 Date Filter")

    min_date = df[cols.date].min().date()
    max_date = df[cols.date].max().date()

    start_date = st.sidebar.date_input("From Date", min_date)
    end_date = st.sidebar.date_input("To Date", max_date)

    if start_date > end_date:
        st.error("❌ From Date cannot be after To Date.")
        st.stop()

    filtered_df = df[
        (df[cols.date].dt.date >= start_date) & (df[cols.date].dt.date <= end_date)
    ].copy()

    if cols.scheme is not None:
        scheme_values = sorted(
            filtered_df["Scheme Group"].dropna().unique(),
            key=lambda x: str(x).upper(),
        )

        selected_schemes = st.sidebar.multiselect(
            "📋 Report Scheme",
            scheme_values,
            default=scheme_values,
        )

        filtered_df = filtered_df[filtered_df["Scheme Group"].isin(selected_schemes)]

        scheme_type_values = sorted(filtered_df["Scheme Type"].dropna().unique())

        selected_scheme_types = st.sidebar.multiselect(
            "💰 Scheme Type",
            scheme_type_values,
            default=scheme_type_values,
        )

        filtered_df = filtered_df[filtered_df["Scheme Type"].isin(selected_scheme_types)]

    branch_values = sort_branches_by_df(filtered_df)

    selected_branches = st.sidebar.multiselect(
        "🏢 Branch",
        branch_values,
        default=branch_values,
    )

    filtered_df = filtered_df[filtered_df["Branch"].isin(selected_branches)]

    return filtered_df.reset_index(drop=True), start_date, end_date


# ============================================================
# KPI
# ============================================================

def _safe_sum(df: pd.DataFrame, col: Optional[str]) -> float:
    if not col:
        return 0.0

    if col not in df.columns:
        return 0.0

    total = df[col].sum()

    if pd.isna(total):
        return 0.0

    return float(total)


def render_kpis(filtered_df: pd.DataFrame, cols: ColumnMap) -> None:
    total_joinings = len(filtered_df)

    total_weight = weight_scheme_total(filtered_df, cols.weight)

    total_installment = _safe_sum(filtered_df, cols.installment)

    st.subheader("📊 Report Summary")

    has_weight_scheme_in_data = False

    if cols.scheme is not None:
        has_weight_scheme_in_data = bool(
            (filtered_df["Scheme Type"] == WEIGHT_SCHEME).any()
        )

    col1, col2, col3 = st.columns(3)

    col1.metric("👥 Total Joinings", f"{total_joinings:,}")

    if has_weight_scheme_in_data:
        col2.metric("⚖️ Booked Weight", f"{total_weight:,.3f}")
    else:
        col2.metric("⚖️ Booked Weight", "N/A (No Weight Scheme)")

    col3.metric("💰 Installment Amount", f"₹{total_installment:,.2f}")

    if "Branch Source" in filtered_df.columns:
        egold_count = int(
            filtered_df["Branch Source"]
            .astype(str)
            .str.contains("e-Gold App", na=False)
            .sum()
        )

        if egold_count:
            st.info(
                f"ℹ️ {egold_count} joining(s) assigned to "
                f"**{EGOLD_APP_BRANCH}** because Sales Man is blank."
            )


# ============================================================
# SCHEME BREAKDOWN
# ============================================================

def render_scheme_breakdown(filtered_df: pd.DataFrame, cols: ColumnMap) -> None:
    if cols.scheme is None:
        return

    st.markdown("### 🧠 Scheme Type Breakdown")

    cash_df = filtered_df[filtered_df["Scheme Type"] == CASH_SCHEME]
    weight_df = filtered_df[filtered_df["Scheme Type"] == WEIGHT_SCHEME]

    cash_delta = None

    if cols.installment:
        cash_amount = cash_df[cols.installment].sum()
        cash_delta = f"₹{cash_amount:,.0f}"

    weight_delta = None

    if cols.weight:
        weight_total = weight_scheme_total(weight_df, cols.weight)
        weight_delta = f"{weight_total:,.3f} g"

    left, right = st.columns(2)

    with left:
        st.metric(
            "🟠 Cash Scheme Joinings",
            f"{len(cash_df):,}",
            delta=cash_delta,
        )

    with right:
        st.metric(
            "🟢 Weight Scheme Joinings",
            f"{len(weight_df):,}",
            delta=weight_delta,
        )


# ============================================================
# SUMMARY TABLES
# ============================================================

def build_summaries(
    filtered_df: pd.DataFrame,
    cols: ColumnMap,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    daily_agg = {"Joinings": (cols.doc, "count")}

    if cols.weight:
        daily_agg["Booked Wt"] = (cols.weight, "sum")

    if cols.installment:
        daily_agg["Installment Amount"] = (cols.installment, "sum")

    daily = (
        filtered_df.assign(__date=filtered_df[cols.date].dt.date)
        .groupby("__date", as_index=False)
        .agg(**daily_agg)
        .rename(columns={"__date": "Date"})
        .sort_values("Date")
        .reset_index(drop=True)
    )

    cumulative = daily.copy()

    cumulative["Cumulative Joinings"] = cumulative["Joinings"].cumsum()

    if "Booked Wt" in cumulative.columns:
        cumulative["Cumulative Wt"] = cumulative["Booked Wt"].cumsum()

    if "Installment Amount" in cumulative.columns:
        cumulative["Cumulative Installment"] = cumulative["Installment Amount"].cumsum()

    branch_agg = {"Joinings": (cols.doc, "count")}

    if cols.weight:
        branch_agg["Booked Wt"] = (cols.weight, "sum")

    if cols.installment:
        branch_agg["Installment Amount"] = (cols.installment, "sum")

    branch_summary = (
        filtered_df.groupby("Branch", dropna=False).agg(**branch_agg).reset_index()
    )

    branch_summary["Branch"] = branch_summary["Branch"].fillna("(Unknown)")

    ordered_names = sort_branches_by_df(filtered_df)

    name_to_rank = {name: i for i, name in enumerate(ordered_names)}

    branch_summary["__order"] = branch_summary["Branch"].map(
        lambda branch: name_to_rank.get(branch, len(name_to_rank))
    )

    branch_summary = (
        branch_summary.sort_values("__order")
        .drop(columns="__order")
        .reset_index(drop=True)
    )

    return cumulative, branch_summary


# ============================================================
# MAPPING VIEWERS
# ============================================================

def render_branch_mapping_viewer(mapping: Dict[str, str]) -> None:
    with st.expander("⚙️ View / Export Branch Mapping"):
        mapping_df = pd.DataFrame(
            [{"Code": code, "Branch Name": branch} for code, branch in mapping.items()]
        ).sort_values("Code")

        st.dataframe(mapping_df, width="stretch", hide_index=True)

        st.download_button(
            "⬇️ Download Branch Mapping (CSV)",
            mapping_df.to_csv(index=False),
            file_name="branch_mapping.csv",
            mime="text/csv",
            key="download_mapping",
        )


def render_scheme_mapping_viewer(df: pd.DataFrame) -> None:
    with st.expander("🧠 View Scheme Type Classification"):
        if "Scheme Label" not in df.columns or df["Scheme Label"].dropna().empty:
            st.info("No schemes were classified.")
            return

        scheme_df = (
            df[["Scheme Label", "Scheme Group", "Scheme Type"]]
            .dropna(subset=["Scheme Label"])
            .drop_duplicates()
            .rename(
                columns={
                    "Scheme Label": "Scheme",
                    "Scheme Group": "Reported As",
                    "Scheme Type": "Type",
                }
            )
            .sort_values(["Type", "Reported As", "Scheme"])
            .reset_index(drop=True)
        )

        st.dataframe(scheme_df, width="stretch", hide_index=True)


# ============================================================
# FOOTER
# ============================================================

def render_footer() -> None:
    st.divider()

    st.caption(
        "📊 New Joining Scheme Report | "
        "Scheme Master is the final source of truth | "
        "Booked Wt from Weight Schemes only | "
        "Sales-Report Style Excel | "
        "Day-wise banner headers per date | "
        "Day-wise % uses that day's own totals | "
        "Empty TOTAL rows show 0.00 % | "
        "Zero weight/amount shown as blank | "
        "Zero counts as '-' | "
        "Scheme names normalized (tier marker preserved) | "
        "Branches ordered by BRANCH_ORDER | "
        "Sales Man blank → e-Gold App"
    )


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    if "branch_mapping" not in st.session_state:
        st.session_state["branch_mapping"] = load_branch_mapping()

    mapping = st.session_state["branch_mapping"]

    if not _master_is_current(st.session_state.get("scheme_master")):
        st.session_state["scheme_master"] = load_scheme_master()

    scheme_master = st.session_state["scheme_master"]

    render_header()

    st.sidebar.header("⚙️ Report Controls")

    uploaded_file = st.sidebar.file_uploader(
        "📁 Upload New Joining Report",
        type=["xlsx", "xls", "csv"],
    )

    st.sidebar.caption("Max file size: 200 MB")

    st.sidebar.caption(
        f"🗂️ Scheme Master: {len(scheme_master)} raw scheme(s) → "
        f"{len(report_registry(scheme_master))} report scheme(s)"
    )

    if uploaded_file is None:
        render_welcome()

        render_scheme_master_editor(scheme_master)

        st.stop()

    # STEP 1 — READ FILE
    df = load_dataframe(uploaded_file)
    df = prepare_dataframe(df)

    # STEP 2 — DETECT COLUMNS
    cols = ColumnMap.detect(df)

    require_columns(df, cols)

    # STEP 3 — CONVERT DATES
    df = convert_dates(df, cols.date)

    # STEP 4 — CLEAN NUMERIC COLUMNS
    for numeric_col in cols.numeric:
        clean_numeric_column(df, numeric_col)

    # STEP 5 — BRANCH DETECTION
    assign_branches(df, cols.doc, mapping)

    # STEP 6 — SALESMAN BLANK
    df = apply_salesman_blank_rule(df, cols.salesman)

    # STEP 7 — UNKNOWN BRANCH MAPPING
    render_unknown_code_mapper(df, cols.doc, mapping)

    # STEP 8 — NORMALIZE SCHEME NAMES
    if cols.scheme is not None:
        df["Scheme (Original)"] = df[cols.scheme]

        df[cols.scheme] = (
            df[cols.scheme].apply(normalize_scheme_name).replace("", pd.NA)
        )

        df = df.dropna(subset=[cols.scheme]).copy()

    # STEP 9 — CLASSIFY FROM THE SCHEME MASTER
    df = classify_from_master(df, cols.scheme, scheme_master)

    if render_new_scheme_mapper(
        df,
        cols.scheme,
        "Scheme (Original)",
        cols.doc,
        scheme_master,
    ):
        render_scheme_master_editor(scheme_master)

        st.stop()

    render_scheme_types_sidebar(df)

    render_scheme_master_editor(scheme_master)

    # STEP 10 — VALIDATE COMPLETE UPLOAD
    validation_issues = validate_scheme_types(
        df=df,
        scheme_col=cols.scheme,
        weight_col=cols.weight,
        doc_col=cols.doc,
    )

    # STEP 11 — SHOW COMPLETE FILE VALIDATION
    st.subheader("🚨 Complete File Scheme Validation")

    render_validation_panel(validation_issues)

    # STEP 12 — CASH SCHEME VALIDATION GATE
    cash_weight_issues = [i for i in validation_issues if i.issue == "CASH_WITH_WEIGHT"]

    if cash_weight_issues and CASH_WEIGHT_POLICY == "STOP":
        st.error("🛑 REPORT PROCESSING STOPPED")

        st.error("Cash Scheme validation failed.")

        st.markdown(
            """
### ❌ Cash Scheme / Weight Mismatch

The uploaded data contains one or more:

**Cash Scheme (from Scheme Master) + Booked Weight > 0**

Booked Wt must come from Weight Schemes only, so a Cash Scheme cannot carry booked weight.

The application will **NOT continue** until the source data or Scheme Master is corrected.
"""
        )

        st.stop()

    # STEP 14 — BOOKED WT ONLY FOR WEIGHT SCHEMES
    if cash_weight_issues:
        st.warning(
            f"⚠️ {len(cash_weight_issues)} Cash Scheme row(s) had booked weight "
            f"in the source file. That weight is ignored and excluded from all "
            f"Booked Weight totals."
        )
    else:
        st.success("✅ ALL CASH SCHEMES CONFIRMED")

        st.success(
            "Every Cash Scheme in the complete uploaded file has ZERO booked weight."
        )

    df = enforce_weight_only_for_weight_schemes(df, cols.weight)

    # STEP 15 — WEIGHT SCHEME WARNING
    weight_zero_issues = [
        i for i in validation_issues if i.issue == "WEIGHT_WITHOUT_WEIGHT"
    ]

    if weight_zero_issues:
        st.warning(
            f"⚠️ {len(weight_zero_issues)} Weight Scheme booking(s) have zero "
            f"booked weight. Please review the source data."
        )

    # STEP 16 — APPLY FILTERS
    filtered_df, start_date, end_date = apply_sidebar_filters(df, cols)

    if filtered_df.empty:
        st.warning("⚠️ No data matches the selected filters.")
        st.stop()

    # STEP 17 — KPI
    render_kpis(filtered_df, cols)

    # STEP 18 — SCHEME BREAKDOWN
    render_scheme_breakdown(filtered_df, cols)

    # STEP 19 — VALIDATION SUMMARY
    st.subheader("🚨 Validation Result")

    if cash_weight_issues:
        st.warning(
            "⚠️ Cash Scheme weight was ignored (CASH_WEIGHT_POLICY = IGNORE)."
        )
    else:
        st.success("✅ Cash Scheme validation passed before report processing.")

    if weight_zero_issues:
        st.warning(
            f"⚠️ Weight Scheme review required: {len(weight_zero_issues)} row(s) "
            f"have zero booked weight."
        )
    else:
        st.success("✅ All Weight Scheme rows have booked weight > 0.")

    # STEP 20 — DAILY / CUMULATIVE
    cumulative, branch_summary = build_summaries(filtered_df, cols)

    st.subheader("📅 Daily & Cumulative")

    st.dataframe(cumulative, width="stretch", hide_index=True)

    # STEP 21 — BRANCH SUMMARY
    st.subheader("🏢 Branch Summary")

    st.dataframe(branch_summary, width="stretch", hide_index=True)

    # STEP 22 — JOINING REGISTER
    register_df = build_register_df(filtered_df, cols)

    st.subheader("📋 Joining Register")

    st.dataframe(register_df, width="stretch", hide_index=True)

    # STEP 23 — EXCEL EXPORT
    excel_data = build_scheme_joining_excel(
        filtered_df=filtered_df,
        report_start=pd.Timestamp(start_date),
        report_end=pd.Timestamp(end_date),
        date_col=cols.date,
        scheme_col=cols.scheme,
        weight_col=cols.weight,
        installment_col=cols.installment,
        validation_issues=validation_issues,
        register_df=register_df,
    )

    st.subheader("⬇️ Download Report")

    def _fmt_report_date(d) -> str:
        """Format a date as DDMonYYYY, e.g. 01Sep2026."""
        return pd.Timestamp(d).strftime("%d%b%Y")

    report_file_name = (
        "New_Joining_Branch_Wise_BSS_Report_Cumulative_Daily_"
        f"{_fmt_report_date(start_date)}_"
        f"{_fmt_report_date(end_date)}"
        ".xlsx"
    )

    st.download_button(
        label="📥 DOWNLOAD FORMATTED EXCEL",
        data=excel_data,
        file_name=report_file_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
    )

    st.caption(f"📄 File will be saved as: `{report_file_name}`")

    # STEP 24 — MAPPINGS
    render_branch_mapping_viewer(mapping)

    render_scheme_mapping_viewer(df)

    render_footer()


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    main()