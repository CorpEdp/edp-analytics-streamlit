import streamlit as st
import pandas as pd
import numpy as np
from io import BytesIO
from calendar import month_name

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)
from openpyxl.utils import get_column_letter


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Employee Referral Performance",
    page_icon="📊",
    layout="wide"
)

st.title("📊 Employee Referral Performance Dashboard")
st.caption(
    "Any Month | Employee-wise | Scheme-wise | Branch-wise"
)


# ============================================================
# CONSTANTS
# ============================================================

EXPECTED_COLUMNS = [
    "Updated Date",
    "Customer Name",
    "Customer Phone",
    "Customer Enrollment Amount",
    "Status",
    "Employee Name",
    "Referral Code",
    "Employee Phone",
    "Employee Code",
    "Branch",
    "Customer Payment",
    "True/False",
    "Scheme Name",
    "Scheme Passbook Number",
    "Category",
    "Month",
    "Not Enrolled",
    "Is Duplicate",
    "Duplicate Group",
    "Week Number",
    "Month Name"
]

ALL_MONTHS = [
    "January", "February", "March", "April",
    "May", "June", "July", "August",
    "September", "October", "November", "December"
]

DEFAULT_MONTHS = ["June", "July", "August"]


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def clean_text(series):
    return (
        series
        .astype(str)
        .str.strip()
        .replace(["nan", "None", "NaN"], "")
    )


def clean_number(series):
    return (
        series
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("₹", "", regex=False)
        .str.replace("$", "", regex=False)
        .str.strip()
        .replace(["", "-", "nan", "None", "NaN"], np.nan)
        .pipe(pd.to_numeric, errors="coerce")
        .fillna(0)
    )


def is_true(series):
    normalized = (
        series
        .astype(str)
        .str.strip()
        .str.lower()
    )
    return normalized.isin(
        ["true", "yes", "1", "1.0", "y", "t", "yes "]
    )


def format_amount(value):
    # Plain number with 2 decimals and thousands separator — no ₹ symbol
    return f"{value:,.2f}"


def format_percent(value):
    try:
        return f"{value:.2f}%"
    except Exception:
        return "0.00%"


def safe_value(value):
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    if isinstance(value, (np.floating, np.integer)):
        try:
            return bool(pd.isna(value))
        except Exception:
            return False
    return False


def add_percentage(df, value_columns, total_columns=None):
    df = df.copy()

    for col in value_columns:
        if col not in df.columns:
            continue

        if total_columns and col in total_columns:
            total = df[total_columns[col]].sum()
        else:
            total = df[col].sum()

        display_name = col.replace("_", " ")
        pct_col = f"{display_name} %"

        if total > 0:
            df[pct_col] = (df[col] / total) * 100
        else:
            df[pct_col] = 0.0

        df[pct_col] = df[pct_col].round(2)

    return df


def select_existing_columns(df, desired):
    cols = [c for c in desired if c in df.columns]
    return df[cols]


def month_percent_columns(df, months, month_prefixes):
    df = df.copy()

    for month in months:
        for prefix in month_prefixes:
            src_col = f"{month} {prefix}"
            if src_col not in df.columns:
                continue

            pct_col = f"{month} {prefix} %"
            total = df[src_col].sum()

            if total > 0:
                df[pct_col] = (df[src_col] / total) * 100
            else:
                df[pct_col] = 0.0

            df[pct_col] = df[pct_col].round(2)

    return df


def build_month_column_order(months, prefixes):
    cols = []
    for month in months:
        for prefix in prefixes:
            cols.append(f"{month} {prefix}")
            cols.append(f"{month} {prefix} %")
    return cols


def build_employee_scheme_wide(df):
    """
    Build the wide Employee + Scheme table with per-scheme % columns
    and a Grand Total % column.
    """
    pivot = (
        df
        .pivot_table(
            index=["Employee Name", "Employee Code", "Branch"],
            columns="Scheme Name",
            values="Customer Enrollment Amount",
            aggfunc="sum",
            fill_value=0
        )
    )
    pivot.columns.name = None
    pivot = pivot.reset_index()

    scheme_columns = sorted(
        [c for c in pivot.columns
         if c not in ("Employee Name", "Employee Code", "Branch")],
        key=lambda x: str(x)
    )

    # ---- Grand Total per employee ----
    pivot["Grand Total"] = pivot[scheme_columns].sum(axis=1)

    # ---- Grand Total % (share of overall total) ----
    grand_total_overall = pivot["Grand Total"].sum()
    if grand_total_overall > 0:
        pivot["Grand Total %"] = (
            (pivot["Grand Total"] / grand_total_overall) * 100
        ).round(2)
    else:
        pivot["Grand Total %"] = 0.0

    # ---- Per-scheme % (share of that scheme's overall total) ----
    for sc in scheme_columns:
        scheme_total = pivot[sc].sum()
        pct_col = f"{sc} %"
        if scheme_total > 0:
            pivot[pct_col] = ((pivot[sc] / scheme_total) * 100).round(2)
        else:
            pivot[pct_col] = 0.0

    # ---- Sort by Grand Total desc + Rank ----
    pivot = pivot.sort_values(
        "Grand Total",
        ascending=False
    ).reset_index(drop=True)

    pivot.insert(0, "Rank", pivot.index + 1)

    # ---- Final column order ----
    final_cols = ["Rank", "Employee Name", "Employee Code", "Branch"]
    for sc in scheme_columns:
        final_cols.append(sc)
        final_cols.append(f"{sc} %")
    final_cols.append("Grand Total")
    final_cols.append("Grand Total %")

    pivot = pivot[final_cols]

    return pivot


# ============================================================
# SINGLE SHEET EXCEL REPORT
# ============================================================

def create_formatted_excel(
    monthly,
    scheme_report,
    employee_report,
    employee_scheme_wide,
    branch_report,
    selected_year,
    selected_months
):

    wb = Workbook()
    ws = wb.active
    ws.title = "Referral Performance"

    # --------------------------------------------------------
    # STYLES
    # --------------------------------------------------------

    dark_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="5B9BD5")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    total_fill = PatternFill("solid", fgColor="E2F0D9")

    title_font = Font(color="FFFFFF", bold=True, size=18)
    section_font = Font(color="FFFFFF", bold=True, size=12)
    header_font = Font(bold=True)
    total_font = Font(bold=True)

    thin_side = Side(style="thin", color="B7B7B7")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    months_label = ", ".join(selected_months).upper()

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    ws.merge_cells("A1:L1")
    ws["A1"] = (
        "EMPLOYEE REFERRAL PERFORMANCE REPORT "
        f"- {months_label} {selected_year}"
    )
    ws["A1"].fill = dark_fill
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 32

    # --------------------------------------------------------
    # SUBTITLE
    # --------------------------------------------------------

    ws.merge_cells("A2:L2")
    ws["A2"] = (
        "Performance calculated using Updated Date | "
        "Duplicate and Not Enrolled records excluded by default | "
        "Percentages shown as share of grand total"
    )
    ws["A2"].alignment = center
    ws["A2"].font = Font(italic=True, size=10)

    current_row = 4

    # ========================================================
    # SECTION TITLE FUNCTION (AUTO SPAN)
    # ========================================================

    def add_section_title(title, row, span):
        span = max(1, int(span))

        ws.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=span
        )
        cell = ws.cell(row=row, column=1)
        cell.value = title
        cell.fill = section_fill
        cell.font = section_font
        cell.alignment = left
        ws.row_dimensions[row].height = 24
        return row + 1

    # ========================================================
    # WRITE DATAFRAME FUNCTION
    # ========================================================

    def write_dataframe(dataframe, start_row, add_total=False):
        dataframe = dataframe.copy()

        # Precompute which columns are numeric so we can apply
        # a numeric format even when the header has no keyword.
        numeric_columns = set()
        for col in dataframe.columns:
            if pd.api.types.is_numeric_dtype(dataframe[col]):
                numeric_columns.add(col)

        # ---- HEADER ----
        for col_index, column in enumerate(dataframe.columns, start=1):
            cell = ws.cell(row=start_row, column=col_index)
            cell.value = str(column)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = center

        # ---- DATA ----
        for row_index, row_data in enumerate(
            dataframe.itertuples(index=False),
            start=start_row + 1
        ):
            for col_index, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_index, column=col_index)

                if safe_value(value):
                    value = ""
                elif isinstance(value, np.integer):
                    value = int(value)
                elif isinstance(value, np.floating):
                    value = float(value)

                cell.value = value
                cell.border = border

                column_name = str(dataframe.columns[col_index - 1])

                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.alignment = center
                else:
                    cell.alignment = left

                # ---- Number format rules (order matters) ----

                # 1. Codes / Passbook → text
                if "Code" in column_name or "Passbook" in column_name:
                    cell.number_format = '@'

                # 2. Percentage columns
                elif "%" in column_name:
                    cell.number_format = '0.00"%"'

                # 3. Amount columns (keyword)
                elif "Amount" in column_name:
                    cell.number_format = '#,##0.00'

                # 4. Count / Rank / Employees / Value
                elif (
                    "Count" in column_name
                    or column_name == "Employees"
                    or column_name == "Value"
                    or column_name == "Rank"
                ):
                    cell.number_format = '#,##0'

                # 5. Fallback: any other numeric column → comma format
                #    This is what fixes Section 5 scheme columns
                #    (headers are scheme names, not "Amount")
                elif column_name in numeric_columns:
                    cell.number_format = '#,##0.00'

        end_row = start_row + len(dataframe)

        # ---- TOTAL ROW ----
        if add_total:
            total_row = end_row + 1

            total_cell = ws.cell(row=total_row, column=1)
            total_cell.value = "TOTAL"
            total_cell.font = total_font
            total_cell.fill = total_fill
            total_cell.border = border
            total_cell.alignment = center

            for col_index, column in enumerate(dataframe.columns, start=1):
                cell = ws.cell(row=total_row, column=col_index)
                cell.fill = total_fill
                cell.border = border
                cell.font = total_font

                if col_index == 1:
                    continue

                column_name = str(column)

                if "%" in column_name:
                    cell.value = 100.00
                    cell.number_format = '0.00"%"'
                    cell.alignment = center
                    continue

                if pd.api.types.is_numeric_dtype(dataframe[column]):
                    col_letter = get_column_letter(col_index)
                    cell.value = (
                        f"=SUM({col_letter}{start_row + 1}:"
                        f"{col_letter}{end_row})"
                    )
                    cell.alignment = center

                    if "Amount" in column_name:
                        cell.number_format = '#,##0.00'
                    elif "Count" in column_name or column_name == "Employees":
                        cell.number_format = '#,##0'
                    else:
                        # Fallback for scheme-name columns in Section 5
                        cell.number_format = '#,##0.00'
                else:
                    cell.value = ""

            return total_row + 2

        return end_row + 2

    # ========================================================
    # 1. MONTH-WISE PERFORMANCE
    # ========================================================

    monthly_export = monthly.copy().rename(columns={
        "Report Month": "Month",
        "Referral_Count": "Referral Count",
        "Referral_Amount": "Referral Amount"
    })

    monthly_desired = [
        "Month",
        "Referral Count",
        "Referral Count %",
        "Referral Amount",
        "Referral Amount %"
    ]
    monthly_export = select_existing_columns(
        monthly_export, monthly_desired
    )

    current_row = add_section_title(
        "1. MONTH-WISE PERFORMANCE",
        current_row,
        span=len(monthly_export.columns)
    )

    current_row = write_dataframe(
        monthly_export,
        current_row,
        add_total=True
    )

    # ========================================================
    # 2. SCHEME-WISE PERFORMANCE
    # ========================================================

    scheme_export = scheme_report.copy().rename(columns={
        "Total_Count": "Total Count",
        "Total_Amount": "Total Amount"
    })

    scheme_desired = [
        "Scheme Name",
        "Total Count",
        "Total Count %",
        "Total Amount",
        "Total Amount %"
    ]
    scheme_export = select_existing_columns(
        scheme_export, scheme_desired
    )

    current_row = add_section_title(
        "2. SCHEME-WISE PERFORMANCE",
        current_row,
        span=len(scheme_export.columns)
    )

    current_row = write_dataframe(
        scheme_export,
        current_row,
        add_total=True
    )

    # ========================================================
    # 3. BRANCH PERFORMANCE (MONTH-WISE)
    # ========================================================

    branch_export = branch_report.copy().rename(columns={
        "Total_Count": "Total Count",
        "Total_Amount": "Total Amount"
    })

    month_block = build_month_column_order(
        selected_months,
        ["Count", "Amount"]
    )

    branch_desired = (
        ["Branch"]
        + month_block
        + [
            "Total Count", "Total Count %",
            "Total Amount", "Total Amount %",
            "Employees"
        ]
    )
    branch_export = select_existing_columns(
        branch_export, branch_desired
    )

    current_row = add_section_title(
        "3. BRANCH PERFORMANCE (MONTH-WISE)",
        current_row,
        span=len(branch_export.columns)
    )

    current_row = write_dataframe(
        branch_export,
        current_row,
        add_total=True
    )

    # ========================================================
    # 4. INDIVIDUAL EMPLOYEE PERFORMANCE
    # ========================================================

    employee_export = employee_report.copy().rename(columns={
        "Total_Count": "Total Count",
        "Total_Amount": "Total Amount"
    })

    employee_month_block = build_month_column_order(
        selected_months,
        ["Count", "Amount"]
    )

    employee_desired = (
        ["Rank", "Employee Name", "Employee Code", "Branch"]
        + employee_month_block
        + [
            "Total Count", "Total Count %",
            "Total Amount", "Total Amount %"
        ]
    )
    employee_export = select_existing_columns(
        employee_export, employee_desired
    )

    current_row = add_section_title(
        "4. INDIVIDUAL EMPLOYEE PERFORMANCE",
        current_row,
        span=len(employee_export.columns)
    )

    current_row = write_dataframe(
        employee_export,
        current_row,
        add_total=True
    )

    # ========================================================
    # 5. EMPLOYEE + SCHEME PERFORMANCE (WIDE FORMAT)
    # ========================================================

    current_row = add_section_title(
        "5. EMPLOYEE + SCHEME PERFORMANCE "
        "(Enrollment Amount by Scheme)",
        current_row,
        span=len(employee_scheme_wide.columns)
    )

    current_row = write_dataframe(
        employee_scheme_wide,
        current_row,
        add_total=True
    )

    # ========================================================
    # COLUMN WIDTH
    # ========================================================

    merged_coordinates = set()
    for rng in ws.merged_cells.ranges:
        for row in ws[rng.coord]:
            for cell in row:
                merged_coordinates.add(cell.coordinate)

    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0

        for cell in column_cells:
            if cell.coordinate in merged_coordinates:
                continue
            if cell.value is None:
                continue
            try:
                max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            35
        )

    # ========================================================
    # FREEZE PANES
    # ========================================================

    ws.freeze_panes = "A4"

    # ========================================================
    # SHEET SETTINGS
    # ========================================================

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:3"

    # ========================================================
    # SAVE TO MEMORY
    # ========================================================

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output


# ============================================================
# FILE UPLOAD
# ============================================================

st.sidebar.header("📂 Upload Data")

uploaded_file = st.sidebar.file_uploader(
    "Upload Excel / CSV file",
    type=["xlsx", "xls", "csv"]
)

if uploaded_file is None:
    st.info(
        "👆 Please upload your employee referral Excel or CSV file."
    )
    st.stop()


# ============================================================
# READ FILE
# ============================================================

try:
    if uploaded_file.name.lower().endswith(".csv"):
        df = pd.read_csv(uploaded_file)
    else:
        df = pd.read_excel(uploaded_file)
except Exception as e:
    st.error(f"❌ Unable to read the file: {e}")
    st.stop()


# ============================================================
# CLEAN COLUMN NAMES
# ============================================================

df.columns = df.columns.astype(str).str.strip()


# ============================================================
# CHECK REQUIRED COLUMNS
# ============================================================

missing_columns = [
    column
    for column in EXPECTED_COLUMNS
    if column not in df.columns
]

if missing_columns:
    st.error("❌ Required columns are missing:")
    for column in missing_columns:
        st.write(f"- {column}")
    st.stop()


# ============================================================
# CLEAN TEXT COLUMNS
# ============================================================

text_columns = [
    "Customer Name",
    "Customer Phone",
    "Status",
    "Employee Name",
    "Referral Code",
    "Employee Phone",
    "Employee Code",
    "Branch",
    "True/False",
    "Scheme Name",
    "Scheme Passbook Number",
    "Category",
    "Month",
    "Not Enrolled",
    "Is Duplicate",
    "Duplicate Group",
    "Month Name"
]

for column in text_columns:
    df[column] = clean_text(df[column])


# ============================================================
# CLEAN AMOUNT COLUMNS
# ============================================================

df["Customer Enrollment Amount"] = clean_number(
    df["Customer Enrollment Amount"]
)


# ============================================================
# DATE CONVERSION
# ============================================================

st.subheader("📅 Date Check")

try:
    df["Updated Date"] = pd.to_datetime(
        df["Updated Date"],
        errors="coerce",
        format="mixed",
        dayfirst=True
    )
except TypeError:
    df["Updated Date"] = pd.to_datetime(
        df["Updated Date"],
        errors="coerce",
        dayfirst=True
    )

valid_dates = df["Updated Date"].notna().sum()
invalid_dates = df["Updated Date"].isna().sum()

d1, d2, d3 = st.columns(3)

d1.metric("Total Records", f"{len(df):,}")
d2.metric("Valid Dates", f"{valid_dates:,}")
d3.metric("Invalid Dates", f"{invalid_dates:,}")


# ============================================================
# DATE RANGE
# ============================================================

if valid_dates > 0:
    min_date = df["Updated Date"].min()
    max_date = df["Updated Date"].max()
    st.info(
        "📅 Date range detected: "
        f"**{min_date.strftime('%d-%m-%Y')}** "
        "to "
        f"**{max_date.strftime('%d-%m-%Y')}**"
    )


# ============================================================
# CREATE REPORT MONTH / YEAR
# ============================================================

df["Report Month"] = df["Updated Date"].dt.month_name()
df["Report Year"] = df["Updated Date"].dt.year


# ============================================================
# DEBUG DATE INFORMATION
# ============================================================

with st.expander("🔍 Check Date Data"):

    st.write("### Months detected from Updated Date")
    month_check = (
        df["Report Month"]
        .value_counts()
        .rename_axis("Month")
        .reset_index(name="Records")
    )
    st.dataframe(month_check, width="stretch", hide_index=True)

    st.write("### Years detected")
    year_check = (
        df["Report Year"]
        .value_counts()
        .rename_axis("Year")
        .reset_index(name="Records")
    )
    st.dataframe(year_check, width="stretch", hide_index=True)

    st.write("### Sample dates")
    st.dataframe(
        df[
            [
                "Updated Date",
                "Month",
                "Month Name",
                "Report Month",
                "Report Year"
            ]
        ].head(30),
        width="stretch",
        hide_index=True
    )


# ============================================================
# AVAILABLE YEARS
# ============================================================

available_years = sorted(df["Report Year"].dropna().unique())

if not available_years:
    st.error("❌ No valid dates were detected in Updated Date.")
    st.stop()


# ============================================================
# YEAR FILTER
# ============================================================

selected_year = st.sidebar.selectbox(
    "📅 Select Year",
    available_years,
    index=len(available_years) - 1
)


# ============================================================
# ALL RECORDS FOR SELECTED YEAR
# ============================================================

df_year = df[df["Report Year"] == selected_year].copy()

if df_year.empty:
    st.error(f"❌ No records found for year {selected_year}.")
    st.stop()


# ============================================================
# DETECT AVAILABLE MONTHS IN SELECTED YEAR
# ============================================================

available_months_in_year = [
    m for m in ALL_MONTHS if m in df_year["Report Month"].unique()
]

if not available_months_in_year:
    st.error(f"❌ No month data found for year {selected_year}.")
    st.stop()


# ============================================================
# MONTH SELECTION (DYNAMIC — ANY MONTH)
# ============================================================

st.sidebar.divider()
st.sidebar.header("📅 Month Selection")

default_selection = [
    m for m in DEFAULT_MONTHS
    if m in available_months_in_year
]

if not default_selection:
    default_selection = available_months_in_year

selected_months = st.sidebar.multiselect(
    "Select Months",
    options=available_months_in_year,
    default=default_selection
)

if not selected_months:
    st.warning("⚠️ Please select at least one month.")
    st.stop()

selected_months = [
    m for m in ALL_MONTHS if m in selected_months
]


# ============================================================
# PERIOD FILTER
# ============================================================

df_period = df_year[
    df_year["Report Month"].isin(selected_months)
].copy()

if df_period.empty:
    st.error(
        f"❌ No records found for "
        f"{', '.join(selected_months)} {selected_year}."
    )
    st.stop()


# ============================================================
# EXCLUDED RECORD COUNTS
# ============================================================

duplicate_count = is_true(df_period["Is Duplicate"]).sum()
not_enrolled_count = is_true(df_period["Not Enrolled"]).sum()


# ============================================================
# SIDEBAR FILTERS
# ============================================================

st.sidebar.divider()
st.sidebar.header("🔎 Filters")

branch_values = sorted(df_period["Branch"].dropna().unique())
selected_branches = st.sidebar.multiselect(
    "Branch",
    branch_values,
    default=branch_values
)

employee_values = sorted(
    df_period["Employee Name"].dropna().unique()
)
selected_employees = st.sidebar.multiselect(
    "Employee",
    employee_values,
    default=[]
)

scheme_values = sorted(
    df_period["Scheme Name"].dropna().unique()
)
selected_schemes = st.sidebar.multiselect(
    "Scheme",
    scheme_values,
    default=[]
)

category_values = sorted(
    df_period["Category"].dropna().unique()
)
selected_categories = st.sidebar.multiselect(
    "Category",
    category_values,
    default=[]
)

status_values = sorted(
    df_period["Status"].dropna().unique()
)
selected_status = st.sidebar.multiselect(
    "Status",
    status_values,
    default=[]
)

include_duplicates = st.sidebar.checkbox(
    "Include duplicate records",
    value=False
)

include_not_enrolled = st.sidebar.checkbox(
    "Include Not Enrolled records",
    value=False
)

with st.sidebar.expander("ℹ️ Excluded records (in period)"):
    st.write(f"Duplicates: **{duplicate_count:,}**")
    st.write(f"Not Enrolled: **{not_enrolled_count:,}**")


# ============================================================
# APPLY FILTERS
# ============================================================

filtered_df = df_period.copy()

if selected_branches:
    filtered_df = filtered_df[
        filtered_df["Branch"].isin(selected_branches)
    ]

if selected_employees:
    filtered_df = filtered_df[
        filtered_df["Employee Name"].isin(selected_employees)
    ]

if selected_schemes:
    filtered_df = filtered_df[
        filtered_df["Scheme Name"].isin(selected_schemes)
    ]

if selected_categories:
    filtered_df = filtered_df[
        filtered_df["Category"].isin(selected_categories)
    ]

if selected_status:
    filtered_df = filtered_df[
        filtered_df["Status"].isin(selected_status)
    ]


# ============================================================
# DUPLICATE FILTER
# ============================================================

filtered_df = filtered_df.assign(
    Duplicate_Record=is_true(filtered_df["Is Duplicate"])
)

if not include_duplicates:
    filtered_df = filtered_df[
        ~filtered_df["Duplicate_Record"]
    ].copy()


# ============================================================
# NOT ENROLLED FILTER
# ============================================================

filtered_df = filtered_df.assign(
    Not_Enrolled_Record=is_true(filtered_df["Not Enrolled"])
)

if not include_not_enrolled:
    filtered_df = filtered_df[
        ~filtered_df["Not_Enrolled_Record"]
    ].copy()


# ============================================================
# CHECK FILTER RESULT
# ============================================================

if filtered_df.empty:
    st.warning("⚠️ No records match the selected filters.")
    st.stop()


# ============================================================
# MAIN DASHBOARD
# ============================================================

months_label = " • ".join(selected_months)

st.divider()
st.header(f"📊 Referral Performance — {months_label} {selected_year}")


# ============================================================
# KPI
# ============================================================

total_referrals = len(filtered_df)

total_enrollment_amount = filtered_df[
    "Customer Enrollment Amount"
].sum()

total_employees = filtered_df["Employee Name"].nunique()
total_schemes = filtered_df["Scheme Name"].nunique()

k1, k2, k3, k4 = st.columns(4)

k1.metric("Total Referrals", f"{total_referrals:,}")
k2.metric(
    "Enrollment Amount",
    format_amount(total_enrollment_amount)
)
k3.metric("Employees", f"{total_employees:,}")
k4.metric("Schemes", f"{total_schemes:,}")


# ============================================================
# MONTHLY PERFORMANCE
# ============================================================

st.header("📅 Month-wise Performance")

monthly = (
    filtered_df
    .groupby("Report Month", as_index=False)
    .agg(
        Referral_Count=("Customer Name", "count"),
        Referral_Amount=("Customer Enrollment Amount", "sum")
    )
)

month_order_map = {m: i for i, m in enumerate(selected_months)}

monthly["Month_Order"] = monthly["Report Month"].map(month_order_map)

monthly = (
    monthly
    .sort_values("Month_Order")
    .drop(columns="Month_Order")
    .reset_index(drop=True)
)

monthly = add_percentage(
    monthly,
    value_columns=[
        "Referral_Count",
        "Referral_Amount"
    ]
)

monthly_display = monthly.copy()

monthly_display["Referral_Amount"] = monthly_display[
    "Referral_Amount"
].map(format_amount)

monthly_display["Referral Count %"] = monthly_display[
    "Referral Count %"
].map(format_percent)
monthly_display["Referral Amount %"] = monthly_display[
    "Referral Amount %"
].map(format_percent)

monthly_display = select_existing_columns(
    monthly_display,
    [
        "Report Month",
        "Referral_Count",
        "Referral Count %",
        "Referral_Amount",
        "Referral Amount %"
    ]
)

monthly_display.columns = [
    "Month",
    "Referral Count",
    "Referral Count %",
    "Referral Amount",
    "Referral Amount %"
]

st.dataframe(monthly_display, width="stretch", hide_index=True)


# ============================================================
# MONTHLY CHARTS
# ============================================================

st.subheader("📈 Referral Count by Month")
st.bar_chart(
    monthly.set_index("Report Month")[["Referral_Count"]]
)

st.subheader("💰 Enrollment Amount by Month")
st.bar_chart(
    monthly.set_index("Report Month")[["Referral_Amount"]]
)


# ============================================================
# SCHEME-WISE PERFORMANCE
# ============================================================

st.header("🏦 Scheme-wise Performance")

scheme_monthly = (
    filtered_df
    .groupby(["Scheme Name", "Report Month"], as_index=False)
    .agg(
        Referral_Count=("Customer Name", "count"),
        Referral_Amount=("Customer Enrollment Amount", "sum")
    )
)

scheme_pivot = (
    scheme_monthly
    .pivot(
        index="Scheme Name",
        columns="Report Month",
        values=["Referral_Count", "Referral_Amount"]
    )
    .fillna(0)
)

scheme_pivot.columns = [
    f"{metric}_{month}"
    for metric, month in scheme_pivot.columns
]

scheme_pivot = scheme_pivot.reset_index()

scheme_total = (
    filtered_df
    .groupby("Scheme Name", as_index=False)
    .agg(
        Total_Count=("Customer Name", "count"),
        Total_Amount=("Customer Enrollment Amount", "sum")
    )
)

scheme_report = scheme_pivot.merge(
    scheme_total,
    on="Scheme Name",
    how="left"
)

scheme_report = scheme_report.sort_values(
    "Total_Amount",
    ascending=False
).reset_index(drop=True)

scheme_report = add_percentage(
    scheme_report,
    value_columns=["Total_Count", "Total_Amount"]
)

st.dataframe(scheme_report, width="stretch", hide_index=True)


# ============================================================
# BRANCH PERFORMANCE (MONTH-WISE)
# ============================================================

st.header("🏢 Branch Performance (Month-wise)")

st.caption(
    "Each month column shows that branch's share (%) of the "
    "grand total for that month."
)

branch_count = (
    filtered_df
    .pivot_table(
        index="Branch",
        columns="Report Month",
        values="Customer Name",
        aggfunc="count",
        fill_value=0
    )
    .reset_index()
)

branch_count.columns.name = None

for month in selected_months:
    if month not in branch_count.columns:
        branch_count[month] = 0

branch_count = branch_count.rename(columns={
    m: f"{m} Count" for m in selected_months
})

branch_amount = (
    filtered_df
    .pivot_table(
        index="Branch",
        columns="Report Month",
        values="Customer Enrollment Amount",
        aggfunc="sum",
        fill_value=0
    )
    .reset_index()
)

branch_amount.columns.name = None

for month in selected_months:
    if month not in branch_amount.columns:
        branch_amount[month] = 0

branch_amount = branch_amount.rename(columns={
    m: f"{m} Amount" for m in selected_months
})

branch_total = (
    filtered_df
    .groupby("Branch", as_index=False)
    .agg(
        Total_Count=("Customer Name", "count"),
        Total_Amount=("Customer Enrollment Amount", "sum"),
        Employees=("Employee Name", "nunique")
    )
)

branch_report = (
    branch_count
    .merge(branch_amount, on="Branch", how="outer")
    .merge(branch_total, on="Branch", how="left")
    .fillna(0)
)

branch_report = add_percentage(
    branch_report,
    value_columns=[
        "Total_Count",
        "Total_Amount"
    ]
)

branch_report = month_percent_columns(
    branch_report,
    months=selected_months,
    month_prefixes=["Count", "Amount"]
)

branch_report = branch_report.sort_values(
    "Total_Amount",
    ascending=False
).reset_index(drop=True)

month_block = build_month_column_order(
    selected_months,
    ["Count", "Amount"]
)

branch_desired = (
    ["Branch"]
    + month_block
    + [
        "Total_Count", "Total Count %",
        "Total_Amount", "Total Amount %",
        "Employees"
    ]
)

branch_report = select_existing_columns(
    branch_report, branch_desired
)

branch_display = branch_report.copy()

amount_cols = [
    c for c in branch_display.columns
    if "Amount" in c and "%" not in c
]

for col in amount_cols:
    if col in branch_display.columns:
        branch_display[col] = branch_display[col].map(format_amount)

percent_cols = [c for c in branch_display.columns if c.endswith("%")]
for col in percent_cols:
    branch_display[col] = branch_display[col].map(format_percent)

branch_display.columns = [
    c.replace("_", " ") for c in branch_display.columns
]

st.dataframe(branch_display, width="stretch", hide_index=True)


# ============================================================
# INDIVIDUAL EMPLOYEE PERFORMANCE
# ============================================================

st.header("👤 Individual Employee Performance")

employee_count = (
    filtered_df
    .pivot_table(
        index=["Employee Name", "Employee Code", "Branch"],
        columns="Report Month",
        values="Customer Name",
        aggfunc="count",
        fill_value=0
    )
    .reset_index()
)

employee_count.columns.name = None

for month in selected_months:
    if month not in employee_count.columns:
        employee_count[month] = 0

employee_count = employee_count.rename(columns={
    m: f"{m} Count" for m in selected_months
})

employee_amount = (
    filtered_df
    .pivot_table(
        index=["Employee Name", "Employee Code", "Branch"],
        columns="Report Month",
        values="Customer Enrollment Amount",
        aggfunc="sum",
        fill_value=0
    )
    .reset_index()
)

employee_amount.columns.name = None

for month in selected_months:
    if month not in employee_amount.columns:
        employee_amount[month] = 0

employee_amount = employee_amount.rename(columns={
    m: f"{m} Amount" for m in selected_months
})

employee_total = (
    filtered_df
    .groupby(
        ["Employee Name", "Employee Code", "Branch"],
        as_index=False
    )
    .agg(
        Total_Count=("Customer Name", "count"),
        Total_Amount=("Customer Enrollment Amount", "sum")
    )
)

employee_report = (
    employee_count
    .merge(
        employee_amount,
        on=["Employee Name", "Employee Code", "Branch"],
        how="outer"
    )
    .merge(
        employee_total,
        on=["Employee Name", "Employee Code", "Branch"],
        how="left"
    )
    .fillna(0)
)

employee_report = add_percentage(
    employee_report,
    value_columns=["Total_Count", "Total_Amount"]
)

employee_report = employee_report.sort_values(
    "Total_Amount",
    ascending=False
).reset_index(drop=True)

employee_report.insert(0, "Rank", employee_report.index + 1)

employee_month_block = build_month_column_order(
    selected_months,
    ["Count", "Amount"]
)

employee_desired = (
    ["Rank", "Employee Name", "Employee Code", "Branch"]
    + employee_month_block
    + [
        "Total_Count", "Total Count %",
        "Total_Amount", "Total Amount %"
    ]
)

employee_report = select_existing_columns(
    employee_report, employee_desired
)

st.dataframe(employee_report, width="stretch", hide_index=True)


# ============================================================
# TOP PERFORMERS
# ============================================================

st.subheader("🏆 Top Performers (by Enrollment Amount)")

top_performers = employee_report.nlargest(3, "Total_Amount")[
    [
        "Rank",
        "Employee Name",
        "Branch",
        "Total_Count",
        "Total_Amount",
        "Total Amount %"
    ]
]

top_performers_display = top_performers.copy()
top_performers_display["Total_Amount"] = top_performers_display[
    "Total_Amount"
].map(format_amount)
top_performers_display["Total Amount %"] = top_performers_display[
    "Total Amount %"
].map(format_percent)

top_performers_display.columns = [
    "Rank",
    "Employee Name",
    "Branch",
    "Referrals",
    "Enrollment Amount",
    "Share %"
]

st.dataframe(
    top_performers_display,
    width="stretch",
    hide_index=True
)


# ============================================================
# EMPLOYEE + SCHEME PERFORMANCE (WIDE — schemes as columns)
# ============================================================

st.header("👤🏦 Employee + Scheme Performance (Wide View)")

st.caption(
    "Each row = one employee | Each scheme = its own column | "
    "Values = Enrollment Amount | % = share of that scheme's total"
)

# ---- Build the wide table (with per-scheme % and Grand Total %) ----
employee_scheme_wide = build_employee_scheme_wide(filtered_df)

# ---- Detect scheme columns (all non-metadata, non-% columns) ----
non_scheme = {
    "Rank", "Employee Name", "Employee Code", "Branch",
    "Grand Total", "Grand Total %"
}
scheme_columns = [
    c for c in employee_scheme_wide.columns
    if c not in non_scheme and not c.endswith("%")
]

# ---- Display version: format amounts and percents ----
employee_scheme_display = employee_scheme_wide.copy()

amount_cols = scheme_columns + ["Grand Total"]
for col in amount_cols:
    if col in employee_scheme_display.columns:
        employee_scheme_display[col] = employee_scheme_display[col].map(
            format_amount
        )

for col in employee_scheme_display.columns:
    if col.endswith("%"):
        employee_scheme_display[col] = employee_scheme_display[col].map(
            format_percent
        )

st.dataframe(
    employee_scheme_display,
    width="stretch",
    hide_index=True
)


# ============================================================
# INDIVIDUAL EMPLOYEE DETAIL
# ============================================================

st.header("🔎 Individual Employee Detail")

employee_list = sorted(
    filtered_df["Employee Name"].dropna().unique()
)

selected_employee = st.selectbox(
    "Select Employee",
    employee_list
)

employee_detail = filtered_df[
    filtered_df["Employee Name"] == selected_employee
].copy()

emp_count = len(employee_detail)
emp_amount = employee_detail["Customer Enrollment Amount"].sum()
emp_schemes = employee_detail["Scheme Name"].nunique()

overall_count = len(filtered_df)
overall_amount = filtered_df["Customer Enrollment Amount"].sum()

emp_count_pct = (
    (emp_count / overall_count) * 100
    if overall_count > 0 else 0
)
emp_amount_pct = (
    (emp_amount / overall_amount) * 100
    if overall_amount > 0 else 0
)

e1, e2, e3 = st.columns(3)

e1.metric(
    "Referrals",
    f"{emp_count:,}",
    f"{emp_count_pct:.2f}% of total"
)
e2.metric(
    "Enrollment Amount",
    format_amount(emp_amount),
    f"{emp_amount_pct:.2f}% of total"
)
e3.metric(
    "Schemes",
    f"{emp_schemes:,}"
)


# ============================================================
# SELECTED EMPLOYEE SCHEME DETAIL
# ============================================================

st.subheader(f"📋 {selected_employee} — Scheme Detail")

employee_detail_summary = (
    employee_detail
    .groupby(["Report Month", "Scheme Name"], as_index=False)
    .agg(
        Referral_Count=("Customer Name", "count"),
        Referral_Amount=("Customer Enrollment Amount", "sum")
    )
)

employee_detail_summary = add_percentage(
    employee_detail_summary,
    value_columns=[
        "Referral_Count",
        "Referral_Amount"
    ]
)

st.dataframe(
    employee_detail_summary,
    width="stretch",
    hide_index=True
)


# ============================================================
# CUSTOMER LEVEL DATA
# ============================================================

with st.expander("👁️ View Customer-level Records"):

    detail_columns = [
        "Updated Date",
        "Customer Name",
        "Customer Phone",
        "Customer Enrollment Amount",
        "Status",
        "Employee Name",
        "Referral Code",
        "Employee Phone",
        "Employee Code",
        "Branch",
        "Scheme Name",
        "Scheme Passbook Number",
        "Category",
        "Not Enrolled",
        "Is Duplicate",
        "Duplicate Group"
    ]

    available_columns = [
        c for c in detail_columns
        if c in employee_detail.columns
    ]

    st.dataframe(
        employee_detail[available_columns],
        width="stretch",
        hide_index=True
    )


# ============================================================
# SINGLE-SHEET FORMATTED EXCEL DOWNLOAD
# ============================================================

st.header("📥 Download Report")

excel_file = create_formatted_excel(
    monthly=monthly,
    scheme_report=scheme_report,
    employee_report=employee_report,
    employee_scheme_wide=employee_scheme_wide,
    branch_report=branch_report,
    selected_year=selected_year,
    selected_months=selected_months
)

months_file = "_".join(selected_months)

st.download_button(
    label="📥 Download Formatted Excel Report",
    data=excel_file,
    file_name=(
        f"Employee_Referral_Performance_"
        f"{months_file}_{selected_year}.xlsx"
    ),
    mime=(
        "application/vnd.openxmlformats-officedocument"
        ".spreadsheetml.sheet"
    )
)


# ============================================================
# FILTERED CSV
# ============================================================

csv_data = filtered_df.to_csv(index=False).encode("utf-8")

st.download_button(
    label="📥 Download Filtered Data CSV",
    data=csv_data,
    file_name=f"Filtered_Referral_Data_{selected_year}.csv",
    mime="text/csv"
)


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "Employee Referral Performance Dashboard | "
    "Date source: Updated Date | "
    "Duplicate records excluded by default | "
    "Percentages shown as share of grand total"
)