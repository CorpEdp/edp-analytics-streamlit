import io
import re
from datetime import datetime, timedelta, date

import numpy as np
import pandas as pd
import streamlit as st

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# PAGE CONFIG
# ============================================================
st.set_page_config(
    page_title="E-GOLD APP Agent Call Summary",
    page_icon="📊",
    layout="wide",
)


# ============================================================
# SESSION STATE INITIALIZATION
# ============================================================
if 'previous_filter_type' not in st.session_state:
    st.session_state.previous_filter_type = "Daily"
if 'filter_changed' not in st.session_state:
    st.session_state.filter_changed = False
if 'manual_mapping' not in st.session_state:
    st.session_state.manual_mapping = None
if 'mapping_applied' not in st.session_state:
    st.session_state.mapping_applied = False


# ============================================================
# REQUIRED COLUMNS - Now all optional with fallbacks
# ============================================================
OPTIONAL_COLUMNS = [
    "Agent Name",
    "Name",
    "Team",
    "Phone Number",
    "Call Start Time",
    "Process Name",
    "DID",
    "Dialer Status",
    "Disposition",
    "Sub Disposition",
    "Call Type",
    "Dialing Mode",
    "Talk Time",
    "List Name",
    "Completed By",
    "Comments",
]

# Column name variations for auto-mapping
COLUMN_VARIATIONS = {
    "Agent Name": ["agent", "agentname", "agent_name", "agent id", "agentid", "rep", "representative", "agent name"],
    "Name": ["name", "customer", "customer name", "client", "client name", "fullname", "full name"],
    "Team": ["team", "group", "department", "dept"],
    "Phone Number": ["phone", "phonenumber", "phone_number", "mobile", "mobilenumber", "contact", "contact number", "phone number"],
    "Call Start Time": ["callstarttime", "call_start_time", "start time", "starttime", "call time", "datetime", "date time", "call start time"],
    "Process Name": ["process", "processname", "process_name", "campaign", "project", "process name"],
    "DID": ["did", "number", "dialed number", "called number"],
    "Dialer Status": ["dialerstatus", "dialer_status", "status", "call status", "dialer status"],
    "Disposition": ["disposition", "result", "outcome", "call result"],
    "Sub Disposition": ["subdisposition", "sub_disposition", "sub result", "subresult", "sub disposition"],
    "Call Type": ["calltype", "call_type", "type", "call category", "call type"],
    "Dialing Mode": ["dialingmode", "dialing_mode", "mode", "dial mode", "dialing mode"],
    "Talk Time": ["talktime", "talk_time", "duration", "call duration", "talk duration", "talk time"],
    "List Name": ["listname", "list_name", "list", "campaign list", "list name"],
    "Completed By": ["completedby", "completed_by", "completed", "completion", "completed by"],
    "Comments": ["comments", "notes", "remarks", "comment"],
}


# ============================================================
# HELPERS
# ============================================================
def clean_text(value):
    if pd.isna(value):
        return ""

    text = str(value)
    text = text.replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def normalize_column_name(value):
    text = clean_text(value)
    return re.sub(r"[^a-z0-9]", "", text.lower())


def find_matching_column(column_name, available_columns):
    """Find matching column based on variations"""
    col_normalized = normalize_column_name(column_name)
    
    # First try exact match
    for avail_col in available_columns:
        if normalize_column_name(avail_col) == col_normalized:
            return avail_col
    
    # Then try variations
    for expected_col, variations in COLUMN_VARIATIONS.items():
        if column_name == expected_col:
            for avail_col in available_columns:
                avail_normalized = normalize_column_name(avail_col)
                for variation in variations:
                    if variation in avail_normalized or avail_normalized in variation:
                        return avail_col
    
    # Try partial match
    for avail_col in available_columns:
        avail_normalized = normalize_column_name(avail_col)
        if col_normalized in avail_normalized or avail_normalized in col_normalized:
            return avail_col
    
    return None


def detect_header_row(df):
    """Detect which row contains the headers by checking for known column patterns"""
    # Check first 15 rows for potential headers
    for row_idx in range(min(15, len(df))):
        row_values = df.iloc[row_idx].astype(str).tolist()
        row_text = " ".join(row_values).lower()
        
        # Check if this row contains any known column names
        found_patterns = 0
        for expected_col, variations in COLUMN_VARIATIONS.items():
            for variation in variations:
                if variation.lower() in row_text:
                    found_patterns += 1
                    break
        
        # If we found at least 3 patterns, this is likely the header row
        if found_patterns >= 3:
            return row_idx
        
        # Also check if any cell contains typical header keywords
        header_keywords = ["name", "phone", "time", "date", "call", "status", "type", "mode", "team"]
        found_keywords = 0
        for cell in row_values:
            cell_lower = cell.lower()
            for keyword in header_keywords:
                if keyword in cell_lower:
                    found_keywords += 1
                    break
        
        if found_keywords >= 2:
            return row_idx
    
    # If no header row found, return 0 as default
    return 0


def detect_data_start_row(df, header_row):
    """Detect where the actual data starts (after headers)"""
    # Start checking from the row after header
    for row_idx in range(header_row + 1, min(header_row + 20, len(df))):
        row_values = df.iloc[row_idx].astype(str).tolist()
        # Check if this row has any non-empty values
        non_empty = [x for x in row_values if x.strip() not in ["", "nan", "None", "null"]]
        if len(non_empty) > 0:
            return row_idx
    
    return header_row + 1


def read_excel_with_headers(uploaded_file):
    """Read Excel file with automatic header detection"""
    try:
        # First, read everything as raw data
        raw_df = pd.read_excel(uploaded_file, header=None)
        
        # Detect header row
        header_row = detect_header_row(raw_df)
        
        # Get headers from the detected row
        headers = raw_df.iloc[header_row].astype(str).tolist()
        # Clean up headers
        headers = [clean_text(h) if h and h.strip() not in ["", "nan", "None"] else f"Column_{i}" for i, h in enumerate(headers)]
        
        # Read the data again with the correct header
        data_start = detect_data_start_row(raw_df, header_row)
        
        # If data starts after header, skip the header row and use detected headers
        if data_start > header_row:
            df = pd.read_excel(uploaded_file, header=None, skiprows=data_start)
            df.columns = headers
        else:
            # If we couldn't detect data start, try reading with header
            df = pd.read_excel(uploaded_file, header=header_row)
        
        # Clean column names
        df.columns = [clean_text(col) if col and str(col).strip() not in ["", "nan", "None"] else f"Column_{i}" for i, col in enumerate(df.columns)]
        
        return df
    
    except Exception as e:
        st.error(f"Error reading Excel file: {str(e)}")
        # Fallback: try reading normally
        return pd.read_excel(uploaded_file)


def normalize_columns(df):
    df = df.copy()
    
    # First try to find matches using variations
    rename_map = {}
    available_columns = list(df.columns)
    
    for expected in OPTIONAL_COLUMNS:
        matching_col = find_matching_column(expected, available_columns)
        if matching_col and matching_col != expected:
            rename_map[matching_col] = expected
    
    # Apply the mapping
    if rename_map:
        df = df.rename(columns=rename_map)
    
    # Fallback: use the original normalized method for remaining columns
    lookup = {}
    for column in df.columns:
        lookup[normalize_column_name(column)] = column
    
    for expected in OPTIONAL_COLUMNS:
        key = normalize_column_name(expected)
        if key in lookup and expected not in df.columns:
            rename_map[lookup[key]] = expected
    
    if rename_map:
        df = df.rename(columns=rename_map)
    
    return df


def find_column(df, *possible_names):
    """Find a column by trying multiple possible names."""
    for name in possible_names:
        if name in df.columns:
            return name
        # Also try case-insensitive
        for col in df.columns:
            if col.lower() == name.lower():
                return col
    return None


# ============================================================
# DATE PARSER
# ============================================================
def parse_datetime_series(series):
    result = pd.to_datetime(
        series,
        errors="coerce",
        dayfirst=True,
    )

    numeric = pd.to_numeric(
        series,
        errors="coerce",
    )

    mask = (
        result.isna()
        & numeric.notna()
    )

    if mask.any():
        result.loc[mask] = pd.to_datetime(
            numeric.loc[mask],
            unit="D",
            origin="1899-12-30",
            errors="coerce",
        )

    return result


# ============================================================
# TALK TIME PARSER
# ============================================================
def parse_duration_seconds(value):
    if pd.isna(value):
        return 0.0

    if isinstance(value, pd.Timedelta):
        return max(0.0, value.total_seconds())

    if isinstance(value, timedelta):
        return max(0.0, value.total_seconds())

    if isinstance(value, (int, float, np.integer, np.floating)):
        numeric = float(value)
        if pd.isna(numeric):
            return 0.0
        # Excel time fraction
        if 0 <= numeric < 1:
            return numeric * 86400
        return numeric

    text = clean_text(value).lower()
    if not text:
        return 0.0

    # h/m/s
    h = re.search(r"(\d+(?:\.\d+)?)\s*h", text)
    m = re.search(r"(\d+(?:\.\d+)?)\s*m", text)
    s = re.search(r"(\d+(?:\.\d+)?)\s*s", text)

    if h or m or s:
        hours = float(h.group(1)) if h else 0
        minutes = float(m.group(1)) if m else 0
        seconds = float(s.group(1)) if s else 0
        return hours * 3600 + minutes * 60 + seconds

    # HH:MM:SS / MM:SS
    if ":" in text:
        try:
            parts = [float(x) for x in text.split(":")]
            if len(parts) == 3:
                return parts[0] * 3600 + parts[1] * 60 + parts[2]
            if len(parts) == 2:
                return parts[0] * 60 + parts[1]
        except Exception:
            pass

    try:
        return float(text)
    except Exception:
        return 0.0


# ============================================================
# FORMAT DURATION
# ============================================================
def format_duration(seconds):
    if pd.isna(seconds):
        return "0s"

    seconds = int(round(max(0, float(seconds))))
    hours, remainder = divmod(seconds, 3600)
    minutes, seconds = divmod(remainder, 60)

    if hours:
        return f"{hours}h {minutes:02d}m {seconds:02d}s"
    if minutes:
        return f"{minutes}m {seconds:02d}s"
    return f"{seconds}s"


# ============================================================
# DATE FILTER FUNCTIONS
# ============================================================
def get_date_range(filter_type, selected_date=None):
    if selected_date is None:
        selected_date = date.today()

    if filter_type == "Daily":
        start_date = selected_date
        end_date = selected_date
        date_label = selected_date.strftime("%d-%m-%Y")

    elif filter_type == "Weekly":
        start_date = selected_date - timedelta(days=selected_date.weekday())
        end_date = start_date + timedelta(days=6)
        date_label = f"{start_date.strftime('%d-%m-%Y')} to {end_date.strftime('%d-%m-%Y')}"

    elif filter_type == "Monthly":
        start_date = selected_date.replace(day=1)
        if selected_date.month == 12:
            end_date = selected_date.replace(year=selected_date.year + 1, month=1, day=1) - timedelta(days=1)
        else:
            end_date = selected_date.replace(month=selected_date.month + 1, day=1) - timedelta(days=1)
        date_label = start_date.strftime("%B %Y")

    else:  # Custom range
        return None, None, "Custom Range"

    return start_date, end_date, date_label


def filter_data_by_date(df, filter_type, start_date=None, end_date=None):
    if df.empty:
        return df

    if "Date" not in df.columns and "Call Start Time" in df.columns:
        df["Date"] = df["Call Start Time"].dt.date

    if "Date" not in df.columns:
        return df

    if filter_type == "Custom Range" and start_date and end_date:
        mask = (df["Date"] >= start_date) & (df["Date"] <= end_date)
        return df[mask].copy()

    elif filter_type == "Daily" and start_date:
        mask = df["Date"] == start_date
        return df[mask].copy()

    elif filter_type == "Weekly" and start_date and end_date:
        mask = (df["Date"] >= start_date) & (df["Date"] <= end_date)
        return df[mask].copy()

    elif filter_type == "Monthly" and start_date and end_date:
        mask = (df["Date"] >= start_date) & (df["Date"] <= end_date)
        return df[mask].copy()

    return df


# ============================================================
# MANUAL COLUMN MAPPING
# ============================================================
def manual_column_mapping(raw_df):
    """Display manual column mapping interface"""
    st.warning("⚠️ Auto-mapping failed. Please manually map the required columns.")
    
    # Show available columns
    available_cols = list(raw_df.columns)
    
    st.write("### Available Columns in Your File:")
    st.write(available_cols)
    
    # Show a preview of the data
    st.write("### Data Preview (first 5 rows):")
    st.dataframe(raw_df.head(5), width="stretch")
    
    st.write("### Map Your Columns to Required Fields:")
    
    mapping = {}
    cols = st.columns(2)
    
    required_fields = [
        "Agent Name",
        "Name", 
        "Phone Number",
        "Call Start Time",
        "Disposition",
        "Call Type",
        "Dialing Mode",
        "Talk Time"
    ]
    
    for idx, field in enumerate(required_fields):
        with cols[idx % 2]:
            mapping[field] = st.selectbox(
                f"Select column for **{field}**",
                options=[""] + available_cols,
                key=f"map_{field}",
                help=f"Select the column that contains {field} data"
            )
    
    # Optional fields
    st.write("### Optional Fields (if available):")
    optional_fields = [
        "Team",
        "Process Name",
        "DID",
        "Dialer Status",
        "Sub Disposition",
        "List Name",
        "Completed By",
        "Comments"
    ]
    
    cols2 = st.columns(2)
    for idx, field in enumerate(optional_fields):
        with cols2[idx % 2]:
            mapping[field] = st.selectbox(
                f"Select column for **{field}** (optional)",
                options=[""] + available_cols,
                key=f"map_opt_{field}",
                help=f"Select the column that contains {field} data (optional)"
            )
    
    col1, col2, col3 = st.columns([1, 1, 1])
    with col2:
        if st.button("✅ Apply Manual Mapping", width="stretch"):
            return mapping
    
    return None


def apply_manual_mapping(raw_df, mapping):
    """Apply manual column mapping to the dataframe"""
    df = raw_df.copy()
    
    # Create rename mapping
    rename_map = {}
    for required_col, source_col in mapping.items():
        if source_col and source_col != required_col:
            rename_map[source_col] = required_col
    
    if rename_map:
        df = df.rename(columns=rename_map)
    
    # Ensure all required columns exist
    for col in OPTIONAL_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    
    return df


# ============================================================
# PREPARE DATA - Now more flexible
# ============================================================
def prepare_data(raw_df):
    df = normalize_columns(raw_df)
    
    # Track which columns are available
    available_cols = {}
    for col in OPTIONAL_COLUMNS:
        available_cols[col] = col in df.columns
    
    # Initialize missing columns with empty values
    for col in OPTIONAL_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    
    # Keep only the columns we need
    df = df[OPTIONAL_COLUMNS].copy()
    
    # Clean text columns
    text_columns = [col for col in OPTIONAL_COLUMNS if col not in ["Call Start Time", "Talk Time"]]
    for column in text_columns:
        if column in df.columns:
            df[column] = df[column].map(clean_text)
    
    # Parse datetime if available
    if "Call Start Time" in df.columns:
        df["Call Start Time"] = parse_datetime_series(df["Call Start Time"])
    
    # Parse talk time if available
    if "Talk Time" in df.columns:
        df["Talk Time Seconds"] = df["Talk Time"].map(parse_duration_seconds)
    else:
        df["Talk Time Seconds"] = 0
    
    # Remove completely blank records (where all key fields are empty)
    useful = []
    if "Agent Name" in df.columns:
        useful.append("Agent Name")
    if "Name" in df.columns:
        useful.append("Name")
    if "Phone Number" in df.columns:
        useful.append("Phone Number")
    if "Call Start Time" in df.columns:
        useful.append("Call Start Time")
    
    if useful:
        valid_rows = df[useful].astype(str).apply(
            lambda row: any(x.strip() not in ["", "nan", "NaT"] for x in row),
            axis=1,
        )
        df = df[valid_rows].copy()
    
    # Add Date column if Call Start Time exists
    if "Call Start Time" in df.columns:
        df["Date"] = df["Call Start Time"].dt.date
    else:
        df["Date"] = None
    
    # Connected = positive talk time
    if "Talk Time Seconds" in df.columns:
        df["Connected"] = df["Talk Time Seconds"] > 0
    else:
        df["Connected"] = False
    
    # Store available columns info for reporting
    df.attrs['available_columns'] = available_cols
    
    return df


# ============================================================
# AGENT DAILY SUMMARY
# ============================================================
def build_agent_summary(df):
    valid = df[df["Call Start Time"].notna()].copy() if "Call Start Time" in df.columns else df.copy()
    
    columns = [
        "Date",
        "Agent Name",
        "Name",
        "Team",
        "Process Name",
        "Total Calls",
        "1st Call Time",
        "Last Call Time",
        "Talk Time",
        "Avg Talk Time",
        "Calls/Hour",
        "Talk %",
        "Unique DIDs",
    ]
    
    if valid.empty:
        return pd.DataFrame(columns=columns)
    
    # Determine grouping columns (only those that exist)
    group_columns = ["Date"]
    for col in ["Agent Name", "Name", "Team", "Process Name"]:
        if col in valid.columns and valid[col].notna().any():
            group_columns.append(col)
    
    rows = []
    
    for keys, group in valid.groupby(group_columns, dropna=False, sort=True):
        if len(group_columns) == 1:
            date = keys
            agent_id = ""
            name = ""
            team = ""
            process = ""
        else:
            date = keys[0]
            agent_id = keys[1] if len(keys) > 1 else ""
            name = keys[2] if len(keys) > 2 else ""
            team = keys[3] if len(keys) > 3 else ""
            process = keys[4] if len(keys) > 4 else ""
        
        first_call = group["Call Start Time"].min() if "Call Start Time" in group.columns else None
        last_call = group["Call Start Time"].max() if "Call Start Time" in group.columns else None
        
        total_calls = len(group)
        unique_dids = group["DID"].nunique() if "DID" in group.columns else 0
        
        talk_seconds = group["Talk Time Seconds"].sum() if "Talk Time Seconds" in group.columns else 0
        connected = int(group["Connected"].sum()) if "Connected" in group.columns else 0
        
        avg_talk = talk_seconds / connected if connected > 0 else 0
        
        if first_call and last_call:
            observed_seconds = (last_call - first_call).total_seconds()
            observed_seconds = max(0, observed_seconds)
            
            calls_per_hour = total_calls / (observed_seconds / 3600) if observed_seconds > 0 else np.nan
            talk_percent = talk_seconds / observed_seconds * 100 if observed_seconds > 0 else np.nan
        else:
            calls_per_hour = np.nan
            talk_percent = np.nan
        
        rows.append({
            "Date": date,
            "Agent Name": agent_id,
            "Name": name,
            "Team": team,
            "Process Name": process,
            "Total Calls": total_calls,
            "1st Call Time": first_call,
            "Last Call Time": last_call,
            "Talk Time": format_duration(talk_seconds),
            "Avg Talk Time": format_duration(avg_talk),
            "Calls/Hour": round(calls_per_hour, 2) if pd.notna(calls_per_hour) else np.nan,
            "Talk %": round(talk_percent, 2) if pd.notna(talk_percent) else np.nan,
            "Unique DIDs": unique_dids,
        })
    
    return pd.DataFrame(rows, columns=columns)


# ============================================================
# CATEGORY SUMMARY
# ============================================================
def build_category_summary(df, category_column):
    if df.empty or category_column not in df.columns:
        return pd.DataFrame()
    
    working = df.copy()
    
    # Use Name if available, otherwise use Agent Name
    working["Summary Name"] = working.get("Name", "").replace("", np.nan).fillna(working.get("Agent Name", ""))
    
    # Get all unique categories
    categories = []
    for value in working[category_column]:
        value = clean_text(value)
        if value and value not in categories:
            categories.append(value)
    
    if not categories:
        return pd.DataFrame({"Name": ["No data found"]})
    
    rows = []
    agent_names = working["Summary Name"].drop_duplicates().tolist()
    
    for name in agent_names:
        group = working[working["Summary Name"] == name]
        total = len(group)
        
        row = {"Name": name}
        
        for category in categories:
            count = int(group[category_column].map(clean_text).eq(category).sum())
            percent = (count / total * 100) if total > 0 else 0
            row[category] = count
            row[f"{category} %"] = round(percent, 2)
        
        row["Total"] = total
        rows.append(row)
    
    # Add grand total
    grand_total = len(working)
    total_row = {"Name": "Total"}
    
    for category in categories:
        count = int(working[category_column].map(clean_text).eq(category).sum())
        percent = (count / grand_total * 100) if grand_total > 0 else 0
        total_row[category] = count
        total_row[f"{category} %"] = round(percent, 2)
    
    total_row["Total"] = grand_total
    rows.append(total_row)
    
    # Create ordered columns
    ordered = ["Name"]
    for category in categories:
        ordered.extend([category, f"{category} %"])
    ordered.append("Total")
    
    df_result = pd.DataFrame(rows)
    existing_columns = [col for col in ordered if col in df_result.columns]
    
    return df_result[existing_columns]


# ============================================================
# TREND ANALYSIS - Updated for Weekly and Monthly aggregation with Agent Name and Name
# ============================================================
def build_trend_summary(df, filter_type="Daily"):
    """
    Build trend summary based on filter type:
    - Daily: Shows daily breakdown by agent
    - Weekly: Shows weekly breakdown by agent
    - Monthly: Shows monthly breakdown by agent
    """
    if df.empty:
        return pd.DataFrame()
    
    # Check if required columns exist
    if "Date" not in df.columns:
        return pd.DataFrame()
    
    # Create a copy of the dataframe with date as datetime
    df_copy = df.copy()
    df_copy["Date"] = pd.to_datetime(df_copy["Date"])
    
    # Determine grouping based on filter type
    if filter_type == "Daily":
        # Daily grouping
        df_copy["Period"] = df_copy["Date"].dt.date
        sort_column = "Period"
        label_format = "%d-%m-%Y"
        
    elif filter_type == "Weekly":
        # Weekly grouping
        df_copy["Year"] = df_copy["Date"].dt.year
        df_copy["Week"] = df_copy["Date"].dt.isocalendar().week
        df_copy["Period"] = df_copy["Year"].astype(str) + "-W" + df_copy["Week"].astype(str).str.zfill(2)
        sort_column = "Period"
        label_format = None
        
    elif filter_type == "Monthly":
        # Monthly grouping
        df_copy["Period"] = df_copy["Date"].dt.strftime("%B %Y")
        sort_column = "Period"
        label_format = None
        
    else:
        # Default to daily for custom range
        df_copy["Period"] = df_copy["Date"].dt.date
        sort_column = "Period"
        label_format = "%d-%m-%Y"
    
    # Get agent name columns
    agent_col = "Agent Name" if "Agent Name" in df_copy.columns else None
    name_col = "Name" if "Name" in df_copy.columns else None
    
    # Build grouping columns
    group_cols = ["Period"]
    
    # Add Year and Week for Weekly sorting
    if filter_type == "Weekly":
        group_cols = ["Year", "Week", "Period"]
    elif filter_type == "Monthly":
        group_cols = ["Period"]
    else:
        group_cols = ["Period"]
    
    # Add Agent Name if available
    if agent_col and agent_col in df_copy.columns:
        group_cols.append(agent_col)
    
    # Add Name if available and different from Agent Name
    if name_col and name_col in df_copy.columns and name_col != agent_col:
        group_cols.append(name_col)
    
    # Aggregate data
    agg_dict = {"Call Start Time": "count"} if "Call Start Time" in df_copy.columns else {}
    if "Connected" in df_copy.columns:
        agg_dict["Connected"] = "sum"
    if "Talk Time Seconds" in df_copy.columns:
        agg_dict["Talk Time Seconds"] = "sum"
    
    if not agg_dict:
        return pd.DataFrame()
    
    # Group by Period and Agent columns
    daily_stats = df_copy.groupby(group_cols).agg(agg_dict).reset_index()
    
    # Rename columns
    col_mapping = {"Call Start Time": "Total Calls"}
    daily_stats = daily_stats.rename(columns=col_mapping)
    
    # Sort by period and agent
    if filter_type == "Daily":
        daily_stats["Date_sort"] = pd.to_datetime(daily_stats["Period"], format="%d-%m-%Y", errors="coerce")
        sort_cols = ["Date_sort"]
        if agent_col:
            sort_cols.append(agent_col)
        if name_col and name_col != agent_col:
            sort_cols.append(name_col)
        daily_stats = daily_stats.sort_values(sort_cols)
        daily_stats = daily_stats.drop(columns=["Date_sort"])
    elif filter_type == "Weekly":
        sort_cols = ["Year", "Week"]
        if agent_col:
            sort_cols.append(agent_col)
        if name_col and name_col != agent_col:
            sort_cols.append(name_col)
        daily_stats = daily_stats.sort_values(sort_cols)
        daily_stats = daily_stats.drop(columns=["Year", "Week"])
    elif filter_type == "Monthly":
        # Sort by month
        month_order = ["January", "February", "March", "April", "May", "June", 
                      "July", "August", "September", "October", "November", "December"]
        daily_stats["Month"] = daily_stats["Period"].str.split(" ").str[0]
        daily_stats["Year"] = daily_stats["Period"].str.split(" ").str[1].astype(int)
        sort_cols = ["Year", "Month"]
        if agent_col:
            sort_cols.append(agent_col)
        if name_col and name_col != agent_col:
            sort_cols.append(name_col)
        daily_stats = daily_stats.sort_values(sort_cols)
        daily_stats = daily_stats.drop(columns=["Month", "Year"])
    
    # Calculate additional metrics
    if "Connected" in daily_stats.columns and "Total Calls" in daily_stats.columns:
        daily_stats["Connect Rate %"] = ((daily_stats["Connected"] / daily_stats["Total Calls"]) * 100).round(2)
    
    if "Talk Time Seconds" in daily_stats.columns:
        daily_stats["Avg Talk Time"] = daily_stats.apply(
            lambda row: format_duration(row["Talk Time Seconds"] / row["Connected"]) 
            if row.get("Connected", 0) > 0 else "0s",
            axis=1
        )
        daily_stats["Talk Time"] = daily_stats["Talk Time Seconds"].apply(format_duration)
        daily_stats = daily_stats.drop(columns=["Talk Time Seconds"])
    
    # Final column order - put Agent Name and Name after Period
    final_cols = ["Period"]
    if agent_col:
        final_cols.append(agent_col)
    if name_col and name_col != agent_col:
        final_cols.append(name_col)
    if "Total Calls" in daily_stats.columns:
        final_cols.append("Total Calls")
    if "Connected" in daily_stats.columns:
        final_cols.append("Connected")
    if "Talk Time" in daily_stats.columns:
        final_cols.append("Talk Time")
    if "Connect Rate %" in daily_stats.columns:
        final_cols.append("Connect Rate %")
    if "Avg Talk Time" in daily_stats.columns:
        final_cols.append("Avg Talk Time")
    
    # Rename Period column based on filter type
    if filter_type == "Daily":
        daily_stats = daily_stats.rename(columns={"Period": "Date"})
    elif filter_type == "Weekly":
        daily_stats = daily_stats.rename(columns={"Period": "Week"})
    elif filter_type == "Monthly":
        daily_stats = daily_stats.rename(columns={"Period": "Month"})
    else:
        daily_stats = daily_stats.rename(columns={"Period": "Date"})
    
    # Return only columns that exist
    return daily_stats[[col for col in final_cols if col in daily_stats.columns]]


# ============================================================
# EXCEL EXPORT FUNCTIONS
# ============================================================
def get_percent_columns(header_values):
    percent_cols = set()
    for idx, header in enumerate(header_values, start=1):
        if header is None:
            continue
        text = str(header).strip()
        if text.endswith("%"):
            percent_cols.add(idx)
    return percent_cols


def write_section_rows(
    final_ws,
    source_ws,
    start_row,
    include_header=True,
    header_row_index=1,
    total_row_index=None,
    header_fill=None,
    white_font=None,
    total_fill=None,
    bold_font=None,
    border=None,
):
    current_row = start_row
    
    header_values = [
        source_ws.cell(row=header_row_index, column=col).value
        for col in range(1, source_ws.max_column + 1)
    ]
    percent_cols = get_percent_columns(header_values)
    
    for source_row in range(1, source_ws.max_row + 1):
        for col in range(1, source_ws.max_column + 1):
            value = source_ws.cell(row=source_row, column=col).value
            cell = final_ws.cell(row=current_row, column=col)
            cell.value = value
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            if source_row == header_row_index:
                cell.fill = header_fill
                cell.font = white_font
            elif total_row_index is not None and source_row == total_row_index:
                cell.fill = total_fill
                cell.font = bold_font
            
            if source_row != header_row_index and col in percent_cols and isinstance(value, (int, float)):
                cell.number_format = '0.00"%"'
        
        current_row += 1
    
    return current_row


def create_single_sheet_excel(
    agent_summary,
    disposition_summary,
    call_type_summary,
    dialing_mode_summary,
    trend_summary=None,
    filter_info="",
    filter_type="Daily",
    raw_data=None
):
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Write all summary dataframes to temporary sheets first
        agent_summary.to_excel(writer, sheet_name="Agent", index=False)
        disposition_summary.to_excel(writer, sheet_name="Disposition", index=False)
        call_type_summary.to_excel(writer, sheet_name="Call Type", index=False)
        dialing_mode_summary.to_excel(writer, sheet_name="Dialing Mode", index=False)
        if trend_summary is not None and not trend_summary.empty:
            trend_summary.to_excel(writer, sheet_name="Trend", index=False)
        # Write Raw Data last (will be Sheet 2 after we move it)
        if raw_data is not None and not raw_data.empty:
            raw_data.to_excel(writer, sheet_name="Raw Data", index=False)
    
    output.seek(0)
    wb = load_workbook(output)
    
    # Get sheets
    agent_ws = wb["Agent"]
    disposition_ws = wb["Disposition"]
    call_type_ws = wb["Call Type"]
    dialing_ws = wb["Dialing Mode"]
    trend_ws = wb["Trend"] if "Trend" in wb.sheetnames else None
    raw_ws = wb["Raw Data"] if "Raw Data" in wb.sheetnames else None
    
    # Create Final Summary sheet (will be Sheet 1)
    final_ws = wb.create_sheet("Final Summary", 0)  # Insert at position 0
    
    # Style the Raw Data sheet (will be Sheet 2)
    if raw_ws is not None:
        # Move Raw Data to position 2
        wb.move_sheet("Raw Data", offset=2 - wb.sheetnames.index("Raw Data"))
        
        # Style headers
        header_fill = PatternFill("solid", fgColor="5B9BD5")
        white_font = Font(color="FFFFFF", bold=True)
        
        for col in range(1, raw_ws.max_column + 1):
            cell = raw_ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Auto-fit columns
        for column_cells in raw_ws.columns:
            column_index = column_cells[0].column
            maximum = 0
            for cell in column_cells:
                if cell.value is not None:
                    maximum = max(maximum, len(str(cell.value)))
            raw_ws.column_dimensions[get_column_letter(column_index)].width = min(max(maximum + 2, 10), 32)
        
        # Freeze the header row
        raw_ws.freeze_panes = "A2"
        
        # Add borders
        thin = Side(style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        
        for row in raw_ws.iter_rows():
            for cell in row:
                cell.border = border
    
    # Delete temporary summary sheets
    for sheet_name in ["Agent", "Disposition", "Call Type", "Dialing Mode", "Trend"]:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
    
    # Styles for final summary
    title_fill = PatternFill("solid", fgColor="17365D")
    section_fill = PatternFill("solid", fgColor="1F4E78")
    header_fill = PatternFill("solid", fgColor="5B9BD5")
    total_fill = PatternFill("solid", fgColor="E2F0D9")
    white_font = Font(color="FFFFFF", bold=True, size=12)
    title_font = Font(color="FFFFFF", bold=True, size=14)
    bold_font = Font(bold=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    # Calculate max columns
    max_cols = max(
        agent_ws.max_column,
        disposition_ws.max_column,
        call_type_ws.max_column,
        dialing_ws.max_column,
        14
    )
    
    # Title
    final_ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_cols)
    title_cell = final_ws.cell(row=1, column=1)
    title_cell.value = f"E-GOLD APP AGENT CALL SUMMARY - {filter_info}"
    title_cell.fill = title_fill
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    final_ws.row_dimensions[1].height = 25
    
    current_row = 3
    
    # Determine trend section title based on filter type
    if filter_type == "Daily":
        trend_title = "DAILY TREND SUMMARY"
    elif filter_type == "Weekly":
        trend_title = "WEEKLY TREND SUMMARY"
    elif filter_type == "Monthly":
        trend_title = "MONTHLY TREND SUMMARY"
    else:
        trend_title = "TREND SUMMARY"
    
    # Trend section
    if trend_ws is not None:
        final_ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
        cell = final_ws.cell(row=current_row, column=1)
        cell.value = trend_title
        cell.fill = section_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="left")
        current_row += 1
        
        current_row = write_section_rows(
            final_ws=final_ws,
            source_ws=trend_ws,
            start_row=current_row,
            header_row_index=1,
            total_row_index=None,
            header_fill=header_fill,
            white_font=white_font,
            total_fill=total_fill,
            bold_font=bold_font,
            border=border,
        )
        current_row += 2
    
    # Agent section
    final_ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    cell = final_ws.cell(row=current_row, column=1)
    cell.value = "E-GOLD APP AGENT DAILY CALL SUMMARY"
    cell.fill = section_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="left")
    current_row += 1
    
    current_row = write_section_rows(
        final_ws=final_ws,
        source_ws=agent_ws,
        start_row=current_row,
        header_row_index=1,
        total_row_index=None,
        header_fill=header_fill,
        white_font=white_font,
        total_fill=total_fill,
        bold_font=bold_font,
        border=border,
    )
    
    # Disposition section
    current_row += 2
    final_ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    cell = final_ws.cell(row=current_row, column=1)
    cell.value = "DISPOSITION SUMMARY"
    cell.fill = section_fill
    cell.font = white_font
    current_row += 1
    
    current_row = write_section_rows(
        final_ws=final_ws,
        source_ws=disposition_ws,
        start_row=current_row,
        header_row_index=1,
        total_row_index=disposition_ws.max_row,
        header_fill=header_fill,
        white_font=white_font,
        total_fill=total_fill,
        bold_font=bold_font,
        border=border,
    )
    
    # Call Type section
    current_row += 2
    final_ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    cell = final_ws.cell(row=current_row, column=1)
    cell.value = "CALL TYPE SUMMARY"
    cell.fill = section_fill
    cell.font = white_font
    current_row += 1
    
    current_row = write_section_rows(
        final_ws=final_ws,
        source_ws=call_type_ws,
        start_row=current_row,
        header_row_index=1,
        total_row_index=call_type_ws.max_row,
        header_fill=header_fill,
        white_font=white_font,
        total_fill=total_fill,
        bold_font=bold_font,
        border=border,
    )
    
    # Dialing Mode section
    current_row += 2
    final_ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    cell = final_ws.cell(row=current_row, column=1)
    cell.value = "DIALING MODE SUMMARY"
    cell.fill = section_fill
    cell.font = white_font
    current_row += 1
    
    current_row = write_section_rows(
        final_ws=final_ws,
        source_ws=dialing_ws,
        start_row=current_row,
        header_row_index=1,
        total_row_index=dialing_ws.max_row,
        header_fill=header_fill,
        white_font=white_font,
        total_fill=total_fill,
        bold_font=bold_font,
        border=border,
    )
    
    # Format cells
    final_ws.sheet_view.showGridLines = False
    final_ws.freeze_panes = "A5"
    
    # Column widths
    for column_cells in final_ws.columns:
        column_index = column_cells[0].column
        maximum = 0
        for cell in column_cells:
            if cell.value is not None:
                maximum = max(maximum, len(str(cell.value)))
        final_ws.column_dimensions[get_column_letter(column_index)].width = min(max(maximum + 2, 10), 32)
    
    # Date/time formats
    for row in final_ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            if isinstance(cell.value, datetime):
                if cell.column == 1:
                    cell.number_format = "dd-mm-yyyy"
                else:
                    cell.number_format = "h:mm AM/PM"
    
    final_output = io.BytesIO()
    wb.save(final_output)
    final_output.seek(0)
    return final_output


# ============================================================
# STREAMLIT UI
# ============================================================
st.title("📊 E-GOLD APP Agent Call Summary")
st.caption("Single-sheet management report with Agent, Disposition, Call Type and Dialing Mode summaries with date filtering.")


# ============================================================
# FILE UPLOAD
# ============================================================
uploaded_file = st.file_uploader(
    "📁 Upload Excel File",
    type=["xlsx", "xls"],
)

if uploaded_file is None:
    st.info("Please upload your call-detail Excel file.")
    st.stop()


# ============================================================
# READ FILE WITH HEADER DETECTION
# ============================================================
try:
    # Use the enhanced reading function that detects headers
    raw_df = read_excel_with_headers(uploaded_file)
except Exception as error:
    st.error(f"Excel reading error: {error}")
    st.stop()


# ============================================================
# COLUMN MAPPING - AUTO OR MANUAL
# ============================================================
# Show what was detected
with st.expander("📋 File Reading Details", expanded=False):
    st.write(f"**Total rows detected:** {len(raw_df)}")
    st.write(f"**Total columns detected:** {len(raw_df.columns)}")
    st.write("**Detected column names:**")
    st.write(list(raw_df.columns))
    st.write("**Data Preview (first 5 rows):**")
    st.dataframe(raw_df.head(5), width="stretch")

# Check if we have a saved manual mapping or need to auto-map
if st.session_state.manual_mapping is not None and st.session_state.mapping_applied:
    # Use saved manual mapping
    st.info("📌 Using previously applied manual mapping")
    df = apply_manual_mapping(raw_df.copy(), st.session_state.manual_mapping)
    mapping_used = "Manual (Saved)"
    
    # Show reset option
    if st.button("🔄 Reset Manual Mapping", help="Clear manual mapping and try auto-mapping again"):
        st.session_state.manual_mapping = None
        st.session_state.mapping_applied = False
        st.rerun()
else:
    # First try auto-mapping
    df_auto = normalize_columns(raw_df.copy())
    
    # Check if critical columns were found
    critical_columns = ["Agent Name", "Call Start Time", "Disposition"]
    found_critical = [col for col in critical_columns if col in df_auto.columns]
    mapping_successful = len(found_critical) >= 2  # At least 2 of 3 critical columns found

    if not mapping_successful:
        # Auto-mapping failed - show manual mapping
        st.warning("⚠️ Auto-mapping could not identify all required columns. Please map them manually.")
        
        mapping = manual_column_mapping(raw_df)
        
        if mapping is None:
            st.info("Please complete the column mapping to continue.")
            st.stop()
        
        # Save mapping to session state
        st.session_state.manual_mapping = mapping
        st.session_state.mapping_applied = True
        
        # Apply manual mapping
        df = apply_manual_mapping(raw_df, mapping)
        mapping_used = "Manual"
    else:
        df = df_auto
        mapping_used = "Auto"
        # Clear any saved manual mapping since auto worked
        st.session_state.manual_mapping = None
        st.session_state.mapping_applied = False
        st.success("✅ Auto-mapping successful! Columns have been mapped automatically.")


# ============================================================
# PREPARE DATA
# ============================================================
try:
    df = prepare_data(df)
except Exception as error:
    st.error(f"Error preparing data: {str(error)}")
    st.stop()

if df.empty:
    st.warning("No usable call records found after processing.")
    
    # Show the data preview to help diagnose
    with st.expander("📊 View Raw Data Preview", expanded=True):
        st.write("Raw data columns:", list(raw_df.columns))
        st.dataframe(raw_df.head(10), width="stretch")
    
    st.stop()

# Show mapping summary
with st.expander("📋 Column Mapping Summary", expanded=False):
    st.write(f"**Mapping Method:** {mapping_used}")
    st.write("**Columns found in data:**")
    
    found_cols = []
    missing_cols = []
    
    for col in OPTIONAL_COLUMNS:
        if col in df.columns:
            # Check if the column has any non-null values
            if not df[col].empty and df[col].notna().any():
                found_cols.append(col)
            else:
                missing_cols.append(col)
        else:
            missing_cols.append(col)
    
    col1, col2 = st.columns(2)
    with col1:
        st.success(f"✅ Found: {len(found_cols)} columns")
        for col in found_cols:
            st.write(f"• {col}")
    
    with col2:
        if missing_cols:
            st.warning(f"⚠️ Missing: {len(missing_cols)} columns")
            for col in missing_cols:
                st.write(f"• {col}")
        else:
            st.success("✅ All columns found!")


# ============================================================
# DATE FILTER UI WITH AUTO-REFRESH
# ============================================================
st.divider()
st.subheader("📅 Date Filter")

# Check if Date column exists
if "Date" not in df.columns or df["Date"].isna().all():
    st.warning("No date information available in the data.")
    st.stop()

min_date = df["Date"].min()
max_date = df["Date"].max()

if min_date is None or max_date is None:
    st.warning("No date information available in the data.")
    st.stop()

col1, col2, col3 = st.columns([2, 2, 1])

with col1:
    filter_type = st.radio(
        "Select Period",
        ["Daily", "Weekly", "Monthly", "Custom Range"],
        horizontal=True,
        index=0,
        key="filter_type_radio"
    )
    
    # Check if filter type changed
    if filter_type != st.session_state.previous_filter_type:
        st.session_state.filter_changed = True
        st.session_state.previous_filter_type = filter_type

with col2:
    if filter_type == "Custom Range":
        start_date = st.date_input("Start Date", value=min_date, min_value=min_date, max_value=max_date)
        end_date = st.date_input("End Date", value=max_date, min_value=min_date, max_value=max_date)
        # Format the date range for display
        date_label = f"{start_date.strftime('%d-%m-%Y')} to {end_date.strftime('%d-%m-%Y')}"
    else:
        selected_date = st.date_input("Select Date", value=max_date, min_value=min_date, max_value=max_date)
        start_date, end_date, date_label = get_date_range(filter_type, selected_date)

with col3:
    st.metric("Period", date_label)

# Apply filter
if filter_type == "Custom Range":
    filtered_df = filter_data_by_date(df, filter_type, start_date, end_date)
    filter_info = f"{start_date.strftime('%d-%m-%Y')} to {end_date.strftime('%d-%m-%Y')}"
else:
    filtered_df = filter_data_by_date(df, filter_type, start_date, end_date)
    filter_info = date_label

if filtered_df.empty:
    st.warning(f"No data available for the selected {filter_type.lower()} period.")
    st.stop()

st.success(f"Showing **{len(filtered_df):,}** records for **{filter_info}**")


# ============================================================
# SUMMARIES
# ============================================================
agent_summary = build_agent_summary(filtered_df)
disposition_summary = build_category_summary(filtered_df, "Disposition")
call_type_summary = build_category_summary(filtered_df, "Call Type")
dialing_mode_summary = build_category_summary(filtered_df, "Dialing Mode")

# Build trend summary with filter type for proper aggregation
trend_summary = build_trend_summary(filtered_df, filter_type)


# ============================================================
# AUTO-REFRESH TRIGGER
# ============================================================
# If filter type changed, trigger a rerun to refresh all data
if st.session_state.filter_changed:
    st.session_state.filter_changed = False
    st.rerun()


# ============================================================
# KPI
# ============================================================
total_calls = len(filtered_df)
connected_calls = int(filtered_df["Connected"].sum()) if "Connected" in filtered_df.columns else 0
talk_seconds = filtered_df["Talk Time Seconds"].sum() if "Talk Time Seconds" in filtered_df.columns else 0
avg_talk_seconds = talk_seconds / connected_calls if connected_calls else 0
unique_agents = filtered_df["Agent Name"].nunique() if "Agent Name" in filtered_df.columns else 0

st.divider()
st.subheader("📌 Overall Summary")

c1, c2, c3, c4, c5 = st.columns(5)
c1.metric("Total Calls", f"{total_calls:,}")
c2.metric("Connected", f"{connected_calls:,}")
c3.metric("Talk Time", format_duration(talk_seconds))
c4.metric("Avg Talk Time", format_duration(avg_talk_seconds))
c5.metric("Unique Agents", f"{unique_agents}")


# ============================================================
# DISPLAY SUMMARIES
# ============================================================
# Determine trend section title based on filter type
if filter_type == "Daily":
    trend_title = "📈 Daily Trend"
elif filter_type == "Weekly":
    trend_title = "📈 Weekly Trend"
elif filter_type == "Monthly":
    trend_title = "📈 Monthly Trend"
else:
    trend_title = "📈 Trend Summary"

if not trend_summary.empty and len(trend_summary) > 1:
    st.divider()
    st.subheader(trend_title)
    st.dataframe(trend_summary, width="stretch", hide_index=True)

st.divider()
st.subheader("1️⃣ E-GOLD APP Agent Daily Call Summary")
st.dataframe(agent_summary, width="stretch", hide_index=True)

st.divider()
st.subheader("2️⃣ Disposition Summary")
st.dataframe(disposition_summary, width="stretch", hide_index=True)

st.divider()
st.subheader("3️⃣ Call Type Summary")
st.dataframe(call_type_summary, width="stretch", hide_index=True)

st.divider()
st.subheader("4️⃣ Dialing Mode Summary")
st.dataframe(dialing_mode_summary, width="stretch", hide_index=True)


# ============================================================
# PROCESSED DATA
# ============================================================
st.divider()
with st.expander("🔎 Processed Data", expanded=False):
    st.write(f"Total processed records: **{len(filtered_df):,}**")
    
    # Show only columns that have data
    display_cols = []
    for col in OPTIONAL_COLUMNS:
        if col in filtered_df.columns:
            if not filtered_df[col].empty and filtered_df[col].notna().any():
                display_cols.append(col)
    
    if display_cols:
        st.dataframe(filtered_df[display_cols], width="stretch", hide_index=True)


# ============================================================
# CREATE EXCEL WITH TWO SHEETS
# ============================================================
try:
    excel_output = create_single_sheet_excel(
        agent_summary,
        disposition_summary,
        call_type_summary,
        dialing_mode_summary,
        trend_summary,
        filter_info,
        filter_type,
        filtered_df  # Pass the filtered raw data
    )
except Exception as error:
    st.error(f"Unable to create Excel report: {error}")
    st.stop()


# ============================================================
# DOWNLOAD
# ============================================================
st.divider()
st.subheader("📥 Download Final Report")

filename = (
    f"E-GOLD_APP_Agent_Summary_{filter_info.replace(' ', '_').replace('/', '-')}_"
    + datetime.now().strftime("%d-%m-%Y_%H-%M-%S")
    + ".xlsx"
)

st.download_button(
    label="📥 Download Report (Final Summary + Raw Data)",
    data=excel_output,
    file_name=filename,
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    width="stretch",
)

st.success("✅ Excel report contains two sheets: `Final Summary` (Sheet 1) and `Raw Data` (Sheet 2)")